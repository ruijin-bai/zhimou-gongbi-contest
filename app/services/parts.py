"""部位档案：把日志、照片、随手记、质控点、隐患按“部位”归集，形成每个部位的时间线与汇总。

- 主体数据全部由确定性查询给出（只读汇总，不改进度、审批、计量语义）；日志与工程量只计技术员已确认（submitted/approved）的日。
- 部位识别：P6 作业的 WBS 层级/名称里的墩台编号（如“3#墩”）优先，其次最末级 WBS 名称；日志里写的部位、照片地点、
  自由文本里的墩台编号也按同一规则归并；识别不出的归“未归部位”，不猜。
- 档案Agent（kind=site.archive）只做两件事：把写法各异的部位名归并到已有部位（如“三号墩东侧”→“3#墩”），
  以及写一段“部位概况”；概况里的数字必须出自统计，由数字守卫校验，不通过则改用规则生成的概况。
- 分包单位只看到本单位（同公司）范围内的记录；港湾岗位、领导看到全项目。
"""
from __future__ import annotations

import hashlib
import re
from collections import Counter, defaultdict
from datetime import date

from ..ai import report_agent
from ..ai.client import AIError, call_json
from ..config import settings
from ..db import all_, audit, dumps, loads, now_iso, one
from .. import roles as R
from . import p6
from .feed import scope_ids
from .media import geo_of
from . import act_names as _an

UNASSIGNED = "未归部位"
PROMPT_VERSION = "site-v1-20261006"
ACTIVE = ("open", "rectifying", "rejected")
CONFIRMED = ("submitted", "approved")
TYPE_NAME = {"photo": "照片", "progress": "施工与工程量", "qc": "质控点", "hazard": "安全质量隐患", "issue": "日志问题", "note": "随手记"}
PURPOSE_NAME = {"log": "日志照片", "safety": "随手拍隐患", "rectify": "整改照片", "ticket": "票据照片", "conceal": "隐蔽验收照片"}

# 3#墩 / 3号桥台 / 1#盖板涵 / 2#承台 → 统一写成 “编号#类型”
_PART_RE = re.compile(r"(\d+)\s*[#＃号]\s*(?:盖板|圆管|箱|拱|桥)?(墩|台|涵|承台|桩|梁)")


def part_from_text(text: str | None) -> str | None:
    m = _PART_RE.search(text or "")
    return f"{int(m.group(1))}#{m.group(2)}" if m else None


def activity_part(a: dict) -> str | None:
    """P6 作业所属部位：WBS 各级名称（由深到浅）里的墩台编号 → 作业名称里的编号 → 最末级 WBS 名称；里程碑不归部位。"""
    segs = [s.strip() for s in (a.get("wbs_path") or "").split("/") if s.strip()]
    for s in reversed(segs):
        p = part_from_text(s)
        if p:
            return p
    p = part_from_text(a.get("name"))
    if p:
        return p
    if a.get("is_milestone"):
        return None
    return segs[-1][:20] if segs else None


def sort_key(name: str):
    return (name == UNASSIGNED, [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", name)])


class Resolver:
    """把作业代码、写法各异的部位名、自由文本统一解析成部位名；解析不了返回 None。"""

    def __init__(self, conn, project_id: int):
        self.acts = {a["code"]: a for a in p6.current_activities(conn, project_id)}
        self.act_part = {c: activity_part(a) for c, a in self.acts.items()}
        self.known = {p for p in self.act_part.values() if p}
        self.alias = {r["key"]: (loads(r["value_json"], {}) or {}).get("to", "")
                      for r in all_(conn, "SELECT key, value_json FROM knowledge WHERE project_id=? AND kind='part_alias'", (project_id,))}

    def of_code(self, code: str | None) -> str | None:
        return self.act_part.get(code) if code else None

    def _known_in(self, text: str) -> str | None:
        hits = [k for k in self.known if len(k) >= 2 and k in text]
        return max(hits, key=len) if hits else None

    def of_label(self, raw: str | None) -> str | None:
        raw = (raw or "").strip()
        if not raw:
            return None
        return part_from_text(raw) or self.alias.get(raw) or self._known_in(raw) or None

    def of_text(self, text: str | None) -> str | None:
        text = (text or "").strip()
        return part_from_text(text) or self._known_in(text) if text else None


# ----------------------------------------------------------------------------- 归集

def _in(ids: list[int]) -> str:
    return ",".join("?" for _ in ids) or "NULL"


def _hhmm(iso: str | None) -> str:
    return iso[11:16] if iso and len(iso) >= 16 else ""


def _photo_view(m: dict, rs_note: str = "") -> dict:
    """照片视图：带拍摄时间与定位（水印相机的拍摄时间/坐标已入库）；没有拍摄时间的如实标注，改用上传时间。"""
    return {"id": m["id"], "thumb": f"/api/media/{m['id']}/thumb", "url": f"/api/media/{m['id']}/file",
            "taken_at": m.get("taken_at"), "uploaded_at": m.get("uploaded_at"), "geo": geo_of(m),
            "purpose": m.get("purpose") or "log", "purpose_name": PURPOSE_NAME.get(m.get("purpose") or "log", "照片"), "log_id": m.get("log_id")}


def collect(conn, user: dict) -> dict:
    """按用户可见范围把所有记录归到部位。返回 {parts: {名: {items, photos}}, rs, unresolved, ...}。"""
    pid = user["project_id"]
    ids = scope_ids(conn, user)  # 分包：本单位人员；其他：None
    rs = Resolver(conn, pid)
    people = {u["id"]: u for u in all_(conn, "SELECT id, name, company, wbs_scope FROM users WHERE project_id=?", (pid,))}
    name_of = lambda uid: people.get(uid, {}).get("name", "已离职")  # noqa: E731
    visible = (lambda uid: True) if ids is None else (lambda uid: uid in ids)
    parts: dict[str, dict] = defaultdict(lambda: {"items": [], "photos": set()})
    unresolved: Counter = Counter()

    def put(part: str, item: dict) -> None:
        parts[part or UNASSIGNED]["items"].append(item)

    # ---- 日志（日报）：施工内容、工程量、问题
    logs = all_(conn, "SELECT * FROM logs WHERE project_id=? AND period='day' AND status<>'returned' ORDER BY log_date, id", (pid,))
    prog_rows = defaultdict(list)
    conf_ids = [lg["id"] for lg in logs if lg["status"] in CONFIRMED]
    for r in all_(conn, f"SELECT * FROM log_progress WHERE voided_at IS NULL AND log_id IN ({_in(conf_ids)}) ORDER BY id", conf_ids):
        prog_rows[r["log_id"]].append(r)
    log_parts: dict[int, list[str]] = {}
    issue_part: dict[tuple, str] = {}
    log_by_id = {lg["id"]: lg for lg in logs}
    for lg in logs:
        c = loads(lg["content_json"], {}) or {}
        works = defaultdict(list)
        in_log: list[str] = []
        for w in c.get("works", []):
            p = rs.of_code(w.get("activity_code")) or rs.of_label(w.get("part")) or rs.of_text(w.get("desc"))
            if not p:
                raw = (w.get("part") or "").strip()
                if raw and visible(lg["author_id"]):
                    unresolved[raw[:20]] += 1
            works[p or UNASSIGNED].append({"process": w.get("process") or "", "desc": w.get("desc") or "", "activity_code": w.get("activity_code") or ""})
            if p and p not in in_log:
                in_log.append(p)
        progs = defaultdict(list)
        for r in prog_rows.get(lg["id"], []):
            p = rs.of_code(r["activity_code"])
            a = rs.acts.get(r["activity_code"]) or {}
            progs[p or UNASSIGNED].append({"code": r["activity_code"], "name": a.get("name") or "（不在当前计划）", "qty": r["qty"], "unit": r["unit"], "pct": r["pct"]})
            if p and p not in in_log:
                in_log.append(p)
        log_parts[lg["id"]] = in_log
        if lg["status"] not in CONFIRMED or not visible(lg["author_id"]):
            continue
        for p in set(works) | set(progs):
            put(p, {"type": "progress", "date": lg["log_date"], "time": "", "title": "施工记录", "log_id": lg["id"], "author": name_of(lg["author_id"]),
                    "weather": lg["weather"] or "", "works": works.get(p, []), "progress": progs.get(p, [])})
        for k, i in enumerate(c.get("issues", [])):
            p = rs.of_code(i.get("activity_code")) or rs.of_text(i.get("desc")) or (in_log[0] if len(in_log) == 1 else None) or UNASSIGNED
            issue_part[(lg["id"], k)] = p
            put(p, {"type": "issue", "date": lg["log_date"], "time": "", "title": f"{i.get('type') or '其他'}问题（{i.get('severity') or '中'}）", "log_id": lg["id"],
                    "author": name_of(lg["author_id"]), "text": i.get("desc") or "", "measure": i.get("measure") or "", "issue_type": i.get("type") or "其他",
                    "activity_code": i.get("activity_code") or "", "k": k})

    # ---- 照片全表（用于解析隐患的位置；可见性另算）
    mrows = {m["id"]: m for m in all_(conn, "SELECT * FROM media WHERE project_id=?", (pid,))}
    for m in mrows.values():
        raw = (m["geo_place"] or "").strip()
        if raw and not rs.of_label(raw) and visible(m["uploader_id"]):
            unresolved[raw[:20]] += 1

    # ---- 质控点与隐患
    fmedia: set[int] = set()
    for f in all_(conn, "SELECT * FROM findings WHERE project_id=? AND agent IN ('safety','quality') AND confirmed_at IS NOT NULL ORDER BY id", (pid,)):
        ref_m = mrows.get(f["ref_id"]) if f["ref_type"] == "media" else None
        rect_m = mrows.get(f["rectify_media_id"]) if f["rectify_media_id"] else None
        ref_lg = log_by_id.get(f["ref_id"]) if f["ref_type"] == "log" else None
        owners = [f["assignee_id"], ref_m["uploader_id"] if ref_m else None, ref_lg["author_id"] if ref_lg else None, rect_m["uploader_id"] if rect_m else None]
        if ids is not None and not any(o in ids for o in owners if o is not None):
            continue
        body = loads(f["body_json"], {}) or {}
        part = None
        if f["kind"] == "reported" and ref_lg:
            fp = (f["fingerprint"] or "").split(":")
            if len(fp) == 3 and fp[0] == "rep":
                part = issue_part.get((ref_lg["id"], int(fp[2]))) if fp[2].isdigit() else None
        lp = log_parts.get(ref_lg["id"], []) if ref_lg else []
        part = (part or rs.of_code(f["activity_code"]) or (ref_m and rs.of_label(ref_m["geo_place"])) or (rect_m and rs.of_label(rect_m["geo_place"]))
                or (lp[0] if len(lp) == 1 else None) or rs.of_text(f["title"])
                or (ref_m and rs.of_text(people.get(ref_m["uploader_id"], {}).get("wbs_scope"))) or UNASSIGNED)
        photos = [_photo_view(m) for m in (ref_m, rect_m) if m]
        for m in (ref_m, rect_m):
            if m:
                fmedia.add(m["id"])
                parts[part]["photos"].add(m["id"])
        status = f["status"]
        if f["kind"] == "qc_reminder":
            pts = body.get("points") or []
            put(part, {"type": "qc", "date": f["created_at"][:10], "time": _hhmm(f["created_at"]), "title": f["title"], "finding_id": f["id"],
                       "points": pts, "status": status, "done": status == "resolved", "due_date": f["due_date"]})
        else:
            put(part, {"type": "hazard", "date": f["created_at"][:10], "time": _hhmm(f["created_at"]), "title": f["title"], "finding_id": f["id"],
                       "agent": f["agent"], "level": f["level"], "status": status, "open": status in ACTIVE, "due_date": f["due_date"],
                       "reported": f["kind"] == "reported", "log_id": ref_lg["id"] if ref_lg else None,
                       "resolved_at": f["resolved_at"], "photos": photos, "category": body.get("category") or ""})

    # ---- 随手记
    note_media_part: dict[int, str] = {}
    for n in all_(conn, "SELECT * FROM notes WHERE project_id=? ORDER BY created_at", (pid,)):
        if not visible(n["author_id"]):
            continue
        mids = [int(x) for x in loads(n["media_json"], [])]
        p = rs.of_text(n["text"]) or rs.of_label(n["geo_place"]) or next((rs.of_label(mrows[i]["geo_place"]) for i in mids if i in mrows and rs.of_label(mrows[i]["geo_place"])), None) or UNASSIGNED
        for i in mids:
            note_media_part.setdefault(i, p)
        ph = [_photo_view(mrows[i]) for i in mids if i in mrows]
        for i in mids:
            parts[p]["photos"].add(i)
        put(p, {"type": "note", "date": n["note_date"], "time": _hhmm(n["created_at"]), "title": "语音随手记" if n["source"] == "voice" else "随手记", "text": n["text"],
                "author": name_of(n["author_id"]), "photos": ph, "geo": geo_of(n)})

    # ---- 照片（隐患/整改照片已并入隐患条目，不重复出时间线，但计入照片数）
    for m in mrows.values():
        if not visible(m["uploader_id"]):
            continue
        mid = m["id"]
        if mid in fmedia:
            continue
        if mid in note_media_part:
            continue  # 已随随手记出现
        geo_p = rs.of_label(m["geo_place"])
        lg_p = log_parts.get(m["log_id"], []) if m["log_id"] else []
        ps = [geo_p] if geo_p else (lg_p or [UNASSIGNED])
        shared = len(ps) > 1  # 同一日志涉及多个部位、照片又没有地点时，在各部位都出现
        for p in ps:
            parts[p]["photos"].add(mid)
            put(p, {"type": "photo", "date": (m["taken_at"] or m["uploaded_at"])[:10], "time": _hhmm(m["taken_at"] or m["uploaded_at"]),
                    "title": PURPOSE_NAME.get(m["purpose"], "照片"), "author": name_of(m["uploader_id"]), "photos": [_photo_view(m)],
                    "time_src": "taken" if m["taken_at"] else "uploaded", "shared": shared})

    # ---- P6 作业清单也算部位（无记录时显示 0，这是真实计数）；分包只看自己有记录的部位
    if ids is None:
        for p in rs.known:
            parts[p]
    return {"parts": parts, "rs": rs, "unresolved": unresolved, "log_parts": log_parts, "pid": pid}


def _stats(name: str, d: dict, rs: Resolver) -> dict:
    items = d["items"]
    hz = [i for i in items if i["type"] == "hazard"]
    # 日志问题里由质量/安全问题生成的待办已在隐患里计数，这里的“问题”只数日志问题条目
    qc = [i for i in items if i["type"] == "qc"]
    days = {i["date"] for i in items if i["type"] in ("progress", "issue")}
    acts: dict[str, dict] = {}
    for code, a in rs.acts.items():
        if rs.act_part.get(code) == name and not a.get("is_milestone"):
            acts[code] = {"code": code, "name": a["name"], "plan_start": a.get("early_start"), "plan_finish": a.get("early_finish"), "status": a.get("status"),
                          "p6_pct": a.get("phys_pct"), "qty": [], "log_pct": None, "log_pct_date": None, "critical": bool(a.get("is_critical"))}
    qty: dict[str, dict] = defaultdict(lambda: defaultdict(float))
    for i in sorted((i for i in items if i["type"] == "progress"), key=lambda x: x["date"]):
        for p in i["progress"]:
            a = acts.setdefault(p["code"], {"code": p["code"], "name": p["name"], "plan_start": None, "plan_finish": None, "status": None, "p6_pct": None,
                                            "qty": [], "log_pct": None, "log_pct_date": None, "critical": False})
            if p["qty"] is not None:
                qty[p["code"]][p["unit"] or ""] += p["qty"]
            if p["pct"] is not None:
                a["log_pct"], a["log_pct_date"] = p["pct"], i["date"]
    for code, a in acts.items():
        a["qty"] = [{"qty": round(v, 2), "unit": u} for u, v in qty.get(code, {}).items()]
    times = [f"{i['date']} {i['time']}".strip() for i in items]
    return {"name": name, "photos": len(d["photos"]), "hazards": len(hz), "hazards_open": sum(1 for i in hz if i["open"]),
            "hazards_closed": sum(1 for i in hz if i["status"] == "resolved"), "qc_total": len(qc), "qc_done": sum(1 for i in qc if i["done"]),
            "issues": sum(1 for i in items if i["type"] == "issue"), "log_days": len(days), "notes": sum(1 for i in items if i["type"] == "note"),
            "last_update": max(times)[:16] if times else None, "activities": sorted(acts.values(), key=lambda a: (a["plan_start"] or "9999", a["code"]))}


def list_parts(conn, user: dict) -> dict:
    data = collect(conn, user)
    out = []
    for name, d in data["parts"].items():
        if name == UNASSIGNED and not d["items"] and not d["photos"]:
            continue
        s = _stats(name, d, data["rs"])
        s["activity_count"] = len(s.pop("activities"))
        out.append(s)
    out.sort(key=lambda s: sort_key(s["name"]))
    return {"parts": out, "unresolved": sum(data["unresolved"].values()), "unresolved_labels": [k for k, _ in data["unresolved"].most_common(30)],
            "scope": (user.get("company") or "本单位") if R.is_sub(user["role"]) else "全项目",
            "note": "日志与工程量只计技术员已确认的日志；识别不出部位的记录归“未归部位”。"}


def detail(conn, user: dict, part: str, types: list[str] | None = None, limit: int = 400) -> dict:
    data = collect(conn, user)
    d = data["parts"].get(part)
    if d is None:
        raise ValueError("部位不存在或无权查看")
    stats = _stats(part, d, data["rs"])
    counts = Counter(i["type"] for i in d["items"])
    items = [i for i in d["items"] if not types or i["type"] in types]
    items.sort(key=lambda i: (i["date"], i["time"]), reverse=True)
    out = {"part": part, "stats": stats, "type_counts": dict(counts), "types": TYPE_NAME, "items": items[:limit], "truncated": len(items) > limit,
           "summary": None, "scope": (user.get("company") or "本单位") if R.is_sub(user["role"]) else "全项目"}
    if not R.is_sub(user["role"]):  # 概况含全项目数字，分包单位不返回
        row = one(conn, "SELECT * FROM site_summaries WHERE project_id=? AND part=?", (user["project_id"], part))
        if row:
            out["summary"] = {"text": row["summary"], "source": row["source"], "at": row["created_at"], "stale": row["facts_hash"] != facts_hash(stats)}
    return out


# ----------------------------------------------------------------------------- 档案Agent：概况与归并

SUMMARY_SYSTEM = """你是工程项目的档案员。请根据给定的某一部位的统计事实，写一段2-4句的“部位概况”，供项目部查档时快速了解该部位的现状。
严格遵守：1. 所有数字（数量、百分比、天数）必须直接取自给定统计，不得计算新数字、不得估计、不得编造；2. 没有给出的情况（如质量结论、完成程度）不要写；
3. 有未闭环隐患时要点明；统计为空的项不要写“无”以外的推断；4. 用简洁规范的中文工程语言。只输出JSON：{"summary":"概况文字"}"""


def facts_text(stats: dict) -> str:
    L = [f"部位：{stats['name']}",
         f"照片 {stats['photos']} 张；有记录的日志日 {stats['log_days']} 天；日志问题 {stats['issues']} 条；随手记 {stats['notes']} 条",
         f"安全质量隐患共 {stats['hazards']} 项，其中未闭环 {stats['hazards_open']} 项，已闭环 {stats['hazards_closed']} 项",
         f"质控提醒 {stats['qc_total']} 项，已完成 {stats['qc_done']} 项",
         f"涉及P6作业 {len(stats['activities'])} 项",
         f"最近更新：{stats['last_update'] or '无记录'}"]
    for a in stats["activities"][:12]:
        q = "、".join(f"{x['qty']:g}{x['unit']}" for x in a["qty"]) or "日志暂无工程量"
        pct = f"日志累计 {a['log_pct']:g}%（{a['log_pct_date']}）" if a["log_pct"] is not None else "日志暂无累计百分比"
        L.append(f"作业 {_an.fmt(a['code'], a['name'])}：计划 {a['plan_start'] or '未知'}~{a['plan_finish'] or '未知'}；日志合计 {q}；{pct}")
    return "\n".join(L)


def facts_hash(stats: dict) -> str:
    return hashlib.sha1(facts_text(stats).encode("utf-8")).hexdigest()[:16]


def rule_summary(stats: dict) -> str:
    """规则概况（离线模拟与数字守卫不通过时使用）：只复述统计。"""
    s = f"{stats['name']}：累计照片 {stats['photos']} 张，有日志记录的日期 {stats['log_days']} 天"
    if stats["activities"]:
        s += f"，涉及 {len(stats['activities'])} 项P6作业"
    s += f"；安全质量隐患 {stats['hazards']} 项，未闭环 {stats['hazards_open']} 项；质控提醒 {stats['qc_total']} 项，已完成 {stats['qc_done']} 项。"
    if stats["last_update"]:
        s += f"最近更新 {stats['last_update']}。"
    return s


def _norm_summary(x) -> dict:
    text = str(x.get("summary") or "").strip() if isinstance(x, dict) else ""
    if not text:
        raise ValueError("缺少summary")
    return {"summary": text[:600]}


def make_summary(conn, user: dict, part: str) -> dict:
    data = collect(conn, user)
    d = data["parts"].get(part)
    if d is None:
        raise ValueError("部位不存在或无权查看")
    stats = _stats(part, d, data["rs"])
    facts = facts_text(stats)
    source, guard, run_id, text = "ai", [], None, ""
    try:
        res = call_json("site.archive", system=SUMMARY_SYSTEM, user=facts, prompt_version=PROMPT_VERSION,
                        mock=lambda: {"summary": "（离线模拟）" + rule_summary(stats)},
                        validate=_norm_summary,
                        project_id=user["project_id"], user_id=user["id"], max_tokens=700, input_meta={"task": "summary", "part_chars": len(part)})
        text, run_id = res.data["summary"], res.run_id
        guard = report_agent.number_guard({"summary": text}, facts)
    except AIError as exc:
        guard = [{"field": "summary", "number": "", "error": str(exc)[:80]}]
    if not text or guard:  # AI 不可用或出现统计里没有的数字：改用规则概况，不放行自造数字
        text, source = rule_summary(stats), "rule"
    conn.execute("INSERT OR REPLACE INTO site_summaries (project_id, part, facts_hash, summary, source, ai_run_id, created_by, created_at) VALUES (?,?,?,?,?,?,?,?)",
                 (user["project_id"], part, facts_hash(stats), text, source, run_id, user["id"], now_iso()))
    audit(conn, user["id"], "site.summary", f"site:{part}", f"source={source}")
    return {"text": text, "source": source, "guard": guard, "at": now_iso(), "stale": False}


ORG_SYSTEM = """你是工程项目的档案员。给你一批现场记录里写法各异的“部位名称”，以及项目已有的规范部位名称。
请把每个写法归并到最匹配的规范部位名称；拿不准、或明显是新部位的，目标留空字符串，不要猜。只输出JSON：{"map":{"写法":"规范部位名或空"}}"""


def organize(conn, user: dict) -> dict:
    """档案Agent归并部位写法：结果存入 knowledge(kind=part_alias)，之后的归集按此识别。"""
    data = collect(conn, user)
    labels = [k for k, _ in data["unresolved"].most_common(40)]
    cands = sorted((n for n, d in data["parts"].items() if n != UNASSIGNED and (d["items"] or d["photos"])), key=sort_key)
    if not labels:
        return {"merged": 0, "labels": 0}
    if not cands:
        return {"merged": 0, "labels": len(labels)}

    def mock():
        out = {}
        for lb in labels:
            hit = [c for c in cands if len(lb) >= 2 and (c in lb or lb in c)]
            out[lb] = hit[0] if len(hit) == 1 else ""
        return {"map": out}

    def norm(x):
        if not isinstance(x, dict) or not isinstance(x.get("map"), dict):
            raise ValueError("缺少map")
        return {"map": {str(k): str(v) for k, v in x["map"].items() if str(k) in labels and str(v) in cands}}

    res = call_json("site.archive", system=ORG_SYSTEM, user="写法：\n" + "\n".join(labels) + "\n\n规范部位：\n" + "\n".join(cands),
                    prompt_version=PROMPT_VERSION, mock=mock, validate=norm, project_id=user["project_id"], user_id=user["id"], max_tokens=1200,
                    input_meta={"task": "organize", "labels": len(labels)})
    n = 0
    for lb, to in res.data["map"].items():
        if to:
            conn.execute("INSERT OR REPLACE INTO knowledge (project_id, kind, key, value_json, ai_run_id, updated_at) VALUES (?,?,?,?,?,?)",
                         (user["project_id"], "part_alias", lb, dumps({"to": to}), res.run_id, now_iso()))
            n += 1
    audit(conn, user["id"], "site.organize", "site", f"labels={len(labels)} merged={n}")
    return {"merged": n, "labels": len(labels)}


# ----------------------------------------------------------------------------- 导出

def _doc_base():
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)
    return doc


def _table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)
    return t


def add_photos(doc, conn, photos: list[dict], per_row: int = 3, limit: int = 12, width_cm: float = 4.6) -> None:
    """照片缩略图表格，每张下方写拍摄时间与地点；文件缺失的只写文字说明。"""
    from docx.shared import Cm
    photos = photos[:limit]
    if not photos:
        return
    rows = (len(photos) + per_row - 1) // per_row
    t = doc.add_table(rows=rows, cols=per_row)
    for idx, p in enumerate(photos):
        cell = t.rows[idx // per_row].cells[idx % per_row]
        m = one(conn, "SELECT thumb FROM media WHERE id=?", (p["id"],))
        para = cell.paragraphs[0]
        try:
            path = (settings.media_dir / m["thumb"]).resolve()
            if m and settings.media_dir.resolve() in path.parents and path.is_file():
                para.add_run().add_picture(str(path), width=Cm(width_cm))
        except Exception:  # noqa: BLE001 — 单张图坏了不影响整份导出
            pass
        g = p.get("geo") or {}
        where = g.get("place") or (f"{g['lat']:.5f},{g['lon']:.5f}" if g.get("lat") is not None else "未取得定位")
        when = (p.get("taken_at") or "").replace("T", " ") or f"拍摄时间未知（上传 {(p.get('uploaded_at') or '').replace('T', ' ')[:16]}）"
        cell.add_paragraph(f"P{p['id']} {when}\n{where}")


def to_docx(conn, user: dict, part: str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = detail(conn, user, part, limit=2000)
    st = d["stats"]
    doc = _doc_base()
    title = doc.add_heading(f"部位档案：{part}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"导出日期：{date.today().isoformat()}    数据范围：{d['scope']}    导出人：{user['name']}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("一、部位概况", level=1)
    if d["summary"]:
        doc.add_paragraph(d["summary"]["text"] + ("（档案Agent起草，数字取自统计）" if d["summary"]["source"] == "ai" else "（规则生成，数字取自统计）"))
    doc.add_paragraph(rule_summary(st))
    _table(doc, ["照片", "隐患", "未闭环", "质控提醒", "质控已完成", "日志问题", "最近更新"],
           [[st["photos"], st["hazards"], st["hazards_open"], st["qc_total"], st["qc_done"], st["issues"], st["last_update"] or "无记录"]])
    doc.add_heading("二、作业与工程量", level=1)
    if st["activities"]:
        _table(doc, ["作业", "计划", "日志工程量合计", "日志累计%"],
               [[f"{a['code']} {a['name']}", f"{a['plan_start'] or '—'}~{a['plan_finish'] or '—'}", "、".join(f"{x['qty']:g}{x['unit']}" for x in a["qty"]) or "—",
                 f"{a['log_pct']:g}%（{a['log_pct_date']}）" if a["log_pct"] is not None else "—"] for a in st["activities"]])
    else:
        doc.add_paragraph("无记录。")
    doc.add_heading("三、时间线", level=1)
    rows = []
    for i in d["items"]:
        if i["type"] == "progress":
            txt = "；".join(w["desc"] or w["process"] for w in i["works"]) or ""
            txt += ("｜" if txt and i["progress"] else "") + "；".join(f"{p['code']} 完成{p['qty']:g}{p['unit']}" if p["qty"] is not None else f"{p['code']}" for p in i["progress"])
        elif i["type"] == "hazard":
            txt = f"{i['title']}（{'高' if i['level'] == 'high' else '中' if i['level'] == 'medium' else '低'}级，{'未闭环' if i['open'] else '已闭环' if i['status'] == 'resolved' else '已忽略'}）"
        elif i["type"] == "qc":
            txt = f"{i['title']}（{'已完成' if i['done'] else '未完成'}，{len(i['points'])}个质控点）"
        elif i["type"] == "issue":
            txt = f"{i['text']}" + (f"｜措施：{i['measure']}" if i["measure"] else "")
        elif i["type"] == "note":
            txt = i["text"]
        else:
            txt = f"{i['title']} P{i['photos'][0]['id']}"
        rows.append([f"{i['date']} {i['time']}".strip(), TYPE_NAME.get(i["type"], i["type"]), i.get("author", ""), txt])
    if rows:
        _table(doc, ["时间", "类型", "记录人", "内容"], rows)
    else:
        doc.add_paragraph("无记录。")
    doc.add_heading("四、照片", level=1)
    photos, seen = [], set()
    for i in d["items"]:
        for ph in i.get("photos") or []:
            if ph["id"] not in seen:
                seen.add(ph["id"])
                photos.append(ph)
    if photos:
        add_photos(doc, conn, photos, limit=18)
        if len(photos) > 18:
            doc.add_paragraph(f"（仅列出最近 18 张，共 {len(photos)} 张）")
    else:
        doc.add_paragraph("无照片。")
    doc.add_heading("附：数据来源说明", level=1)
    doc.add_paragraph("本档案由系统从技术员确认的施工日志、现场照片（含水印拍摄时间与定位）、随手记、质控提醒和安全质量隐患记录自动归集，"
                      "部位按P6作业的WBS层级及名称识别，识别不出的归“未归部位”。工程量、进度、整改闭环均以原记录为准，AI 不改变其结论。")
    safe = hashlib.sha1(part.encode("utf-8")).hexdigest()[:8]
    out = settings.export_dir / f"site_{user['project_id']}_{safe}.docx"
    _an.humanize_doc(doc)   # 作业代码 → 中文简称（services/act_names.py）
    doc.save(out)
    audit(conn, user["id"], "site.export", f"site:{part}", f"items={len(d['items'])}")
    return out

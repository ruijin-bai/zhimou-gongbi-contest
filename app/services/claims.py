"""工期/费用索赔证据包：从施工日志里的停工、缺料、设备故障、恶劣天气、变更/指令等问题事件出发，汇集证据。

- 事件、证据、受影响作业与天数都由确定性查询得出；受影响天数只用日志工程量与P6已有数据估算，缺数据留空标“待补”，不补零；
- 天气只取日志里已记录的天气，不联网查；
- 档案Agent（kind=claim.pack）只把已汇集的证据整理成“事件→证据清单→待补充项”的说明文字，数字守卫不通过时改用规则文字；
- 证据包一律注明“AI整理，待商务/项目经理核实后使用”，生成与导出都写 audit；本模块不改变工程量、进度、审批、计量语义。
"""
from __future__ import annotations

import hashlib
from datetime import date, timedelta

from ..ai import report_agent
from ..ai.client import AIError, call_json
from ..config import settings
from ..db import all_, audit, dumps, insert, loads, now_iso, one
from . import p6
from .parts import PROMPT_VERSION, _doc_base, _table, add_photos
from .role_views import company_label
from . import act_names as _an

NOTICE = "AI整理，待商务/项目经理核实后使用"
CONFIRMED = ("submitted", "approved")

HALT_KW = ("停工", "停止施工", "暂停施工", "无法施工", "被迫停", "窝工", "停产")
KW = {
    "weather": ("暴雨", "降雨", "大雨", "强降雨", "阵雨", "台风", "大风", "雷暴", "高温", "洪水", "积水", "沙尘"),
    "material": ("缺料", "断料", "材料不足", "进场不足", "未到场", "延误到场", "供应不足", "物资短缺", "材料短缺", "未进场"),
    "equipment": ("故障", "损坏", "抛锚", "停机", "维修", "机械失效", "设备失效"),
    "change": ("变更", "指令", "业主要求", "监理要求", "设计修改", "图纸修改", "图纸变更", "暂停令", "通知"),
}
CATS = {
    "stop": {"name": "停工", "todo": ["停工的起止时间（精确到时分）——日志只记到日", "业主/监理的停工通知或书面确认", "停工期间窝工的人员、机械清单及台班（当日投入见下）"]},
    "material": {"name": "缺料/材料延误", "todo": ["采购订单、催货函及供应商回复", "材料实际到场时间与验收记录", "缺料导致的窝工人员、机械清单"]},
    "equipment": {"name": "设备故障", "todo": ["设备故障报告及维修记录（含起止时间）", "备用设备租赁/调配记录", "设备停机期间的窝工人员清单"]},
    "weather": {"name": "恶劣天气", "todo": ["当地气象部门出具的气象证明（本系统不联网查天气）", "监理对恶劣天气停工的签认", "天气影响时段的现场照片（带水印时间）"]},
    "change": {"name": "变更/指令/外部因素", "todo": ["业主/监理书面指令或变更单（编号、日期）", "指令前后的工作范围、工程量对比", "指令引起的返工/窝工记录"]},
}


# ----------------------------------------------------------------------------- 识别事件

def _classify(issue: dict) -> tuple[str | None, bool]:
    """日志问题 → (索赔类别, 是否含停工字样)。类型字段优先，其次关键词；与索赔无关的问题返回 None。"""
    t, desc = issue.get("type") or "", f"{issue.get('desc') or ''}{issue.get('measure') or ''}"
    halt = any(k in desc for k in HALT_KW)
    by_type = {"材料": "material", "设备": "equipment", "天气": "weather", "图纸": "change"}
    if t in by_type:
        return by_type[t], halt
    hit = next((c for c in ("weather", "material", "equipment", "change") if any(k in desc for k in KW[c])), None)
    if t == "外部":
        return (hit or ("stop" if halt else "change")), halt
    if hit:
        return hit, halt
    return ("stop" if halt else None), halt


def _event_key(cat: str, start: str, company: str) -> str:
    return f"{cat}-{start}-{hashlib.sha1(company.encode('utf-8')).hexdigest()[:6]}"


def all_events(conn, project_id: int, merge_gap: int = 2) -> list[dict]:
    """全项目事件：同类别、同单位、记录日期相隔不超过 merge_gap 天的合并为一个事件。"""
    users = {u["id"]: u for u in all_(conn, "SELECT id, name, company FROM users WHERE project_id=?", (project_id,))}
    rows = all_(conn, "SELECT * FROM logs WHERE project_id=? AND period='day' AND status<>'returned' ORDER BY log_date, id", (project_id,))
    cand = []
    for lg in rows:
        c = loads(lg["content_json"], {}) or {}
        for k, i in enumerate(c.get("issues", [])):
            cat, halt = _classify(i)
            if not cat:
                continue
            u = users.get(lg["author_id"], {})
            cand.append({"cat": cat, "company": company_label(u.get("company")), "date": lg["log_date"], "log_id": lg["id"], "k": k, "author_id": lg["author_id"],
                         "author": u.get("name", "已离职"), "type": i.get("type") or "其他", "severity": i.get("severity") or "中", "desc": (i.get("desc") or "").strip(),
                         "measure": (i.get("measure") or "").strip(), "activity_code": i.get("activity_code") or "", "halt": halt, "confirmed": lg["status"] in CONFIRMED})
    events: list[dict] = []
    open_ev: dict[tuple, dict] = {}
    for x in sorted(cand, key=lambda x: (x["cat"], x["company"], x["date"], x["log_id"], x["k"])):
        g = (x["cat"], x["company"])
        ev = open_ev.get(g)
        if ev and (date.fromisoformat(x["date"]) - date.fromisoformat(ev["date_to"])).days <= merge_gap:
            ev["date_to"] = max(ev["date_to"], x["date"])
        else:
            ev = {"cat": x["cat"], "cat_name": CATS[x["cat"]]["name"], "company": x["company"], "date_from": x["date"], "date_to": x["date"], "issues": []}
            events.append(ev)
            open_ev[g] = ev
        ev["issues"].append(x)
    for ev in events:
        ev["key"] = _event_key(ev["cat"], ev["date_from"], ev["company"])
        ev["dates"] = sorted({i["date"] for i in ev["issues"]})
        ev["halt"] = any(i["halt"] for i in ev["issues"])
        ev["codes"] = sorted({i["activity_code"] for i in ev["issues"] if i["activity_code"]})
        ev["unconfirmed"] = any(not i["confirmed"] for i in ev["issues"])
    return sorted(events, key=lambda e: (e["date_from"], e["key"]), reverse=True)


def list_events(conn, user: dict, date_from: str | None = None, date_to: str | None = None) -> list[dict]:
    today = date.today()
    df = date_from or (today - timedelta(days=120)).isoformat()
    dt = date_to or today.isoformat()
    acts = {a["code"]: a for a in p6.current_activities(conn, user["project_id"])}
    out = []
    for ev in all_events(conn, user["project_id"]):
        if ev["date_to"] < df or ev["date_from"] > dt:
            continue
        out.append({"key": ev["key"], "cat": ev["cat"], "cat_name": ev["cat_name"], "company": ev["company"], "date_from": ev["date_from"], "date_to": ev["date_to"],
                    "log_days": len(ev["dates"]), "n_issues": len(ev["issues"]), "halt": ev["halt"], "unconfirmed": ev["unconfirmed"],
                    "descs": [i["desc"] for i in ev["issues"][:2]],
                    "activities": [{"code": c, "name": (acts.get(c) or {}).get("name") or "（不在当前计划）"} for c in ev["codes"]]})
    return out


# ----------------------------------------------------------------------------- 汇集证据

def _impact(conn, pid: int, ev: dict, acts: dict) -> dict:
    """受影响作业与工期：天数只用日志工程量对比估算；基线不足 3 个正常施工日或事件日无工程量记录则留空（待补）。"""
    dates = set(ev["dates"])
    items = []
    total, has = 0.0, False
    for code in ev["codes"]:
        a = acts.get(code) or {}
        rows = all_(conn, """SELECT lp.log_date, lp.qty, lp.unit FROM log_progress lp JOIN logs l ON l.id=lp.log_id
                             WHERE lp.project_id=? AND lp.activity_code=? AND lp.qty IS NOT NULL AND lp.voided_at IS NULL AND l.status IN ('submitted','approved') AND l.period='day'
                             ORDER BY lp.log_date""", (pid, code))
        units = [r["unit"] for r in rows]
        unit = max(set(units), key=units.count) if units else ""
        base = [r["qty"] for r in rows if r["unit"] == unit and r["log_date"] not in dates and r["qty"] > 0]
        evq = [{"date": r["log_date"], "qty": r["qty"]} for r in rows if r["unit"] == unit and r["log_date"] in dates]
        baseline = round(sum(base) / len(base), 2) if len(base) >= 3 else None
        loss = round(sum(max(0.0, 1 - e["qty"] / baseline) for e in evq), 1) if baseline and evq else None
        if loss is not None:
            total, has = total + loss, True
        items.append({"code": code, "name": a.get("name") or "（不在当前计划）", "plan_start": a.get("early_start"), "plan_finish": a.get("early_finish"),
                      "float_d": a.get("total_float_d"), "critical": bool(a.get("is_critical")) if a else None, "unit": unit, "baseline": baseline,
                      "baseline_days": len(base), "event_qty": evq, "loss_days": loss})
    span = (date.fromisoformat(ev["date_to"]) - date.fromisoformat(ev["date_from"])).days + 1
    halt_dates = sorted({i["date"] for i in ev["issues"] if i["halt"]})
    return {"span_days": span, "log_days": len(ev["dates"]), "halt_dates": halt_dates, "activities": items, "est_loss_days": round(total, 1) if has else None,
            "basis": "按日志工程量与同作业其他施工日的日均量对比估算（基线至少3个施工日），仅供参考，须商务核实"}


def evidence(conn, user: dict, ev: dict) -> dict:
    pid = user["project_id"]
    acts = {a["code"]: a for a in p6.current_activities(conn, pid)}
    log_ids = sorted({i["log_id"] for i in ev["issues"]})
    marks = ",".join("?" for _ in log_ids)
    logs = all_(conn, f"SELECT * FROM logs WHERE id IN ({marks}) ORDER BY log_date, id", log_ids)
    users = {u["id"]: u["name"] for u in all_(conn, "SELECT id, name FROM users WHERE project_id=?", (pid,))}
    ev_logs, labor, weather = [], [], []
    for lg in logs:
        c = loads(lg["content_json"], {}) or {}
        mine = [i for i in ev["issues"] if i["log_id"] == lg["id"]]
        ev_logs.append({"id": lg["id"], "date": lg["log_date"], "author": users.get(lg["author_id"], ""), "status": lg["status"], "confirmed": lg["status"] in CONFIRMED,
                        "weather": lg["weather"] or "", "voice_text": lg["voice_text"] or "", "summary": c.get("summary") or "",
                        "works": [w.get("desc") or w.get("process") or "" for w in c.get("works", [])][:6],
                        "issues": [{"desc": i["desc"], "measure": i["measure"], "type": i["type"], "severity": i["severity"], "halt": i["halt"]} for i in mine]})
        weather.append({"date": lg["log_date"], "log_id": lg["id"], "weather": lg["weather"] or ""})
        labor.append({"date": lg["log_date"], "log_id": lg["id"],
                      "labor": [{"trade": x.get("trade") or "", "count": x.get("count")} for x in c.get("labor", [])],
                      "machinery": [{"name": x.get("name") or "", "count": x.get("count"), "status": x.get("status") or ""} for x in c.get("machinery", [])]})
    # 照片：日志照片 + 当事人事件日的随手记照片（均带拍摄时间/定位）
    photos: dict[int, dict] = {}
    for m in all_(conn, f"SELECT * FROM media WHERE log_id IN ({marks}) ORDER BY COALESCE(taken_at, uploaded_at), id", log_ids):
        photos[m["id"]] = _photo(m, f"日志 L{m['log_id']}")
    authors = sorted({i["author_id"] for i in ev["issues"]})
    notes = []
    for n in all_(conn, f"SELECT * FROM notes WHERE project_id=? AND author_id IN ({','.join('?' for _ in authors)}) AND note_date IN ({','.join('?' for _ in ev['dates'])}) ORDER BY created_at",
                  [pid] + authors + ev["dates"]):
        mids = [int(i) for i in loads(n["media_json"], [])]
        notes.append({"id": n["id"], "date": n["note_date"], "author": users.get(n["author_id"], ""), "text": n["text"], "media": mids})
        for m in all_(conn, f"SELECT * FROM media WHERE id IN ({','.join('?' for _ in mids) or 'NULL'}) AND project_id=?", mids + [pid]):
            photos.setdefault(m["id"], _photo(m, f"随手记 {n['note_date']}"))
    # 相关指令/整改：挂在这些日志上的事项，或同作业在事件前后几天内产生的事项
    fq = f"SELECT * FROM findings WHERE project_id=? AND kind<>'qc_reminder' AND confirmed_at IS NOT NULL AND ((ref_type='log' AND ref_id IN ({marks}))"
    fp: list = [pid] + log_ids
    if ev["codes"]:
        fq += f" OR (activity_code IN ({','.join('?' for _ in ev['codes'])}) AND substr(created_at,1,10) BETWEEN ? AND ?)"
        fp += ev["codes"] + [(date.fromisoformat(ev["date_from"]) - timedelta(days=1)).isoformat(), (date.fromisoformat(ev["date_to"]) + timedelta(days=3)).isoformat()]
    findings = [{"id": f["id"], "title": f["title"], "level": f["level"], "status": f["status"], "agent": f["agent"], "due_date": f["due_date"],
                 "created": f["created_at"][:10], "resolved_at": (f["resolved_at"] or "")[:10]}
                for f in all_(conn, fq + ") ORDER BY id LIMIT 12", fp)]
    imp = _impact(conn, pid, ev, acts)
    out = {"key": ev["key"], "cat": ev["cat"], "cat_name": ev["cat_name"], "company": ev["company"], "date_from": ev["date_from"], "date_to": ev["date_to"],
           "dates": ev["dates"], "halt": ev["halt"], "issues": ev["issues"], "logs": ev_logs, "labor_machinery": labor, "weather": weather,
           "photos": sorted(photos.values(), key=lambda p: p["taken_at"] or p["uploaded_at"] or ""), "notes": notes, "findings": findings, "impact": imp}
    out["gaps"] = gaps(out)
    return out


def _photo(m: dict, src: str) -> dict:
    g = None
    if m.get("geo_lat") is not None:
        g = {"lat": m["geo_lat"], "lon": m["geo_lon"], "acc": m.get("geo_acc"), "place": m.get("geo_place") or ""}
    elif m.get("geo_place"):
        g = {"lat": None, "lon": None, "acc": None, "place": m["geo_place"]}
    return {"id": m["id"], "thumb": f"/api/media/{m['id']}/thumb", "url": f"/api/media/{m['id']}/file", "taken_at": m.get("taken_at"),
            "uploaded_at": m.get("uploaded_at"), "geo": g, "source": src}


def gaps(e: dict) -> list[dict]:
    """待补充项：先列该类事件常见但系统里没有的材料，再按实际缺失逐条检查。"""
    out = [{"item": t, "why": "系统内无此类材料，需人工补充", "auto": False} for t in CATS[e["cat"]]["todo"]]

    def add(item, why):
        out.append({"item": item, "why": why, "auto": True})
    imp = e["impact"]
    if not e["photos"]:
        add("现场照片", "事件日没有挂接带拍摄时间的照片")
    else:
        no_time = [p for p in e["photos"] if not p["taken_at"]]
        if no_time:
            add("照片拍摄时间", f"{len(no_time)} 张照片没有拍摄时间，仅有上传时间")
        no_geo = [p for p in e["photos"] if not p["geo"]]
        if no_geo:
            add("照片定位", f"{len(no_geo)} 张照片未取得定位")
    if not any(x["labor"] or x["machinery"] for x in e["labor_machinery"]):
        add("事件日人机投入", "日志未记录人员和机械投入")
    if e["cat"] == "weather" and not any(w["weather"] for w in e["weather"]):
        add("天气记录", "事件日日志未填写天气")
    if not imp["activities"]:
        add("受影响作业", "日志问题未关联P6作业")
    else:
        for a in imp["activities"]:
            if a["loss_days"] is None:
                add(f"{a['code']} 的工期影响天数", "日志工程量基线不足（正常施工日少于3天）或事件日无工程量记录，无法估算")
    if imp["est_loss_days"] is None:
        add("受影响工期天数", "待补：无法由日志数据估算，请商务按实际核定")
    if e["halt"] and not imp["halt_dates"]:
        add("停工日期", "未能从日志识别")
    if any(not i["measure"] for i in e["issues"]):
        add("采取的措施", "部分问题记录未写措施")
    if any(not l["confirmed"] for l in e["logs"]):
        add("日志确认", "部分日志尚未经技术员周确认")
    return out


# ----------------------------------------------------------------------------- 档案Agent整理说明

SYSTEM = """你是工程项目的档案员，正在把施工日志里的停工、缺料、设备故障、天气、变更/指令事件整理成工期/费用索赔的“证据包说明”，供商务和项目经理核实。
严格遵守：1. 只用给定的证据，不得编造事实；所有数字（天数、数量、人数）必须直接取自给定事实，不得计算新数字或估计；
2. 对每个事件输出：statement（事件经过，客观陈述，≤150字）、evidence（证据清单，每条一句，指明来源，如“日志L12”“照片P5”，≤8条）、gaps（在给定“系统已识别待补项”之外你认为还需人补充的证据，≤3条，没有则空）；
3. 受影响天数写“待补”或引用给定估算值，并说明是估算；不得做法律结论或索赔金额；4. 天气只用日志已记录的天气。
只输出JSON：{"overview":"总说明≤120字","events":[{"ref":"E1","statement":"","evidence":[""],"gaps":[""]}]}"""


def facts_text(evs: list[dict]) -> str:
    L = []
    for n, e in enumerate(evs, 1):
        imp = e["impact"]
        L.append(f"[E{n}] 类别：{e['cat_name']}；单位：{e['company']}；时间：{e['date_from']}~{e['date_to']}（日志记录 {len(e['dates'])} 天）{'；含停工记录' if e['halt'] else ''}")
        for i in e["issues"]:
            L.append(f"  问题 L{i['log_id']} {i['date']} {i['type']}/{i['severity']}：{i['desc']}（措施：{i['measure'] or '未写'}）")
        for a in imp["activities"]:
            L.append(f"  受影响作业 {_an.fmt(a['code'], a['name'])}：计划 {a['plan_start'] or '未知'}~{a['plan_finish'] or '未知'}；总浮时 "
                     f"{('%g天' % a['float_d']) if a['float_d'] is not None else '未知'}；{'关键作业' if a['critical'] else '非关键' if a['critical'] is False else '关键性未知'}；"
                     + (f"估算损失 {a['loss_days']:g} 天（基线日均 {a['baseline']:g}{a['unit']}，样本 {a['baseline_days']} 天）" if a["loss_days"] is not None else "工期影响天数待补"))
        L.append("  工期影响：" + (f"估算合计 {imp['est_loss_days']:g} 天（估算，待核实）" if imp["est_loss_days"] is not None else "待补"))
        for l in e["logs"]:
            L.append(f"  日志 L{l['id']} {l['date']} 天气：{l['weather'] or '未记录'}；口述：{(l['voice_text'] or l['summary'])[:120]}")
        for x in e["labor_machinery"]:
            lab = "、".join(f"{p['trade']}{p['count']:g}人" if p["count"] is not None else p["trade"] for p in x["labor"]) or "未记录"
            mac = "、".join(f"{p['name']}{p['count']:g}台{p['status']}" if p["count"] is not None else f"{p['name']}{p['status']}" for p in x["machinery"]) or "未记录"
            L.append(f"  人机投入 {x['date']}（L{x['log_id']}）：人员 {lab}；机械 {mac}")
        for p in e["photos"][:8]:
            L.append(f"  照片 P{p['id']}（{p['source']}）：拍摄 {p['taken_at'] or '时间未知'}；定位 {(p['geo'] or {}).get('place') or ('有坐标' if p['geo'] else '无')}")
        for f in e["findings"][:5]:
            L.append(f"  整改/指令 F{f['id']} {f['title']}（{f['status']}）")
        L.append("  系统已识别待补项：" + "；".join(g["item"] for g in e["gaps"] if g["auto"]) if any(g["auto"] for g in e["gaps"]) else "  系统已识别待补项：无")
    return "\n".join(L)


def rule_narrative(evs: list[dict]) -> dict:
    """规则整理：不经模型，只复述证据（离线模拟与数字守卫不通过时使用）。数字全部取自 facts_text。"""
    out = {"overview": f"本证据包汇集所选事件的日志记录、照片、人机投入和相关整改，{NOTICE}。", "events": []}
    for n, e in enumerate(evs, 1):
        imp = e["impact"]
        first = e["issues"][0]
        st = f"{e['company']}在 {e['date_from']} 至 {e['date_to']} 期间，施工日志共 {len(e['dates'])} 天记录了{e['cat_name']}事件" \
             + f"，首条记录：{first['desc']}。" + (f"受影响工期估算 {imp['est_loss_days']:g} 天（估算，待核实）。" if imp["est_loss_days"] is not None else "受影响工期待补。")
        ev = [f"施工日志 L{l['id']}（{l['date']}）原文及天气记录" for l in e["logs"][:5]]
        ev += [f"现场照片 P{p['id']}（拍摄 {p['taken_at'] or '时间未知'}）" for p in e["photos"][:5]]
        if any(x["labor"] or x["machinery"] for x in e["labor_machinery"]):
            ev.append("事件日人员机械投入记录")
        ev += [f"相关整改/指令 F{f['id']}：{f['title']}" for f in e["findings"][:3]]
        out["events"].append({"ref": f"E{n}", "statement": st, "evidence": ev[:8], "gaps": []})
    return out


def _norm(data, n: int) -> dict:
    if not isinstance(data, dict) or not isinstance(data.get("events"), list):
        raise ValueError("缺少events")
    byref = {str(x.get("ref")): x for x in data["events"] if isinstance(x, dict)}
    evs = []
    for k in range(1, n + 1):
        x = byref.get(f"E{k}")
        if not x or not str(x.get("statement") or "").strip():
            raise ValueError(f"缺少事件E{k}")
        evs.append({"ref": f"E{k}", "statement": str(x["statement"]).strip()[:400],
                    "evidence": [str(s).strip() for s in (x.get("evidence") or []) if str(s).strip()][:8],
                    "gaps": [str(s).strip() for s in (x.get("gaps") or []) if str(s).strip()][:3]})
    return {"overview": str(data.get("overview") or "").strip()[:300], "events": evs}


def narrate(conn, user: dict, evs: list[dict]) -> dict:
    facts = facts_text(evs)
    note, run_id, source = "", None, "ai"
    try:
        res = call_json("claim.pack", system=SYSTEM, user=facts, prompt_version=PROMPT_VERSION, mock=lambda: rule_narrative(evs),
                        validate=lambda d: _norm(d, len(evs)), project_id=user["project_id"], user_id=user["id"], max_tokens=2500,
                        input_meta={"events": len(evs)})
        nar, run_id = res.data, res.run_id
        flat = {"overview": nar["overview"]}
        for e in nar["events"]:
            flat[e["ref"]] = "\n".join([e["statement"]] + e["evidence"] + e["gaps"])
        guard = report_agent.number_guard(flat, facts)
        if guard:
            note = "AI 叙述中出现统计里没有的数字（" + "、".join(g["number"] for g in guard[:5]) + "），已改用规则整理"
    except AIError as exc:
        nar, guard, note = None, [], f"AI 不可用（{str(exc)[:60]}），已用规则整理"
    if nar is None or guard:
        nar, source = rule_narrative(evs), "rule"
    nar = {**nar, "source": source, "note": note, "notice": NOTICE, "ai_run_id": run_id if source == "ai" else None}
    if not nar["overview"]:
        nar["overview"] = rule_narrative(evs)["overview"]
    return nar


# ----------------------------------------------------------------------------- 证据包

def create_pack(conn, user: dict, keys: list[str], title: str = "") -> dict:
    keys = list(dict.fromkeys(k for k in keys if k))
    if not keys:
        raise ValueError("请先勾选至少一个事件")
    if len(keys) > 8:
        raise ValueError("一份证据包最多合并 8 个事件")
    by_key = {e["key"]: e for e in all_events(conn, user["project_id"])}
    miss = [k for k in keys if k not in by_key]
    if miss:
        raise ValueError("事件不存在或已变化，请刷新后重选")
    evs = [evidence(conn, user, by_key[k]) for k in keys]
    nar = narrate(conn, user, evs)
    title = (title or "").strip()[:60] or f"索赔证据包（{min(e['date_from'] for e in evs)} 起，{len(evs)} 个事件）"
    pid = insert(conn, "claim_packs", {"project_id": user["project_id"], "title": title, "event_keys_json": dumps(keys), "snapshot_json": dumps(evs),
                                       "narrative_json": dumps(nar), "ai_run_id": nar.get("ai_run_id"), "created_by": user["id"], "created_at": now_iso()})
    audit(conn, user["id"], "claim.generate", f"claim:{pid}", f"events={len(keys)} source={nar['source']}")
    return pack_view(conn, one(conn, "SELECT * FROM claim_packs WHERE id=?", (pid,)), full=True)


def pack_view(conn, row: dict, full: bool = False) -> dict:
    nar = loads(row["narrative_json"], {}) or {}
    names = {u["id"]: u["name"] for u in all_(conn, "SELECT id, name FROM users WHERE project_id=?", (row["project_id"],))}
    out = {"id": row["id"], "title": row["title"], "created_at": row["created_at"], "created_by": names.get(row["created_by"], ""), "exported_at": row["exported_at"],
           "n_events": len(loads(row["event_keys_json"], [])), "source": nar.get("source"), "notice": NOTICE}
    if full:
        out.update({"events": loads(row["snapshot_json"], []), "narrative": nar})
    return out


def to_docx(conn, user: dict, row: dict):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    evs, nar = loads(row["snapshot_json"], []), loads(row["narrative_json"], {}) or {}
    by_ref = {e["ref"]: e for e in nar.get("events", [])}
    doc = _doc_base()
    t = doc.add_heading(row["title"], level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(f"【{NOTICE}】")
    r.bold, r.font.color.rgb, r.font.size = True, RGBColor(0xC6, 0x28, 0x28), Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"生成日期：{row['created_at'][:10]}    导出日期：{date.today().isoformat()}    导出人：{user['name']}    "
                      f"说明文字：{'档案Agent整理' if nar.get('source') == 'ai' else '规则整理'}" + (f"（{nar['note']}）" if nar.get("note") else ""))
    doc.add_heading("一、总说明", level=1)
    doc.add_paragraph(nar.get("overview") or "")
    _table(doc, ["序号", "类别", "单位", "起止", "日志记录天数", "受影响工期（估算）"],
           [[n, e["cat_name"], e["company"], f"{e['date_from']}~{e['date_to']}", len(e["dates"]),
             f"{e['impact']['est_loss_days']:g} 天（估算，待核实）" if e["impact"]["est_loss_days"] is not None else "待补"] for n, e in enumerate(evs, 1)])
    for n, e in enumerate(evs, 1):
        x = by_ref.get(f"E{n}", {})
        doc.add_heading(f"事件{n}：{e['cat_name']}（{e['company']}，{e['date_from']}~{e['date_to']}）", level=1)
        doc.add_heading("事件说明", level=2)
        doc.add_paragraph(x.get("statement") or "")
        doc.add_heading("证据清单", level=2)
        for s in x.get("evidence") or []:
            doc.add_paragraph(s, style="List Bullet")
        doc.add_heading("受影响作业与工期", level=2)
        imp = e["impact"]
        if imp["activities"]:
            _table(doc, ["作业", "P6计划", "总浮时(天)", "关键", "事件日工程量", "估算损失(天)"],
                   [[f"{a['code']} {a['name']}", f"{a['plan_start'] or '—'}~{a['plan_finish'] or '—'}", f"{a['float_d']:g}" if a["float_d"] is not None else "未知",
                     "是" if a["critical"] else "否" if a["critical"] is False else "未知",
                     "、".join(f"{q['date']} {q['qty']:g}{a['unit']}" for q in a["event_qty"]) or "无记录",
                     f"{a['loss_days']:g}" if a["loss_days"] is not None else "待补"] for a in imp["activities"]])
        else:
            doc.add_paragraph("受影响作业：待补（日志问题未关联P6作业）。")
        doc.add_paragraph(f"口径：{imp['basis']}。" + (f"日志记录含停工字样的日期：{'、'.join(imp['halt_dates'])}。" if imp["halt_dates"] else ""))
        doc.add_heading("当日日志原文", level=2)
        for l in e["logs"]:
            doc.add_paragraph(f"L{l['id']}  {l['date']}  {l['author']}  天气：{l['weather'] or '日志未记录'}" + ("" if l["confirmed"] else "  （尚未周确认）"))
            doc.add_paragraph(l["voice_text"] or l["summary"] or "（无口述原文）")
            for i in l["issues"]:
                doc.add_paragraph(f"问题（{i['type']}/{i['severity']}）：{i['desc']}｜措施：{i['measure'] or '未写'}", style="List Bullet")
        doc.add_heading("当日人机投入（费用依据）", level=2)
        rows = []
        for lm in e["labor_machinery"]:
            rows.append([lm["date"], "、".join(f"{p['trade']} {p['count']:g}人" if p["count"] is not None else p["trade"] for p in lm["labor"]) or "未记录",
                         "、".join(f"{p['name']} {p['count']:g}台（{p['status']}）" if p["count"] is not None else f"{p['name']}（{p['status']}）" for p in lm["machinery"]) or "未记录"])
        _table(doc, ["日期", "人员", "机械"], rows)
        doc.add_heading("相关照片（带水印拍摄时间）", level=2)
        if e["photos"]:
            add_photos(doc, conn, e["photos"], limit=12)
        else:
            doc.add_paragraph("无挂接照片，待补。")
        if e["findings"]:
            doc.add_heading("相关指令/整改", level=2)
            _table(doc, ["编号", "事项", "状态", "日期"], [[f"F{f['id']}", f["title"], f["status"], f["created"]] for f in e["findings"]])
        doc.add_heading("待补充项", level=2)
        extra = [{"item": g, "why": "档案Agent建议补充"} for g in x.get("gaps") or []]
        for g in e["gaps"] + extra:
            doc.add_paragraph(f"□ {g['item']}（{g['why']}）")
    doc.add_heading("附：说明", level=1)
    doc.add_paragraph(f"{NOTICE}。事件、日期、照片、人机投入均取自系统内已记录的施工日志、随手记与整改记录；天气只取日志已记录内容，未联网查询；"
                      "受影响工期为按日志工程量对比的估算，缺数据处标“待补”，不代表索赔结论。")
    out = settings.export_dir / f"claim_{row['id']}.docx"
    _an.humanize_doc(doc)   # 作业代码 → 中文简称（services/act_names.py）
    doc.save(out)
    conn.execute("UPDATE claim_packs SET exported_at=? WHERE id=?", (now_iso(), row["id"]))
    audit(conn, user["id"], "claim.export", f"claim:{row['id']}", f"events={len(evs)}")
    return out

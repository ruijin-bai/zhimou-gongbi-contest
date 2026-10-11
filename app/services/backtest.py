"""历史回测（P0-2）：用多期 P6 更新 XER 检验进度预警规则在 P6 实绩上的表现。

口径（D1、D2，10-07 用户确认）：
- 分组：第 i 期的计划与总浮时 + 第 i+1 期的完成%，在第 i+1 期数据日期按线上同一规则（progress.assess_full）判定，i=0…n-2；
- 事实取能拿到的最晚一期：已完工用实际完成日，未完工用最晚一期的 P6 预测完成日（标注“预测”）；
  与判定所用计划完成日相比推迟超过当时总浮时 + 容差，即为“确实延误”。另给“只算已完工”的严格口径；
- 主指标“命中率”＝TP/(TP+FP)，同时给召回率、实际延误占比、样本数；(TP+TN)/N 只作参考；
  主表按“作业×期次”统计，另附按作业去重；
- 完成%按 P6 作业自身的类型读取（p6.typed_pct），工期%换算；
- 数据缺口（完成%未知等“没报/证据不足”）与末期缺失的作业不进入混淆矩阵，单列数量；分母不足 MIN_RATE_DENOM 不给百分比。

回测是确定性计算，不调用 AI；XER 原文件只在内存中解析，不落盘，只记文件名与 sha256 前缀；只读白名单字段，费用字段一律不读。
结果存 backtests 表（建表见 routers/backtest.py），带来源标识：真实项目数据 / 仿真/测试数据（G-03）。
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path
from statistics import median

from ..config import settings
from ..db import audit, dumps, insert, loads, now_iso, one, tx
from . import p6
from . import progress as prog
from . import act_names as _an

SOURCES = {"real": "真实项目数据", "test": "仿真/测试数据"}
SOURCE_NOTES = {"real": "真实数据，不在仿真时间线内", "test": "仿真/测试数据，不计入应用成效"}
_SOURCE_ALIASES = {"real": "real", "真实": "real", "真实项目数据": "real",
                   "test": "test", "sim": "test", "仿真": "test", "测试": "test", "仿真/测试数据": "test"}
READ_ROLES = ("eng_sup", "admin")   # 运行、查看、导出；项目经理、领导等 403（G-13）
ADOPT_ROLES = ("eng_sup",)          # 采用阈值是业务决定，管理员不参与

RED_GRID = (1.0, 2.0, 3.0)          # 敏感性表：红色阈值（天）
TOL_GRID = (0.0, 3.0, 7.0)          # 敏感性表：容差（天）
MIN_RATE_DENOM = 5                  # 分母少于此数不给百分比，只给 k/n
MAX_FILES = 24
MAX_BYTES = 60 * 1024 * 1024
MAX_TOLERANCE = 60.0
LABEL_MAX = 40
# 期次校验：疑似大范围重排（与上一期相比，共有的未完工作业中计划完成日变动 ≥ RESCHEDULE_DAYS 天的占比 ≥ RESCHEDULE_SHARE）
RESCHEDULE_DAYS = 30
RESCHEDULE_SHARE = 0.3
OVERLAP_MIN = 0.9                   # 作业代码与上一期重合率低于此值单独标出
DOCX_ROWS_MAX = 300


class BacktestError(ValueError):
    def __init__(self, msg: str, status: int = 400):
        super().__init__(msg)
        self.msg, self.status = msg, status


# ----------------------------------------------------------------------------- 输入

def normalize_source(value: str | None) -> str:
    text = (value or "").strip()
    if not text:
        raise BacktestError("请选择数据来源：真实项目数据，或仿真/测试数据")
    src = _SOURCE_ALIASES.get(text.lower()) or _SOURCE_ALIASES.get(text)
    if not src:
        raise BacktestError("数据来源只能是“真实项目数据”或“仿真/测试数据”")
    return src


def parse_exclude(value) -> list[str]:
    """要剔除的期次：数据日期，逗号分隔，或 JSON 数组。"""
    if value in (None, ""):
        return []
    if isinstance(value, (list, tuple, set)):
        items = [str(v).strip() for v in value]
    else:
        text = str(value).strip()
        if text.startswith("["):
            try:
                items = [str(v).strip() for v in json.loads(text)]
            except (ValueError, TypeError):
                raise BacktestError("要剔除的期次格式不对，应为数据日期，如 2026-01-12") from None
        else:
            items = [x.strip() for x in re.split(r"[,，;；\s]+", text)]
    out = []
    for x in items:
        if not x:
            continue
        try:
            out.append(date.fromisoformat(x).isoformat())
        except ValueError:
            raise BacktestError(f"要剔除的期次应写数据日期（如 2026-01-12），不能识别：{x[:20]}") from None
    return sorted(set(out))


def _short_name(name: str) -> str:
    base = Path((name or "").replace("\\", "/")).name.strip()
    return base[:80] or "未命名.xer"


def parse_files(files: list[tuple[str, bytes]]) -> list[dict]:
    """[(文件名, 原始字节)] → 按数据日期升序的期次；完成%按类型读取。原文件不落盘。"""
    if len(files) < 2:
        raise BacktestError("回测至少需要两期 P6 更新 XER（建议三期以上）")
    if len(files) > MAX_FILES:
        raise BacktestError(f"一次最多上传 {MAX_FILES} 期")
    snaps = []
    for name, raw in files:
        short = _short_name(name)
        if not short.lower().endswith(".xer"):
            raise BacktestError("回测需要 P6 导出的 .xer 更新文件")
        if len(raw) > MAX_BYTES:
            raise BacktestError(f"{short} 超过 60MB")
        try:
            meta, acts, _ = p6.activities_from_xer(p6.decode(raw), pct_mode="typed")
        except p6.P6Error as exc:
            raise BacktestError(f"{short}：{exc}") from None
        if not meta.get("data_date"):
            raise BacktestError(f"{short} 未读取到数据日期（Data Date）")
        snaps.append({"file": short, "sha12": hashlib.sha256(raw).hexdigest()[:12], "data_date": meta["data_date"],
                      "proj": hashlib.sha256((meta.get("proj_name") or "").encode()).hexdigest()[:8], "acts": acts})
    dates = Counter(s["data_date"] for s in snaps)
    dup = sorted(d for d, n in dates.items() if n > 1)
    if dup:
        raise BacktestError(f"有多个文件的数据日期相同（{'、'.join(dup)}），请检查是否重复上传")
    return sorted(snaps, key=lambda s: s["data_date"])


# ----------------------------------------------------------------------------- 计算

def _rate(k: int, n: int) -> dict:
    return {"k": k, "n": n, "rate": round(k / n, 4) if n >= MIN_RATE_DENOM else None}


def _counts(rows: list[dict]) -> dict:
    tp = sum(1 for r in rows if r["predicted"] and r["actual"])
    fp = sum(1 for r in rows if r["predicted"] and not r["actual"])
    fn = sum(1 for r in rows if not r["predicted"] and r["actual"])
    tn = len(rows) - tp - fp - fn
    n = len(rows)
    return {"samples": n, "tp": tp, "fp": fp, "fn": fn, "tn": tn,
            "precision": _rate(tp, tp + fp), "recall": _rate(tp, tp + fn),
            "delay_share": _rate(tp + fn, n), "accuracy_ref": _rate(tp + tn, n)}


def _outcome(predicted: bool, actual: bool) -> str:
    return ("TP" if actual else "FP") if predicted else ("FN" if actual else "TN")


def _days(a: str, b: str) -> int:
    return (date.fromisoformat(a) - date.fromisoformat(b)).days


def _evaluate(used: list[dict], th: dict, tol: float) -> dict:
    """按组判定，返回样本行、数据缺口行与跳过计数。"""
    last = used[-1]
    last_by = {a["code"]: a for a in last["acts"]}
    rows, gaps, skipped = [], [], Counter()
    for i in range(len(used) - 1):
        plan, judge = used[i], used[i + 1]
        d0, d1 = date.fromisoformat(plan["data_date"]), date.fromisoformat(judge["data_date"])
        judge_by = {a["code"]: a for a in judge["acts"]}
        same = judge is last
        for a in plan["acts"]:
            if a.get("act_finish") or a.get("is_milestone"):
                continue
            if not a.get("early_start") or not a.get("early_finish"):
                skipped["no_dates"] += 1
                continue
            if date.fromisoformat(a["early_start"]) > d1:  # 判定时点尚未计划开工的不评
                continue
            later, outcome = judge_by.get(a["code"]), last_by.get(a["code"])
            if not later:
                skipped["missing_in_judge"] += 1
                continue
            if not outcome:
                skipped["missing_in_last"] += 1
                continue
            base = {"plan_date": plan["data_date"], "judge_date": judge["data_date"], "period": f"{plan['data_date']}→{judge['data_date']}",
                    "code": a["code"], "name": a.get("name") or ""}
            pct = later.get("phys_pct")
            started = bool(later.get("act_start")) or (pct or 0) > 0
            if started and pct is None:  # 完成%未知：不补零，单列为数据缺口
                gaps.append({**base, "reason": "完成%未知", "idle_days": None})
                continue
            st = {"pct": pct, "pct_date": d1.isoformat(), "last_date": d1.isoformat(), "reports": 1} if started else None
            res, gap = prog.assess_full(a, st, d1, d0, th)
            if gap:  # 没报/证据不足：既不是预测也不是事实
                gaps.append({**base, "reason": "无进度记录", "idle_days": gap["idle_days"]})
                continue
            tf = a.get("total_float_d")
            if tf is None:  # 浮时未知，无法判断是否“超出浮时”
                skipped["float_unknown"] += 1
                continue
            fin = outcome.get("act_finish") or outcome.get("early_finish")
            if not fin:
                skipped["no_fact_date"] += 1
                continue
            delay = _days(fin, a["early_finish"])
            actual = delay > max(tf, 0.0) + tol
            predicted = bool(res and res["level"] == "red")
            if res:
                slip = res["facts"]["slip_days"]
            elif st:
                slip = prog.earned_schedule_slip(a, pct, d1, d0)
            else:
                slip = None
            plan_pct = prog.planned_pct(a, d1, d0)
            rows.append({**base, "critical": bool(a.get("is_critical")) or tf <= th["critical_float_days"], "float": tf,
                         "plan_finish": a["early_finish"], "pct_start": a.get("phys_pct"), "pct_plan": round(plan_pct, 1) if plan_pct is not None else None,
                         "pct_actual": pct, "slip": slip, "level": res["level"] if res else None, "rule": res["rule"] if res else None,
                         "predicted": predicted, "fact_finish": fin, "fact_kind": "actual" if outcome.get("act_finish") else "forecast",
                         "true_delay": delay, "actual": actual, "outcome": _outcome(predicted, actual), "same_period": same})
    return {"rows": rows, "gaps": gaps, "skipped": dict(skipped)}


def _by_activity(rows: list[dict]) -> list[dict]:
    """每个作业只计一次：任一期判红即为预测预警；是否延误以该作业首次进入样本那一期的计划与浮时衡量。"""
    first: dict[str, dict] = {}
    pred: dict[str, bool] = {}
    for r in rows:
        first.setdefault(r["code"], r)
        pred[r["code"]] = pred.get(r["code"], False) or r["predicted"]
    return [{"code": c, "predicted": pred[c], "actual": r["actual"]} for c, r in first.items()]


def _lead_time(rows: list[dict], used: list[dict]) -> dict:
    """对确实延误的作业：P6 总浮时首次转负的期次 − 本规则首次判红的期次（天，正数＝本规则更早）。"""
    first_row: dict[str, dict] = {}
    first_red: dict[str, str] = {}
    for r in sorted(rows, key=lambda x: x["judge_date"]):
        first_row.setdefault(r["code"], r)
        if r["predicted"]:
            first_red.setdefault(r["code"], r["judge_date"])
    p6_neg: dict[str, str] = {}
    for s in used:
        for a in s["acts"]:
            tf = a.get("total_float_d")
            if tf is not None and tf < 0:
                p6_neg.setdefault(a["code"], s["data_date"])
    cats: Counter = Counter()
    leads, items, never_p6 = [], [], 0
    for code, r in first_row.items():
        if not r["actual"]:
            continue
        ours, p6d = first_red.get(code), p6_neg.get(code)
        lead = None
        if not ours:
            cat = "never"
            never_p6 += 1 if p6d else 0
        elif not p6d:
            cat = "only_ours"
        else:
            lead = _days(p6d, ours)
            leads.append(lead)
            cat = "earlier" if lead > 0 else ("same" if lead == 0 else "later")
        cats[cat] += 1
        items.append({"code": code, "name": r["name"], "ours": ours, "p6": p6d, "lead_days": lead, "category": cat})
    return {"delayed": len(items), "compared": len(leads),
            "mean_days": round(sum(leads) / len(leads), 1) if leads else None,
            "median_days": median(leads) if leads else None,
            "earlier": cats["earlier"], "same": cats["same"], "later": cats["later"],
            "only_ours": cats["only_ours"], "never": cats["never"], "never_p6_flagged": never_p6, "items": items}


def _periods_report(snaps: list[dict], excluded: set[str]) -> list[dict]:
    """期次校验：期数、数据日期、作业代码重合率、完成%类型、已完工数；疑似大范围重排的期次单独标出。
    比较对象是上一个未被剔除的期次。"""
    out, prev = [], None
    used_list = [s for s in snaps if s["data_date"] not in excluded]
    for k, s in enumerate(snaps):
        acts = s["acts"]
        codes = [a["code"] for a in acts]
        uniq = set(codes)
        flags = []
        overlap = moved_share = None
        if prev is not None:
            pcodes = {a["code"] for a in prev["acts"]}
            overlap = round(len(uniq & pcodes) / len(pcodes), 4) if pcodes else None
            if overlap is not None and overlap < OVERLAP_MIN:
                flags.append(f"作业代码与上一期重合率 {overlap * 100:.0f}%，低于 {OVERLAP_MIN * 100:.0f}%")
            pby = {a["code"]: a for a in prev["acts"]}
            common = [(pby[a["code"]], a) for a in acts if a["code"] in pby and not pby[a["code"]].get("act_finish")
                      and pby[a["code"]].get("early_finish") and a.get("early_finish")]
            if common:
                moved = sum(1 for pa, a in common if abs(_days(a["early_finish"], pa["early_finish"])) >= RESCHEDULE_DAYS)
                moved_share = round(moved / len(common), 4)
                if moved_share >= RESCHEDULE_SHARE:
                    flags.append(f"疑似大范围重排：{moved_share * 100:.0f}% 的未完工作业计划完成日变动 {RESCHEDULE_DAYS} 天以上")
        if len(uniq) != len(codes):
            flags.append(f"作业代码重复 {len(codes) - len(uniq)} 个")
        types = Counter((a.get("pct_type") or "未标注") for a in acts if not a.get("is_milestone"))
        used = s["data_date"] not in excluded
        role = ""
        if used:
            role = "首期计划" if s is used_list[0] else ("事实（最晚一期）" if s is used_list[-1] else "计划/判定")
        out.append({
            "index": k, "data_date": s["data_date"], "file": s["file"], "sha12": s["sha12"],
            "activities": len(acts), "milestones": sum(1 for a in acts if a.get("is_milestone")),
            "completed": sum(1 for a in acts if a.get("act_finish")),
            "in_progress": sum(1 for a in acts if a.get("act_start") and not a.get("act_finish")),
            "pct_types": dict(types), "overlap_prev": overlap, "moved_share": moved_share,
            "flags": flags, "suspect": any(f.startswith(("疑似", "作业代码与上一期")) for f in flags),
            "used": used, "excluded": not used, "role": role,
        })
        if used:
            prev = s
    return out


def _definitions(tol: float) -> dict:
    return {
        "样本": "判定时点已计划开工、尚未完工的非里程碑作业 × 期次；数据缺口、末期缺失、浮时未知的不计入，单列数量",
        "判定": "用第 i 期的计划与总浮时、第 i+1 期的完成%，在第 i+1 期数据日期按线上同一规则判定；红色＝预测为预警",
        "事实为延误": "取最晚一期：已完工用实际完成日，未完工用 P6 预测完成日（标“预测”）；与判定所用计划完成日相比推迟超过当时总浮时"
                    + (f" + {tol:g} 天容差" if tol else ""),
        "命中率（主指标）": "TP/(TP+FP)：发出的红色预警中后来确实延误的比例",
        "召回率": "TP/(TP+FN)：确实延误的样本中被提前判红的比例",
        "实际延误占比": "(TP+FN)/样本数",
        "(TP+TN)/N": "只作参考：大多数作业本来就不延误，此数天然偏高",
        "只算已完工": "严格口径：只统计最晚一期已实际完工的样本",
        "按作业去重": "每个作业只计一次：任一期判红即为预测预警；是否延误以该作业首次进入样本那一期的计划与浮时衡量",
        "提前量": "对确实延误的作业：P6 总浮时首次转负的期次数据日期 − 本规则首次判红的期次数据日期（天，正数＝本规则更早）",
        "数据缺口": prog.GAP_NOTE + "：完成%未知或没有进度记录的作业不进入混淆矩阵，只单列数量",
        "样本不足": f"分母少于 {MIN_RATE_DENOM} 不给百分比，只给 k/n",
        "完成%口径": "按 P6 作业自身的完成%类型读取：工期%＝（原定工期−尚需工期）/原定工期；数量%＝实际数量/（实际+尚需数量）；实物%照读",
        "说明": "回测验证的是判定规则在 P6 实绩上的表现，不验证“随手记→工程量”的提取；确定性计算，不调用 AI，数据不出本机，不读费用字段",
    }


def run(snaps: list[dict], th: dict | None = None, tolerance_days: float = 0.0, exclude=()) -> dict:
    """纯计算：返回 {params, summary, periods_report, rows, gap_rows, history, sensitivity, lead_time}，不写库。"""
    th = {**prog.DEFAULT_THRESHOLDS, **(th or {})}
    try:
        tol = float(tolerance_days or 0.0)
    except (TypeError, ValueError):
        raise BacktestError("容差天数应为数字") from None
    if not 0 <= tol <= MAX_TOLERANCE:
        raise BacktestError(f"容差天数应在 0～{MAX_TOLERANCE:g} 之间")
    if len(snaps) < 2:
        raise BacktestError("回测至少需要两期 P6 更新 XER（建议三期以上）")
    excl = set(parse_exclude(exclude))
    dates = {s["data_date"] for s in snaps}
    unknown = sorted(excl - dates)
    if unknown:
        raise BacktestError(f"要剔除的期次不在上传文件中：{'、'.join(unknown)}")
    used = [s for s in snaps if s["data_date"] not in excl]
    if len(used) < 2:
        raise BacktestError("剔除后不足两期，无法回测")

    report = _periods_report(snaps, excl)
    ev = _evaluate(used, th, tol)
    rows, gaps = ev["rows"], ev["gaps"]
    counts = _counts(rows)
    by_period = []
    for i in range(len(used) - 1):
        jd = used[i + 1]["data_date"]
        c = _counts([r for r in rows if r["judge_date"] == jd])
        by_period.append({"plan_date": used[i]["data_date"], "judge_date": jd, "same_period": i + 1 == len(used) - 1,
                          "data_gaps": sum(1 for g in gaps if g["judge_date"] == jd), **c})
    lead = _lead_time(rows, used)
    red_now = float(th["red_slip_days"])
    sensitivity = []
    for r in RED_GRID:
        for t in TOL_GRID:
            e = _evaluate(used, {**th, "red_slip_days": r}, t)
            sensitivity.append({"red_slip_days": r, "tolerance_days": t, "current": abs(r - red_now) < 1e-9 and abs(t - tol) < 1e-9,
                                "data_gaps": len(e["gaps"]), **_counts(e["rows"])})
    pct_types: Counter = Counter()
    for s in used:
        pct_types.update((a.get("pct_type") or "未标注") for a in s["acts"] if not a.get("is_milestone"))

    fact_actual = sum(1 for r in rows if r["fact_kind"] == "actual")
    notes = []
    if len(used) == 2:
        notes.append("只有 2 期：事实与判定取自同一期（未完工用 P6 预测完成日），属弱口径，仅供参考")
    if counts["tp"] + counts["fp"] == 0:
        notes.append("无预警样本：回测期内本规则没有判出红色预警，命中率不适用")
    elif any(counts[k]["rate"] is None and counts[k]["n"] > 0 for k in ("precision", "recall", "delay_share")):
        notes.append(f"样本不足：分母少于 {MIN_RATE_DENOM} 的比例只给 k/n，不给百分比")
    if rows and fact_actual == 0:
        notes.append("最晚一期没有已完工的样本，事实全部来自 P6 预测完成日")
    if excl:
        notes.append(f"已剔除 {len(excl)} 期：{'、'.join(sorted(excl))}")
    if len({s["proj"] for s in snaps}) > 1:
        notes.append("各期 P6 项目名称不一致，请确认是同一项目的更新")
    suspect = [p["data_date"] for p in report if p["suspect"] and p["used"]]
    if suspect:
        notes.append(f"有 {len(suspect)} 期疑似大范围重排或代码变动较大，未剔除：{'、'.join(suspect)}")
    if pct_types.get("CP_Drtn"):
        notes.append("完成%按 P6 工期%换算：（原定工期−尚需工期）/原定工期")
    summary = {
        "periods": len(used), "periods_uploaded": len(snaps), "excluded_periods": sorted(excl),
        "data_date_from": used[0]["data_date"], "data_date_to": used[-1]["data_date"], "fact_date": used[-1]["data_date"],
        "weak": len(used) == 2, **counts,
        "no_alert_samples": counts["tp"] + counts["fp"] == 0,
        "data_gaps": len(gaps), "skipped": ev["skipped"],
        "fact_actual": fact_actual, "fact_forecast": len(rows) - fact_actual,
        "same_period_samples": sum(1 for r in rows if r["same_period"]),
        "strict": _counts([r for r in rows if r["fact_kind"] == "actual"]),
        "by_activity": _counts(_by_activity(rows)),
        "by_period": by_period, "lead_mean_days": lead["mean_days"],
        "pct_types": dict(pct_types), "min_rate_denom": MIN_RATE_DENOM,
        "notes": notes, "definitions": _definitions(tol),
    }
    odd = {r["code"] for r in rows if r["outcome"] in ("FP", "FN")}
    history: dict[str, list[dict]] = {}
    for s in used:
        for a in s["acts"]:
            if a["code"] in odd:
                history.setdefault(a["code"], []).append({
                    "data_date": s["data_date"], "status": a.get("status"), "pct": a.get("phys_pct"), "pct_type": a.get("pct_type") or "",
                    "total_float_d": a.get("total_float_d"), "early_start": a.get("early_start"), "early_finish": a.get("early_finish"),
                    "act_start": a.get("act_start"), "act_finish": a.get("act_finish")})
    params = {"thresholds": {k: th[k] for k in prog.DEFAULT_THRESHOLDS}, "tolerance_days": tol,
              "exclude_periods": sorted(excl), "pct_mode": "typed", "red_grid": list(RED_GRID), "tolerance_grid": list(TOL_GRID)}
    return {"params": params, "summary": summary, "periods_report": report, "rows": rows, "gap_rows": gaps,
            "history": history, "sensitivity": sensitivity, "lead_time": lead}


# ----------------------------------------------------------------------------- 存取

def run_files(conn, user: dict, files: list[tuple[str, bytes]], source: str | None, label: str | None = "",
              tolerance_days: float = 0.0, exclude=None) -> dict:
    """上传的多期 XER → 回测 → 存 backtests 表 → 返回 view。"""
    pid = user.get("project_id")
    if not pid:
        raise BacktestError("账号未关联项目")
    src = normalize_source(source)
    excl = parse_exclude(exclude)
    snaps = parse_files(files)
    res = run(snaps, prog.thresholds(conn, pid), tolerance_days, excl)
    s = res["summary"]
    with tx(conn):
        bid = insert(conn, "backtests", {
            "project_id": pid, "source": src, "label": (label or "").strip()[:LABEL_MAX], "created_by": user["id"], "created_at": now_iso(),
            "params_json": dumps(res["params"]), "files_json": dumps(res["periods_report"]),
            "summary_json": dumps({"summary": s, "lead_time": res["lead_time"]}),
            "rows_json": dumps({"rows": res["rows"], "gap_rows": res["gap_rows"], "history": res["history"]}),
            "sensitivity_json": dumps(res["sensitivity"]),
        })
        audit(conn, user["id"], "backtest.run", f"backtest:{bid}",
              f"source={src} periods={s['periods']} samples={s['samples']} tp={s['tp']} fp={s['fp']} fn={s['fn']} tn={s['tn']} gaps={s['data_gaps']}")
    return view(conn, get(conn, pid, bid))


def get(conn, project_id: int, bid: int) -> dict:
    row = one(conn, "SELECT * FROM backtests WHERE id=? AND project_id=?", (bid, project_id))
    if not row:
        raise BacktestError("回测记录不存在", 404)
    return row


def latest(conn, project_id: int, source: str | None = None) -> dict | None:
    if source:
        return one(conn, "SELECT * FROM backtests WHERE project_id=? AND source=? ORDER BY id DESC LIMIT 1", (project_id, source))
    return one(conn, "SELECT * FROM backtests WHERE project_id=? ORDER BY id DESC LIMIT 1", (project_id,))


def view(conn, row: dict) -> dict:
    sj = loads(row["summary_json"], {})
    rj = loads(row["rows_json"], {})
    params = loads(row["params_json"], {})
    who = one(conn, "SELECT name FROM users WHERE id=?", (row["created_by"],)) if row["created_by"] else None
    adopted = None
    if row["adopted_at"]:
        by = one(conn, "SELECT name FROM users WHERE id=?", (row["adopted_by"],)) if row["adopted_by"] else None
        adopted = {**(params.get("adopted") or {}), "at": row["adopted_at"], "by": by["name"] if by else None}
    return {
        "id": row["id"], "created_at": row["created_at"], "created_by": who["name"] if who else None,
        "source": row["source"], "source_name": SOURCES.get(row["source"], row["source"]), "source_note": SOURCE_NOTES.get(row["source"], ""),
        "label": row["label"], "params": params,
        "summary": sj.get("summary", {}), "lead_time": sj.get("lead_time", {}),
        "periods_report": loads(row["files_json"], []), "sensitivity": loads(row["sensitivity_json"], []),
        "rows": rj.get("rows", []), "gap_rows": rj.get("gap_rows", []), "history": rj.get("history", {}),
        "adopted": adopted,
    }


def adopt(conn, user: dict, row: dict, red_slip_days: float, tolerance_days: float) -> dict:
    """由工程主管从敏感性表里选定一组阈值：红色阈值写入项目设置（线上预警判定随之改变），容差记入回测记录；留审计。"""
    cells = loads(row["sensitivity_json"], [])
    red, tol = float(red_slip_days), float(tolerance_days)
    if not any(abs(c["red_slip_days"] - red) < 1e-9 and abs(c["tolerance_days"] - tol) < 1e-9 for c in cells):
        raise BacktestError("只能采用敏感性表中的一组阈值")
    with tx(conn):
        proj = one(conn, "SELECT config_json FROM projects WHERE id=?", (row["project_id"],))
        cfg = loads(proj["config_json"], {}) if proj else {}
        ths = dict(cfg.get("thresholds") or {})
        old = ths.get("red_slip_days", prog.DEFAULT_THRESHOLDS["red_slip_days"])
        ths["red_slip_days"] = red
        cfg["thresholds"] = ths
        conn.execute("UPDATE projects SET config_json=? WHERE id=?", (dumps(cfg), row["project_id"]))
        params = loads(row["params_json"], {})
        params["adopted"] = {"red_slip_days": red, "tolerance_days": tol, "previous_red_slip_days": old}
        at = now_iso()
        conn.execute("UPDATE backtests SET params_json=?, adopted_at=?, adopted_by=? WHERE id=?", (dumps(params), at, user["id"], row["id"]))
        audit(conn, user["id"], "backtest.adopt", f"backtest:{row['id']}",
              f"source={row['source']} 红色阈值 {old:g}→{red:g} 天，回测容差 {tol:g} 天")
    return view(conn, get(conn, row["project_id"], row["id"]))


def metrics_brief(conn, project_id: int) -> dict | None:
    """应用成效页只引用来源为“真实项目数据”的最新一次回测（G-03、AC-2.14）；没有就为 None。"""
    row = latest(conn, project_id, "real")
    if not row:
        return None
    s = loads(row["summary_json"], {}).get("summary", {})
    return {"id": row["id"], "at": row["created_at"], "source": row["source"], "source_name": SOURCES["real"], "source_note": SOURCE_NOTES["real"],
            "label": row["label"], "periods": s.get("periods"), "data_date_from": s.get("data_date_from"), "data_date_to": s.get("data_date_to"),
            "samples": s.get("samples"), "tp": s.get("tp"), "fp": s.get("fp"), "fn": s.get("fn"), "tn": s.get("tn"),
            "precision": (s.get("precision") or {}).get("rate"), "precision_k": (s.get("precision") or {}).get("k"),
            "precision_n": (s.get("precision") or {}).get("n"),
            "recall": (s.get("recall") or {}).get("rate"), "delay_share": (s.get("delay_share") or {}).get("rate"),
            "data_gaps": s.get("data_gaps"), "lead_mean_days": s.get("lead_mean_days"), "weak": s.get("weak"),
            "link": f"#/backtest?id={row['id']}"}


# ----------------------------------------------------------------------------- Word 导出

def fmt_rate(r: dict | None, zero: str = "—") -> str:
    """百分比（k/n）；分母不足给“样本不足（k/n）”，分母为 0 给 zero（命中率用“无预警样本”）。"""
    if not r:
        return "—"
    kn = f"（{r['k']}/{r['n']}）"
    if r.get("rate") is not None:
        return f"{r['rate'] * 100:.1f}%{kn}"
    return (zero if not r.get("n") else "样本不足") + kn


def _cnt(v) -> str:
    return "—" if v is None else str(v)


def to_docx(v: dict, project_id: int) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, val in enumerate(r):
                cells[i].text = "" if val is None else str(val)
        return t

    s, lt = v["summary"], v["lead_time"]
    title = doc.add_heading("进度预警规则历史回测报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"数据来源：{v['source_name']}（{v['source_note']}）    标签：{v['label'] or '未命名'}    "
                          f"回测于 {(v['created_at'] or '')[:10]}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("一、结论", level=1)
    lead = f"，平均提前 {lt['mean_days']:g} 天（{lt['compared']} 个作业可比）" if lt.get("mean_days") is not None else ""
    doc.add_paragraph(f"共 {s['periods']} 期（数据日期 {s['data_date_from']} 至 {s['data_date_to']}），样本 {s['samples']}，"
                      f"命中率 {fmt_rate(s['precision'], '无预警样本')}，"
                      f"召回率 {fmt_rate(s['recall'])}，实际延误占比 {fmt_rate(s['delay_share'])}{lead}；"
                      f"数据缺口 {s['data_gaps']} 个不计入。")
    for n in s.get("notes") or []:
        doc.add_paragraph(n, style="List Bullet")

    doc.add_heading("二、主要指标", level=1)
    st, ba = s["strict"], s["by_activity"]
    table(["指标", "作业×期次", "只算已完工", "按作业去重"], [
        ["样本数", s["samples"], st["samples"], ba["samples"]],
        ["TP（命中）", s["tp"], st["tp"], ba["tp"]],
        ["FP（误报）", s["fp"], st["fp"], ba["fp"]],
        ["FN（漏报）", s["fn"], st["fn"], ba["fn"]],
        ["TN", s["tn"], st["tn"], ba["tn"]],
        ["命中率（主指标）", fmt_rate(s["precision"], "无预警样本"), fmt_rate(st["precision"], "无预警样本"), fmt_rate(ba["precision"], "无预警样本")],
        ["召回率", fmt_rate(s["recall"]), fmt_rate(st["recall"]), fmt_rate(ba["recall"])],
        ["实际延误占比", fmt_rate(s["delay_share"]), fmt_rate(st["delay_share"]), fmt_rate(ba["delay_share"])],
        ["(TP+TN)/N（仅作参考）", fmt_rate(s["accuracy_ref"]), fmt_rate(st["accuracy_ref"]), fmt_rate(ba["accuracy_ref"])],
        ["数据缺口（不计入）", s["data_gaps"], "—", "—"],
    ])

    doc.add_heading("三、期次校验", level=1)
    table(["期次", "数据日期", "作业数", "完成%类型", "已完工", "与上期代码重合率", "说明"], [
        [p["index"] + 1, p["data_date"], p["activities"], "、".join(f"{k} {n}" for k, n in p["pct_types"].items()), p["completed"],
         "—" if p["overlap_prev"] is None else f"{p['overlap_prev'] * 100:.0f}%",
         "；".join((["已剔除"] if p["excluded"] else []) + p["flags"] + ([p["role"]] if p.get("role") else []))]
        for p in v["periods_report"]])

    doc.add_heading("四、分期混淆矩阵", level=1)
    table(["计划期", "判定期", "样本", "TP", "FP", "FN", "TN", "数据缺口"],
          [[b["plan_date"], b["judge_date"] + ("（与事实同期）" if b["same_period"] else ""), b["samples"], b["tp"], b["fp"], b["fn"], b["tn"],
            b["data_gaps"]] for b in s["by_period"]])

    doc.add_heading("五、提前量（与 P6 总浮时转负相比）", level=1)
    doc.add_paragraph(f"确实延误的作业 {lt['delayed']} 个；可比 {lt['compared']} 个，平均 {_cnt(lt['mean_days'])} 天，中位数 {_cnt(lt['median_days'])} 天。")
    table(["更早", "同期", "更晚", "仅本规则判出（P6 浮时未转负）", "本规则从未判出"],
          [[lt["earlier"], lt["same"], lt["later"], lt["only_ours"], lt["never"]]])

    doc.add_heading("六、阈值敏感性", level=1)
    table(["红色阈值（天）", "容差（天）", "样本", "命中率", "召回率", "当前"],
          [[f"{c['red_slip_days']:g}", f"{c['tolerance_days']:g}", c["samples"], fmt_rate(c["precision"], "无预警样本"), fmt_rate(c["recall"]),
            "当前" if c["current"] else ""] for c in v["sensitivity"]])
    if v.get("adopted"):
        a = v["adopted"]
        doc.add_paragraph(f"已采用：红色阈值 {a['red_slip_days']:g} 天、容差 {a['tolerance_days']:g} 天（{a.get('by') or ''} {(a.get('at') or '')[:10]}）")

    doc.add_heading("七、误报与漏报清单", level=1)
    odd = [r for r in v["rows"] if r["outcome"] in ("FP", "FN")]
    if odd:
        table(["类别", "作业代码", "作业名称", "计划期→判定期", "滞后（天）", "总浮时", "计划完成", "事实完成", "事实来源"],
              [["误报" if r["outcome"] == "FP" else "漏报", r["code"], r["name"], r["period"], _cnt(r["slip"]), _cnt(r["float"]),
                r["plan_finish"], r["fact_finish"], "实际完成" if r["fact_kind"] == "actual" else "P6 预测"] for r in odd[:DOCX_ROWS_MAX]])
        if len(odd) > DOCX_ROWS_MAX:
            doc.add_paragraph(f"共 {len(odd)} 条，此处列出前 {DOCX_ROWS_MAX} 条，其余见系统页面。")
    else:
        doc.add_paragraph("无。")

    doc.add_heading("八、口径说明", level=1)
    for k, d in s["definitions"].items():
        doc.add_paragraph(f"{k}：{d}", style="List Bullet")

    out = settings.export_dir / f"backtest_{project_id}_{v['id']}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    _an.humanize_doc(doc)   # 作业代码 → 中文简称（services/act_names.py）
    doc.save(str(out))
    return out

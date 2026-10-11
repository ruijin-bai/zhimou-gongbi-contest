"""Excel 日报导入（RJ 2026-10-09 意见 7：技术员习惯用 Excel 编日报）：上传 → 识别 → 预览核对 → 技术员确认后生成系统日志草稿。

- 按项目现用的“生产日报/Daily Report”模板识别（表头标签找位置，不写死行列）：单位、施工日期、天气、中方/属地人数、进场设备、
  填报人/复核人、当日施工内容简述、现场进度表（施工区域/内容/单位/设计量/桩号位置/当日完成/累计完成/累计占比）、明日施工计划；
- 识别是规则解析，不调用模型、不补数：表格里是“\\”、空白、文字的数量一律记为“未填”（None），不当 0；识别不了的列和行写进 warnings；
- 进度表每一行给出候选 P6 作业（按桩号重叠与关键词打分），只是建议：每行由技术员决定“计入工程量（选定作业与数量）/只记施工内容/不导入”，
  全部决定后才生成日志草稿；草稿照常走日志核对、提交、审核，导入本身不提交、不改进度；
- 只存识别结果与文件指纹（sha256），不保存原始 Excel。
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import date, datetime

from ..db import dumps, insert, loads, now_iso, one, tx
from . import logs as log_svc
from . import p6

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS, MAX_COLS = 400, 30
SRC = "X"  # 日志里的依据代号：Excel 日报


class XlsxError(ValueError):
    def __init__(self, msg: str, status: int = 400):
        super().__init__(msg)
        self.status = status


# ----------------------------------------------------------------------------- 读表

def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


_BLANK = {"", "\\", "/", "-", "—", "--", "无", "n/a", "N/A", "/ ", "\\\\"}


def num(v) -> float | None:
    """数量：数字原样；“\\”、“/”、空白、文字一律 None（未填，不当 0）。"""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _text(v).replace(",", "").replace("，", "")
    if s in _BLANK:
        return None
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def _count(v) -> int | None:
    """人数、台数：取开头的整数（“74台（发电机15台…）”→ 74）；没有数字为 None。"""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v)
    m = re.match(r"\s*(\d+)", _text(v))
    return int(m.group(1)) if m else None


def parse_day(v, filename: str = "") -> str | None:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    for s in (_text(v), filename or ""):
        m = re.search(r"(20\d{2})\s*[.\-/年]\s*(\d{1,2})\s*[.\-/月]\s*(\d{1,2})", s)
        if m:
            try:
                return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
            except ValueError:
                continue
    return None


def _load(raw: bytes):
    if not raw:
        raise XlsxError("文件是空的")
    if len(raw) > MAX_BYTES:
        raise XlsxError("文件不超过 10MB", 413)
    if raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise XlsxError("这是旧版 .xls 文件，请在 Excel 里“另存为 .xlsx”后再导入")
    if not raw.startswith(b"PK"):
        raise XlsxError("只支持 Excel .xlsx 文件")
    import openpyxl
    try:
        return openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise XlsxError(f"Excel 文件打不开：{exc.__class__.__name__}") from None


def _grid(ws) -> list[list]:
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, MAX_ROWS), max_col=min(ws.max_column or 0, MAX_COLS)):
        rows.append([c.value for c in r])
    return rows


def _has(cells: list, *keys: str) -> bool:
    return any(any(k in _text(c) for k in keys) for c in cells)


def _is_table_head(cells: list) -> bool:
    return _has(cells, "施工内容", "Construction content") and _has(cells, "当日完成", "Daily completed", "累计完成", "Accumulated")


# 表头标签 → 字段；值取同一行标签右边第一个非空、且不是另一个标签的格子
_LABELS = [
    ("section", ("单位名称", "Section Name")),
    ("date", ("施工日期", "Construction date")),
    ("reporter", ("填报人", "Reporter")),
    ("reviewer", ("复核人", "Reviewer")),
    ("weather", ("天气", "Weather")),
    ("staff_cn", ("中方人员", "Chinese personnel")),
    ("staff_local", ("属地员工", "Nigerian employees", "当地员工")),
    ("equipment", ("进场设备", "Equipment Onsite")),
]


def _label_of(v) -> str | None:
    t = _text(v)
    if not t:
        return None
    for key, words in _LABELS:
        if any(w in t for w in words):
            return key
    return None


_SEC_OVERVIEW = ("当日施工内容", "Overview of construction")
_SEC_PROGRESS = ("现场进度情况", "On site Progress")
_SEC_TOMORROW = ("明日施工计划", "Tomorrow")
_SEC_PHOTO = ("现场形象进度照片", "Progress Photograph")

_COLS = [
    ("no", ("序号", "No")),
    ("area", ("施工区域", "Construction Area")),
    ("content", ("施工内容", "Construction content")),
    ("design", ("设计总工程量", "Total design")),
    ("location", ("桩号", "位置", "Location")),
    ("daily", ("当日完成", "Daily completed")),
    ("pct", ("累计完成占比", "completion percentage", "占比")),
    ("cum", ("累计完成数量", "Accumulated completed quantity", "累计完成")),
    ("unit", ("单位", "unit")),
]


def _colmap(cells: list) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, c in enumerate(cells):
        t = _text(c)
        if not t:
            continue
        for key, words in _COLS:
            if key in out:
                continue
            if key == "no" and not (t.startswith("序号") or t.lower().startswith("no")):
                continue
            if key == "unit" and not (t.startswith("单位") or t.lower().startswith("unit")):
                continue
            if any(w in t for w in words):
                out[key] = i
                break
    return out


def _first(cells: list) -> str:
    for c in cells:
        t = _text(c)
        if t:
            return t
    return ""


def _line_items(rows: list[list]) -> list[dict]:
    """简述/计划的文字行：“标段A:”“半岛段：”这类小标题作为分组带到下面各行；“无”跳过；去掉行首编号。"""
    out, group = [], ""
    for cells in rows:
        t = _first(cells)
        if not t or t in ("无", "/", "\\"):
            continue
        if re.fullmatch(r"[^，。；;、]{1,20}[:：]", t):
            group = t.rstrip(":：").strip()
            continue
        desc = re.sub(r"^\s*\d{1,2}\s*[、.．)）]\s*", "", t).strip()
        if desc:
            out.append({"group": group, "text": desc})
    return out


def _pct_text(v) -> str:
    n = num(v)
    if n is None:
        return _text(v) if _text(v) not in _BLANK else ""
    return f"{round(n * 100 if n <= 1 else n, 1):g}%"


def parse(raw: bytes, filename: str = "") -> dict:
    """识别一份 Excel 日报，返回结构化结果（不入库）。找不到进度表也照样返回已识别的部分，warnings 写明。"""
    wb = _load(raw)
    sheets = wb.worksheets
    ws = next((s for s in sheets if any(_is_table_head(r) for r in _grid(s))), None) or \
        next((s for s in sheets if any(_has(r, "日报", "Daily Report") for r in _grid(s)[:5])), None) or sheets[0]
    grid = _grid(ws)
    head: dict = {}
    warnings: list[str] = []

    # 表头字段
    for cells in grid[:12]:
        for i, c in enumerate(cells):
            key = _label_of(c)
            if not key or key in head:
                continue
            val = None
            for v in cells[i + 1:]:
                if _text(v) == "":
                    continue
                if _label_of(v):
                    break
                val = v
                break
            head[key] = val

    # 分段
    def find(words, start=0):
        for i in range(start, len(grid)):
            if any(w in _first(grid[i]) for w in words):
                return i
        return None

    i_over = find(_SEC_OVERVIEW)
    i_head = next((i for i, r in enumerate(grid) if _is_table_head(r)), None)
    i_prog = find(_SEC_PROGRESS)
    i_tom = find(_SEC_TOMORROW)
    i_photo = find(_SEC_PHOTO)
    ends = [x for x in (i_prog, i_head, i_tom, i_photo, len(grid)) if x is not None]

    overview = []
    if i_over is not None:
        stop = min(x for x in ends if x > i_over)
        overview = _line_items(grid[i_over + 1:stop])
    else:
        warnings.append("没找到“当日施工内容简述”")

    rows = []
    if i_head is not None:
        cols = _colmap(grid[i_head])
        missing = [n for k, n in (("content", "施工内容"), ("daily", "当日完成数量")) if k not in cols]
        if missing:
            warnings.append("进度表缺少列：" + "、".join(missing))
        stop = min([x for x in (i_tom, i_photo, len(grid)) if x is not None and x > i_head])
        prev = {"area": "", "content": "", "unit": ""}
        for r_i in range(i_head + 1, stop):
            cells = grid[r_i]
            cell = lambda k: cells[cols[k]] if k in cols and cols[k] < len(cells) else None  # noqa: E731
            vals = {k: cell(k) for k in ("area", "content", "unit", "design", "location", "daily", "cum", "pct")}
            if not any(_text(v) not in ("",) for v in vals.values()):
                continue
            area, content, unit = _text(vals["area"]), _text(vals["content"]), _text(vals["unit"])
            # 合并单元格的续行：施工区域、内容（以及单位）沿用上一行
            area = area or prev["area"]
            if not content:
                content = prev["content"]
                unit = unit or prev["unit"]
            elif not unit and content == prev["content"]:  # 同一内容分几行写、后面几行没填单位（四航局模板“软基清淤”）
                unit = prev["unit"]
            prev = {"area": area, "content": content, "unit": unit}
            daily, cum = num(vals["daily"]), num(vals["cum"])
            item = {"i": len(rows), "excel_row": r_i + 1, "area": area, "content": content, "unit": unit,
                    "design": num(vals["design"]), "location": _text(vals["location"]), "daily": daily, "cum": cum,
                    "pct_text": _pct_text(vals["pct"])}
            if daily is None and _text(vals["daily"]) not in _BLANK:
                warnings.append(f"第 {r_i + 1} 行当日完成数量“{_text(vals['daily'])}”不是数字，按未填处理")
            if not content:
                warnings.append(f"第 {r_i + 1} 行没有施工内容，已跳过")
                continue
            rows.append(item)
        for n, r in enumerate(rows):
            r["i"] = n
    else:
        warnings.append("没找到现场进度表（需有“施工内容”和“当日完成数量”列）")

    tomorrow = []
    if i_tom is not None:
        stop = min([x for x in (i_photo, len(grid)) if x is not None and x > i_tom])
        tomorrow = _line_items(grid[i_tom + 1:stop])

    if not rows and not overview and not tomorrow:  # 如同包里的设备监控表等其他报表：不是生产日报，直接说明，不给空预览
        raise XlsxError("没识别出生产日报的内容（施工内容简述、现场进度表、明日计划都没有），请确认上传的是“生产日报/Daily Report”")
    day = parse_day(head.get("date"), filename)
    if not day:
        warnings.append("施工日期没识别出来，请在下面选择")
    equipment = _text(head.get("equipment"))
    return {
        "sheet": ws.title, "date": day, "section": _text(head.get("section")), "weather": _text(head.get("weather")),
        "reporter": _text(head.get("reporter")), "reviewer": _text(head.get("reviewer")),
        "staff_cn": _count(head.get("staff_cn")), "staff_local": _count(head.get("staff_local")),
        "equipment": equipment, "equipment_count": _count(head.get("equipment")), "machines": machines_of(equipment),
        "overview": overview, "rows": rows, "tomorrow": tomorrow, "warnings": warnings,
    }


def machines_of(text: str) -> list[dict]:
    """“74台（发电机15台、加油车1台…）”里括号内的分项 → [{name, count}]；没有分项为空。"""
    inner = re.search(r"[（(](.*)[)）]?", text or "")
    if not inner:
        return []
    out = []
    for name, n in re.findall(r"([一-龥A-Za-z][一-龥A-Za-z0-9\-]{0,11}?)\s*(\d+)\s*(?:台|辆|套|部)", inner.group(1)):
        out.append({"name": name, "count": int(n)})
    return out


# ----------------------------------------------------------------------------- 候选 P6 作业（只是建议）

_CH = re.compile(r"(?:CH|K)\s*(\d+)\s*\+\s*(\d+(?:\.\d+)?)", re.I)
# 中文工序词 → P6 作业名里常见的英文词（P6 作业为英文名时用）
_KW = {
    "清表": ["clearance", "topsoil", "clearing"], "清荒": ["clearance", "clearing"], "清淤": ["mud", "removal", "dredg"],
    "换填": ["replacement", "replace"], "填砂": ["sand", "fill"], "填方": ["fill", "embankment"], "路基": ["subgrade", "embankment", "fill"],
    "涵洞": ["culvert"], "圆管涵": ["pipe", "culvert"], "箱涵": ["box", "culvert"], "桩基": ["pile"], "钻孔": ["bored", "pile", "drill"],
    "钢筋笼": ["cage", "pile", "reinforcement"], "灌注": ["pile", "concret"], "承台": ["cap", "pile cap"], "墩": ["pier"], "盖梁": ["cap", "pier cap"],
    "系梁": ["tie beam"], "梁场": ["girder", "precast"], "预制": ["precast"], "台座": ["precast", "girder"], "拌合站": ["batching"],
    "搅拌站": ["batching"], "营地": ["camp"], "围墙": ["fenc", "wall"], "排水": ["drain"], "水沟": ["drain"], "边沟": ["side drain"],
    "碎石": ["crushed stone", "gravel"], "水稳": ["cement-stabilized", "stabilized"], "沥青": ["asphalt"], "路面": ["pavement"],
    "土工布": ["geotextile"], "砂袋": ["sand bag", "sandbag"], "强夯": ["compaction"], "冲击碾压": ["impact compaction"], "碾压": ["compaction"],
    "栈桥": ["trestle"], "平台": ["platform"], "护栏": ["barrier"], "路缘石": ["kerb"], "沼泽": ["swamp"], "便道": ["access road", "haul road"],
}
_STOP = {"and", "the", "of", "in", "on", "for", "with", "side", "hand", "left", "right", "left-hand", "right-hand", "construction", "works", "work"}


def _ranges(text: str) -> list[tuple[float, float]]:
    pts = [int(a) * 1000 + float(b) for a, b in _CH.findall(text or "")]
    if not pts:
        return []
    if len(pts) == 1:
        return [(pts[0], pts[0])]
    return [(min(a, b), max(a, b)) for a, b in zip(pts[0::2], pts[1::2])] + ([(pts[-1], pts[-1])] if len(pts) % 2 else [])


def _overlap(a: list, b: list) -> bool:
    return any(x0 <= y1 + 1 and y0 <= x1 + 1 for x0, x1 in a for y0, y1 in b)


def _score(row: dict, act: dict) -> int:
    name = (act.get("name") or "").lower()
    text = f"{row['area']} {row['content']}"
    s = 0
    words = {w for w in re.findall(r"[a-z][a-z\-]{2,}", text.lower()) if w not in _STOP}
    s += 2 * sum(1 for w in words if w in name)
    for cn, ens in _KW.items():
        if cn in text and any(e in name for e in ens):
            s += 3
    s += min(5, sum(1 for ch in set(re.sub(r"[^一-龥]", "", f"{text} {row['location']}")) if ch in name))  # 中文作业名
    s += 2 * sum(1 for t in re.findall(r"\d+#[一-龥]", f"{text} {row['location']}") if t in name)  # 3#墩 这类部位
    rr, ar = _ranges(f"{row['location']} {row['content']}"), _ranges(act.get("name") or "")
    if rr and ar:
        s += 5 if _overlap(rr, ar) else -2
    if act.get("code") and act["code"].lower() in text.lower():
        s += 10
    return s


def candidates(acts: list[dict], row: dict, limit: int = 3) -> list[dict]:
    scored = sorted(((_score(row, a), a) for a in acts), key=lambda x: (-x[0], x[1]["code"]))
    return [{"code": a["code"], "name": a["name"], "score": sc} for sc, a in scored[:limit] if sc >= 3]


def _acts(conn, user: dict) -> list[dict]:
    acts = [a for a in p6.current_activities(conn, user["project_id"]) if not a.get("is_milestone")]
    mine = [a for a in acts if p6.in_scope(a, user.get("wbs_scope", ""))]
    return mine or acts


# ----------------------------------------------------------------------------- 导入记录

def preview(conn, user: dict, raw: bytes, filename: str) -> dict:
    """识别并登记一次导入（只存识别结果与文件指纹），返回预览：识别结果 + 每行候选作业 + 可选作业清单 + 当天日志状态。"""
    parsed = parse(raw, filename)
    acts = _acts(conn, user)
    for r in parsed["rows"]:
        r["candidates"] = candidates(acts, r)
    sha = hashlib.sha256(raw).hexdigest()
    with tx(conn):
        iid = insert(conn, "xlsx_imports", {"project_id": user["project_id"], "user_id": user["id"], "filename": (filename or "")[:120],
                                             "sha256": sha, "parsed_json": dumps(parsed), "created_at": now_iso()})
    return view(conn, user, iid)


def _mine(conn, user: dict, iid: int) -> dict:
    row = one(conn, "SELECT * FROM xlsx_imports WHERE id=? AND user_id=?", (iid, user["id"]))
    if not row:
        raise XlsxError("导入记录不存在", 404)
    return row


def view(conn, user: dict, iid: int) -> dict:
    row = _mine(conn, user, iid)
    parsed = loads(row["parsed_json"], {})
    same = one(conn, "SELECT id, log_id, confirmed_at FROM xlsx_imports WHERE user_id=? AND sha256=? AND id<>? AND confirmed_at IS NOT NULL ORDER BY id DESC LIMIT 1",
               (user["id"], row["sha256"], iid))
    acts = _acts(conn, user)
    day = parsed.get("date")
    lg = one(conn, "SELECT id, status FROM logs WHERE project_id=? AND author_id=? AND log_date=? AND mode='ai' AND period='day'",
             (user["project_id"], user["id"], day)) if day else None
    return {"import_id": iid, "filename": row["filename"], "parsed": parsed, "confirmed_at": row["confirmed_at"], "log_id": row["log_id"],
            "imported_before": {"log_id": same["log_id"], "at": same["confirmed_at"]} if same else None,
            "activities": [{"code": a["code"], "name": a["name"]} for a in acts],
            "day_log": lg}


def _works_from(parsed: dict) -> list[dict]:
    return [{"part": x["group"], "process": "", "desc": x["text"], "activity_code": "", "sources": [SRC]} for x in parsed.get("overview", [])]


def confirm(conn, user: dict, iid: int, log_date: str, decisions: list[dict], keep: dict | None = None) -> dict:
    """技术员逐行决定后生成（或并入）当天的日志草稿。decisions: [{i, action: p6|text|skip, activity_code?, qty?}]，必须覆盖进度表每一行。
    keep: {overview, tomorrow, staff, machinery} 各段是否带入（默认都带）。只写草稿，不提交、不改进度计划。"""
    row = _mine(conn, user, iid)
    if row["confirmed_at"]:
        raise XlsxError("这份导入已生成过日志，如需重来请重新上传", 409)
    parsed = loads(row["parsed_json"], {})
    try:
        day = date.fromisoformat(log_date or "").isoformat()
    except ValueError:
        raise XlsxError("请选择施工日期") from None
    if day > date.today().isoformat():
        raise XlsxError("施工日期不能晚于今天")
    keep = {"overview": True, "tomorrow": True, "staff": True, "machinery": True, **(keep or {})}
    rows = parsed.get("rows", [])
    by_i = {}
    for d in decisions or []:
        try:
            by_i[int(d.get("i"))] = d
        except (TypeError, ValueError):
            continue
    left = [r["i"] for r in rows if (by_i.get(r["i"]) or {}).get("action") not in ("p6", "text", "skip")]
    if left:
        raise XlsxError(f"还有 {len(left)} 行没决定（计入工程量 / 只记施工内容 / 不导入）")
    codes = {a["code"]: a for a in p6.current_activities(conn, user["project_id"])}
    progress, works = [], _works_from(parsed) if keep["overview"] else []
    for r in rows:
        d = by_i[r["i"]]
        label = "".join(x for x in (r["area"] + " " if r["area"] else "", r["content"]))
        where = f"（{r['location']}）" if r["location"] else ""
        if d["action"] == "skip":
            continue
        if d["action"] == "text":
            qty = "" if r["daily"] is None else f"，当日完成 {r['daily']:g}{r['unit']}"
            works.append({"part": r["area"], "process": r["content"], "desc": f"{label}{where}{qty}".strip(), "activity_code": "", "sources": [SRC]})
            continue
        code = str(d.get("activity_code") or "").strip()
        if code not in codes:
            raise XlsxError(f"第 {r['excel_row']} 行“{r['content']}”请选择 P6 作业")
        q = d.get("qty", r["daily"])
        try:
            qty = None if q in (None, "") else float(q)
        except (TypeError, ValueError):
            raise XlsxError(f"第 {r['excel_row']} 行数量不是数字") from None
        if qty is not None and qty < 0:
            raise XlsxError(f"第 {r['excel_row']} 行数量不能为负")
        if qty is None:
            raise XlsxError(f"第 {r['excel_row']} 行“{r['content']}”没有当日完成数量，不能计入工程量（可改为只记施工内容）")
        note = f"Excel：{label}{where}" + (f"，累计 {r['cum']:g}{r['unit']}" if r["cum"] is not None else "") + (f"，累计占比 {r['pct_text']}" if r["pct_text"] else "")
        progress.append({"activity_code": code, "qty": qty, "unit": r["unit"], "pct": None, "note": note[:200], "sources": [SRC]})
        works.append({"part": r["area"], "process": r["content"], "desc": f"{label}{where}，当日完成 {qty:g}{r['unit']}", "activity_code": code, "sources": [SRC]})
    labor = []
    if keep["staff"]:
        labor = [{"trade": t, "count": parsed.get(k), "sources": [SRC]} for k, t in (("staff_cn", "中方人员"), ("staff_local", "属地员工")) if parsed.get(k) is not None]
    machinery = []
    if keep["machinery"]:
        machinery = [{"name": m["name"], "count": m["count"], "status": "正常", "sources": [SRC]} for m in parsed.get("machines", [])]
        if not machinery and parsed.get("equipment_count") is not None:
            machinery = [{"name": "进场设备（合计）", "count": parsed["equipment_count"], "status": "正常", "sources": [SRC]}]
    tomorrow = [{"desc": (f"{x['group']}：" if x["group"] else "") + x["text"], "activity_code": "", "sources": [SRC]}
                for x in parsed.get("tomorrow", [])] if keep["tomorrow"] else []

    lg = log_svc.ensure_log(conn, user, day, "ai", "day")
    log_svc._editable(lg, user)
    content = loads(lg["content_json"], {}) or {}
    base = {k: content.get(k, v) for k, v in log_svc.log_agent.empty_content().items()}
    # 并入当天已有草稿：工程量按作业替换（同一作业以 Excel 为准），其余各段追加；人员、机械以 Excel 为准
    new_codes = {p["activity_code"] for p in progress}
    base["progress"] = [p for p in base["progress"] if p.get("activity_code") not in new_codes] + progress
    base["works"] = base["works"] + works
    base["tomorrow"] = base["tomorrow"] + tomorrow
    if labor:
        base["labor"] = labor
    if machinery:
        base["machinery"] = machinery
    src_note = f"已从 Excel 日报导入（{row['filename'] or '未命名'}，填报人 {parsed.get('reporter') or '未填'}）"
    if parsed.get("equipment") and keep["machinery"]:
        src_note += f"；进场设备原文：{parsed['equipment']}"
    base["notes"] = (base.get("notes") + "\n" if base.get("notes") else "") + src_note
    with tx(conn):
        conn.execute("""UPDATE logs SET content_json=?, weather=CASE WHEN weather='' THEN ? ELSE weather END, version=version+1 WHERE id=?""",
                     (dumps(base), parsed.get("weather") or "", lg["id"]))
        conn.execute("UPDATE xlsx_imports SET confirmed_at=?, log_id=?, decisions_json=? WHERE id=?",
                     (now_iso(), lg["id"], dumps(list(by_i.values())), iid))
    lg = log_svc.get_log(conn, lg["id"])
    log_svc.recheck(conn, lg, user)
    return {"log_id": lg["id"], "log_date": day, "progress": len(progress), "works": len(works)}

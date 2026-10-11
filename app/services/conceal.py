"""隐蔽工程验收资料（资料Agent）：验收模板、记录、缺项检查、缺资料扫描、Word 成册。

分工（AI 提取、人工确认、监理结果分开存放）：
- 确定性：验收模板（内置，项目可增删）、必拍照片与数据项的缺项检查、状态流转、缺资料扫描、导出；
- 资料Agent：识别每张照片拍的是什么、对应模板的哪一项，并把技术员口述整理成数据项候选；只写 ai_json，不改状态，也不会自动成为人工值；
- 人工：归类照片、录入实测值（fields_json）、确认齐全、报验、录入监理意见。
不编造规范条文号和允许偏差数值：模板只列“要拍什么、要记什么”，设计值/要求值由人填写。
给其他模块的只读口径：conceal_records 的 activity_code / part / template_key / status / inspection_date 可直接 SQL 读取；
status ∈ (ready, submitted, passed) 视为“该部位该工序已有验收资料”，可用 covered() 判断。
"""
from __future__ import annotations

import re
import sqlite3
from collections import defaultdict
from datetime import date, datetime, timedelta

from .. import roles as R
from ..ai import conceal_agent
from ..ai.client import AIError, call_json
from ..config import settings
from ..db import all_, audit, dumps, insert, loads, now_iso, one
from . import p6
from . import parts as part_svc
from .media import MediaError, geo_of, save_photo
from .role_views import in_scope
from . import act_names as _an

STATUS_NAME = {"draft": "待补资料", "ready": "已齐待报验", "submitted": "已报验", "passed": "监理通过", "rejected": "监理退回"}
OK_STATUS = ("ready", "submitted", "passed")
READ_ROLES = frozenset({"manager", "eng_sup", "tech", "leader", "safety_sup"})
WRITE_ROLES = frozenset({"manager", "eng_sup", "tech"})
RESULT_ROLES = frozenset({"manager", "eng_sup"})
NOTICE = "验收要求、规格与允许偏差以设计文件、项目验收规范和监理要求为准；本系统只核对“拍了没有、记了没有”，不判定合格。"
MEASURE_NOTE = "须现场实测、由人录入"
FOOTER = "AI 预填，经人工确认；签认以纸质/监理系统为准。"
FIELD_SRC = ("人工", "口述", "照片读数")
PHOTO_SRC = ("当场拍摄", "部位已有照片")


class ConcealError(Exception):
    def __init__(self, status: int, msg: str):
        super().__init__(msg)
        self.status, self.msg = status, msg


# ----------------------------------------------------------------------------- 内置验收模板

def _s(key, label, hint="", required=True):
    return {"key": key, "label": label, "hint": hint, "required": required}


def _n(key, label, unit, required=True, lo=None, hi=None, kw=()):
    return {"key": key, "label": label, "type": "number", "unit": unit, "required": required, "measure": True, "min": lo, "max": hi, "kw": list(kw)}


def _t(key, label, required=True, measure=False):
    return {"key": key, "label": label, "type": "text", "unit": "", "required": required, "measure": measure, "min": None, "max": None, "kw": []}


BUILTIN: list[dict] = [
    {"key": "rebar", "name": "钢筋隐蔽验收", "group": "钢筋",
     "keywords": ["钢筋绑扎", "钢筋安装", "钢筋制安", "钢筋工程", "钢筋隐蔽", "钢筋网", "绑扎钢筋", "钢筋骨架"],
     "shots": [_s("overview", "钢筋整体绑扎全貌", "带部位标识牌"), _s("spacing", "钢筋规格与间距", "带尺，读数清楚"),
               _s("cover", "保护层垫块特写", "带尺"), _s("lap", "搭接/锚固长度", "带尺"), _s("tie", "绑扎节点与扎丝"),
               _s("embed", "预埋件/插筋位置", "有则拍", False)],
     "items": [_t("spec", "钢筋规格型号"), _n("spacing", "钢筋间距实测", "mm", True, 20, 1000, ("间距",)),
               _n("cover", "保护层厚度实测", "mm", True, 5, 200, ("保护层",)), _n("lap_len", "搭接/锚固长度实测", "mm", False, 50, 3000, ("搭接", "锚固"))],
     "roles": ["施工单位", "监理"]},
    {"key": "formwork", "name": "模板验收", "group": "模板",
     "keywords": ["模板", "支模", "立模", "模板安装", "模板支护"],
     "shots": [_s("overview", "模板安装全貌", "带部位标识牌"), _s("dimension", "截面/轴线尺寸", "带尺"), _s("vertical", "垂直度检查", "靠尺或吊线"),
               _s("support", "支撑与加固"), _s("joint", "拼缝与封堵"), _s("release", "脱模剂涂刷")],
     "items": [_n("dim", "截面尺寸实测", "mm", True, 50, 20000, ("截面", "尺寸")), _n("vertical", "垂直度偏差实测", "mm", True, 0, 200, ("垂直度",)),
               _n("joint_gap", "拼缝宽度实测", "mm", False, 0, 50, ("拼缝",)), _t("support", "支撑形式及间距", False)],
     "roles": ["施工单位", "监理"]},
    {"key": "foundation_pit", "name": "基槽/基坑验槽", "group": "基槽基坑",
     "keywords": ["基槽", "基坑", "验槽", "基底", "基础开挖", "承台开挖"],
     "shots": [_s("overview", "基槽全貌", "带部位标识牌"), _s("size", "开挖尺寸", "带尺"), _s("elevation", "基底标高", "水准尺读数"),
               _s("bearing", "持力层土质/岩性特写"), _s("drain", "排水与集水措施")],
     "items": [_n("length", "开挖长度实测", "m", False, 0.1, 500, ("长度",)), _n("width", "开挖宽度实测", "m", False, 0.1, 500, ("宽度",)),
               _n("elev", "基底标高实测", "m", True, -500, 9000, ("标高",)), _t("strata", "持力层描述（须经勘察/设计/监理确认）"), _t("water", "渗水及排水情况", False)],
     "roles": ["施工单位", "监理"]},
    {"key": "pile", "name": "桩基验收", "group": "桩基",
     "keywords": ["桩基", "成孔", "钻孔桩", "灌注桩", "钻孔灌注", "旋挖", "沉桩", "打桩", "钢筋笼"],
     "shots": [_s("position", "桩位复核", "桩位放样/桩位偏差"), _s("depth", "成孔深度", "测绳读数"), _s("diameter", "孔径检测", "探孔器"),
               _s("cage", "钢筋笼", "含长度、加劲箍"), _s("sediment", "沉渣厚度检测")],
     "items": [_t("pile_no", "桩号"), _n("depth", "成孔深度实测", "m", True, 0.5, 200, ("深度", "孔深")), _n("dia", "孔径实测", "mm", True, 100, 5000, ("孔径", "直径")),
               _n("sediment", "沉渣厚度实测", "mm", True, 0, 2000, ("沉渣",)), _n("pos_dev", "桩位偏差实测", "mm", False, 0, 1000, ("桩位偏差", "偏位"))],
     "roles": ["施工单位", "监理"]},
    {"key": "embedded", "name": "预埋件与预留孔洞", "group": "预埋预留",
     "keywords": ["预埋", "预留孔", "预留洞", "预埋件", "预埋管", "套管", "预留孔洞"],
     "shots": [_s("layout", "预埋件/预留孔洞布置全貌", "带部位标识牌"), _s("position", "位置与尺寸", "带尺"), _s("fix", "固定与封堵措施"), _s("label", "型号规格标识", "有标牌则拍", False)],
     "items": [_t("spec", "预埋件/孔洞型号规格"), _n("count", "数量", "个", False, 1, 10000, ("数量",)), _n("pos_dev", "位置偏差实测", "mm", False, 0, 1000, ("偏差",))],
     "roles": ["施工单位", "监理"]},
    {"key": "waterproof", "name": "防水层验收", "group": "防水",
     "keywords": ["防水", "防水层", "防水卷材", "涂膜防水", "防水涂料"],
     "shots": [_s("base", "基层处理", "找平、清洁"), _s("layer", "防水层整体铺设"), _s("lap", "搭接宽度", "带尺"), _s("corner", "阴阳角与细部节点"),
               _s("protect", "保护层/试水", "有则拍", False)],
     "items": [_t("material", "防水材料品种及规格"), _n("lap_w", "搭接宽度实测", "mm", False, 10, 1000, ("搭接",)),
               _n("layers", "铺设层数", "层", False, 1, 10, ("层数",)), _n("thick", "涂膜/卷材厚度实测", "mm", False, 0.1, 50, ("厚度",))],
     "roles": ["施工单位", "监理"]},
]
BUILTIN_KEYS = {t["key"] for t in BUILTIN}
_KEY = re.compile(r"^[a-z][a-z0-9_]{0,23}$")


def _full(key: str, body: dict, builtin: bool, customized: bool = False) -> dict:
    return {"key": key, "name": body["name"], "group": body.get("group") or "", "keywords": list(body.get("keywords") or []),
            "shots": [dict(s) for s in body["shots"]], "items": [dict(i) for i in body.get("items") or []],
            "roles": list(body.get("roles") or ["施工单位", "监理"]), "builtin": builtin, "customized": customized}


def normalize_template(key: str, body: dict) -> dict:
    """自定义模板校验：不含任何允许偏差/规范条文字段；多余字段丢弃。"""
    if not _KEY.match(key or ""):
        raise ConcealError(400, "模板标识只能用小写字母、数字、下划线，字母开头，最长24位")
    if not isinstance(body, dict):
        raise ConcealError(400, "模板格式不对")
    name = str(body.get("name") or "").strip()
    if not 2 <= len(name) <= 30:
        raise ConcealError(400, "模板名称2-30字")
    kws = [str(k).strip()[:20] for k in (body.get("keywords") or []) if str(k).strip()][:30]
    shots, seen = [], set()
    for s in body.get("shots") or []:
        k, label = str((s or {}).get("key") or "").strip(), str((s or {}).get("label") or "").strip()
        if not _KEY.match(k) or not label or k in seen:
            raise ConcealError(400, "必拍照片项需要不重复的标识和名称")
        seen.add(k)
        shots.append({"key": k, "label": label[:30], "hint": str(s.get("hint") or "").strip()[:40], "required": bool(s.get("required", True))})
    if not 1 <= len(shots) <= 20:
        raise ConcealError(400, "必拍照片项 1-20 项")
    items, seen = [], set()
    for i in body.get("items") or []:
        k, label = str((i or {}).get("key") or "").strip(), str((i or {}).get("label") or "").strip()
        typ = i.get("type") if i.get("type") in ("number", "text") else None
        if not _KEY.match(k) or not label or k in seen or not typ:
            raise ConcealError(400, "数据项需要不重复的标识、名称和类型(number/text)")
        seen.add(k)
        lo, hi = conceal_agent.num(i.get("min")), conceal_agent.num(i.get("max"))
        items.append({"key": k, "label": label[:30], "type": typ, "unit": str(i.get("unit") or "").strip()[:8] if typ == "number" else "",
                      "required": bool(i.get("required", True)), "measure": bool(i.get("measure", typ == "number")),
                      "min": lo if typ == "number" else None, "max": hi if typ == "number" else None,
                      "kw": [str(x)[:10] for x in (i.get("kw") or [])][:6] if typ == "number" else []})
    if len(items) > 20:
        raise ConcealError(400, "数据项最多 20 项")
    roles = [r for r in (body.get("roles") or ["施工单位", "监理"]) if r in ("施工单位", "监理")] or ["施工单位", "监理"]
    return {"name": name, "group": str(body.get("group") or "").strip()[:12], "keywords": kws, "shots": shots, "items": items, "roles": roles}


def templates(conn, pid: int, include_hidden: bool = False) -> list[dict]:
    rows = {r["key"]: r for r in all_(conn, "SELECT * FROM conceal_templates WHERE project_id=?", (pid,))}
    out = []
    for b in BUILTIN:
        r = rows.get(b["key"])
        if r and r["hidden"]:
            if include_hidden:
                out.append({**_full(b["key"], b, True), "hidden": True})
            continue
        out.append(_full(b["key"], loads(r["body_json"], b) if r else b, True, bool(r)))
    for k, r in rows.items():
        if k not in BUILTIN_KEYS and not r["hidden"]:
            out.append(_full(k, loads(r["body_json"], {}), False, True))
    return out


def get_template(conn, pid: int, key: str) -> dict | None:
    return next((t for t in templates(conn, pid) if t["key"] == key), None)


def save_template(conn, user: dict, key: str, body: dict) -> dict:
    tpl = normalize_template(key, body)
    conn.execute("INSERT OR REPLACE INTO conceal_templates (project_id, key, body_json, hidden, updated_by, updated_at) VALUES (?,?,?,?,?,?)",
                 (user["project_id"], key, dumps(tpl), 0, user["id"], now_iso()))
    audit(conn, user["id"], "conceal.template", f"tpl:{key}", f"save {tpl['name']}")
    return get_template(conn, user["project_id"], key)


def delete_template(conn, user: dict, key: str, restore: bool = False) -> dict:
    pid = user["project_id"]
    row = one(conn, "SELECT * FROM conceal_templates WHERE project_id=? AND key=?", (pid, key))
    if key in BUILTIN_KEYS:
        if restore:
            conn.execute("DELETE FROM conceal_templates WHERE project_id=? AND key=?", (pid, key))
        else:
            conn.execute("INSERT OR REPLACE INTO conceal_templates (project_id, key, body_json, hidden, updated_by, updated_at) VALUES (?,?,?,?,?,?)",
                         (pid, key, row["body_json"] if row else "{}", 1, user["id"], now_iso()))
    else:
        if not row:
            raise ConcealError(404, "模板不存在")
        conn.execute("DELETE FROM conceal_templates WHERE project_id=? AND key=?", (pid, key))
    audit(conn, user["id"], "conceal.template", f"tpl:{key}", "restore" if restore else "delete")
    return {"ok": True}


def match_score(tpl: dict, text: str) -> int:
    return max((len(k) for k in tpl["keywords"] if k and k in (text or "")), default=0)


def match_templates(tpls: list[dict], text: str) -> list[dict]:
    hits = [(match_score(t, text), t) for t in tpls]
    return [t for s, t in sorted((h for h in hits if h[0]), key=lambda h: -h[0])]


# ----------------------------------------------------------------------------- 范围与读取

def _today() -> str:
    return date.today().isoformat()


def _np(part: str | None) -> str:
    p = (part or "").strip()
    return part_svc.part_from_text(p) or p


def _act(conn, pid: int, code: str | None) -> dict | None:
    return next((a for a in p6.current_activities(conn, pid) if a["code"] == code), None) if code else None


SCOPE_MSG = "所选作业不在您的负责范围内；不关联作业时可填写任意具体部位"


def _scope_ok(user: dict, act: dict | None) -> bool:
    """D20：关联作业时，技术员只能关联自己负责范围（wbs_scope）内的作业；范围为空表示不限。
    不关联作业时可填任意具体部位，只记录作者（可见范围仍按作者/单位过滤，见 get_row、list_records）。"""
    if not act or user["role"] != "tech" or not (user.get("wbs_scope") or "").strip():
        return True
    return p6.in_scope(act, user["wbs_scope"])


def get_row(conn, user: dict, rid: int) -> dict:
    row = one(conn, "SELECT * FROM conceal_records WHERE id=? AND project_id=?", (rid, user["project_id"]))
    if not row:
        raise ConcealError(404, "验收记录不存在")
    if user["role"] == "tech" and not in_scope(conn, user, row["created_by"]):
        raise ConcealError(403, "不在您的范围内")
    return row


def _tpl_for(conn, row: dict) -> dict:
    """草稿用最新模板（模板被改后下一次检查生效）；已确认及之后用建档时的快照，保证报验依据不变。"""
    snap = loads(row.get("template_json"), None)
    live = get_template(conn, row["project_id"], row["template_key"])
    tpl = (live or snap) if row["status"] == "draft" else (snap or live)
    return tpl or {"key": row["template_key"], "name": row["template_key"], "group": "", "keywords": [], "shots": [], "items": [], "roles": [], "builtin": False, "customized": False}


def _meta(conn, ids: list[int]) -> dict[int, dict]:
    if not ids:
        return {}
    return {m["id"]: m for m in all_(conn, f"SELECT * FROM media WHERE id IN ({','.join('?' for _ in ids)})", ids)}


def _hist(row: dict, user: dict, action: str, note: str = "") -> str:
    h = loads(row.get("history_json"), []) or []
    h.append({"at": now_iso(), "by": user["id"], "name": user["name"], "action": action, "note": note[:200]})
    return dumps(h[-60:])


# ----------------------------------------------------------------------------- 缺项检查（确定性）

def check(tpl: dict, photos: list[dict], fields: dict, meta: dict[int, dict], ai: dict | None = None) -> dict:
    by_shot: dict[str, list[dict]] = defaultdict(list)
    for p in photos:
        if p.get("shot"):
            by_shot[p["shot"]].append(p)
    miss_shots = [{"key": s["key"], "label": s["label"], "hint": s.get("hint", "")} for s in tpl["shots"] if s.get("required", True) and not by_shot.get(s["key"])]
    pend = [{"key": i["key"], "label": i["label"], "unit": i.get("unit", "")} for i in tpl["items"]
            if i.get("required", True) and (fields.get(i["key"]) or {}).get("value") in (None, "")]
    warns = []
    if photos:
        ms = [meta.get(p["media_id"]) or {} for p in photos]
        no_t = [m for m in ms if not m.get("taken_at")]
        no_g = [m for m in ms if not geo_of(m)] if ms else []
        if len(no_t) == len(ms):
            warns.append({"kind": "no_time", "text": "拍摄时间缺失：所有照片都没有水印拍摄时间"})
        elif no_t:
            warns.append({"kind": "no_time", "text": f"{len(no_t)} 张照片缺拍摄时间"})
        if len(no_g) == len(ms):
            warns.append({"kind": "no_geo", "text": "定位缺失：所有照片都没有手机定位或地点"})
        elif no_g:
            warns.append({"kind": "no_geo", "text": f"{len(no_g)} 张照片缺定位"})
        unassigned = sum(1 for p in photos if not p.get("shot"))
        if unassigned:
            warns.append({"kind": "unassigned", "text": f"{unassigned} 张照片还没有归入清单项"})
        aip = (ai or {}).get("photos") or {}
        for s in tpl["shots"]:
            ps = by_shot.get(s["key"]) or []
            if ps and all((aip.get(str(p["media_id"])) or {}).get("clear") is False for p in ps):
                warns.append({"kind": "unclear", "text": f"“{s['label']}”的照片经识别不够清晰，建议重拍"})
    complete = not miss_shots and not pend
    return {"missing_shots": miss_shots, "pending_items": pend, "warnings": warns, "complete": complete,
            "conclusion": "资料齐全，可报验" if complete else "有缺项",
            "counts": {"shots_required": sum(1 for s in tpl["shots"] if s.get("required", True)), "shots_missing": len(miss_shots),
                       "items_required": sum(1 for i in tpl["items"] if i.get("required", True)), "items_pending": len(pend), "photos": len(photos)}}


def _snapshot(chk: dict) -> str:
    return dumps({"shots": [m["key"] for m in chk["missing_shots"]], "fields": [m["key"] for m in chk["pending_items"]],
                  "warnings": [w["kind"] for w in chk["warnings"]], "complete": chk["complete"], "conclusion": chk["conclusion"], "at": now_iso()})


def _refresh(conn, row: dict) -> dict:
    """重算并写回 missing_json 快照（供其他模块直接 SQL 读取），返回检查结果。"""
    tpl = _tpl_for(conn, row)
    photos = loads(row["photos_json"], []) or []
    chk = check(tpl, photos, loads(row["fields_json"], {}) or {}, _meta(conn, [p["media_id"] for p in photos]), loads(row["ai_json"], {}) or {})
    conn.execute("UPDATE conceal_records SET missing_json=?, updated_at=? WHERE id=?", (_snapshot(chk), now_iso(), row["id"]))
    return chk


def _save(conn, rid: int, **cols) -> dict:
    cols["updated_at"] = now_iso()
    conn.execute(f"UPDATE conceal_records SET {', '.join(k + '=?' for k in cols)} WHERE id=?", (*cols.values(), rid))
    return one(conn, "SELECT * FROM conceal_records WHERE id=?", (rid,))


# ----------------------------------------------------------------------------- 视图

def _users(conn, pid: int) -> dict[int, dict]:
    return {u["id"]: u for u in all_(conn, "SELECT id, name, company FROM users WHERE project_id=?", (pid,))}


def _photo_view(p: dict, m: dict | None, ai: dict | None) -> dict:
    m = m or {}
    return {"media_id": p["media_id"], "shot": p.get("shot"), "shot_src": p.get("shot_src"), "source": p.get("source"),
            "thumb": f"/api/media/{p['media_id']}/thumb", "url": f"/api/media/{p['media_id']}/file", "taken_at": m.get("taken_at"),
            "uploaded_at": m.get("uploaded_at"), "geo": geo_of(m) if m else None, "ai": ai}


def can_flags(user: dict, status: str) -> dict:
    w, r = user["role"] in WRITE_ROLES, user["role"] in RESULT_ROLES
    return {"edit": w and status == "draft", "analyze": w and status == "draft", "ready": w and status == "draft",
            "reopen": w and status in ("ready", "rejected"), "submit": w and status == "ready", "result": r and status == "submitted",
            "delete": w and status == "draft", "template": r}


def view(conn, user: dict, row: dict, users: dict | None = None) -> dict:
    tpl = _tpl_for(conn, row)
    users = users or _users(conn, row["project_id"])
    photos = loads(row["photos_json"], []) or []
    meta = _meta(conn, [p["media_id"] for p in photos])
    ai = loads(row["ai_json"], {}) or {}
    fields = loads(row["fields_json"], {}) or {}
    chk = check(tpl, photos, fields, meta, ai)
    aip = ai.get("photos") or {}
    act = _act(conn, row["project_id"], row["activity_code"])
    fv = {}
    for k, f in fields.items():
        it = next((i for i in tpl["items"] if i["key"] == k), None)
        doubt = None
        if it and it["type"] == "number" and isinstance(f.get("value"), (int, float)) and not conceal_agent._in_range(it, f["value"]):
            doubt = "数值超出常见范围，请核对"
        fv[k] = {**f, "doubt": doubt}
    return {"id": row["id"], "project_id": row["project_id"], "title": row["title"], "part": row["part"], "activity_code": row["activity_code"],
            "activity_name": act["name"] if act else None, "template_key": row["template_key"], "template": tpl,
            "status": row["status"], "status_name": STATUS_NAME[row["status"]], "inspection_date": row["inspection_date"],
            "fields": fv, "narration": row.get("narration") or "", "photos": [_photo_view(p, meta.get(p["media_id"]), aip.get(str(p["media_id"]))) for p in photos],
            "ai": {"analyzed": bool(ai.get("analyzed_at")), "analyzed_at": ai.get("analyzed_at"), "mock": bool(ai.get("mock")), "provider": ai.get("provider"),
                   "model": ai.get("model"), "fields": ai.get("fields") or [], "notice": "离线模拟识别，仅供开发测试" if ai.get("mock") else "AI 识别结果仅供参考，须人工核对后采用"},
            "check": chk, "ready_note": row.get("ready_note") or "", "forced": row["status"] != "draft" and bool(row.get("ready_note")),
            "supervisor_name": row.get("supervisor_name") or "", "submitted_at": row.get("submitted_at"), "result_note": row.get("result_note") or "",
            "result_at": row.get("result_at"), "created_by": users.get(row["created_by"], {}).get("name", "已离职"), "created_at": row["created_at"],
            "history": [{"at": h["at"], "name": h["name"], "action": h["action"], "note": h.get("note", "")} for h in (loads(row.get("history_json"), []) or [])],
            "can": can_flags(user, row["status"]), "notice": NOTICE, "measure_note": MEASURE_NOTE, "status_names": STATUS_NAME}


def _brief(conn, row: dict, tpls: dict, users: dict) -> dict:
    ms = loads(row.get("missing_json"), {}) or {}
    t = tpls.get(row["template_key"])
    photos = loads(row["photos_json"], []) or []
    return {"id": row["id"], "title": row["title"], "part": row["part"], "activity_code": row["activity_code"], "template_key": row["template_key"],
            "template_name": (t or loads(row.get("template_json"), {}) or {}).get("name", row["template_key"]), "status": row["status"],
            "status_name": STATUS_NAME[row["status"]], "inspection_date": row["inspection_date"], "photos": len(photos),
            "missing_shots": len(ms.get("shots") or []), "pending_items": len(ms.get("fields") or []), "complete": bool(ms.get("complete")),
            "forced": row["status"] != "draft" and bool(row.get("ready_note")), "created_by": users.get(row["created_by"], {}).get("name", "已离职"),
            "company": users.get(row["created_by"], {}).get("company", ""), "submitted_at": row.get("submitted_at"), "result_at": row.get("result_at"),
            "result_note": row.get("result_note") or "", "updated_at": row.get("updated_at") or row["created_at"]}


def list_records(conn, user: dict, status: str = "", part: str = "", activity: str = "", template: str = "") -> dict:
    pid = user["project_id"]
    sql, args = "SELECT * FROM conceal_records WHERE project_id=?", [pid]
    if status:
        sql += " AND status=?"
        args.append(status)
    if part:
        sql += " AND part=?"
        args.append(part)
    if activity:
        sql += " AND activity_code=?"
        args.append(activity)
    if template:
        sql += " AND template_key=?"
        args.append(template)
    rows = [r for r in all_(conn, sql + " ORDER BY id DESC LIMIT 500", args) if in_scope(conn, user, r["created_by"])]
    tpls = {t["key"]: t for t in templates(conn, pid)}
    users = _users(conn, pid)
    items = [_brief(conn, r, tpls, users) for r in rows]
    counts = {k: sum(1 for i in items if i["status"] == k) for k in STATUS_NAME}
    return {"records": items, "counts": counts, "status_names": STATUS_NAME,
            "can_write": user["role"] in WRITE_ROLES, "can_result": user["role"] in RESULT_ROLES, "notice": NOTICE}


# ----------------------------------------------------------------------------- 创建、编辑

def _date(s: str | None, default: str | None = None) -> str:
    s = (s or "").strip() or (default or _today())
    try:
        return date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        raise ConcealError(400, "日期格式应为 YYYY-MM-DD") from None


def create(conn, user: dict, part: str, template_key: str, activity_code: str | None = None, inspection_date: str | None = None, title: str = "") -> dict:
    pid = user["project_id"]
    tpl = get_template(conn, pid, template_key)
    if not tpl:
        raise ConcealError(400, "验收模板不存在")
    act = None
    if activity_code:
        act = _act(conn, pid, activity_code)
        if not act:
            raise ConcealError(400, "作业不在当前 P6 计划里")
    part = (part or "").strip() or ((part_svc.activity_part(act) or "") if act else "")
    if not part:
        raise ConcealError(400, "请选择或填写部位")
    if not _scope_ok(user, act):
        raise ConcealError(403, SCOPE_MSG)
    rid = insert(conn, "conceal_records", {
        "project_id": pid, "activity_code": activity_code or None, "part": part[:40], "template_key": template_key,
        "title": (title or "").strip()[:60] or f"{part}{tpl['name']}", "status": "draft", "inspection_date": _date(inspection_date),
        "fields_json": "{}", "photos_json": "[]", "ai_json": "{}", "missing_json": "{}", "template_json": dumps(tpl), "history_json": "[]",
        "created_by": user["id"], "created_at": now_iso(), "updated_at": now_iso()})
    row = one(conn, "SELECT * FROM conceal_records WHERE id=?", (rid,))
    row = _save(conn, rid, history_json=_hist(row, user, "创建"))
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.create", f"conceal:{rid}", f"{part} {template_key}")
    return view(conn, user, one(conn, "SELECT * FROM conceal_records WHERE id=?", (rid,)))


def _draft(row: dict) -> None:
    if row["status"] != "draft":
        raise ConcealError(409, "记录已确认或报验，不能再改；请先“退回补资料”")


def update_meta(conn, user: dict, row: dict, title=None, inspection_date=None, narration=None, part=None) -> dict:
    _draft(row)
    cols = {}
    if title is not None and title.strip():
        cols["title"] = title.strip()[:60]
    if inspection_date:
        cols["inspection_date"] = _date(inspection_date)
    if narration is not None:
        cols["narration"] = narration.strip()[:1500]
    if part is not None and part.strip():
        if not _scope_ok(user, _act(conn, row["project_id"], row["activity_code"])):
            raise ConcealError(403, SCOPE_MSG)
        cols["part"] = part.strip()[:40]
    row = _save(conn, row["id"], **cols) if cols else row
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.edit", f"conceal:{row['id']}", ",".join(cols))
    return view(conn, user, one(conn, "SELECT * FROM conceal_records WHERE id=?", (row["id"],)))


def delete(conn, user: dict, row: dict) -> dict:
    _draft(row)
    conn.execute("DELETE FROM conceal_records WHERE id=?", (row["id"],))
    audit(conn, user["id"], "conceal.delete", f"conceal:{row['id']}", row["title"])
    return {"ok": True}


def set_fields(conn, user: dict, row: dict, payload: dict) -> dict:
    """人工录入数据项（设计值/要求值可选）。来源：人工 / 口述（采纳AI候选）/ 照片读数（采纳照片读数），采纳也算人工确认。"""
    _draft(row)
    tpl = _tpl_for(conn, row)
    item_by = {i["key"]: i for i in tpl["items"]}
    fields = loads(row["fields_json"], {}) or {}
    for k, f in (payload or {}).items():
        it = item_by.get(k)
        if not it:
            raise ConcealError(400, f"数据项不存在：{k}")
        f = f if isinstance(f, dict) else {"value": f}
        raw = f.get("value")
        if raw is None or str(raw).strip() == "":
            fields.pop(k, None)
            continue
        if it["type"] == "number":
            v = conceal_agent.num(raw if not isinstance(raw, str) else raw.strip())
            if v is None:
                raise ConcealError(400, f"“{it['label']}”需要填数字")
        else:
            v = str(raw).strip()[:80]
        src = f.get("src") if f.get("src") in FIELD_SRC else "人工"
        fields[k] = {"value": v, "design": str(f.get("design") or "").strip()[:40], "src": src, "by": user["id"], "at": now_iso()}
    row = _save(conn, row["id"], fields_json=dumps(fields))
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.fields", f"conceal:{row['id']}", ",".join((payload or {}).keys()))
    return view(conn, user, one(conn, "SELECT * FROM conceal_records WHERE id=?", (row["id"],)))


# ----------------------------------------------------------------------------- 照片

def _valid_shot(tpl: dict, shot: str | None) -> str | None:
    if shot in (None, ""):
        return None
    if shot not in {s["key"] for s in tpl["shots"]}:
        raise ConcealError(400, "清单项不存在")
    return shot


def _add_photos(conn, user: dict, row: dict, entries: list[dict]) -> dict:
    photos = loads(row["photos_json"], []) or []
    have = {p["media_id"] for p in photos}
    for e in entries:
        if e["media_id"] not in have:
            photos.append(e)
            have.add(e["media_id"])
    row = _save(conn, row["id"], photos_json=dumps(photos))
    _refresh(conn, row)
    return row


def add_uploads(conn, user: dict, row: dict, files: list[dict], shot: str | None) -> dict:
    _draft(row)
    tpl = _tpl_for(conn, row)
    shot = _valid_shot(tpl, shot)
    if len(files) > 20:
        raise ConcealError(400, "一次最多上传20张")
    entries = []
    for f in files:
        try:
            m = save_photo(conn, user["project_id"], user["id"], f["name"], f["raw"], "conceal", taken_hint=f.get("taken_at"), geo=f.get("geo"))
        except MediaError as exc:
            raise ConcealError(400, f"{f['name']}: {exc}") from None
        entries.append({"media_id": m["id"], "shot": shot, "shot_src": "人工" if shot else None, "source": "当场拍摄", "added_by": user["id"], "added_at": now_iso()})
    row = _add_photos(conn, user, row, entries)
    audit(conn, user["id"], "conceal.photo", f"conceal:{row['id']}", f"upload={len(entries)} shot={shot or '-'}")
    return view(conn, user, row)


def _media_day(m: dict) -> str:
    return (m.get("taken_at") or m.get("uploaded_at") or "")[:10]


def attach_existing(conn, user: dict, row: dict, media_ids: list[int], shot: str | None) -> dict:
    """从部位已有照片里引用：限验收日期前后 1 天，且是本人/本范围可见的照片。"""
    _draft(row)
    tpl = _tpl_for(conn, row)
    shot = _valid_shot(tpl, shot)
    day = date.fromisoformat(row["inspection_date"])
    ok_days = {(day + timedelta(days=d)).isoformat() for d in (-1, 0, 1)}
    meta = _meta(conn, media_ids)
    entries, skipped = [], []
    for mid in media_ids:
        m = meta.get(mid)
        if not m or m["project_id"] != row["project_id"] or not in_scope(conn, user, m["uploader_id"]):
            skipped.append({"media_id": mid, "why": "照片不存在或无权引用"})
        elif _media_day(m) not in ok_days:
            skipped.append({"media_id": mid, "why": "照片不在验收日期前后 1 天内"})
        else:
            entries.append({"media_id": mid, "shot": shot, "shot_src": "人工" if shot else None, "source": "部位已有照片", "added_by": user["id"], "added_at": now_iso()})
    if not entries:
        raise ConcealError(400, "没有可引用的照片：" + "；".join(s["why"] for s in skipped[:3]))
    row = _add_photos(conn, user, row, entries)
    audit(conn, user["id"], "conceal.photo", f"conceal:{row['id']}", f"attach={len(entries)} shot={shot or '-'}")
    out = view(conn, user, row)
    out["skipped"] = skipped
    return out


def set_photo_shot(conn, user: dict, row: dict, mid: int, shot: str | None) -> dict:
    _draft(row)
    shot = _valid_shot(_tpl_for(conn, row), shot)
    photos = loads(row["photos_json"], []) or []
    p = next((x for x in photos if x["media_id"] == mid), None)
    if not p:
        raise ConcealError(404, "记录里没有这张照片")
    p["shot"], p["shot_src"] = shot, ("人工" if shot else None)
    row = _save(conn, row["id"], photos_json=dumps(photos))
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.photo", f"conceal:{row['id']}", f"shot media={mid} -> {shot or '-'}")
    return view(conn, user, row)


def remove_photo(conn, user: dict, row: dict, mid: int) -> dict:
    """只是把照片从记录里取下，照片原件仍留在照片库（原件不可改写）。"""
    _draft(row)
    photos = [p for p in (loads(row["photos_json"], []) or []) if p["media_id"] != mid]
    row = _save(conn, row["id"], photos_json=dumps(photos))
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.photo", f"conceal:{row['id']}", f"remove media={mid}")
    return view(conn, user, row)


def accept_ai_shots(conn, user: dict, row: dict, media_ids: list[int] | None = None) -> dict:
    """人确认采用AI对照片的归类建议：只处理还没归类的照片；采用后才计入缺项检查。"""
    _draft(row)
    ai = (loads(row["ai_json"], {}) or {}).get("photos") or {}
    photos = loads(row["photos_json"], []) or []
    keys = {s["key"] for s in _tpl_for(conn, row)["shots"]}
    n = 0
    for p in photos:
        sug = (ai.get(str(p["media_id"])) or {}).get("shot")
        if not p.get("shot") and sug in keys and (not media_ids or p["media_id"] in media_ids):
            p["shot"], p["shot_src"] = sug, "AI建议·人确认"
            n += 1
    row = _save(conn, row["id"], photos_json=dumps(photos))
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.photo", f"conceal:{row['id']}", f"accept_ai={n}")
    out = view(conn, user, row)
    out["accepted"] = n
    return out


def candidates(conn, user: dict, part: str, day: str | None, record_id: int | None = None) -> dict:
    """该部位已有的照片（日志照片、随手拍、随手记照片等），限验收日期前后 1 天；另列同日未归部位的照片供人挑选。"""
    day = _date(day)
    d0 = date.fromisoformat(day)
    ok_days = {(d0 + timedelta(days=d)).isoformat() for d in (-1, 0, 1)}
    have = set()
    if record_id:
        r = one(conn, "SELECT photos_json FROM conceal_records WHERE id=? AND project_id=?", (record_id, user["project_id"]))
        have = {p["media_id"] for p in (loads(r["photos_json"], []) if r else [])}

    def pick(name: str, unassigned: bool) -> list[dict]:
        try:
            d = part_svc.detail(conn, user, name, ["photo", "note"], limit=3000)
        except ValueError:
            return []
        out, seen = [], set()
        for it in d["items"]:
            for p in it.get("photos") or []:
                if p["id"] in seen or p.get("purpose") == "ticket":
                    continue
                seen.add(p["id"])
                m = _meta(conn, [p["id"]]).get(p["id"])
                if not m or not in_scope(conn, user, m["uploader_id"]):
                    continue
                if ((p.get("taken_at") or p.get("uploaded_at") or "")[:10]) in ok_days:
                    out.append({"media_id": p["id"], "thumb": p["thumb"], "url": p["url"], "taken_at": p.get("taken_at"), "uploaded_at": p.get("uploaded_at"),
                                "geo": p.get("geo"), "purpose_name": p.get("purpose_name"), "in_record": p["id"] in have, "unassigned": unassigned})
        return out

    own = pick(part, False)
    ids = {p["media_id"] for p in own}
    extra = [p for p in pick(part_svc.UNASSIGNED, True) if p["media_id"] not in ids]
    key = lambda p: p["taken_at"] or p["uploaded_at"] or ""  # noqa: E731
    return {"part": part, "date": day, "photos": sorted(own, key=key), "unassigned": sorted(extra, key=key),
            "note": "限验收日期前后 1 天；照片归部位按拍照时填写的地点/日志挂接识别，识别不出的在“未归部位”里。"}


# ----------------------------------------------------------------------------- 资料Agent识别

def analyze(conn, user: dict, row: dict, narration: str | None = None) -> dict:
    _draft(row)
    if narration is not None:
        row = _save(conn, row["id"], narration=narration.strip()[:1500])
    narr = (row.get("narration") or "").strip()
    tpl = _tpl_for(conn, row)
    photos = loads(row["photos_json"], []) or []
    if not photos and not narr:
        raise ConcealError(400, "先拍几张照片或写一段口述，再让资料Agent整理")
    meta = _meta(conn, [p["media_id"] for p in photos])
    refs = [(f"P{i + 1}", p) for i, p in enumerate(photos)]
    assign = conceal_agent.mock_assign(tpl, [{"ref": r, "shot": p.get("shot")} for r, p in refs])
    merged_photos: dict[str, dict] = {}
    fields: list[dict] = []
    run_ids, provider, model = [], settings.ai_provider, ""
    chunks = [refs[i:i + conceal_agent.MAX_PER_CALL] for i in range(0, len(refs), conceal_agent.MAX_PER_CALL)] or [[]]
    for ci, chunk in enumerate(chunks):
        names = [r for r, _ in chunk]
        paths = []
        for _r, p in chunk:
            m = meta.get(p["media_id"])
            path = (settings.media_dir / m["ai_file"]) if m else None
            if path and path.is_file():
                paths.append(str(path))
        text_narr = narr if ci == 0 else ""
        sub_assign = {r: assign.get(r) for r in names}
        try:
            res = call_json("conceal.pack", system=conceal_agent.SYSTEM, user=conceal_agent.build_user(tpl, chunk, text_narr),
                            prompt_version=conceal_agent.PROMPT_VERSION, images=paths,
                            mock=lambda a=sub_assign, n=text_narr: conceal_agent.mock_response(tpl, a, n),
                            validate=lambda d, nm=names: conceal_agent.normalize(d, nm, tpl),
                            project_id=user["project_id"], user_id=user["id"], max_tokens=2500,
                            input_meta={"task": "conceal", "template": tpl["key"], "photos": len(chunk), "narration_chars": len(text_narr)})
        except AIError as exc:
            raise ConcealError(502, f"资料Agent暂不可用：{str(exc)[:120]}") from None
        run_ids.append(res.run_id)
        provider, model = res.provider, res.model
        for ref, info in res.data["photos"].items():
            merged_photos[ref] = info
        if ci == 0:
            fields = res.data["fields"]
    norm = conceal_agent.apply_rules(tpl, {"photos": merged_photos, "fields": fields}, narr)
    by_media = {str(photos[int(r[1:]) - 1]["media_id"]): v for r, v in norm["photos"].items()}
    ai = {"analyzed_at": now_iso(), "provider": provider, "model": model, "mock": provider == "mock", "run_ids": run_ids,
          "prompt_version": conceal_agent.PROMPT_VERSION, "photos": by_media, "fields": norm["fields"], "narration": narr}
    row = _save(conn, row["id"], ai_json=dumps(ai))
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.analyze", f"conceal:{row['id']}", f"photos={len(photos)} runs={len(run_ids)} provider={provider}")
    return view(conn, user, one(conn, "SELECT * FROM conceal_records WHERE id=?", (row["id"],)))


# ----------------------------------------------------------------------------- 消息（F-04）

def link(rid: int) -> dict:
    """消息跳转目标：隐蔽验收记录页 #/conceal?id=。"""
    return {"route": "conceal", "params": {"id": str(rid)}}


def _notify(conn, row: dict, uids: list[int], title: str, body: str) -> None:
    """站内消息，带记录跳转。收件人按岗位职责定，不发项目经理（AGENTS.md：项目经理只看进度、质量、安全）。"""
    for uid in uids:
        insert(conn, "notifications", {"project_id": row["project_id"], "user_id": uid, "alert_id": None, "title": title[:80], "body": body,
                                       "channels_json": dumps({"inapp": "sent"}), "created_at": now_iso(), "link_json": dumps(link(row["id"]))})


def _what(row: dict, tpl_name: str) -> str:
    """消息里的“哪份验收”：标题已含部位/模板名就不重复。"""
    t = row["title"] or ""
    s = t if row["part"] in t else f"{row['part']} · {t}"
    return s if tpl_name in t else f"{s}（{tpl_name}）"


def _notify_submitted(conn, user: dict, row: dict) -> None:
    """报验 → 本项目工程主管（操作人本人除外）：待录入监理结果。"""
    to = [r["id"] for r in all_(conn, "SELECT id FROM users WHERE project_id=? AND role='eng_sup' AND active=1 AND id<>?", (row["project_id"], user["id"]))]
    sup = f"，监理：{row['supervisor_name']}" if row.get("supervisor_name") else ""
    forced = "（带缺项说明确认）" if row.get("ready_note") else ""
    _notify(conn, row, to, f"隐蔽验收待录入监理结果：{row['title']}",
            f"{user['name']}（{R.role_name(user['role'])}）已报验 {_what(row, _tpl_for(conn, row)['name'])}{forced}{sup}。"
            "监理验收后请在记录里录入结果；退回须写明原因。")


def _notify_result(conn, user: dict, row: dict, result: str, note: str) -> None:
    """监理结果 → 编制人与报验人（操作人本人除外）：退回写明原因，通过附监理意见（如有）。"""
    ids = [u for u in dict.fromkeys((row.get("created_by"), row.get("submitted_by"))) if u and u != user["id"]]
    to = [r["id"] for r in all_(conn, f"SELECT id FROM users WHERE project_id=? AND active=1 AND id IN ({','.join('?' for _ in ids)})",
                                (row["project_id"], *ids))] if ids else []
    who = f"{user['name']}（{R.role_name(user['role'])}）"
    sup = f"监理 {row['supervisor_name']}" if row.get("supervisor_name") else "监理"
    what = _what(row, _tpl_for(conn, row)["name"])
    if result == "rejected":
        _notify(conn, row, to, f"隐蔽验收被监理退回：{row['title']}",
                f"{who}录入{sup}退回 {what}。退回原因：{note}。请点“退回补资料”改好后重新报验。")
    else:
        _notify(conn, row, to, f"隐蔽验收监理通过：{row['title']}",
                f"{who}录入{sup}通过 {what}" + (f"。监理意见：{note}" if note else "") + "。")


# ----------------------------------------------------------------------------- 状态流转

def set_ready(conn, user: dict, row: dict, note: str = "") -> dict:
    _draft(row)
    chk = _refresh(conn, row)
    note = (note or "").strip()
    if not chk["complete"] and not note:
        raise ConcealError(400, "还有缺项：请补齐，或写明缺项说明后再确认")
    row = _save(conn, row["id"], status="ready", ready_note=note[:300] if not chk["complete"] else "", ready_by=user["id"], ready_at=now_iso(),
                template_json=dumps(_tpl_for(conn, row)), history_json=_hist(row, user, "确认" + ("（带缺项）" if not chk["complete"] else "齐全"), note))
    audit(conn, user["id"], "conceal.ready", f"conceal:{row['id']}", "forced" if not chk["complete"] else "complete")
    return view(conn, user, row)


def reopen(conn, user: dict, row: dict) -> dict:
    if row["status"] not in ("ready", "rejected"):
        raise ConcealError(409, "只有“已齐待报验”或“监理退回”的记录可退回补资料")
    row = _save(conn, row["id"], status="draft", ready_note="", ready_by=None, ready_at=None, submitted_at=None, supervisor_name="", result_note="", result_at=None,
                history_json=_hist(row, user, "退回补资料", row.get("result_note") or ""))
    _refresh(conn, row)
    audit(conn, user["id"], "conceal.reopen", f"conceal:{row['id']}", "")
    return view(conn, user, one(conn, "SELECT * FROM conceal_records WHERE id=?", (row["id"],)))


def submit(conn, user: dict, row: dict, supervisor_name: str = "") -> dict:
    if row["status"] != "ready":
        raise ConcealError(409, "需先确认资料（已齐待报验）才能报验")
    row = _save(conn, row["id"], status="submitted", submitted_at=now_iso(), submitted_by=user["id"], supervisor_name=(supervisor_name or "").strip()[:30],
                result_note="", result_at=None, history_json=_hist(row, user, "报验", supervisor_name or ""))
    audit(conn, user["id"], "conceal.submit", f"conceal:{row['id']}", row["supervisor_name"])
    _notify_submitted(conn, user, row)
    return view(conn, user, row)


def record_result(conn, user: dict, row: dict, result: str, note: str = "", day: str | None = None, supervisor_name: str = "") -> dict:
    """人工录入监理验收结果与日期；系统不判定合格。"""
    if row["status"] != "submitted":
        raise ConcealError(409, "只有已报验的记录可以录入监理结果")
    if result not in ("passed", "rejected"):
        raise ConcealError(400, "结果只能是通过或退回")
    note = (note or "").strip()
    if result == "rejected" and not note:
        raise ConcealError(400, "退回须写明原因")
    sup = (supervisor_name or "").strip()[:30] or row.get("supervisor_name") or ""
    row = _save(conn, row["id"], status=result, result_note=note[:300], result_at=_date(day), result_by=user["id"], supervisor_name=sup,
                history_json=_hist(row, user, "监理通过" if result == "passed" else "监理退回", note))
    audit(conn, user["id"], "conceal.result", f"conceal:{row['id']}", result)
    _notify_result(conn, user, row, result, note[:300])
    return view(conn, user, row)


# ----------------------------------------------------------------------------- 其他模块只读口径

def covered(conn, project_id: int, activity_code: str | None = None, part: str | None = None, template_key: str | None = None) -> bool:
    """是否已有 ready/submitted/passed 的隐蔽验收记录（供放行清单等引用）。按作业或同部位匹配；可限定模板。"""
    sql = f"SELECT activity_code, part FROM conceal_records WHERE project_id=? AND status IN ({','.join('?' for _ in OK_STATUS)})"
    args = [project_id, *OK_STATUS]
    if template_key:
        sql += " AND template_key=?"
        args.append(template_key)
    for r in all_(conn, sql, args):
        if (activity_code and r["activity_code"] == activity_code) or (part and _np(r["part"]) == _np(part)):
            return True
    return False


# ----------------------------------------------------------------------------- 创建选项与缺资料扫描

def _started(a: dict) -> bool:
    return bool(a.get("act_start") or a.get("act_finish") or a.get("status") in ("进行中", "已完成") or (a.get("phys_pct") or 0) > 0)


_GATE_TPL = {"rebar": "rebar", "form": "formwork", "hole": "pile", "cage": "pile"}   # 放行清单 conceal:<kind> → 验收模板


def _gate_templates(conn, pid: int, acts: list[dict], tpls: list[dict]) -> dict[str, list[dict]]:
    """F-11①：放行清单按作业类别要核验的隐蔽验收（如混凝土浇筑要钢筋、模板验收）对应的模板。
    只读放行Agent 的规则与已缓存归类（不调用模型）；放行清单不可用时返回空，不影响按名称匹配的结果。"""
    try:
        from . import gate
        cls = gate.classify_all(conn, pid, acts, allow_ai=False)
    except (ImportError, sqlite3.Error):
        return {}
    by_key = {t["key"]: t for t in tpls}
    out: dict[str, list[dict]] = {}
    for a in acts:
        cat = (cls.get(a["code"]) or {}).get("category")
        kinds = [c["auto"].split(":", 1)[1] for c in gate.SPECIFIC.get(cat, []) if (c.get("auto") or "").startswith("conceal:")]
        keys = [k for k in dict.fromkeys(_GATE_TPL[k] for k in kinds if k in _GATE_TPL) if k in by_key]
        if keys:
            out[a["code"]] = [by_key[k] for k in keys]
    return out


def options(conn, user: dict) -> dict:
    """新建验收的部位/作业/模板选项：作业名命中模板关键词的，加上放行清单要核验隐蔽验收的作业（含未开工）；技术员只列自己范围。"""
    pid = user["project_id"]
    tpls = templates(conn, pid)
    acts = [a for a in p6.current_activities(conn, pid) if not a.get("is_milestone")]
    if user["role"] == "tech":
        acts = [a for a in acts if p6.in_scope(a, user.get("wbs_scope") or "")]
    by_gate = _gate_templates(conn, pid, acts, tpls)
    by_part: dict[str, list[dict]] = defaultdict(list)
    for a in acts:
        ms = match_templates(tpls, f"{a['name']} {a['wbs_path']}")
        ms += [t for t in by_gate.get(a["code"], []) if t["key"] not in {m["key"] for m in ms}]
        if not ms:
            continue
        by_part[part_svc.activity_part(a) or part_svc.UNASSIGNED].append(
            {"code": a["code"], "name": a["name"], "status": a["status"], "started": _started(a), "templates": [t["key"] for t in ms]})
    parts = [{"name": n, "activities": by_part[n]} for n in sorted(by_part, key=part_svc.sort_key)]
    return {"parts": parts, "templates": [{"key": t["key"], "name": t["name"], "group": t["group"]} for t in tpls], "has_plan": bool(acts),
            "note": "" if acts else "尚未导入 P6 计划（或没有在您范围内的作业），可直接填写部位名称新建。"}


def gaps(conn, user: dict) -> dict:
    """已开始/已完成的作业，按作业名与 WBS 匹配验收模板；找不到 ready/submitted/passed 的对应记录就列为“疑似缺验收资料”。"""
    pid = user["project_id"]
    acts = p6.current_activities(conn, pid)
    note = "“疑似”：依据 = P6 作业状态 + 系统内没找到对应的已齐/已报验/已通过记录；是否确需验收，以设计、规范和监理要求为准。"
    empty = {"gaps": [], "parts": [], "groups": [], "summary": {"total": 0, "done": 0, "pending": 0, "missing": 0}, "has_plan": False,
             "note": "尚未导入 P6 计划，无法扫描缺验收资料。", "generated_at": now_iso()}
    if not acts:
        return empty
    if user["role"] == "tech":
        acts = [a for a in acts if p6.in_scope(a, user.get("wbs_scope") or "")]
    tpls = templates(conn, pid)
    recs = all_(conn, "SELECT id, activity_code, part, template_key, status, created_by FROM conceal_records WHERE project_id=?", (pid,))
    out, by_part, by_group = [], defaultdict(lambda: {"done": 0, "pending": 0, "missing": 0}), defaultdict(lambda: {"done": 0, "pending": 0, "missing": 0})
    for a in acts:
        if a.get("is_milestone") or not _started(a):
            continue
        part = part_svc.activity_part(a) or part_svc.UNASSIGNED
        for t in match_templates(tpls, f"{a['name']} {a['wbs_path']}"):
            mine = [r for r in recs if r["template_key"] == t["key"] and (r["activity_code"] == a["code"] or (part != part_svc.UNASSIGNED and _np(r["part"]) == _np(part)))]
            if any(r["status"] in OK_STATUS for r in mine):
                state = "done"
            elif mine:
                state = "pending"
            else:
                state = "missing"
            by_part[part][state] += 1
            by_group[t["group"] or "其他"][state] += 1
            if state == "done":
                continue
            draft = next((r for r in mine if r["status"] == "draft" and (user["role"] != "tech" or in_scope(conn, user, r["created_by"]))), None)
            began = a.get("act_start") or a.get("early_start")
            out.append({"activity_code": a["code"], "activity_name": a["name"], "part": part, "wbs_path": a["wbs_path"], "template_key": t["key"], "template_name": t["name"],
                        "p6_status": a["status"], "act_start": a.get("act_start"), "act_finish": a.get("act_finish"), "state": state,
                        "state_name": "有记录但未齐（草稿/监理退回）" if state == "pending" else "未找到验收记录",
                        "basis": f"P6 作业“{a['status'] or '已开始'}”" + (f"（实际开始 {a['act_start']}）" if a.get("act_start") else "") + "，" +
                                 ("有对应记录但还没确认齐全或被监理退回" if state == "pending" else "系统内没有找到对应的验收记录"),
                        "record_ids": [r["id"] for r in mine if user["role"] != "tech" or in_scope(conn, user, r["created_by"])], "draft_id": draft["id"] if draft else None,
                        "_k": (0 if a.get("act_finish") else 1, a.get("act_start") or began or "9999")})
    out.sort(key=lambda g: g.pop("_k"))
    for g in out:
        g.pop("_k", None)
    tot = {k: sum(v[k] for v in by_part.values()) for k in ("done", "pending", "missing")}
    tot["total"] = sum(tot.values())
    plist = [{"part": n, "total": sum(by_part[n].values()), **by_part[n]} for n in sorted(by_part, key=part_svc.sort_key)]
    glist = [{"group": n, "total": sum(v.values()), **v} for n, v in sorted(by_group.items())]
    return {"gaps": out, "parts": plist, "groups": glist, "summary": tot, "has_plan": True, "note": note, "generated_at": now_iso()}


# ----------------------------------------------------------------------------- Word 成册

def _fmt_when(p: dict) -> str:
    return (p.get("taken_at") or "").replace("T", " ") or "拍摄时间缺失"


def _record_doc_section(doc, conn, v: dict, level: int = 1) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401
    tpl = v["template"]
    doc.add_heading(f"{v['title']}", level=level)
    part_svc._table(doc, ["部位", "作业", "验收项目", "验收日期", "编制人", "记录状态"],
                    [[v["part"], f"{v['activity_code']} {v['activity_name'] or ''}".strip() if v["activity_code"] else "—", tpl["name"], v["inspection_date"],
                      v["created_by"], v["status_name"] + ("（带缺项确认）" if v["forced"] else "")]])
    chk = v["check"]
    doc.add_paragraph("资料检查：" + chk["conclusion"] + ("" if chk["complete"] else
                      "；缺拍：" + ("、".join(m["label"] for m in chk["missing_shots"]) or "无") + "；待补数据：" + ("、".join(m["label"] for m in chk["pending_items"]) or "无")))
    if v["ready_note"]:
        doc.add_paragraph("缺项说明（人工填写）：" + v["ready_note"])
    for w in chk["warnings"]:
        doc.add_paragraph("提示：" + w["text"])
    doc.add_heading("一、验收数据（现场实测，由人录入）", level=level + 1)
    rows = []
    for it in tpl["items"]:
        f = v["fields"].get(it["key"]) or {}
        rows.append([it["label"] + ("（必填）" if it.get("required", True) else ""), f.get("design") or "", "待补" if f.get("value") in (None, "") else f["value"], it.get("unit", ""),
                     f.get("src", "") if f.get("value") not in (None, "") else ""])
    if rows:
        part_svc._table(doc, ["数据项", "设计值/要求值", "实测/记录值", "单位", "来源"], rows)
    else:
        doc.add_paragraph("该模板未设数据项。")
    doc.add_paragraph("设计值/要求值、允许偏差以设计文件、项目验收规范和监理要求为准，由人填写。")
    doc.add_heading("二、验收照片", level=level + 1)
    for s in tpl["shots"]:
        ps = [p for p in v["photos"] if p["shot"] == s["key"]]
        doc.add_paragraph(f"{s['label']}" + ("" if s.get("required", True) else "（可选）") + (f"：{s['hint']}" if s.get("hint") else "") + ("" if ps else "　【缺拍】"))
        if ps:
            part_svc.add_photos(doc, conn, [{"id": p["media_id"], "taken_at": p["taken_at"], "uploaded_at": p["uploaded_at"], "geo": p["geo"]} for p in ps], limit=9)
            notes = [f"P{p['media_id']}：{p['ai']['scene']}" for p in ps if p.get("ai") and p["ai"].get("scene")]
            if notes:
                doc.add_paragraph("（AI 识别画面，仅供参考）" + "；".join(notes))
    left = [p for p in v["photos"] if not p["shot"]]
    if left:
        doc.add_paragraph("未归入清单项的照片：")
        part_svc.add_photos(doc, conn, [{"id": p["media_id"], "taken_at": p["taken_at"], "uploaded_at": p["uploaded_at"], "geo": p["geo"]} for p in left], limit=9)
    doc.add_heading("三、验收意见", level=level + 1)
    if v["result_at"]:
        doc.add_paragraph(f"系统登记的监理结果（仅供核对，以签认为准）：{v['status_name']}，{v['result_at']}" + (f"，{v['result_note']}" if v["result_note"] else ""))
    part_svc._table(doc, ["施工单位意见", "监理单位验收意见", "监理工程师签字", "日期"], [["", "", "", ""], ["", "", "", ""]])
    doc.add_paragraph(FOOTER)


def _new_doc():
    doc = part_svc._doc_base()
    doc.sections[0].footer.paragraphs[0].text = FOOTER
    return doc


def _title(doc, text: str, sub: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc.add_heading(text, level=0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(sub).alignment = WD_ALIGN_PARAGRAPH.CENTER


def to_docx(conn, user: dict, row: dict):
    v = view(conn, user, row)
    doc = _new_doc()
    _title(doc, "隐蔽工程验收记录", f"导出日期：{_today()}    导出人：{user['name']}" + ("    （照片归类含离线模拟识别）" if v["ai"]["mock"] else ""))
    _record_doc_section(doc, conn, v, level=1)
    out = settings.export_dir / f"conceal_{user['project_id']}_{row['id']}.docx"
    _an.humanize_doc(doc)   # 作业代码 → 中文简称（services/act_names.py）
    doc.save(out)
    audit(conn, user["id"], "conceal.export", f"conceal:{row['id']}", "single")
    return out


def part_docx(conn, user: dict, part: str, include_draft: bool = False):
    sts = ("draft", *OK_STATUS) if include_draft else OK_STATUS
    rows = [r for r in all_(conn, f"SELECT * FROM conceal_records WHERE project_id=? AND part=? AND status IN ({','.join('?' for _ in sts)}) ORDER BY inspection_date, id",
                            (user["project_id"], part, *sts)) if in_scope(conn, user, r["created_by"])]
    if not rows:
        raise ConcealError(404, "该部位还没有已确认的隐蔽验收记录")
    users = _users(conn, user["project_id"])
    views = [view(conn, user, r, users) for r in rows]
    doc = _new_doc()
    _title(doc, f"隐蔽工程验收资料册：{part}", f"导出日期：{_today()}    导出人：{user['name']}    共 {len(views)} 份记录")
    part_svc._table(doc, ["序号", "验收项目", "作业", "验收日期", "记录状态", "资料检查", "监理结果"],
                    [[i + 1, v["template"]["name"], v["activity_code"] or "—", v["inspection_date"], v["status_name"], v["check"]["conclusion"],
                      (v["result_at"] or "") + (f" {v['result_note']}" if v["result_note"] else "") if v["result_at"] else "未登记"] for i, v in enumerate(views)])
    doc.add_paragraph(NOTICE)
    for v in views:
        doc.add_page_break()
        _record_doc_section(doc, conn, v, level=1)
    safe = re.sub(r"[^0-9A-Za-z一-鿿]+", "_", part)[:20]
    out = settings.export_dir / f"conceal_{user['project_id']}_part_{safe}.docx"
    _an.humanize_doc(doc)   # 作业代码 → 中文简称（services/act_names.py）
    doc.save(out)
    audit(conn, user["id"], "conceal.export", f"conceal-part:{part}", f"records={len(views)}")
    return out

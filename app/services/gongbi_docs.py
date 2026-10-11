"""工笔文书（周月报以外的类型）：事实底稿冻结入库 → 交融起草 → 逐栏核对（不过的栏退回规则版）→ 人工核对修改 → 确认冻结 → 导出 Word。

- 类型定义见 ai/docs_registry.py；周报 / 月报仍由 services/reports.py 处理，行为不变；
- 事实底稿只放系统里已确认的数据：不放人名、不放产值与金额（所有岗位一致）；叙述里写进“产值”由核对标出，各岗位看到同一份文字（RJ 10-10 项目经理只读产值后不再单独裁剪）；
- 确认后的文书冻结，不能再改；再次起草同一对象 = 新版本（version + 1，revision_of 指向上一版）；
- 对外文书由人发送，系统不对外发送任何信函；通知函的“已发出”在索赔证据页由人登记（claim_notice.decide）。
"""
from __future__ import annotations

import time
from datetime import date, timedelta
from pathlib import Path

from ..ai import docs_registry as reg
from ..ai.client import AIError, call_json, cause_of
from ..config import settings
from ..db import all_, audit, dumps, event, insert, loads, now_iso, one, tx
from . import claim_notice, claims, p6, reports, schedule, sub_letters

SCHEMA = """
CREATE TABLE IF NOT EXISTS gongbi_docs (
    id INTEGER PRIMARY KEY,
    project_id INTEGER NOT NULL,
    kind TEXT NOT NULL,
    ref TEXT NOT NULL,
    title TEXT NOT NULL,
    lang TEXT NOT NULL,
    period_start TEXT,
    period_end TEXT,
    facts_json TEXT NOT NULL,
    narrative_json TEXT NOT NULL,
    check_json TEXT,
    source TEXT,
    ai_run_id INTEGER,
    status TEXT NOT NULL DEFAULT 'draft' CHECK (status IN ('draft','confirmed')),
    version INTEGER NOT NULL DEFAULT 1,
    revision_of INTEGER,
    created_by INTEGER,
    created_at TEXT NOT NULL,
    confirmed_by INTEGER,
    confirmed_at TEXT,
    gen_seconds REAL
);
CREATE INDEX IF NOT EXISTS ix_gongbi_docs ON gongbi_docs(project_id, kind, ref);
"""

LEVEL_EN = {"red": "Red", "orange": "Orange", "yellow": "Yellow"}
ALERT_STATUS_EN = {"open": "open", "confirmed": "confirmed", "false_alarm": "false alarm", "closed": "closed"}
NOTICE_STATUS_EN = {"pending": "pending", "sent": "notice issued", "waived": "no notice (decided)"}


def _type(kind: str) -> dict:
    t = reg.TYPES.get(kind)
    if not t or t.get("legacy"):
        raise ValueError("不支持的文书类型")
    return t


# ----------------------------------------------------------------------------- 事实底稿

def period_facts(conn, pid: int, kind: str, ref: str | None) -> dict:
    """英文周报 / 月报：与中文周报 / 月报同一份统计（reports.build_facts），去掉产值、人名、填报率分母与自报 / 核定拆分。"""
    start, end = reports.period_for(reg.TYPES[kind]["ref"], ref)
    f = reports.build_facts(conn, pid, start, end)
    if not f["logs"]["count"]:
        raise ValueError("该周期没有已提交的日志，无法起草")
    out = {
        "project": f["project"], "period_start": start, "period_end": end, "p6": f.get("p6"), "photos": f["photos"],
        "logs": {"count": f["logs"]["count"], "ids": f["logs"]["ids"]},
        "progress": [{"code": p["code"], "name": p["name"], "qty": [{"qty": q["qty"], "unit": q["unit"]} for q in p["qty"]],
                      "end_pct": p["end_pct"], "plan_pct": p["plan_pct"], "gap": p["gap"], "critical": p.get("critical", False)} for p in f["progress"]],
        "milestones": [{k: m.get(k) for k in ("code", "name", "p6_date", "forecast", "p80", "prob_pct")} for m in f.get("milestones", [])],
        "safety": {**{k: (f.get("safety") or {}).get(k, 0) for k in ("found", "high", "resolved", "open")},
                   "by_category": (f.get("safety") or {}).get("by_category") or {},
                   "open_high": [{k: h[k] for k in ("id", "title", "due_date")} for h in (f.get("safety") or {}).get("open_high", [])]},
        "quality": f.get("quality") or {}, "labor": f.get("labor", []), "machinery": f.get("machinery", []),
        "issues": [{k: i.get(k, "") for k in ("log_id", "date", "type", "severity", "desc", "measure", "activity_code")} for i in f["issues"]],
        "alerts": [{"id": a["id"], "level": LEVEL_EN.get(a["level"], a["level"]), "title": a["title"], "status": ALERT_STATUS_EN.get(a["status"], a["status"])} for a in f["alerts"]],
        "works": [{"log_id": w["log_id"], "date": w["date"], "text": w["text"]} for w in f["works"]],
        "next": f["next"],
    }
    if kind == "monthly_en":
        st = claim_notice.states(conn, pid)
        out["notices"] = sorted(({"key": k, "cat_en": reg.CAT_EN[n["event"]["cat"]], "cat_name": n["event"]["cat_name"], "date_from": n["event"]["date_from"],
                                  "date_to": n["event"]["date_to"], "deadline": n["deadline"], "status_en": NOTICE_STATUS_EN[n["status"]], "status_name": n["status_name"]}
                                 for k, n in st.items() if n["event"]["date_to"] >= start and n["event"]["date_from"] <= end), key=lambda x: x["date_from"])
    return out


def _downstream_milestones(conn, pid: int, codes: list[str]) -> list[dict] | None:
    """按当前 P6 计划的逻辑关系，从受影响作业往后找里程碑（最新进度预测里的里程碑），附当前预测日期。没有预测返回 None。"""
    fc = schedule.latest_forecast(conn, pid)
    imp = p6.current_import(conn, pid)
    if not fc or not imp:
        return None
    succ: dict[str, list[str]] = {}
    for r in all_(conn, "SELECT pred_code, succ_code FROM relations WHERE import_id=?", (imp["id"],)):
        succ.setdefault(r["pred_code"], []).append(r["succ_code"])
    seen, stack = set(codes), list(codes)
    while stack:
        for s in succ.get(stack.pop(), []):
            if s not in seen:
                seen.add(s)
                stack.append(s)
    return [{k: m.get(k) for k in ("code", "name", "p6_date", "forecast", "p80")} for m in fc.get("milestones", []) if m["code"] in seen and m["code"] not in codes]


def event_facts(conn, user: dict, kind: str, key: str) -> dict:
    """通知函 / 延误简报：单个事件的 claims.evidence + 通知时效。不放受影响天数估算（不出天数结论）、不放人名与单位。"""
    pid = user["project_id"]
    ev = claim_notice._event(conn, pid, key)
    evd = claims.evidence(conn, user, ev)
    st = claim_notice.view_one(conn, pid, key)
    project = one(conn, "SELECT name FROM projects WHERE id=?", (pid,))["name"]
    affected = [{k: a.get(k) for k in ("code", "name", "plan_start", "plan_finish", "float_d", "critical")} for a in evd["impact"]["activities"]]
    out = {
        "project": project, "key": key,
        "event": {"key": key, "cat": ev["cat"], "cat_name": ev["cat_name"], "cat_en": reg.CAT_EN[ev["cat"]], "date_from": ev["date_from"], "date_to": ev["date_to"],
                  "log_days": len(ev["dates"]), "halt_dates": evd["impact"]["halt_dates"],
                  "issues": [{k: i[k] for k in ("log_id", "date", "type", "severity", "desc", "measure")} for i in ev["issues"]],
                  "weather": [w for w in evd["weather"] if w["weather"]]},
        "notice": {**{k: st[k] for k in ("aware_date", "deadline", "days", "status", "status_name")}, "clause": claim_notice.CLAUSE, "clause_en": claim_notice.CLAUSE_EN},
        "affected": affected, "logs": {"ids": sorted({i["log_id"] for i in ev["issues"]})},
    }
    if kind == "delay_brief":
        out["milestones"] = _downstream_milestones(conn, pid, [a["code"] for a in affected])
        out["gaps"] = [f"{g['item']}（{g['why']}）" for g in evd["gaps"] if g["auto"]][:6]
    return out


def build_facts(conn, user: dict, kind: str, ref: str | None) -> dict:
    rt = reg.TYPES[kind]["ref"]
    if reg.TYPES[kind].get("facts"):   # 增补类型（例会底稿、三周滚动计划、月度复盘）自带事实底稿
        return reg.TYPES[kind]["facts"](conn, user, kind, ref)
    if rt in ("week", "month"):
        return period_facts(conn, user["project_id"], kind, ref)
    if rt == "source":
        return {**sub_letters.source_facts(conn, user["project_id"], ref or ""), "ref": (ref or "").partition("|")[0]}
    if rt == "company":
        return {**sub_letters.company_facts(conn, user["project_id"], ref or ""), "ref": ref}
    return event_facts(conn, user, kind, ref or "")


# ----------------------------------------------------------------------------- 起草

def finalize(kind: str, narrative: dict, facts: dict, ai: bool) -> tuple[dict, list[dict], list[str]]:
    """逐栏核对；AI 稿里核对不过的栏目退回规则版文字。返回（最终文字，剩余提示，退回规则版的栏目）。"""
    ftext = reg.facts_text(kind, facts)
    issues = reg.check(kind, narrative, ftext)
    back: list[str] = []
    if ai and issues:
        rule = reg.rule(kind, facts)
        back = sorted({x["field"] for x in issues if x["field"] in rule})
        narrative = {**narrative, **{k: rule[k] for k in back}}
        issues = reg.check(kind, narrative, ftext)
    return narrative, issues, back


def fallback_note(kind: str, back: list[str]) -> str:
    """“AI 稿数字核对未通过、换成规则版”的提示：写栏目标题，不露字段名（验收 I17）。"""
    names = {s["key"]: s["title"] for s in reg.template(kind)["sections"]}
    return "这些栏目的 AI 稿数字与系统事实对不上，已换成规则版文字：" + "、".join(f"“{names.get(k, k)}”" for k in back)


_OLD_BACK = "以下栏目数字核对未通过，已退回规则版："


def _note(kind: str, chk: dict) -> str:
    """旧记录的提示里带字段名：显示时按 fallback 重写。"""
    note = chk.get("note", "")
    if _OLD_BACK in note:
        note = "；".join(x for x in (note.split(_OLD_BACK)[0].rstrip("；"), fallback_note(kind, chk.get("fallback") or [])) if x)
    return note


def _ref_of(kind: str, facts: dict) -> str:
    rt = reg.TYPES[kind]["ref"]
    return facts["period_start"] if rt in ("week", "month") else facts["ref"] if rt in ("source", "company") else facts["key"]


def _title(kind: str, facts: dict) -> str:
    t = reg.TYPES[kind]["title"]
    ev = facts.get("event")
    t = t.replace("{company}", facts.get("company", "")).replace("{title}", (facts.get("title") or "")[:40]).replace("{area}", facts.get("area_name", ""))
    return t.replace("{project}", facts["project"]).replace("{event}", f"{ev['cat_en'] if reg.TYPES[kind]['lang'] == 'en' else ev['cat_name']} {ev['date_from']}" if ev else "")


def generate(conn, user: dict, kind: str, ref: str | None) -> dict:
    t = _type(kind)
    if user["role"] not in t["roles"]:
        raise PermissionError("当前岗位不能起草这类文书")
    t0 = time.monotonic()
    facts = build_facts(conn, user, kind, ref)
    if kind == "work_contact" and user["role"] not in sub_letters.ISSUER[facts["area"]]:
        raise PermissionError(f"{facts['area_name']}类工作联系单由{'安全主管' if facts['area'] == 'safety' else '工程主管'}签发")
    ftext = reg.facts_text(kind, facts)
    source, run_id, note = "ai", None, ""
    try:
        r = call_json(f"report.{kind}", system=reg.system_prompt(kind), user=ftext, prompt_version=reg.PROMPT_VERSION,
                      mock=lambda: reg.rule(kind, facts), validate=lambda d: reg.normalize(kind, d), project_id=user["project_id"], user_id=user["id"],
                      max_tokens=3500, input_meta={"kind": kind, "ref": _ref_of(kind, facts)})
        nar, run_id = r.data, r.run_id
        if r.provider == "mock":
            source, note = "rule", "离线模拟模型：文字为规则版"
    except AIError as e:
        nar, source, note = reg.rule(kind, facts), "rule", f"模型暂不可用（{cause_of(e)}），已用规则版文字"
    nar, issues, back = finalize(kind, nar, facts, source == "ai")
    if back:
        note = "；".join(x for x in (note, fallback_note(kind, back)) if x)
    ref_key = _ref_of(kind, facts)
    prev = one(conn, "SELECT id, version FROM gongbi_docs WHERE project_id=? AND kind=? AND ref=? ORDER BY version DESC LIMIT 1", (user["project_id"], kind, ref_key))
    with tx(conn):
        did = insert(conn, "gongbi_docs", {
            "project_id": user["project_id"], "kind": kind, "ref": ref_key, "title": _title(kind, facts), "lang": t["lang"],
            "period_start": facts.get("period_start") or facts.get("event", {}).get("date_from") or facts.get("date") or facts.get("since"), "period_end": facts.get("period_end") or facts.get("event", {}).get("date_to") or facts.get("due_date"),
            "facts_json": dumps(facts), "narrative_json": dumps(nar), "check_json": dumps({"issues": issues, "fallback": back, "note": note}),
            "source": source, "ai_run_id": run_id, "status": "draft", "version": (prev["version"] + 1) if prev else 1, "revision_of": prev["id"] if prev else None,
            "created_by": user["id"], "created_at": now_iso(), "gen_seconds": round(time.monotonic() - t0, 1)})
        event(conn, "gongbi.generate", project_id=user["project_id"], user_id=user["id"], ref_id=did, value=round(time.monotonic() - t0, 1), data={"kind": kind})
    return get(conn, user, did)


# ----------------------------------------------------------------------------- 查看、修改、确认

def _row(conn, user: dict, did: int) -> dict:
    r = one(conn, "SELECT * FROM gongbi_docs WHERE id=? AND project_id=?", (did, user["project_id"]))
    if not r or (user["role"] == "leader" and r["status"] != "confirmed"):
        raise LookupError("文书不存在")
    return r


def tables(kind: str, f: dict) -> dict:
    """表格由程序从事实底稿直接填（数字不经模型）。"""
    g = reg._g
    en = reg.TYPES[kind]["lang"] == "en"
    if reg.TYPES[kind].get("tables"):
        return reg.TYPES[kind]["tables"](f)
    out = {}
    if f.get("progress") is not None:
        out["progress"] = {"headers": ["Code", "Activity", "Qty this period", "Cumulative %", "Planned %", "Variance"],
                           "rows": [[p["code"], p["name"], ", ".join(f"{q['qty']:g} {q['unit']}" for q in p["qty"]) or "—", g(p["end_pct"]), g(p["plan_pct"]), g(p["gap"])] for p in f["progress"]]}
        out["safety"] = {"headers": ["Category", "Count"], "rows": [[k, v] for k, v in (f.get("safety") or {}).get("by_category", {}).items()]}
        out["issues"] = {"headers": ["Date", "Type", "Severity", "Issue (original)", "Measure (original)", "Ref."],
                         "rows": [[i["date"], i["type"], i["severity"], i["desc"], i["measure"], f"L{i['log_id']}"] for i in f["issues"]]}
        out["next"] = {"headers": ["Activity", "Planned start", "Planned finish", "Critical"],
                       "rows": [[f"{n['code']} {n['name']}", n["start"], n["finish"], "Yes" if n["critical"] else ""] for n in f["next"]]}
        out["resources"] = {"headers": ["Resource", "Average daily"], "rows": [[x["name"], g(x["avg"])] for x in f.get("labor", []) + f.get("machinery", [])]}
        if f.get("notices") is not None:
            out["notices"] = {"headers": ["Event", "From", "To", "Notice deadline", "Status"],
                              "rows": [[n["cat_en"], n["date_from"], n["date_to"], n["deadline"], n["status_en"]] for n in f["notices"]]}
    if f.get("milestones"):
        out["milestones"] = ({"headers": ["Milestone", "Programme", "Forecast", "P80"] + (["On-time %"] if "prob_pct" in f["milestones"][0] else []),
                              "rows": [[f"{m['code']} {m['name']}", m["p6_date"], m["forecast"], m["p80"]] + ([g(m["prob_pct"])] if "prob_pct" in m else []) for m in f["milestones"]]}
                             if en else {"headers": ["里程碑", "P6 日期", "当前预测", "P80"], "rows": [[f"{m['code']} {m['name']}", m["p6_date"], m["forecast"], m["p80"]] for m in f["milestones"]]})
    if f.get("problems") is not None:
        out["problems"] = {"headers": ["编号", "日期", "问题", "状态"], "rows": [[p["ref"], p["date"], p["title"], p["status"]] for p in f["problems"]]}
    if f.get("affected") is not None:
        out["affected"] = {"headers": ["作业", "计划开始", "计划完成", "总浮时（天）", "关键"],
                           "rows": [[f"{a['code']} {a['name']}", a["plan_start"] or "—", a["plan_finish"] or "—", g(a["float_d"]), "是" if a["critical"] else ""] for a in f["affected"]]}
    return out


def view(conn, r: dict, role: str = "") -> dict:
    t = reg.TYPES[r["kind"]]
    facts, nar, chk = loads(r["facts_json"]), loads(r["narrative_json"], {}), loads(r["check_json"], {}) or {}
    tpl = reg.template(r["kind"])
    tpl["title"] = r["title"]
    tb = tables(r["kind"], facts)
    return {**{k: r[k] for k in ("id", "kind", "ref", "title", "lang", "period_start", "period_end", "status", "version", "revision_of", "source", "created_at", "confirmed_at", "gen_seconds")},
            "type_name": t["name"], "type_zh": t.get("zh") or t["name"], "outward": bool(t.get("outward")), "template": tpl, "narrative": nar,
            "tables": {s["key"]: tb[s["table"]] for s in tpl["sections"] if s.get("table") and s["table"] in tb},
            "check": chk.get("issues", []), "fallback": chk.get("fallback", []), "note": _note(r["kind"], chk), "can_confirm": role in t["confirm"],
            "can_edit": r["status"] == "draft" and role in t["roles"] | t["confirm"],
            "notice": facts.get("notice"), "event_key": facts.get("key"), "company": facts.get("company"), "area": facts.get("area"),
            "letter": _letter_brief(conn, r["id"]) if r["kind"] in reg.SUB_KINDS else None, "project_title": facts.get("project", ""), "logs": len((facts.get("logs") or {}).get("ids", []))}


def _letter_brief(conn, did: int) -> dict | None:
    x = one(conn, "SELECT * FROM sub_letters WHERE doc_id=?", (did,))
    return {"id": x["id"], **{k: v for k, v in sub_letters.state(conn, x).items() if k != "finding"}, "acked_at": x["acked_at"]} if x else None


def get(conn, user: dict, did: int) -> dict:
    return view(conn, _row(conn, user, did), user["role"])


def listing(conn, user: dict, kind: str | None = None, ref: str | None = None) -> list[dict]:
    q, a = "SELECT * FROM gongbi_docs WHERE project_id=?", [user["project_id"]]
    if kind:
        q, a = q + " AND kind=?", a + [kind]
    if ref:
        q, a = q + " AND ref=?", a + [ref]
    rows = all_(conn, q + " ORDER BY id DESC LIMIT 200", a)
    if user["role"] == "leader":
        rows = [r for r in rows if r["status"] == "confirmed"]
    return [{**{k: r[k] for k in ("id", "kind", "ref", "title", "lang", "period_start", "period_end", "status", "version", "source", "created_at", "confirmed_at")},
             "type_name": reg.TYPES[r["kind"]]["name"], "warnings": len((loads(r["check_json"], {}) or {}).get("issues", [])),
             "fallback": len((loads(r["check_json"], {}) or {}).get("fallback", []))} for r in rows if r["kind"] in reg.TYPES]


def edit(conn, user: dict, did: int, narrative: dict, confirm: bool = False) -> dict:
    r = _row(conn, user, did)
    t = _type(r["kind"])
    if r["status"] == "confirmed":
        raise ValueError("已确认的文书已冻结，不能修改；如需修改请重新起草新版本")
    if user["role"] not in t["roles"] | t["confirm"]:
        raise PermissionError("当前岗位不能修改这类文书")
    if confirm and user["role"] not in t["confirm"]:
        raise PermissionError("当前岗位不能确认这类文书")
    facts = loads(r["facts_json"])
    if confirm and r["kind"] == "work_contact" and user["role"] not in sub_letters.ISSUER[facts["area"]]:
        raise PermissionError(f"{facts['area_name']}类工作联系单由{'安全主管' if facts['area'] == 'safety' else '工程主管'}签发")
    old = loads(r["narrative_json"], {})
    nar = reg.normalize(r["kind"], {**old, **(narrative or {})}) if narrative else old
    issues = reg.check(r["kind"], nar, reg.facts_text(r["kind"], facts))
    chk = {**(loads(r["check_json"], {}) or {}), "issues": issues}
    fields = {"narrative_json": dumps(nar), "check_json": dumps(chk)}
    if confirm:
        fields |= {"status": "confirmed", "confirmed_by": user["id"], "confirmed_at": now_iso()}
    with tx(conn):
        conn.execute(f"UPDATE gongbi_docs SET {', '.join(k + '=?' for k in fields)} WHERE id=?", [*fields.values(), did])
        if confirm and r["kind"] in reg.SUB_KINDS:      # 确认即下发：只进分包账号的待办和消息，不对外发送；找不到接收人则整体回滚
            sub_letters.issue(conn, user, r, facts)
        audit(conn, user["id"], "gongbi.confirm" if confirm else "gongbi.edit", f"gongbi:{did}", r["kind"])
    return get(conn, user, did)


# ----------------------------------------------------------------------------- Word

def to_docx(conn, user: dict, did: int) -> tuple[Path, str]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    v = get(conn, user, did)
    en = v["lang"] == "en"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman" if en else "宋体"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    st.font.size = Pt(11 if en else 10.5)
    draft = v["status"] != "confirmed"
    if v["kind"] in reg.SUB_KINDS:
        h = doc.add_heading("工作联系单" if v["kind"] == "work_contact" else "函", level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for line in (f"致：{v['company']}", f"发文：{v['project_title']} 项目部", "编号：［由项目部填写］", "日期：［签发日期］"):
            doc.add_paragraph(line)
        doc.add_paragraph("事由：" + (v["narrative"].get("subject") or "")).runs[0].bold = True
    elif v["kind"] == "notice_of_claim":
        for line in ("To: The Engineer", "From: The Contractor", "Our Ref.: [to be filled in]", "Date: [date of issue]"):
            doc.add_paragraph(line)
        doc.add_paragraph("Subject: " + (v["narrative"].get("subject") or "")).runs[0].bold = True
    else:
        h = doc.add_heading(v["title"], level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if v["period_start"] and reg.TYPES[v["kind"]]["ref"] != "event":
            p = doc.add_paragraph((f"Reporting period: {v['period_start']} to {v['period_end']}    Based on {v['logs']} confirmed daily logs") if en
                                  else f"期间：{v['period_start']} 至 {v['period_end']}")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if draft:
        doc.add_paragraph("DRAFT — for internal review, not issued" if en else "草稿 — 内部核对用，未确认").runs[0].italic = True

    def table(tb):
        t = doc.add_table(rows=1, cols=len(tb["headers"]))
        t.style = "Table Grid"
        for i, h in enumerate(tb["headers"]):
            t.rows[0].cells[i].text = str(h)
        for row in tb["rows"]:
            cells = t.add_row().cells
            for i, x in enumerate(row):
                cells[i].text = "" if x is None else str(x)

    for sec in v["template"]["sections"]:
        if v["kind"] in ("notice_of_claim",) + reg.SUB_KINDS and sec["key"] == "subject":
            continue
        doc.add_heading(sec["title"], level=1)
        val = v["narrative"].get(sec["key"])
        if sec["type"] == "list":
            for c in val or []:
                doc.add_paragraph(c, style="List Bullet")
            if not val:
                doc.add_paragraph("None" if en else "无")
        elif val:
            for para in str(val).split("\n"):
                doc.add_paragraph(para)
        if sec["key"] in v["tables"] and v["tables"][sec["key"]]["rows"]:
            table(v["tables"][sec["key"]])
    if v["kind"] in reg.SUB_KINDS:
        doc.add_paragraph("")
        doc.add_paragraph("签发：［签发人、岗位］" + ("    项目经理确认：［签字］" if v["kind"] == "letter_to_sub" else ""))
        doc.add_paragraph("签收：［分包单位签收人、日期］")
    if v["kind"] == "notice_of_claim":
        doc.add_paragraph("")
        doc.add_paragraph("For and on behalf of the Contractor")
        doc.add_paragraph("[Name and position of the authorised representative]")
    doc.add_heading("Source of data" if en else "附：数据来源说明", level=1)
    doc.add_paragraph("Figures are compiled by the system from confirmed site daily logs and the P6 programme; references [L] are daily log numbers, [W] programme warnings, "
                      "[S] safety hazards. The text was drafted with AI assistance and checked against the system facts before review by the project team." if en else
                      "数字由系统从已确认的施工日志及 P6 计划统计；L 为施工日志编号，W 为进度预警编号，S 为安全隐患编号。文字由交融大模型起草、经数字核对后由项目部审核。")
    if v["kind"] in reg.SUB_KINDS:
        doc.add_paragraph("本文书由系统依据已确认的隐患、质量问题或进度预警记录起草，经签发人确认后下发；不含罚款、扣款及处罚结论。")
    if v["kind"] in ("notice_of_claim", "delay_brief"):
        doc.add_paragraph("是否构成索赔事件以商务和合同部门判断为准；本文件不含金额与工期天数结论。")
    nums = [x["number"] for x in v["check"] if x.get("number")]
    if nums and draft:
        doc.add_paragraph("核对提示：以下内容请审核人核实：" + "、".join(dict.fromkeys(nums)))
    out = settings.export_dir / f"gongbi_{did}_v{v['version']}.docx"
    _an.humanize_doc(doc)   # 作业代码 → 中文简称（services/act_names.py）
    doc.save(out)
    audit(conn, user["id"], "gongbi.docx", f"gongbi:{did}", v["kind"])
    safe = "".join(c for c in v["title"] if c not in '\\/:*?"<>|').strip()[:60]
    return out, f"{safe}_v{v['version']}.docx"


def ref_suggest(conn, user: dict, kind: str) -> list[dict]:
    """可选对象。事件类：按通知时效排序（待处理、剩余天数少的在前）；工作联系单：本人能签发的已确认问题与属实预警；致分包函：单位 × 类别。"""
    rt = reg.TYPES.get(kind, {}).get("ref")
    if reg.TYPES.get(kind, {}).get("refs"):   # 增补类型（预控提醒、资料补齐通知）自带可选对象
        return reg.TYPES[kind]["refs"](conn, user)
    if rt == "source":
        return _source_refs(conn, user)
    if rt == "company":
        return _company_refs(conn, user)
    if rt != "event":
        return []
    st = claim_notice.states(conn, user["project_id"])
    cut = (date.today() - timedelta(days=365)).isoformat()
    rows = [n for n in st.values() if n["event"]["date_to"] >= cut]
    return [{"key": n["event"]["key"], "label": f"{n['event']['cat_name']} {n['event']['date_from']}～{n['event']['date_to'][5:]}", "notice": n["label"], "tone": n["tone"]}
            for n in sorted(rows, key=claim_notice.sort_key)][:60]


def _source_refs(conn, user: dict) -> list[dict]:
    pid, out = user["project_id"], []
    areas = [a for a, roles in sub_letters.ISSUER.items() if user["role"] in roles and a != "progress"]
    if areas:
        for f in all_(conn, f"""SELECT f.id, f.agent, f.title, f.level, f.due_date, u.company FROM findings f LEFT JOIN users u ON u.id=f.assignee_id
                               WHERE f.project_id=? AND f.agent IN ({','.join('?' for _ in areas)}) AND f.kind<>'qc_reminder' AND f.confirmed_at IS NOT NULL
                               AND f.status IN ('open','rectifying','rejected') ORDER BY f.id DESC LIMIT 60""", [pid] + areas):
            if f["company"]:
                out.append({"key": f"finding:{f['id']}", "label": f"{sub_letters.AREA_NAME[f['agent']]} · {f['title'][:40]}", "notice": f"{f['company']} · 期限 {(f['due_date'] or '—')[:10]}",
                            "tone": "red" if f["level"] == "high" else "orange"})
    if user["role"] in sub_letters.ISSUER["progress"]:
        for a in all_(conn, "SELECT id, title, level, activity_code, as_of FROM alerts WHERE project_id=? AND status='confirmed' AND rule<>'stale' ORDER BY as_of DESC, id DESC LIMIT 40", (pid,)):
            out.append({"key": f"alert:{a['id']}", "label": f"进度 · {a['title'][:40]}", "notice": "已确认属实" + (f" · 判定 {a['as_of'][5:10]}" if a.get("as_of") else ""),
                        "tone": {"red": "red", "orange": "orange"}.get(a["level"], "blue")})
    return out


def _company_refs(conn, user: dict) -> list[dict]:
    pid = user["project_id"]
    since = (date.today() - timedelta(days=sub_letters.REPEAT_DAYS)).isoformat()
    rows = all_(conn, """SELECT u.company, f.agent, COUNT(*) n, SUM(f.status IN ('open','rectifying','rejected')) o FROM findings f JOIN users u ON u.id=f.assignee_id
                         WHERE f.project_id=? AND f.agent IN ('safety','quality') AND f.kind<>'qc_reminder' AND f.confirmed_at IS NOT NULL AND u.company<>''
                         AND substr(f.created_at,1,10)>=? GROUP BY u.company, f.agent ORDER BY n DESC""", (pid, since))
    return [{"key": f"{r['company']}|{r['agent']}", "label": f"{r['company']} · {sub_letters.AREA_NAME[r['agent']]}类问题", "notice": f"近 {sub_letters.REPEAT_DAYS} 天 {r['n']} 项，未闭环 {r['o'] or 0} 项",
             "tone": "red" if (r["o"] or 0) >= 2 else "orange"} for r in rows]


from . import gongbi_more  # noqa: E402,F401  注册增补文书类型
from . import act_names as _an

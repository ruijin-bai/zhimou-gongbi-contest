"""P6 计划接入：XER（首选）或 P6 导出的 Excel。

一数一源：作业代码、WBS、计划日期、总浮时、关键标记均以 P6 导入批次为唯一来源，
系统内不另建一套计划；日报进度按 P6 作业代码挂接，并可导出回写 P6 的进度更新表。
"""
from __future__ import annotations

import codecs
import hashlib
import io
from datetime import date, datetime
from typing import Any

from ..db import all_, audit, insert, now_iso, one, tx


class P6Error(ValueError):
    pass


# --------------------------------------------------------------------------- XER

def decode(raw: bytes) -> str:
    for bom, enc in ((codecs.BOM_UTF16_LE, "utf-16-le"), (codecs.BOM_UTF16_BE, "utf-16-be"), (codecs.BOM_UTF8, "utf-8")):
        if raw.startswith(bom):
            return raw[len(bom):].decode(enc, errors="replace")
    for enc in ("utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def parse_xer(text: str) -> dict[str, list[dict[str, str]]]:
    tables: dict[str, list[dict[str, str]]] = {}
    current, fields = None, []
    for line in text.splitlines():
        parts = line.rstrip("\r").split("\t")
        tag = parts[0]
        if tag == "%T":
            current = parts[1].strip() if len(parts) > 1 else None
            fields = []
            if current:
                tables.setdefault(current, [])
        elif tag == "%F":
            fields = [p.strip() for p in parts[1:]]
        elif tag == "%R" and current:
            tables[current].append({fields[i] if i < len(fields) else f"col{i}": v for i, v in enumerate(parts[1:])})
        elif tag == "%E":
            current, fields = None, []
    return tables


def _d(value: Any) -> str | None:
    """P6 日期 → YYYY-MM-DD；无法识别返回 None（不猜）。"""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%b-%y %H:%M", "%d-%b-%y", "%Y/%m/%d", "%d/%m/%Y", "%Y/%m/%d %H:%M"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _f(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError:
        return None


def _hours_to_days(hours: float | None, hpd: float) -> float | None:
    return None if hours is None else round(hours / (hpd or 8.0), 2)


_STATUS = {"TK_NotStart": "未开始", "TK_Active": "进行中", "TK_Complete": "已完成"}
_REL = {"PR_FS": "FS", "PR_SS": "SS", "PR_FF": "FF", "PR_SF": "SF"}
PCT_MODES = ("phys", "typed")


def _sum_qty(*values: Any) -> float | None:
    nums = [v for v in (_f(x) for x in values) if v is not None]
    return sum(nums) if nums else None


def typed_pct(row: dict) -> float | None:
    """按 P6 作业自身的完成%类型（complete_pct_type）给出完成%；无法确定时返回 None（不补零）。

    - CP_Drtn 工期%＝（原定工期 − 尚需工期）/ 原定工期，限制在 0～100；
    - CP_Units 数量%＝实际数量 /（实际数量 + 尚需数量），人工与机械数量合计（只读数量，不读任何费用字段）；
    - CP_Phys 与未标注：照读实物%（phys_complete_pct），与旧行为一致。
    已完工按 100、未开始按 0（P6 自身口径）。
    """
    kind = (row.get("complete_pct_type") or "").strip()
    if kind not in ("CP_Drtn", "CP_Units"):
        return _f(row.get("phys_complete_pct"))
    if row.get("status_code") == "TK_Complete" or _d(row.get("act_end_date")):
        return 100.0
    if row.get("status_code") == "TK_NotStart" and not _d(row.get("act_start_date")):
        return 0.0
    if kind == "CP_Drtn":
        orig, remain = _f(row.get("target_drtn_hr_cnt")), _f(row.get("remain_drtn_hr_cnt"))
        if orig and orig > 0 and remain is not None:
            return round(min(100.0, max(0.0, (orig - remain) / orig * 100.0)), 2)
        return None
    done = _sum_qty(row.get("act_work_qty"), row.get("act_equip_qty"))
    left = _sum_qty(row.get("remain_work_qty"), row.get("remain_equip_qty"))
    if done is not None and left is not None and done + left > 0:
        return round(min(100.0, max(0.0, done / (done + left) * 100.0)), 2)
    return None


def activities_from_xer(text: str, pct_mode: str = "phys") -> tuple[dict, list[dict], list[dict]]:
    """解析 XER。pct_mode="phys"（默认）照旧读实物%；"typed" 按作业的完成%类型读取（见 typed_pct），并记录 pct_type。

    两种模式对不带 complete_pct_type 的 XER 结果一致（typed 只多一个 pct_type=""）。
    """
    if pct_mode not in PCT_MODES:
        raise ValueError(f"未知的完成%读取方式：{pct_mode}")
    t = parse_xer(text)
    tasks = t.get("TASK") or []
    if not tasks:
        raise P6Error("文件中没有 TASK 表，请确认是 P6 导出的 .xer 文件")
    # 多项目 XER：取作业最多的项目
    counts: dict[str, int] = {}
    for row in tasks:
        counts[row.get("proj_id", "")] = counts.get(row.get("proj_id", ""), 0) + 1
    proj_id = max(counts, key=counts.get)
    project = next((p for p in t.get("PROJECT", []) if p.get("proj_id") == proj_id), {})
    calendars = {c.get("clndr_id"): _f(c.get("day_hr_cnt")) or 8.0 for c in t.get("CALENDAR", [])}
    wbs = {w.get("wbs_id"): w for w in t.get("PROJWBS", [])}

    def wbs_path(wbs_id: str | None) -> str:
        names, seen = [], set()
        while wbs_id and wbs_id in wbs and wbs_id not in seen:
            seen.add(wbs_id)
            node = wbs[wbs_id]
            if node.get("proj_node_flag") == "Y":
                break
            names.append(node.get("wbs_name") or node.get("wbs_short_name") or "")
            wbs_id = node.get("parent_wbs_id")
        return " / ".join(reversed([n for n in names if n]))

    acts, by_id = [], {}
    for row in tasks:
        if row.get("proj_id") != proj_id or row.get("task_type") in ("TT_WBS", "TT_LOE"):
            continue
        hpd = calendars.get(row.get("clndr_id"), 8.0)
        tf = _hours_to_days(_f(row.get("total_float_hr_cnt")), hpd)
        ttype = row.get("task_type") or ""
        code = (row.get("task_code") or "").strip()
        if not code:
            continue
        act = {
            "task_id": row.get("task_id", ""), "code": code, "name": (row.get("task_name") or "").strip(),
            "wbs_path": wbs_path(row.get("wbs_id")), "task_type": ttype,
            "status": _STATUS.get(row.get("status_code", ""), row.get("status_code") or ""),
            "target_start": _d(row.get("target_start_date")), "target_finish": _d(row.get("target_end_date")),
            "early_start": _d(row.get("early_start_date") or row.get("restart_date")),
            "early_finish": _d(row.get("early_end_date") or row.get("reend_date")),
            "late_start": _d(row.get("late_start_date")), "late_finish": _d(row.get("late_end_date")),
            "act_start": _d(row.get("act_start_date")), "act_finish": _d(row.get("act_end_date")),
            "orig_dur_d": _hours_to_days(_f(row.get("target_drtn_hr_cnt")), hpd),
            "remain_dur_d": _hours_to_days(_f(row.get("remain_drtn_hr_cnt")), hpd),
            "total_float_d": tf,
            "phys_pct": _f(row.get("phys_complete_pct")),
            "is_critical": 1 if (row.get("driving_path_flag") == "Y" or (tf is not None and tf <= 0)) else 0,
            "is_milestone": 1 if ttype in ("TT_Mile", "TT_FinMile") else 0,
        }
        if pct_mode == "typed":
            act["phys_pct"] = typed_pct(row)
            act["pct_type"] = (row.get("complete_pct_type") or "").strip()
        acts.append(act)
        by_id[act["task_id"]] = code
    rels = []
    for r in t.get("TASKPRED", []):
        pred, succ = by_id.get(r.get("pred_task_id")), by_id.get(r.get("task_id"))
        if pred and succ:
            rels.append({"pred_code": pred, "succ_code": succ, "kind": _REL.get(r.get("pred_type", ""), "FS"),
                         "lag_d": _hours_to_days(_f(r.get("lag_hr_cnt")), 8.0) or 0.0})
    meta = {"proj_name": project.get("proj_short_name") or project.get("proj_name") or "",
            "data_date": _d(project.get("last_recalc_date") or project.get("next_data_date"))}
    return meta, acts, rels


# --------------------------------------------------------------------------- Excel

_XLSX_ALIASES = {
    "code": ("task_code", "activity id", "作业代码", "作业id", "作业编号"),
    "name": ("task_name", "activity name", "作业名称"),
    "wbs_path": ("wbs", "wbs_name", "wbs名称", "wbs路径"),
    "early_start": ("start", "early_start_date", "开始", "开始时间", "计划开始", "最早开始"),
    "early_finish": ("finish", "early_end_date", "完成", "完成时间", "计划完成", "最早完成"),
    "act_start": ("actual start", "act_start_date", "实际开始"),
    "act_finish": ("actual finish", "act_end_date", "实际完成"),
    "total_float_d": ("total float", "total_float_hr_cnt", "总浮时", "总浮动"),
    "phys_pct": ("physical % complete", "phys_complete_pct", "实物完成百分比", "完成百分比", "完成%"),
    "orig_dur_d": ("original duration", "原定工期", "计划工期"),
    "remain_dur_d": ("remaining duration", "尚需工期", "剩余工期"),
    "predecessors": ("predecessors", "紧前作业"),
}


def activities_from_xlsx(raw: bytes) -> tuple[dict, list[dict], list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = None, {}
    for i, row in enumerate(rows[:5]):
        labels = [str(c).strip().lower() if c is not None else "" for c in row]
        found = {}
        for key, aliases in _XLSX_ALIASES.items():
            for j, label in enumerate(labels):
                if label in aliases:
                    found.setdefault(key, j)
        if "code" in found and "name" in found:
            header_idx, mapping = i, found
            # P6 导出 Excel 第二行常是中文/英文标签，跳过
            break
    if header_idx is None:
        raise P6Error("Excel 中未找到“作业代码/作业名称”表头（支持 P6 导出的 task_code/task_name 或中文列名）")
    acts, rels = [], []
    for row in rows[header_idx + 1:]:
        def cell(key):
            j = mapping.get(key)
            return row[j] if j is not None and j < len(row) else None
        code = str(cell("code") or "").strip()
        if not code or code.lower() in _XLSX_ALIASES["code"]:
            continue
        tf = _f(cell("total_float_d"))
        acts.append({
            "task_id": code, "code": code, "name": str(cell("name") or "").strip(), "wbs_path": str(cell("wbs_path") or ""),
            "task_type": "", "status": "已完成" if cell("act_finish") else ("进行中" if cell("act_start") else "未开始"),
            "target_start": _d(cell("early_start")), "target_finish": _d(cell("early_finish")),
            "early_start": _d(cell("early_start")), "early_finish": _d(cell("early_finish")),
            "late_start": None, "late_finish": None,
            "act_start": _d(cell("act_start")), "act_finish": _d(cell("act_finish")),
            "orig_dur_d": _f(cell("orig_dur_d")), "remain_dur_d": _f(cell("remain_dur_d")),
            "total_float_d": tf, "phys_pct": _f(cell("phys_pct")),
            "is_critical": 1 if tf is not None and tf <= 0 else 0, "is_milestone": 0,
        })
        for pred in str(cell("predecessors") or "").replace("，", ",").split(","):
            pred = pred.strip()
            if pred:
                rels.append({"pred_code": pred.split(":")[0].strip(), "succ_code": code, "kind": "FS", "lag_d": 0.0})
    if not acts:
        raise P6Error("Excel 中没有作业行")
    return {"proj_name": ws.title, "data_date": None}, acts, rels


# --------------------------------------------------------------------------- 入库

def import_schedule(conn, project_id: int, user_id: int, file_name: str, raw: bytes, data_date: str | None = None) -> dict:
    if len(raw) > 60 * 1024 * 1024:
        raise P6Error("文件超过 60MB")
    lower = file_name.lower()
    tables = None
    if lower.endswith(".xer"):
        # D2（10-07 用户确认）：完成%按 P6 作业自身的类型读取，工期%类型按工期换算；不带类型的 XER（含仿真）结果不变
        text = decode(raw)
        meta, acts, rels = activities_from_xer(text, pct_mode="typed")
        tables = parse_xer(text)
    elif lower.endswith((".xlsx", ".xlsm")):
        meta, acts, rels = activities_from_xlsx(raw)
    else:
        raise P6Error("仅支持 P6 导出的 .xer 或 .xlsx 文件")
    meta["data_date"] = data_date or meta.get("data_date")
    if not meta["data_date"]:
        raise P6Error("未读取到数据日期（Data Date），请在导入时填写")
    codes = {a["code"] for a in acts}
    if len(codes) != len(acts):
        raise P6Error("作业代码重复，请检查导出范围")
    rels = [r for r in rels if r["pred_code"] in codes and r["succ_code"] in codes]
    sha = hashlib.sha256(raw).hexdigest()
    with tx(conn):
        conn.execute("UPDATE p6_imports SET is_current=0 WHERE project_id=?", (project_id,))
        import_id = insert(conn, "p6_imports", {
            "project_id": project_id, "file_name": file_name[:200], "sha256": sha, "proj_name": meta.get("proj_name"),
            "data_date": meta["data_date"], "activity_count": len(acts), "relation_count": len(rels),
            "is_current": 1, "imported_by": user_id, "imported_at": now_iso(),
        })
        conn.executemany(
            "INSERT INTO activities (import_id, task_id, code, name, wbs_path, task_type, status, target_start, target_finish,"
            " early_start, early_finish, late_start, late_finish, act_start, act_finish, orig_dur_d, remain_dur_d,"
            " total_float_d, phys_pct, is_critical, is_milestone) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(import_id, a["task_id"], a["code"], a["name"], a["wbs_path"], a["task_type"], a["status"], a["target_start"],
              a["target_finish"], a["early_start"], a["early_finish"], a["late_start"], a["late_finish"], a["act_start"],
              a["act_finish"], a["orig_dur_d"], a["remain_dur_d"], a["total_float_d"], a["phys_pct"], a["is_critical"],
              a["is_milestone"]) for a in acts],
        )
        conn.executemany("INSERT INTO relations (import_id, pred_code, succ_code, kind, lag_d) VALUES (?,?,?,?,?)",
                         [(import_id, r["pred_code"], r["succ_code"], r["kind"], r["lag_d"]) for r in rels])
        ext = None
        if tables is not None:  # 日历、约束、UDF（数量/计划工效）另存扩展表，见 p6ext
            from . import p6ext
            ext = p6ext.save(conn, import_id, tables, codes)
        audit(conn, user_id, "p6.import", f"import:{import_id}", f"{file_name} sha256={sha[:12]} acts={len(acts)}")
    pct_types: dict[str, int] = {}
    for a in acts:
        if "pct_type" in a:
            key = a["pct_type"] or "未标注"
            pct_types[key] = pct_types.get(key, 0) + 1
    return {"import_id": import_id, "activities": len(acts), "relations": len(rels), **meta,
            "critical": sum(a["is_critical"] for a in acts), "has_float": any(a["total_float_d"] is not None for a in acts),
            "pct_types": pct_types, "ext": ext}


def current_import(conn, project_id: int) -> dict | None:
    return one(conn, "SELECT * FROM p6_imports WHERE project_id=? AND is_current=1 ORDER BY id DESC LIMIT 1", (project_id,))


def current_activities(conn, project_id: int) -> list[dict]:
    imp = current_import(conn, project_id)
    if not imp:
        return []
    return all_(conn, "SELECT * FROM activities WHERE import_id=? ORDER BY early_start, code", (imp["id"],))


def in_scope(act: dict, scope: str) -> bool:
    """技术员负责范围：WBS 路径或作业代码前缀，多个用逗号分隔；空表示全部。"""
    parts = [p.strip() for p in (scope or "").replace("，", ",").split(",") if p.strip()]
    if not parts:
        return True
    return any(p in (act.get("wbs_path") or "") or (act.get("code") or "").startswith(p) for p in parts)


def window_activities(acts: list[dict], day: str, scope: str = "", horizon_days: int = 7, limit: int = 40) -> list[dict]:
    """当天应关注的作业：进行中、计划窗口覆盖当天、或未来一周计划开工；优先关键作业。"""
    target = date.fromisoformat(day)
    picked = []
    for a in acts:
        if a.get("act_finish") or not in_scope(a, scope):
            continue
        es = date.fromisoformat(a["early_start"]) if a.get("early_start") else None
        ef = date.fromisoformat(a["early_finish"]) if a.get("early_finish") else None
        active = bool(a.get("act_start")) or (es and es <= target and (not ef or ef >= target))
        upcoming = es and 0 < (es - target).days <= horizon_days
        overdue = ef and ef < target
        if active or upcoming or overdue:
            picked.append(a)
    picked.sort(key=lambda a: (0 if a.get("is_critical") else 1, a.get("early_start") or "9999", a["code"]))
    return picked[:limit]


def range_activities(acts: list[dict], start: str, end: str, scope: str = "", horizon_days: int = 7, limit: int = 60) -> list[dict]:
    """某一时段内应关注的作业：时段内计划进行、已开工未完成、或时段结束后一周内计划开工；关键作业优先。"""
    s, e = date.fromisoformat(start), date.fromisoformat(end)
    picked = []
    for a in acts:
        if a.get("act_finish") or a.get("is_milestone") or not in_scope(a, scope):
            continue
        es = date.fromisoformat(a["early_start"]) if a.get("early_start") else None
        ef = date.fromisoformat(a["early_finish"]) if a.get("early_finish") else None
        overlap = es and es <= e and (not ef or ef >= s)
        upcoming = es and 0 < (es - e).days <= horizon_days
        if a.get("act_start") or overlap or upcoming or (ef and ef < s):
            picked.append(a)
    picked.sort(key=lambda a: (0 if a.get("is_critical") else 1, a.get("early_start") or "9999", a["code"]))
    return picked[:limit]

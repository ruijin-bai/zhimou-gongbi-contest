"""周报/月报：统计事实冻结入库 → 交融起草叙述 → 数字守卫 → Word 导出。"""
from __future__ import annotations

import re
import time
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from ..ai import report_agent
from ..ai.client import call_json
from ..config import settings
from ..db import all_, dumps, event, insert, loads, now_iso, one, tx
from . import p6
from .progress import planned_pct
from . import act_names as _an

LEVEL_NAMES = {"red": "红色", "orange": "橙色", "yellow": "黄色"}
STATUS_NAMES = {"open": "待处理", "confirmed": "已确认", "false_alarm": "误报", "closed": "已关闭"}


def period_for(kind: str, ref: str | None) -> tuple[str, str]:
    d = date.fromisoformat(ref) if ref else date.today()
    if kind == "week":
        start = d - timedelta(days=d.weekday())
        return start.isoformat(), (start + timedelta(days=6)).isoformat()
    start = d.replace(day=1)
    nxt = (start + timedelta(days=32)).replace(day=1)
    return start.isoformat(), (nxt - timedelta(days=1)).isoformat()


def build_facts(conn, project_id: int, start: str, end: str) -> dict:
    project = one(conn, "SELECT * FROM projects WHERE id=?", (project_id,))
    # 施工内容/人机料/问题取已确认的每日日志；工程量与累计%只取正式进度记录（周确认），避免重复计量
    logs = all_(conn, """SELECT l.*, u.name author FROM logs l JOIN users u ON u.id=l.author_id WHERE l.project_id=? AND l.mode='ai'
                         AND l.period='day' AND l.status IN ('submitted','approved') AND l.log_date BETWEEN ? AND ? ORDER BY l.log_date, l.id""",
                      (project_id, start, end))
    weeks = all_(conn, """SELECT * FROM logs WHERE project_id=? AND mode='ai' AND period='week' AND status IN ('submitted','approved')
                          AND log_date<=? AND period_end>=?""", (project_id, end, start))
    techs = one(conn, "SELECT COUNT(*) n FROM users WHERE project_id=? AND role='tech' AND active=1", (project_id,))["n"]
    days = (min(date.fromisoformat(end), date.today()) - date.fromisoformat(start)).days + 1
    n_weeks = max(1, round(((date.fromisoformat(end) - date.fromisoformat(start)).days + 1) / 7))
    on_time = sum(1 for w in weeks if w["submitted_at"] and w["submitted_at"][:10] <= (date.fromisoformat(w["period_end"]) + timedelta(days=1)).isoformat())
    imp = p6.current_import(conn, project_id)
    acts = {a["code"]: a for a in p6.current_activities(conn, project_id)}

    works, issues = [], []
    qty: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    last_pct: dict[str, float] = {}
    labor_sum, mach_sum, labor_days, mach_days = defaultdict(float), defaultdict(float), defaultdict(set), defaultdict(set)
    materials: dict[tuple, float] = defaultdict(float)
    for lg in logs:
        c = loads(lg["content_json"], {})
        for w in c.get("works", []):
            works.append({"log_id": lg["id"], "date": lg["log_date"], "code": w.get("activity_code"),
                          "text": f"{w.get('part', '')}{w.get('process', '')}：{w.get('desc', '')}".strip("：")})
        for x in c.get("labor", []):
            if x.get("count") is not None:
                labor_sum[x["trade"]] += x["count"]
                labor_days[x["trade"]].add(lg["log_date"])
        for x in c.get("machinery", []):
            if x.get("count") is not None:
                mach_sum[x["name"]] += x["count"]
                mach_days[x["name"]].add(lg["log_date"])
        for x in c.get("materials", []):
            if x.get("qty") is not None:
                materials[(x["name"], x.get("unit") or "", x.get("action") or "")] += x["qty"]
        for i in c.get("issues", []):
            issues.append({"log_id": lg["id"], "date": lg["log_date"], "author": lg["author"], **{k: i.get(k, "") for k in ("type", "severity", "desc", "measure", "activity_code")}})

    # 自报 = 日志待审核；核定 = 港湾已审核；失效（退回）行不计。合计 = 自报 + 核定，口径不变
    qty_by: dict[str, dict[str, dict[str, float]]] = {"self": defaultdict(lambda: defaultdict(float)), "approved": defaultdict(lambda: defaultdict(float))}
    prog_logs: dict[str, dict[str, list[int]]] = {"self": defaultdict(list), "approved": defaultdict(list)}
    for r in all_(conn, """SELECT lp.*, l.status lstatus FROM log_progress lp JOIN logs l ON l.id=lp.log_id WHERE lp.project_id=? AND l.status IN ('submitted','approved')
                          AND lp.voided_at IS NULL AND lp.log_date BETWEEN ? AND ? ORDER BY lp.log_date, lp.id""", (project_id, start, end)):
        tk = "approved" if r["lstatus"] == "approved" else "self"
        if r["qty"] is not None:
            qty[r["activity_code"]][r["unit"] or ""] += r["qty"]
            qty_by[tk][r["activity_code"]][r["unit"] or ""] += r["qty"]
        if r["log_id"] not in prog_logs[tk][r["activity_code"]]:
            prog_logs[tk][r["activity_code"]].append(r["log_id"])
        if r["pct"] is not None:
            last_pct[r["activity_code"]] = r["pct"]
    end_day = date.fromisoformat(end)
    dd = date.fromisoformat(imp["data_date"]) if imp else None
    progress = []
    for code in sorted(set(qty) | set(last_pct)):
        a = acts.get(code, {})
        plan = round(planned_pct(a, min(end_day, date.today()), dd), 1) if a and dd else None
        endp = last_pct.get(code)
        progress.append({"code": code, "name": a.get("name", "（不在当前计划）"),
                         "qty": [{"qty": round(v, 2), "unit": u, "self": round(qty_by["self"][code].get(u, 0.0), 2),
                                  "approved": round(qty_by["approved"][code].get(u, 0.0), 2)} for u, v in qty[code].items()],
                         "log_ids_self": prog_logs["self"].get(code, []), "log_ids_approved": prog_logs["approved"].get(code, []),
                         "end_pct": endp, "plan_pct": plan, "gap": round(endp - plan, 1) if endp is not None and plan is not None else None,
                         "critical": bool(a.get("is_critical"))})
    alerts = all_(conn, """SELECT id, level, title, status FROM alerts WHERE project_id=? AND rule<>'stale' AND substr(last_detected,1,10)>=? AND substr(first_detected,1,10)<=?
                           ORDER BY CASE level WHEN 'red' THEN 0 WHEN 'orange' THEN 1 ELSE 2 END, id""", (project_id, start, end))
    nxt_start, nxt_end = end_day + timedelta(days=1), end_day + timedelta(days=7 if (end_day - date.fromisoformat(start)).days < 8 else 31)
    nxt = [{"code": a["code"], "name": a["name"], "start": a["early_start"], "finish": a["early_finish"], "critical": bool(a["is_critical"])}
           for a in acts.values() if not a.get("act_finish") and a.get("early_start") and a.get("early_finish")
           and date.fromisoformat(a["early_start"]) <= nxt_end and date.fromisoformat(a["early_finish"]) >= nxt_start]
    nxt.sort(key=lambda n: (not n["critical"], n["start"]))
    photos = one(conn, """SELECT COUNT(*) n FROM media m JOIN logs l ON l.id=m.log_id WHERE l.project_id=? AND l.mode='ai'
                          AND l.status IN ('submitted','approved') AND l.log_date BETWEEN ? AND ?""", (project_id, start, end))["n"]
    from . import output as output_svc, schedule
    fc = schedule.latest_forecast(conn, project_id)
    milestones = [{**m, "prob_pct": round(m["on_time_prob"] * 100) if m["on_time_prob"] is not None else None} for m in schedule.with_prob_text(dict(fc or {})).get("milestones", [])][:12]
    out = output_svc.compute(conn, project_id, start, end)
    rec = one(conn, "SELECT SUM(actual) a, SUM(planned) p, SUM(baseline) b FROM output_records WHERE project_id=? AND period_start BETWEEN ? AND ?", (project_id, start, end))
    out_facts = {**out, "record_actual": rec["a"], "planned": rec["p"], "baseline": rec["b"]} if (out["by_activity"] or rec["a"] is not None) else {"total": None}
    fs = all_(conn, """SELECT * FROM findings WHERE project_id=? AND agent='safety' AND confirmed_at IS NOT NULL AND substr(created_at,1,10) BETWEEN ? AND ?""", (project_id, start, end))
    cats: dict[str, int] = defaultdict(int)
    for f in fs:
        cats[(loads(f["body_json"], {}) or {}).get("category") or "日志上报"] += 1
    open_high = all_(conn, """SELECT id, title, due_date FROM findings WHERE project_id=? AND agent='safety' AND level='high' AND confirmed_at IS NOT NULL
                              AND status IN ('open','rectifying','rejected') ORDER BY id""", (project_id,))
    safety = {"found": len(fs), "high": sum(1 for f in fs if f["level"] == "high"), "resolved": sum(1 for f in fs if f["status"] == "resolved"),
              "open": sum(1 for f in fs if f["status"] in ("open", "rectifying", "rejected")), "by_category": dict(cats), "open_high": open_high}
    q = one(conn, """SELECT SUM(kind<>'qc_reminder') i, SUM(kind='qc_reminder') r FROM findings WHERE project_id=? AND agent='quality'
                     AND substr(created_at,1,10) BETWEEN ? AND ?""", (project_id, start, end))
    return {
        "milestones": milestones, "output": out_facts, "safety": safety, "quality": {"issues": q["i"] or 0, "reminders": q["r"] or 0},
        "project": project["name"], "period_start": start, "period_end": end,
        "p6": {"file": imp["file_name"], "data_date": imp["data_date"]} if imp else None,
        "logs": {"count": len(logs), "expected": techs * max(days, 0), "on_time": on_time, "authors": len({l["author_id"] for l in logs}),
                 "ids": [l["id"] for l in logs], "weeks": len(weeks), "weeks_expected": techs * n_weeks},
        "photos": photos, "works": works, "issues": issues, "progress": progress,
        "labor": [{"name": k, "avg": round(v / len(labor_days[k]), 1)} for k, v in sorted(labor_sum.items(), key=lambda x: -x[1])],
        "machinery": [{"name": k, "avg": round(v / len(mach_days[k]), 1)} for k, v in sorted(mach_sum.items(), key=lambda x: -x[1])],
        "materials": [{"name": k[0], "unit": k[1], "action": k[2], "qty": round(v, 2)} for k, v in materials.items()],
        "alerts": [{**a, "level_name": LEVEL_NAMES[a["level"]], "status_name": STATUS_NAMES[a["status"]]} for a in alerts],
        "next": nxt,
    }


def get_template(conn, project_id: int, kind: str) -> dict:
    row = one(conn, "SELECT value_json FROM knowledge WHERE project_id=? AND kind='report_template' AND key=?", (project_id, kind))
    return loads(row["value_json"]) if row else report_agent.default_template(kind)


def save_template(conn, project_id: int, kind: str, tpl: dict, user_id: int) -> dict:
    if kind not in ("week", "month"):
        raise ValueError("kind 只能是 week 或 month")
    clean = report_agent.validate_template(tpl, kind)
    conn.execute("INSERT OR REPLACE INTO knowledge (project_id, kind, key, value_json, ai_run_id, updated_at) VALUES (?,?,?,?,?,?)",
                 (project_id, "report_template", kind, dumps(clean), None, now_iso()))
    return clean


def generate(conn, project_id: int, user: dict, kind: str, ref: str | None) -> dict:
    start, end = period_for(kind, ref)
    t0 = time.monotonic()
    facts = build_facts(conn, project_id, start, end)
    if not facts["logs"]["count"]:
        raise ValueError("该周期没有已提交的日志，无法生成报表")
    tpl = get_template(conn, project_id, kind)
    facts["template"] = tpl
    prompt = report_agent.build_prompt(facts)
    result = call_json(
        f"report.{kind}", system=report_agent.system_prompt(tpl), user=prompt,
        prompt_version=report_agent.PROMPT_VERSION, mock=lambda: report_agent.mock(facts, tpl), validate=lambda d: report_agent.normalize(d, tpl),
        project_id=project_id, user_id=user["id"], max_tokens=3500, input_meta={"logs": facts["logs"]["count"], "period": [start, end]},
    )
    stance = report_agent.stance_fix(result.data, facts)   # 定性与表格/预警矛盾的句子先改掉，再做数字守卫（验收 I11）
    guard = stance + report_agent.number_guard(result.data, prompt)
    tcheck = report_agent.template_check(result.data, tpl)
    with tx(conn):
        rid = insert(conn, "reports", {
            "project_id": project_id, "kind": kind, "period_start": start, "period_end": end, "facts_json": dumps(facts),
            "narrative_json": dumps(result.data), "number_check_json": dumps(guard + [{"field": x["field"], "number": "", "issue": x["issue"]} for x in tcheck]),
            "ai_run_id": result.run_id, "status": "draft", "created_by": user["id"], "created_at": now_iso(), "gen_seconds": round(time.monotonic() - t0, 1),
        })
        event(conn, "report.generate", project_id=project_id, user_id=user["id"], ref_id=rid, value=round(time.monotonic() - t0, 1),
              data={"kind": kind, "logs": facts["logs"]["count"]})
    return one(conn, "SELECT * FROM reports WHERE id=?", (rid,))


def _split_text(qty: list[dict], key: str) -> str:
    """某一层（self/approved）的完成量文字；没有这一层的数量写“0”，没有任何数量写“—”。旧报表事实里没有拆分字段，写“—”不补零。"""
    if not qty:
        return "—"
    if any(key not in q for q in qty):
        return "—"
    return "、".join(f"{q[key]:g}{q['unit']}" for q in qty)


# ----------------------------------------------------------------------------- 项目经理口径：报表不含产值节与自报/核定拆分

def output_keys(tpl: dict) -> set[str]:
    """模板里属于产值的栏目（键为 output 或挂产值表）。"""
    return {s["key"] for s in tpl.get("sections", []) if s.get("key") == "output" or s.get("table") == "output"}


_SENT_RE = re.compile(r"[^。；;！？!?\n]*(?:[。；;！？!?\n]|$)")


def strip_output(value):
    """其他栏目里顺带写到产值的句子（按中文句号/分号切句）去掉；字符串数组逐条处理，整条去空即删。"""
    if isinstance(value, list):
        return [x for x in (strip_output(v) for v in value) if x]
    if not isinstance(value, str) or "产值" not in value:
        return value
    return "".join(s for s in _SENT_RE.findall(value) if s and "产值" not in s).strip()


_SEC_NO = re.compile(r"^[一二三四五六七八九十]+、")


def _cn_no(n: int) -> str:
    d = "一二三四五六七八九"
    return d[n - 1] if n < 10 else "十" + (d[n - 11] if n > 10 else "") if n < 20 else str(n)


def renumber(sections: list[dict]) -> list[dict]:
    """剔除栏目后按显示顺序重新编号（“三、关键节点预测”后面接“四、安全管理”，不跳号）；标题不带编号的不动。网页与 Word 导出同一口径。"""
    out, n = [], 0
    for s in sections:
        if _SEC_NO.match(s.get("title") or ""):
            n += 1
            s = {**s, "title": _SEC_NO.sub(_cn_no(n) + "、", s["title"], count=1)}
        out.append(s)
    return out


def pm_scope(facts: dict, tpl: dict, narrative: dict, checks: list[dict], drop_output: bool = False) -> tuple[dict, dict, dict, list[dict]]:
    """给项目经理看的报表：去掉工程量的自报/核定拆分（日志审核数据，进度表与产值表同口径）；产值节只读照常显示（RJ 10-10）。
    drop_output=True（问答引用报表时）：另剔除产值栏目（叙述、表格、数字核验）、产值事实，其他栏目里含“产值”的句子同样裁掉——
    问答不向项目经理推产值。只裁剪返回给项目经理的视图，冻结入库的事实与叙述不改。"""
    hidden = output_keys(tpl) if drop_output else set()
    tpl = {**tpl, "sections": renumber([s for s in tpl["sections"] if s["key"] not in hidden])}
    facts = {k: v for k, v in facts.items() if not (drop_output and k == "output")}
    facts["progress"] = [{**{k: v for k, v in p.items() if k not in ("log_ids_self", "log_ids_approved")},
                          "qty": [{k: v for k, v in q.items() if k not in ("self", "approved")} for q in p.get("qty", [])]} for p in facts.get("progress", [])]
    if isinstance(facts.get("output"), dict):
        split_keys = {"self_total", "approved_total", "log_ids_self", "log_ids_approved"}
        o = {k: v for k, v in facts["output"].items() if k not in split_keys}
        if o.get("by_activity"):
            o["by_activity"] = [{k: v for k, v in a.items() if not k.endswith(("_self", "_approved"))} for a in o["by_activity"]]
        facts["output"] = o
    if "template" in facts:
        facts["template"] = tpl
    narrative = {k: (strip_output(v) if drop_output else v) for k, v in (narrative or {}).items() if k not in hidden}
    checks = [x for x in checks or [] if x.get("field") not in hidden]
    return facts, tpl, narrative, checks


def table_rows(facts: dict, name: str, split: bool = True) -> tuple[list[str], list[list]]:
    """split=False：进度表、产值表不列“分包自报/港湾核定”拆分（项目经理口径）。"""
    g = lambda v: "未知" if v is None else f"{v:g}"  # noqa: E731
    if name == "progress":
        if not split:
            return (["作业代码", "作业名称", "本期完成量", "期末累计%", "期末计划%", "偏差"],
                    [[x["code"], x["name"], "、".join(f"{q['qty']:g}{q['unit']}" for q in x["qty"]) or "—", g(x["end_pct"]), g(x["plan_pct"]), g(x["gap"])]
                     for x in facts["progress"]])
        return (["作业代码", "作业名称", "本期完成量（合计）", "其中分包自报", "其中港湾核定", "期末累计%", "期末计划%", "偏差"],
                [[x["code"], x["name"], "、".join(f"{q['qty']:g}{q['unit']}" for q in x["qty"]) or "—", _split_text(x["qty"], "self"), _split_text(x["qty"], "approved"),
                  g(x["end_pct"]), g(x["plan_pct"]), g(x["gap"])] for x in facts["progress"]])
    if name == "milestones":
        return (["节点", "P6日期", "预测", "P80", "按期概率"],
                [[f"{m['code']} {m['name']}", m["p6_date"], m["forecast"], m["p80"], m.get("prob_label") or (f"{m['prob_pct']}%" if m["prob_pct"] is not None else "—")] for m in facts.get("milestones", [])])
    if name == "output":
        o = facts.get("output") or {}
        if not split:
            return (["作业", "清单", "工程量", "内部产值统计（非计量认证）"],
                    [[f"{a['activity_code']} {a['name']}", a["boq"], f"{a['qty']:g}{a['unit']}", f"{a['value']:,.2f}"] for a in o.get("by_activity", [])])
        return (["作业", "清单", "工程量（合计）", "分包自报 / 港湾核定", "内部产值统计（非计量认证）", "其中自报 / 核定"],
                [[f"{a['activity_code']} {a['name']}", a["boq"], f"{a['qty']:g}{a['unit']}",
                  f"{a['qty_self']:g} / {a['qty_approved']:g}{a['unit']}" if "qty_self" in a else "—", f"{a['value']:,.2f}",
                  f"{a['value_self']:,.2f} / {a['value_approved']:,.2f}" if "value_self" in a else "—"] for a in o.get("by_activity", [])])
    if name == "safety":
        s = facts.get("safety") or {}
        return (["类别", "数量"], [[k, v] for k, v in (s.get("by_category") or {}).items()])
    if name == "issues":
        return (["日期", "类型", "等级", "问题", "措施", "来源"], [[i["date"], i["type"], i["severity"], i["desc"], i["measure"], f"L{i['log_id']}"] for i in facts["issues"]])
    if name == "next":
        return (["作业", "计划开始", "计划完成", "关键"], [[f"{n['code']} {n['name']}", n["start"], n["finish"], "是" if n["critical"] else ""] for n in facts["next"]])
    return [], []


def to_docx(report: dict, hide_split: bool = False) -> Path:
    """hide_split=True：项目经理导出的版本，不含自报/核定拆分与相关说明（产值节照常，只读）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    facts, nar = loads(report["facts_json"]), loads(report["narrative_json"], {})
    guard = loads(report["number_check_json"], [])
    tpl = facts.get("template") or report_agent.default_template(report["kind"])
    if hide_split:
        facts, tpl, nar, guard = pm_scope(facts, tpl, nar, guard)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)
    tpl0 = facts.get("template") or {}
    title = doc.add_heading((tpl0.get("title") or "{project} 施工报告").replace("{project}", facts["project"]), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"报告期：{facts['period_start']} 至 {facts['period_end']}    "
                          f"依据日报 {facts['logs']['count']} 份 / 照片 {facts['photos']} 张"
                          + (f"    P6数据日期：{facts['p6']['data_date']}" if facts.get("p6") else ""))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = "" if v is None else str(v)
        return t

    for sec in tpl["sections"]:
        doc.add_heading(sec["title"], level=1)
        val = nar.get(sec["key"])
        if sec["type"] == "list":
            for c in val or []:
                doc.add_paragraph(c, style="List Bullet")
            if not val:
                doc.add_paragraph("无")
        elif val:
            doc.add_paragraph(val)
        if sec.get("table"):
            headers, rows = table_rows(facts, sec["table"], split=not hide_split)
            if rows:
                table(headers, rows)
    doc.add_heading("附：数据来源说明", level=1)
    doc.add_paragraph("表格数字由系统从技术员确认提交的施工日志及 P6 计划自动统计；文字部分由中国交建交融大模型起草，经项目部审核。"
                      f"引用编号 L 为施工日志编号（共 {len(facts['logs']['ids'])} 份），W 为进度预警编号。")
    if not hide_split:
        doc.add_paragraph("工程量分两层：分包自报＝技术员已提交、待港湾审核的数量；港湾核定＝港湾已审核通过的数量；二者相加为合计。"
                          "被退回的数量不计入。产值为内部产值统计（非计量认证），不是业主计量或结算金额。")
    else:
        doc.add_paragraph("被退回的数量不计入。产值为内部产值统计（非计量认证），不是业主计量或结算金额。")
    if guard:
        nums = [x["number"] for x in guard if x.get("number")]
        if nums:
            doc.add_paragraph("数字核验提示：叙述中以下数字未在统计事实中找到，请审核人核实：" + "、".join(nums))
    out = settings.export_dir / f"report_{report['id']}{'_pm' if hide_split else ''}.docx"
    _an.humanize_doc(doc)   # 作业代码 → 中文简称（services/act_names.py）
    doc.save(out)
    return out

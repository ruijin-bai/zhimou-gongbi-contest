"""产值Agent：清单只在后台使用。AI 建立 P6 作业→清单单价映射，商务确认一次；此后按日志工程量自动算产值。

产值 = Σ 日志确认工程量 × 映射系数 × 清单单价（单位一致才计入；不一致或未映射的工程量单列，不补零）。
这是内部产值统计（非计量认证），不是业主计量或结算金额。工程量分两层：分包自报（日志待审核）与港湾核定（日志已审核），合计=两者之和。
"""
from __future__ import annotations

import hashlib
import io
import re
from collections import defaultdict
from datetime import date, timedelta

from ..ai.client import AIError, call_json
from ..db import all_, audit, dumps, loads, now_iso, one, tx
from . import p6
from . import act_names as _an

_BOQ_ALIASES = {
    "code": ("编码", "清单编号", "清单编码", "项目编码", "item", "item no", "item no.", "code", "编号"),
    "name": ("名称", "项目名称", "清单名称", "description", "name", "工作内容"),
    "unit": ("单位", "unit"),
    "unit_price": ("单价", "综合单价", "rate", "unit price", "unit rate"),
    "contract_qty": ("数量", "合同数量", "工程量", "qty", "quantity"),
    "currency": ("币种", "currency"),
}
_REC_ALIASES = {
    "start": ("开始", "周起", "起始日期", "开始日期", "start", "week start", "from"),
    "end": ("结束", "周止", "截止日期", "结束日期", "end", "week end", "to"),
    "planned": ("计划", "计划产值", "更新计划", "更新产值计划", "planned", "plan"),
    "baseline": ("基准计划", "基准产值计划", "baseline"),
    "actual": ("实际", "实际产值", "完成产值", "产值", "actual", "value"),
    "currency": ("币种", "currency"),
}


def _sheet_rows(raw: bytes, aliases: dict) -> tuple[list[tuple], dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[:8]):
            labels = [str(c).strip().lower() if c is not None else "" for c in row]
            found = {}
            for key, names in aliases.items():
                for j, lab in enumerate(labels):
                    if lab in names or any(lab.startswith(n) for n in names if len(n) >= 2):
                        found.setdefault(key, j)
            if len(found) >= 2:
                return rows[i + 1:], found
    raise ValueError("未找到可识别的表头")


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def import_boq(conn, project_id: int, user_id: int, raw: bytes) -> dict:
    rows, m = _sheet_rows(raw, _BOQ_ALIASES)
    if "code" not in m or "name" not in m:
        raise ValueError("清单需包含“编码”和“名称”列")
    n = 0
    with tx(conn):
        for r in rows:
            code = str(r[m["code"]] or "").strip() if m["code"] < len(r) else ""
            name = str(r[m["name"]] or "").strip() if m["name"] < len(r) else ""
            if not code or not name:
                continue
            g = lambda k: r[m[k]] if k in m and m[k] < len(r) else None  # noqa: E731
            conn.execute("""INSERT INTO boq_items (project_id, code, name, unit, unit_price, currency, contract_qty) VALUES (?,?,?,?,?,?,?)
                            ON CONFLICT(project_id, code) DO UPDATE SET name=excluded.name, unit=excluded.unit, unit_price=excluded.unit_price,
                            currency=excluded.currency, contract_qty=excluded.contract_qty""",
                         (project_id, code[:60], name[:200], str(g("unit") or "").strip(), _num(g("unit_price")), str(g("currency") or "").strip(), _num(g("contract_qty"))))
            n += 1
        audit(conn, user_id, "boq.import", f"project:{project_id}", f"{n} items")
    return {"items": n}


def import_records(conn, project_id: int, user_id: int, raw: bytes, source: str = "record") -> dict:
    """导入周/月产值记录（计划与实际）。真实记录 source=record；仿真生成 source=sim。
    计划分两列（与项目周报一致）：“基准产值计划”→ baseline，“更新产值计划/计划产值”→ planned；缺列即为空，不补零。"""
    from .p6 import _d
    rows, m = _sheet_rows(raw, _REC_ALIASES)
    if "start" not in m or not {"actual", "planned", "baseline"} & set(m):
        raise ValueError("产值记录需包含开始日期及实际/计划产值列")
    n = 0
    with tx(conn):
        for r in rows:
            g = lambda k: r[m[k]] if k in m and m[k] < len(r) else None  # noqa: E731
            s = _d(g("start"))
            if not s:
                continue
            e = _d(g("end")) or (date.fromisoformat(s) + timedelta(days=6)).isoformat()
            # 只更新本次文件里有的列：单独导入一份计划表不会把已有实际值清空
            vals = {k: _num(g(k)) for k in ("planned", "actual", "baseline") if k in m}
            sets = ", ".join(f"{k}=excluded.{k}" for k in vals)
            conn.execute(f"""INSERT INTO output_records (project_id, period_start, period_end, {", ".join(vals)}, currency, source)
                             VALUES (?,?,?,{",".join("?" for _ in vals)},?,?)
                             ON CONFLICT(project_id, period_start, source) DO UPDATE SET period_end=excluded.period_end, {sets},
                             currency=CASE WHEN excluded.currency<>'' THEN excluded.currency ELSE output_records.currency END""",
                         (project_id, s, e, *vals.values(), str(g("currency") or ""), source))
            n += 1
        audit(conn, user_id, "output.records", f"project:{project_id}", f"{n} rows {source}")
    return {"rows": n}


# ----------------------------------------------------------------------------- AI 映射

MAP_SYSTEM = """你是工程商务人员，负责把P6进度作业对应到工程量清单子目，以便按完成工程量计算产值。
对每个作业，从给定候选清单中选出最匹配的一项（工作内容与计量单位都要匹配），没有合适的填空字符串；factor为作业工程量换算到清单单位的系数，通常为1。
只输出JSON：{"maps":[{"activity_code":"作业代码","boq_code":"清单编码或空","factor":1,"reason":"简要理由"}]}"""


def _grams(s: str) -> set[str]:
    s = re.sub(r"[\s\d#（）()\-_/、，,.]+", "", s or "")
    return {s[i:i + 2] for i in range(len(s) - 1)} | set(s)


def suggest_map(conn, project_id: int, user_id: int | None, limit: int = 120) -> dict:
    boq = all_(conn, "SELECT * FROM boq_items WHERE project_id=?", (project_id,))
    if not boq:
        raise ValueError("请先导入工程量清单")
    done = {r["activity_code"] for r in all_(conn, "SELECT activity_code FROM boq_map WHERE project_id=? AND status IN ('confirmed','rejected')", (project_id,))}
    acts = [a for a in p6.current_activities(conn, project_id) if not a["is_milestone"] and a["code"] not in done][:limit]
    bg = [(b, _grams(b["name"])) for b in boq]
    by_code = {b["code"]: b for b in boq}
    made = 0
    for i in range(0, len(acts), 25):
        batch = acts[i:i + 25]
        lines, cands = [], {}
        for a in batch:
            ag = _grams(a["name"] + (a["wbs_path"] or "").split(" / ")[-1])
            ranked = sorted(bg, key=lambda x: -len(ag & x[1]))[:6]
            cands[a["code"]] = {b["code"] for b, _ in ranked}
            lines.append(f"作业 {a['code']} {a['name']}（{a['wbs_path']}）候选：" + "；".join(f"{b['code']} {b['name']} [{b['unit']}]" for b, _ in ranked))

        def mock(batch=batch, cands=cands):
            out = []
            for a in batch:
                ag = _grams(a["name"])
                best = max((by_code[c] for c in cands[a["code"]]), key=lambda b: len(ag & _grams(b["name"])), default=None)
                ok = best is not None and len(ag & _grams(best["name"])) >= 3
                out.append({"activity_code": a["code"], "boq_code": best["code"] if ok else "", "factor": 1, "reason": "（离线模拟）名称相似"})
            return {"maps": out}

        def norm(d, cands=cands):
            if not isinstance(d, dict) or not isinstance(d.get("maps"), list):
                raise ValueError("缺少maps")
            return [m for m in d["maps"] if isinstance(m, dict) and m.get("activity_code") in cands]
        try:
            res = call_json("output.map", system=MAP_SYSTEM, user="\n".join(lines), prompt_version="map-v1-20261006", mock=mock,
                            validate=norm, project_id=project_id, user_id=user_id, max_tokens=3000, input_meta={"acts": len(batch)})
        except AIError:
            continue
        with tx(conn):
            for mp in res.data:
                code = str(mp.get("boq_code") or "")
                if code not in cands[mp["activity_code"]]:
                    continue
                f = _num(mp.get("factor")) or 1.0
                conn.execute("""INSERT INTO boq_map (project_id, activity_code, boq_id, factor, status, reason, ai_run_id, updated_at)
                                VALUES (?,?,?,?,?,?,?,?) ON CONFLICT(project_id, activity_code) DO UPDATE SET boq_id=excluded.boq_id,
                                factor=excluded.factor, reason=excluded.reason, ai_run_id=excluded.ai_run_id, updated_at=excluded.updated_at
                                WHERE boq_map.status='suggested'""",
                             (project_id, mp["activity_code"], by_code[code]["id"], f, "suggested", str(mp.get("reason") or "")[:200], res.run_id, now_iso()))
                made += 1
    return {"suggested": made, "activities": len(acts)}


def mappings(conn, project_id: int) -> list[dict]:
    acts = {a["code"]: a for a in p6.current_activities(conn, project_id)}
    rows = all_(conn, """SELECT m.*, b.code boq_code, b.name boq_name, b.unit, b.unit_price, b.currency FROM boq_map m
                         JOIN boq_items b ON b.id=m.boq_id WHERE m.project_id=? ORDER BY m.status, m.activity_code""", (project_id,))
    return [{**r, "activity_name": acts.get(r["activity_code"], {}).get("name", "（不在当前计划）")} for r in rows]


def set_map(conn, project_id: int, user_id: int, ids: list[int], status: str) -> int:
    if status not in ("confirmed", "rejected", "suggested"):
        raise ValueError("状态无效")
    n = 0
    with tx(conn):
        for i in ids:
            n += conn.execute("UPDATE boq_map SET status=?, confirmed_by=?, updated_at=? WHERE id=? AND project_id=?",
                              (status, user_id, now_iso(), i, project_id)).rowcount
        audit(conn, user_id, f"boq_map.{status}", f"project:{project_id}", f"{n}")
    return n


# ----------------------------------------------------------------------------- 产值计算

def compute(conn, project_id: int, start: str, end: str) -> dict:
    """产值统计。total = 自报 + 核定（原口径不变）；另给出拆分：
    自报 = 来自“待审核”日志的工程量；核定 = 来自“港湾已审核”日志的工程量。失效（退回）行不计。"""
    maps = {m["activity_code"]: m for m in mappings(conn, project_id) if m["status"] == "confirmed"}
    rows = all_(conn, """SELECT lp.activity_code, lp.log_date, lp.qty, lp.unit, lp.log_id, l.status FROM log_progress lp JOIN logs l ON l.id=lp.log_id
                         WHERE lp.project_id=? AND l.status IN ('submitted','approved') AND lp.voided_at IS NULL
                         AND lp.log_date BETWEEN ? AND ? AND lp.qty IS NOT NULL ORDER BY lp.log_date, lp.id""",
                      (project_id, start, end))
    by_act: dict[str, dict] = {}
    by_day: dict[str, float] = defaultdict(float)
    unmapped, unit_mismatch, currency = defaultdict(lambda: {"self": 0.0, "approved": 0.0}), [], ""
    src: dict[str, list[int]] = {"self": [], "approved": []}
    for r in rows:
        k = "approved" if r["status"] == "approved" else "self"
        if r["log_id"] not in src[k]:
            src[k].append(r["log_id"])
        m = maps.get(r["activity_code"])
        if not m or m["unit_price"] is None:
            unmapped[f"{r['activity_code']}|{r['unit']}"][k] += r["qty"]
            continue
        if (r["unit"] or "").lower() != (m["unit"] or "").lower() and m["factor"] == 1:
            unit_mismatch.append({"activity_code": r["activity_code"], "log_unit": r["unit"], "boq_unit": m["unit"]})
            continue
        v = r["qty"] * m["factor"] * m["unit_price"]
        currency = currency or m["currency"]
        a = by_act.setdefault(r["activity_code"], {"activity_code": r["activity_code"], "name": m["activity_name"], "boq": m["boq_code"],
                                                    "qty": 0.0, "unit": m["unit"], "value": 0.0, "qty_self": 0.0, "qty_approved": 0.0,
                                                    "value_self": 0.0, "value_approved": 0.0, "log_ids_self": [], "log_ids_approved": []})
        a["qty"] += r["qty"] * m["factor"]
        a["value"] += v
        a["qty_" + k] += r["qty"] * m["factor"]
        a["value_" + k] += v
        if r["log_id"] not in a["log_ids_" + k]:
            a["log_ids_" + k].append(r["log_id"])
        by_day[r["log_date"]] += v
    total = round(sum(a["value"] for a in by_act.values()), 2)
    r2 = lambda x: round(x, 2)  # noqa: E731
    return {"start": start, "end": end, "total": total, "currency": currency,
            "self_total": r2(sum(a["value_self"] for a in by_act.values())), "approved_total": r2(sum(a["value_approved"] for a in by_act.values())),
            "log_ids_self": src["self"], "log_ids_approved": src["approved"],
            "by_activity": sorted(({**a, "qty": r2(a["qty"]), "value": r2(a["value"]), "qty_self": r2(a["qty_self"]), "qty_approved": r2(a["qty_approved"]),
                                    "value_self": r2(a["value_self"]), "value_approved": r2(a["value_approved"])} for a in by_act.values()), key=lambda a: -a["value"]),
            "by_day": [{"day": d, "value": round(v, 2)} for d, v in sorted(by_day.items())],
            "unmapped": [{"activity_code": k.split("|")[0], "unit": k.split("|")[1], "qty": r2(v["self"] + v["approved"]),
                          "qty_self": r2(v["self"]), "qty_approved": r2(v["approved"])} for k, v in unmapped.items()],
            "unit_mismatch": unit_mismatch[:20],
            "note": "内部产值统计（非计量认证）：日志工程量×清单单价，不是业主计量或结算金额；自报=分包提交待审核，核定=港湾已审核"}


def weekly_series(conn, project_id: int, weeks: int = 16, end: date | None = None) -> list[dict]:
    end = end or date.today()
    monday = end - timedelta(days=end.weekday())
    recs = {r["period_start"]: r for r in all_(conn, "SELECT * FROM output_records WHERE project_id=? ORDER BY period_start", (project_id,))}
    out = []
    for i in range(weeks - 1, -1, -1):
        s = monday - timedelta(days=7 * i)
        e = s + timedelta(days=6)
        c = compute(conn, project_id, s.isoformat(), e.isoformat())
        rec = recs.get(s.isoformat())
        out.append({"week": s.isoformat(), "system": c["total"] or None, "record_actual": rec["actual"] if rec else None,
                    "planned": rec["planned"] if rec else None, "baseline": rec["baseline"] if rec else None, "source": rec["source"] if rec else None,
                    # 系统统计的自报/核定拆分（自报+核定=system）与来源日志；无统计时为 None（不补零）
                    "system_self": c["self_total"] if c["total"] else None, "system_approved": c["approved_total"] if c["total"] else None,
                    "log_ids_self": c["log_ids_self"], "log_ids_approved": c["log_ids_approved"]})
    return out


def month_forecast(series: list[dict], today: date | None = None) -> dict:
    """月末预测：近4周平均周产值 × 本月剩余周数 + 本月已完成（只用实际值，不足4周不预测）。"""
    today = today or date.today()
    vals = [w["record_actual"] if w["record_actual"] is not None else w["system"] for w in series]
    recent = [v for v in vals[-5:-1] if v is not None]
    if len(recent) < 4:
        return {"forecast": None, "note": "近4周产值数据不足，不做预测"}
    avg = sum(recent) / 4
    month_start = today.replace(day=1)
    mtd = sum((v or 0) for w, v in zip(series, vals) if date.fromisoformat(w["week"]) >= month_start)
    nxt = (month_start + timedelta(days=32)).replace(day=1)
    remain_weeks = (nxt - today).days / 7
    return {"forecast": round(mtd + avg * remain_weeks, 2), "avg_week": round(avg, 2), "mtd": round(mtd, 2),
            "note": "近4周平均周产值×本月剩余周数+本月已完成"}


# ----------------------------------------------------------------------------- 产值分析：比较由系统算，模型只解释原因

def fmt_money(v, cur: str = "") -> str:
    """与前端 money() 同一格式：≥1亿 两位小数“亿”，≥1万 一位小数“万”，否则取整。"""
    if v is None:
        return "—"
    a = abs(v)
    s = f"{v / 1e8:.2f}亿" if a >= 1e8 else f"{v / 1e4:.1f}万" if a >= 1e4 else f"{v:.0f}"
    return s + (f" {cur}" if cur else "")


def compare(a, b) -> dict | None:
    """a 相对 b：高于/低于/持平、差额（a-b）、百分比（相对 b）。任一缺值或 b 为 0 不比较（不补零）。"""
    if a is None or b is None or not b:
        return None
    diff = round(a - b, 2)
    return {"word": "高于" if diff > 0 else "低于" if diff < 0 else "持平", "diff": diff, "pct": round(diff / b * 100, 1)}


_CMP_KW = {"计划": "计划", "前一周": "前一周|上上周|前周"}
_NUM_RE = re.compile(r"\d+(?:\.\d+)?")
_SIGN_RE = re.compile(r"([+\-−])\s*(\d+(?:\.\d+)?)\s*(万|亿)")


def _cmp_sentence(label: str, c: dict | None, ref, cur: str) -> str:
    if not c:
        return f"无{label}值，不作比较"
    if c["word"] == "持平":
        return f"与{label} {fmt_money(ref, cur)} 持平"
    sign = "+" if c["diff"] > 0 else "-"
    return f"{c['word']}{label} {fmt_money(ref, cur)}，差额 {sign}{fmt_money(abs(c['diff']), cur)}（{sign}{abs(c['pct']):g}%）"


def check_summary(text: str, comps: dict[str, dict | None], allowed: str, words: bool = True) -> list[str]:
    """产值分析出口校验：与计划/前一周的比较词必须与系统计算一致；差额正负号一致；数字只能取自系统给出的（按万显示的）数字。
    比较词只认明确的比较结构（“高于/低于/超过/未达 计划”“比计划多/少”），“未完成计划的作业”“钢筋不足”不算；
    words=False（措施）不查比较词。"""
    problems = []
    ok_nums = set(_NUM_RE.findall(allowed))
    ok_nums |= {str(int(float(x))) for x in ok_nums}
    for n in _NUM_RE.findall(text):
        if n not in ok_nums:
            problems.append(f"数字 {n} 不在系统给出的数字里（金额须按“万”显示，不得自行计算）")
    for label, c in (comps.items() if words else ()):
        if not c:
            continue
        kw = _CMP_KW[label]
        pos = re.compile(rf"(?:高于|超过|超出|多于|超额完成)\s*(?:{kw})|[比较]\s*(?:{kw})值?\s*(?:多|高|增|超)")
        neg = re.compile(rf"(?:低于|少于|未达到|未达|不及|落后于)\s*(?:{kw})|[比较]\s*(?:{kw})值?\s*(?:少|低|减|下降|差)")
        for clause in re.split(r"[，。；;,\n]", text):
            bad = (neg if c["word"] != "低于" else pos).search(clause)
            if c["word"] == "持平":
                bad = pos.search(clause) or neg.search(clause)
            if bad:
                problems.append(f"与{label}的比较写成“{bad.group()}”，系统计算为{c['word']}（差额 {c['diff']:g}）")
    for m in _SIGN_RE.finditer(text):
        v = m.group(2) + m.group(3)
        for label, c in comps.items():
            if c and c["diff"] and fmt_money(abs(c["diff"])).startswith(v) and (m.group(1) == "+") != (c["diff"] > 0):
                problems.append(f"与{label}的差额写成 {m.group(1)}{v}，正负号与系统计算相反")
    return list(dict.fromkeys(problems))


ANALYZE_SYSTEM = """你是项目商务经理。系统已算好上周产值与计划、前一周的比较（高于/低于、差额、百分比），写在【系统已算好的比较】里。
你只负责解释原因：主要贡献或拖累作业、偏差原因（只能引用给定日志问题）、需要采取的措施。
硬性要求：
1. 高于/低于、差额、百分比只能照抄【系统已算好的比较】，不得改写大小关系，不得自行计算新数字；
2. 金额一律照抄给定的“万”格式（如 40.3万 USD），不要写成原始长数字；
3. 周期叫“上周”（与页面卡片一致），不要写“本周”。
只输出JSON：{"summary":"两句话结论","drivers":["主要贡献或拖累作业"],"actions":["建议措施"]}"""


def _currency(conn, project_id: int) -> str:
    r = one(conn, "SELECT MAX(currency) c FROM boq_items WHERE project_id=?", (project_id,))
    if r and r["c"]:
        return r["c"]
    r = one(conn, "SELECT MAX(currency) c FROM output_records WHERE project_id=? AND currency<>''", (project_id,))
    return (r or {}).get("c") or ""


def analyze(conn, project_id: int, user_id: int | None) -> dict:
    from . import facts as F
    series = weekly_series(conn, project_id, 8)
    w, pw = series[-2], series[-3]  # 上一个完整周：与产值页卡片“上周产值”同一周期
    start = w["week"]
    end = (date.fromisoformat(start) + timedelta(days=6)).isoformat()
    week = compute(conn, project_id, start, end)
    cur = _currency(conn, project_id)
    val = lambda x: x["record_actual"] if x["record_actual"] is not None else x["system"]  # noqa: E731
    actual, prev = val(w), val(pw)
    label = f"上周（{start[5:]}～{end[5:]}）"
    comps = {"计划": compare(actual, w["planned"]), "前一周": compare(actual, prev)}
    fc = month_forecast(series)
    M = lambda v: fmt_money(v, cur)  # noqa: E731
    if actual is None:
        cmp_text = f"{label}无产值数据，不作比较"
    else:
        cmp_text = (f"{label}产值 {M(actual)}（{'周记录' if w['record_actual'] is not None else '系统统计'}），"
                    f"{_cmp_sentence('计划', comps['计划'], w['planned'], cur)}；{_cmp_sentence('前一周', comps['前一周'], prev, cur)}")
    issues = []
    for r in all_(conn, "SELECT id, content_json FROM logs WHERE project_id=? AND log_date BETWEEN ? AND ? AND status IN ('submitted','approved')", (project_id, start, end)):
        issues += [f"L{r['id']} {i.get('desc')}" for i in loads(r["content_json"], {}).get("issues", [])[:2]]
    facts = ("【系统已算好的比较】" + cmp_text +
             "\n近8周（周起 系统统计 记录实际 计划）：\n" + "\n".join(f"{x['week']} {M(x['system'])} {M(x['record_actual'])} {M(x['planned'])}" for x in series) +
             f"\n上周按作业：" + ("；".join(f"{_an.fmt(a['activity_code'], a['name'])} {M(a['value'])}" for a in week["by_activity"][:8]) or "无") +
             f"\n月末预测：{M(fc.get('forecast'))}（{fc['note']}）\n上周日志问题：" + ("；".join(issues[:10]) or "无"))
    template = lambda: {"summary": cmp_text + "。", "drivers": [a["name"] for a in week["by_activity"][:2]], "actions": []}  # noqa: E731

    def call(feedback: str):
        return call_json("output.analyze", system=ANALYZE_SYSTEM, user=facts + feedback, prompt_version="output-v2-20261006",
                         mock=template,
                         validate=lambda d: {"summary": str(d.get("summary") or ""), "drivers": [str(x) for x in d.get("drivers") or []][:4],
                                             "actions": [str(x) for x in d.get("actions") or []][:4]},
                         project_id=project_id, user_id=user_id, max_tokens=800)
    # 比较词/正负号/数字与系统计算不一致：带原因重试一次，仍不过用系统模板
    g = F.guarded(call, lambda d: "", None, template,
                  check=lambda d: check_summary(" ".join([d["summary"], *d["drivers"]]), comps, facts)
                  + check_summary(" ".join(d["actions"]), comps, facts + " " + " ".join(re.findall(r"\d+(?:\.\d+)?", " ".join(d["actions"]))), words=False))  # 措施只查数字与正负号，不查比较词
    out = {**g.data, "source": g.source, "period_label": label, "comparisons": comps, "currency": cur,
           "week": week, "series": series, "month_forecast": fc, "at": now_iso()}
    if g.source == "template":
        out["note"] = "AI 文字与系统计算不一致（比较方向/数字），已改用系统模板文字"
        out["rejected"] = g.problems[:3]
    conn.execute("INSERT OR REPLACE INTO knowledge (project_id, kind, key, value_json, ai_run_id, updated_at) VALUES (?,?,?,?,?,?)",
                 (project_id, "output_analysis", "latest", dumps(out), g.run_id, now_iso()))
    return out


def overview(conn, project_id: int) -> dict:
    series = weekly_series(conn, project_id, 16)
    boq = one(conn, "SELECT COUNT(*) n, MAX(currency) cur FROM boq_items WHERE project_id=?", (project_id,))
    mp = one(conn, "SELECT SUM(status='confirmed') c, SUM(status='suggested') s FROM boq_map WHERE project_id=?", (project_id,))
    latest = one(conn, "SELECT value_json FROM knowledge WHERE project_id=? AND kind='output_analysis' AND key='latest'", (project_id,))
    cur_rec = [w for w in series if w["record_actual"] is not None]
    return {"series": series, "month_forecast": month_forecast(series), "boq_items": boq["n"], "currency": boq["cur"] or "",
            "maps_confirmed": mp["c"] or 0, "maps_suggested": mp["s"] or 0,
            "analysis": loads(latest["value_json"]) if latest else None, "has_records": bool(cur_rec),
            "report": report_kpis(conn, project_id, series), "kpi10": kpi10(conn, project_id, series=series)}


_SPLIT_KEYS = ("system_self", "system_approved", "log_ids_self", "log_ids_approved", "week_self", "week_approved")


def strip_review_split(series: list[dict]) -> list[dict]:
    """项目经理口径：周序列去掉自报/核定拆分与来源日志编号（日志审核数据），合计值照常。"""
    return [{k: v for k, v in w.items() if k not in _SPLIT_KEYS} for w in series or []]


# ----------------------------------------------------------------------------- 周报指标（与项目周报“产值完成情况”同一套口径）

def plan_config(conn, project_id: int) -> dict:
    """项目级产值计划参数（存 projects.config_json 的 output_plan 段，缺即为空，不补零）：
    合同额（只作参考显示）、清单总产值（项目整体完成百分比的唯一分母，与项目产值台账一致；缺失即不算）、年度计划产值、
    report_lag_days（周产值按“上报日 = 周末 + N 天”归年度，与台账“记录日期”一致；缺省 0 即按周末日期）。"""
    p = one(conn, "SELECT config_json FROM projects WHERE id=?", (project_id,))
    raw = (loads(p["config_json"], {}) if p else {}).get("output_plan") or {}
    years = {str(k): _num(v) for k, v in (raw.get("year_plan") or {}).items() if _num(v) is not None}
    return {"contract_amount": _num(raw.get("contract_amount")), "boq_total": _num(raw.get("boq_total")), "year_plan": years,
            "report_lag_days": int(_num(raw.get("report_lag_days")) or 0), "source": str(raw.get("source") or "")}


def save_plan_config(conn, project_id: int, user_id: int, contract_amount=None, year: str | None = None, year_plan=None, source: str = "",
                     boq_total=None, report_lag_days: int | None = None) -> dict:
    """工程主管录入合同额/清单总产值/年度计划产值；合同额、年度计划传空值即清空该项（显示“未录入”），清单总产值不传则不改。"""
    p = one(conn, "SELECT config_json FROM projects WHERE id=?", (project_id,))
    cfg = loads(p["config_json"], {}) if p else {}
    plan = dict(cfg.get("output_plan") or {})
    ca = _num(contract_amount)
    if ca is not None and ca <= 0:
        raise ValueError("合同额须大于 0")
    plan["contract_amount"] = ca
    if boq_total is not None:
        bt = _num(boq_total)
        if bt is None or bt <= 0:
            raise ValueError("清单总产值须大于 0")
        plan["boq_total"] = bt
    if report_lag_days is not None:
        plan["report_lag_days"] = max(0, min(6, int(report_lag_days)))
    if year:
        if not re.fullmatch(r"20\d\d", str(year)):
            raise ValueError("年份格式不对")
        yp = _num(year_plan)
        if yp is not None and yp < 0:
            raise ValueError("年度计划产值不能为负")
        years = dict(plan.get("year_plan") or {})
        years[str(year)] = yp
        plan["year_plan"] = {k: v for k, v in years.items() if v is not None}
    if source:
        plan["source"] = source[:200]
    cfg["output_plan"] = plan
    with tx(conn):
        conn.execute("UPDATE projects SET config_json=? WHERE id=?", (dumps(cfg), project_id))
        audit(conn, user_id, "output.plan", f"project:{project_id}", dumps(plan))
    return plan_config(conn, project_id)


def _pct(a, b):
    return round(a / b * 100, 2) if a is not None and b else None


def _money_currencies(conn, project_id: int, last_end: str) -> list[str]:
    """开累产值（周产值记录）与清单总产值（工程量清单）涉及的币种；多于一种时不算完成百分比（不混币种、不换算）。"""
    rows = all_(conn, """SELECT DISTINCT currency c FROM output_records WHERE project_id=? AND period_end<=? AND actual IS NOT NULL AND currency<>''
                         UNION SELECT DISTINCT currency FROM boq_items WHERE project_id=? AND currency<>''""", (project_id, last_end, project_id))
    return sorted({str(r["c"]).strip().upper() for r in rows if str(r["c"] or "").strip()})


def report_kpis(conn, project_id: int, series: list[dict] | None = None, today: date | None = None) -> dict:
    """按项目周报“产值完成情况”的指标给出一组数：上周（基准计划/更新计划/完成/实际与计划差值）、本周计划、
    开累产值与完成百分比、本年度计划与完成。完成值优先取周产值记录，缺失时取系统统计（与周图同一取法）。
    项目整体完成百分比按产值口径（RJ 10-10 决定 2）：开累产值 ÷ 清单总产值；清单总产值缺失或为零、或币种不一致时为 None（显示“—”），
    不用合同额顶替。项目总体计划完成百分比仅取已确认权威台账同周期的原值；来源未核实则为 None，不套用该算法、不推导。
    任何一项缺数据就为 None 并列进 missing（前端显示“未报送/—”），不补零、不估算。"""
    today = today or date.today()
    series = series if series is not None else weekly_series(conn, project_id, 16, today)
    cur = _currency(conn, project_id)
    plan = plan_config(conn, project_id)
    val = lambda w: w["record_actual"] if w["record_actual"] is not None else w["system"]  # noqa: E731
    lw, tw = (series[-2] if len(series) >= 2 else {}), (series[-1] if series else {})
    last_end = (date.fromisoformat(lw["week"]) + timedelta(days=6)).isoformat() if lw else today.isoformat()
    actual = val(lw) if lw else None
    last = {"week": lw.get("week"), "end": last_end if lw else None, "baseline": lw.get("baseline"), "planned": lw.get("planned"), "actual": actual,
            "actual_src": ("周记录" if lw.get("record_actual") is not None else "系统统计") if actual is not None else None,
            "diff": round(actual - lw["planned"], 2) if actual is not None and lw.get("planned") is not None else None}
    this = {"week": tw.get("week"), "baseline": tw.get("baseline"), "planned": tw.get("planned")}
    # 开累与本年：周产值记录逐周累计；没有任何记录时按日志统计（两者不混加）
    by_week: dict[str, dict] = {}
    for r in all_(conn, """SELECT period_start, period_end, actual, baseline, source FROM output_records WHERE project_id=? AND period_end<=?
                           ORDER BY period_start, source<>'record'""", (project_id, last_end)):
        by_week.setdefault(r["period_start"], r)   # 同一周有多个来源时取真实记录（record）
    with_actual = [r for r in by_week.values() if r["actual"] is not None]
    lag = timedelta(days=plan["report_lag_days"])
    year_of = lambda end: str((date.fromisoformat(end) + lag).year)  # noqa: E731
    year = year_of(last_end)
    if with_actual:
        cum = round(sum(r["actual"] for r in with_actual), 2)
        ytd = round(sum(r["actual"] for r in with_actual if year_of(r["period_end"]) == year), 2)
        cum_src = "周产值记录累计"
    else:
        c_all = compute(conn, project_id, "1900-01-01", last_end)["total"]
        c_ytd = compute(conn, project_id, f"{year}-01-01", last_end)["total"]
        cum, ytd = (c_all or None), (c_ytd or None)
        cum_src = "按日志统计累计" if cum else None
    contract, boq_total = plan["contract_amount"], plan["boq_total"]
    curs = _money_currencies(conn, project_id, last_end)
    pct, pct_why = None, None
    if cum is None:
        pct_why = "开累产值未报送"
    elif not boq_total or boq_total <= 0:
        pct_why = "清单总产值未录入"
    elif len(curs) > 1:
        pct_why = f"开累产值与清单总产值币种不一致（{'、'.join(curs)}），不换算"
    else:
        pct = _pct(cum, boq_total)
    base_cum = None
    if with_actual and all(r["baseline"] is not None for r in with_actual):
        base_cum = round(sum(r["baseline"] for r in with_actual), 2)
    year_plan = plan["year_plan"].get(year)
    # RJ 2026-10-11 指定周产值台账为权威：总体计划百分比直接读取同周期来源值，不推导。
    project_cfg = loads((one(conn, "SELECT config_json FROM projects WHERE id=?", (project_id,)) or {}).get("config_json"), {})
    ledger = project_cfg.get("output_ledger") or {}
    source_row = next((r for r in ledger.get("rows", []) if r.get("start") == lw.get("week") and r.get("end") == last_end), None)
    source_matches = (source_row is not None and ledger.get("confirmed_by") and ledger.get("sha256")
                      and ledger.get("currency") == cur and len(curs) <= 1 and cum is not None
                      and _num(source_row.get("cum")) is not None and abs(_num(source_row["cum"]) - cum) <= 0.02
                      and source_row.get("actual") == lw.get("record_actual")
                      and source_row.get("planned") == lw.get("planned") and source_row.get("baseline") == lw.get("baseline"))
    plan_pct = _num(source_row.get("plan_pct")) if source_matches else None
    plan_pct = round(plan_pct, 2) if plan_pct is not None else None
    plan_pct_why = None if plan_pct is not None else "来源未核实，暂不计算"
    out = {
        "currency": cur, "last_week": last, "this_week": this,
        "cumulative": {"actual": cum, "src": cum_src, "contract": contract, "boq_total": boq_total, "denom_name": "清单总产值",
                       "pct": pct, "pct_basis": "按产值口径：开累产值 ÷ 清单总产值", "pct_why": pct_why, "currencies": curs,
                       "baseline_cum": base_cum, "plan_pct": plan_pct, "plan_pct_why": plan_pct_why,
                       "plan_pct_source": ledger.get("file") if plan_pct is not None else None,
                       "plan_pct_sha256": ledger.get("sha256") if plan_pct is not None else None,
                       "gap": round(pct - plan_pct, 2) if pct is not None and plan_pct is not None else None},
        "year": {"year": year, "plan": year_plan, "actual": ytd, "pct": _pct(ytd, year_plan)},
        "source": plan["source"],
        "note": "项目整体完成百分比按产值口径：开累产值 ÷ 清单总产值（不替代关键路径判断）；项目总体计划完成百分比来源未核实，暂不计算。"
                "内部产值统计，不是业主计量或结算金额。",
    }
    if plan_pct is not None:
        out["note"] = ("项目整体完成百分比按产值口径：开累产值 ÷ 清单总产值（不替代关键路径判断）；"
                       f"项目总体计划完成百分比直接取权威周台账《{ledger['file']}》，记录日期 {source_row['reported']}，不由模型推导。"
                       "内部产值统计，不是业主计量或结算金额。")
    labels = [("上周基准产值计划", last["baseline"]), ("上周更新产值计划", last["planned"]), ("本周基准产值计划", this["baseline"]),
              ("本周更新产值计划", this["planned"]), ("清单总产值", boq_total), (f"{year} 年度计划产值", year_plan)]
    out["missing"] = [k for k, v in labels if v is None]
    return out


# ----------------------------------------------------------------------------- 公司周报十项指标（23 号约定）：产值页查询与周报发布快照同一来源

KPI10 = [  # (键, 名称（用户原名，完整显示）, 单位类型, 分组)
    ("last_baseline", "上周基准产值计划", "money", "上周产值"),
    ("last_planned", "上周更新产值计划", "money", "上周产值"),
    ("last_actual", "上周完成产值", "money", "上周产值"),
    ("last_diff", "上周实际与计划差值", "money", "上周产值"),
    ("pct", "项目整体完成百分比", "pct", "整体进度"),
    ("plan_pct", "项目总体计划完成百分比", "pct", "整体进度"),
    ("gap", "项目实际进度与总体计划进度差", "pp", "整体进度"),
    ("year_plan", "本年度计划产值", "money", "年度与开累"),
    ("year_actual", "本年度已完成产值", "money", "年度与开累"),
    ("cum", "项目开累产值", "money", "年度与开累"),
]
_KPI10_NOTE = {
    "last_diff": "完成产值－更新计划：正值超更新计划，负值低于更新计划；不等于进度正常或滞后",
    "pct": "按产值口径：开累产值 ÷ 清单总产值；不替代关键路径判断",
    "gap": "实际百分比－总体计划百分比",
}


def money_unit(cur: str) -> str:
    c = (cur or "").strip().upper()
    return "万美元" if c in ("USD", "US$", "美元") else f"万{cur.strip()}" if c else "万（币种未注明）"


def kpi10(conn, project_id: int, week: str | None = None, today: date | None = None, series: list[dict] | None = None) -> dict:
    """十项指标。week=报告所属周的周一（周报用）；缺省为上一个完整周（产值页用）。金额按“万+币种”给出，原币不换算。"""
    if week:
        today = date.fromisoformat(week) + timedelta(days=7)   # 让“上周”= 所属周
    today = today or date.today()
    rep = report_kpis(conn, project_id, series if series is not None and not week else weekly_series(conn, project_id, 16, today), today)
    lw, cum, yr = rep["last_week"], rep["cumulative"], rep["year"]
    wan = lambda v: None if v is None else round(v / 1e4, 2)  # noqa: E731
    raw = {"last_baseline": wan(lw["baseline"]), "last_planned": wan(lw["planned"]), "last_actual": wan(lw["actual"]),
           "last_diff": wan(lw["diff"]), "pct": cum["pct"], "plan_pct": cum["plan_pct"],
           "gap": round(cum["pct"] - cum["plan_pct"], 2) if cum["pct"] is not None and cum["plan_pct"] is not None else None,
           "year_plan": wan(yr["plan"]), "year_actual": wan(yr["actual"]), "cum": wan(cum["actual"])}
    why = {"pct": cum["pct_why"], "plan_pct": cum["plan_pct_why"], "gap": "总体计划完成百分比暂不计算" if cum["plan_pct"] is None else None,
           "last_diff": "缺更新产值计划或完成产值" if lw["diff"] is None else None}
    unit = money_unit(rep["currency"])
    units = {"money": unit, "pct": "%", "pp": "个百分点"}
    items = [{"key": k, "label": lab, "group": grp, "value": raw[k], "unit": units[t], "note": _KPI10_NOTE.get(k, ""),
              "why": (why.get(k) or "未报送") if raw[k] is None else None} for k, lab, t, grp in KPI10]
    end = lw["end"]
    # 来源版本：所用周产值记录与计划参数的摘要（数据不变则版本号不变），用于追溯快照依据
    recs = all_(conn, """SELECT period_start, period_end, planned, actual, baseline, currency, source FROM output_records
                         WHERE project_id=? AND period_end<=? ORDER BY period_start, source""", (project_id, end or today.isoformat()))
    ver = hashlib.sha1(dumps({"r": recs, "p": plan_config(conn, project_id),
                             "plan_pct": cum["plan_pct"], "plan_pct_source": cum.get("plan_pct_source"),
                             "plan_pct_sha256": cum.get("plan_pct_sha256")}).encode("utf-8")).hexdigest()[:10]
    return {"period": {"week": lw["week"], "start": lw["week"], "end": end}, "year": yr["year"], "currency": rep["currency"], "unit": unit,
            "items": items, "year_pct": yr["pct"], "actual_src": lw["actual_src"], "cum_src": cum["src"], "basis": rep["note"],
            "source": {"version": ver, "records": len(recs), "plan_source": rep["source"], "data_date": end}, "kind": "live", "at": now_iso()}


def snapshot_report(conn, report: dict) -> dict | None:
    """周报发布时冻结十项指标（期间、单位、口径、来源版本一并存）；已有快照不覆盖，月报不存。"""
    if report["kind"] != "week":
        return None
    key = str(report["id"])
    old = one(conn, "SELECT value_json FROM knowledge WHERE project_id=? AND kind='report_kpi10' AND key=?", (report["project_id"], key))
    if old:
        return loads(old["value_json"])
    snap = {**kpi10(conn, report["project_id"], week=report["period_start"]), "kind": "snapshot"}
    conn.execute("INSERT INTO knowledge (project_id, kind, key, value_json, ai_run_id, updated_at) VALUES (?,?,?,?,?,?)",
                 (report["project_id"], "report_kpi10", key, dumps(snap), None, now_iso()))
    return snap


def report_kpi10(conn, report: dict) -> dict | None:
    """报表里的十项指标：已发布读快照；发布于快照上线前的没有快照，如实说明，不回填；草稿为实时查询。"""
    if report["kind"] != "week":
        return None
    if report["status"] == "published":
        row = one(conn, "SELECT value_json FROM knowledge WHERE project_id=? AND kind='report_kpi10' AND key=?", (report["project_id"], str(report["id"])))
        if row:
            return loads(row["value_json"])
        return {"kind": "none", "note": "该周报发布时尚未保存指标快照，历史报告不回填"}
    return {**kpi10(conn, report["project_id"], week=report["period_start"]), "kind": "live"}

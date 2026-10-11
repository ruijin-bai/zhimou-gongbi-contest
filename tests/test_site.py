"""档案Agent：部位档案与索赔证据包。数据直接写表构造，避免依赖其他 Agent 的行为。"""
import io
from datetime import timedelta

from conftest import TODAY, jpeg, login, sample_xer
from test_roles import _team

from app import db as dbmod
from app.ai.client import AIError, AIResult
from app.services import claims as claim_svc
from app.services import findings as find_svc
from app.services import media as media_svc
from app.services import parts as part_svc


def D(n: int) -> str:
    return (TODAY - timedelta(days=n)).isoformat()


def _mklog(conn, env, user, day, content, status="approved", weather="晴 30℃", voice="", progress=()):
    pid, uid = env["project_id"], env["users"][user]
    lid = dbmod.insert(conn, "logs", {"project_id": pid, "author_id": uid, "log_date": day, "mode": "ai", "status": status, "weather": weather, "voice_text": voice,
                                      "content_json": dbmod.dumps(content), "started_at": day + "T07:00:00", "period": "day", "auto": 1,
                                      "submitted_at": (day + "T20:00:00") if status != "draft" else None})
    for code, qty, unit, pct in progress:
        dbmod.insert(conn, "log_progress", {"log_id": lid, "project_id": pid, "log_date": day, "activity_code": code, "qty": qty, "unit": unit, "pct": pct, "note": ""})
    return lid


def _photo(conn, env, user, color, log_id=None, purpose="log", taken=None, place="", lat=None):
    geo = media_svc.parse_geo(lat, 113.0 if lat else None, 8, place) if (lat or place) else None
    m = media_svc.save_photo(conn, env["project_id"], env["users"][user], "p.jpg", jpeg(color), purpose, taken_hint=taken, geo=geo)
    if log_id:
        conn.execute("UPDATE media SET log_id=? WHERE id=?", (log_id, m["id"]))
    return m["id"]


def _world(env):
    """甲(tech1)做3#墩，乙(tech2)做路基；3#墩连续两天设备故障、一条缺料(未周确认)；乙有暴雨停工。"""
    _team(env)
    _add_boss(env)
    pm = login("pm")
    assert pm.post("/api/p6/import", files={"file": ("plan.xer", sample_xer(), "application/octet-stream")}).status_code == 200
    conn = dbmod.connect()
    pid = env["project_id"]
    work = lambda d, code="A1010": {"part": "3#墩", "process": "桩基", "desc": f"3#墩桩基施工 {d}", "activity_code": code}  # noqa: E731
    lab = [{"trade": "钢筋工", "count": 12}, {"trade": "普工", "count": 6}]
    mac = [{"name": "旋挖钻", "count": 2, "status": "正常"}]
    ids = {}
    for n, qty in ((D(6), 10), (D(5), 10), (D(2), 10)):
        ids[n] = _mklog(conn, env, "tech1", n, {"works": [work(n)], "labor": lab, "machinery": mac, "issues": []}, progress=[("A1010", qty, "根", None)])
    ids["e1"] = _mklog(conn, env, "tech1", D(4), {"summary": "旋挖钻故障", "works": [work(D(4))], "labor": lab, "machinery": [{"name": "旋挖钻", "count": 2, "status": "故障"}],
                       "issues": [{"type": "设备", "severity": "中", "desc": "旋挖钻故障停机半天", "measure": "已联系维修", "activity_code": "A1010"}]},
                       voice="旋挖钻故障停机半天，下午恢复", progress=[("A1010", 2, "根", 20)])
    ids["e2"] = _mklog(conn, env, "tech1", D(3), {"works": [work(D(3))], "labor": lab, "machinery": mac,
                       "issues": [{"type": "设备", "severity": "中", "desc": "旋挖钻液压故障，继续维修", "measure": "", "activity_code": "A1010"}]},
                       progress=[("A1010", 4, "根", 30)])
    ids["m"] = _mklog(conn, env, "tech1", D(1), {"works": [work(D(1))], "labor": [], "machinery": [], "issues": [
        {"type": "材料", "severity": "高", "desc": "钢筋进场不足，缺料", "measure": "", "activity_code": "A1010"}]}, status="draft")
    ids["w"] = _mklog(conn, env, "tech2", D(3), {"works": [{"part": "路基工程", "process": "开挖", "desc": "路基土方开挖", "activity_code": "A3010"}], "labor": [], "machinery": [],
                       "issues": [{"type": "天气", "severity": "中", "desc": "暴雨停工一天", "measure": "排水", "activity_code": "A3010"}]},
                       weather="暴雨", progress=[("A3010", 100, "m3", None)])
    ids["x"] = _mklog(conn, env, "tech1", D(7), {"works": [], "labor": [], "machinery": [], "issues": [
        {"type": "安全", "severity": "低", "desc": "个别工人未戴手套", "measure": "已纠正", "activity_code": ""}]})
    # 照片：日志照片（带拍摄时间与地点）、随手拍隐患（地点写3#墩东侧）、随手记照片
    taken = f"{D(4)}T14:20:00"
    ids["p_log"] = _photo(conn, env, "tech1", (90, 100, 110), ids["e1"], taken=taken, place="3#墩", lat=22.5)
    ids["p_hz"] = _photo(conn, env, "tech1", (200, 50, 50), purpose="safety", place="3#墩东侧")
    ids["p_note"] = _photo(conn, env, "tech1", (50, 120, 60))
    ids["p_b"] = _photo(conn, env, "tech2", (10, 20, 220), ids["w"])
    # 隐患（挂在随手拍照片上）+ 质控提醒
    ids["hz"] = find_svc.create(conn, pid, "safety", "hazard", "high", "3#墩临边防护缺失", {"category": "临边"}, ref_type="media", ref_id=ids["p_hz"],
                                activity_code=None, fingerprint="t:hz1", assignee_id=env["users"]["tech1"], due_days=2)
    ids["qc"] = find_svc.create(conn, pid, "quality", "qc_reminder", "info", "A1010 开工前质控", {"points": [{"stage": "开工前", "item": "桩位复核", "record": "测量记录"}]},
                                activity_code="A1010", fingerprint="t:qc1", assignee_id=env["users"]["tech1"], due_days=1)
    dbmod.insert(conn, "notes", {"project_id": pid, "author_id": env["users"]["tech1"], "note_date": D(4), "text": "3号墩钢筋笼吊装完成", "media_json": dbmod.dumps([ids["p_note"]]),
                                 "source": "voice", "created_at": f"{D(4)}T16:00:00"})
    conn.close()
    return ids


def _add_boss(env):
    from test_roles import _add
    _add(env, "boss", "leader", "领导")


def test_part_resolution():
    assert part_svc.part_from_text("3#墩桩基施工") == "3#墩"
    assert part_svc.part_from_text("三期 03号桥台") == "3#台"
    assert part_svc.part_from_text("1#盖板涵施工") == "1#涵"
    assert part_svc.part_from_text("路基土方") is None
    assert part_svc.activity_part({"wbs_path": "桥涵工程 / 立交桥", "name": "立交桥桩基施工"}) == "立交桥"
    assert part_svc.activity_part({"wbs_path": "示例大桥 / 3#墩", "name": "承台"}) == "3#墩"
    assert part_svc.activity_part({"wbs_path": "", "name": "开工", "is_milestone": 1}) is None
    assert sorted(["10#墩", "2#墩", part_svc.UNASSIGNED, "1#墩"], key=part_svc.sort_key) == ["1#墩", "2#墩", "10#墩", part_svc.UNASSIGNED]
    assert claim_svc._classify({"type": "其他", "desc": "午后强降雨停工约3小时"}) == ("weather", True)
    assert claim_svc._classify({"type": "安全", "desc": "未系安全带"}) == (None, False)
    assert claim_svc._classify({"type": "外部", "desc": "社区道路占用协调，短暂停工"}) == ("stop", True)
    assert claim_svc._classify({"type": "质量", "desc": "业主要求变更配筋"}) == ("change", False)


def test_parts_archive_and_scope(env):
    ids = _world(env)
    pm, boss = login("pm"), login("boss")
    data = pm.get("/api/sites").json()
    names = [p["name"] for p in data["parts"]]
    assert "3#墩" in names and "4#墩" in names and "路基工程" in names and data["scope"] == "全项目" and data["can_write"]
    p3 = next(p for p in data["parts"] if p["name"] == "3#墩")
    # 日志照片+随手拍+随手记照片都归3#墩；隐患未闭环1；质控1/0；只计已确认日志的日期
    assert p3["photos"] == 3 and p3["hazards"] == 1 and p3["hazards_open"] == 1 and p3["qc_total"] == 1 and p3["qc_done"] == 0
    assert p3["log_days"] == 5 and p3["issues"] == 2 and p3["last_update"] >= D(3)
    assert next(p for p in data["parts"] if p["name"] == "4#墩")["log_days"] == 0  # 无记录如实为0
    # 无法判断部位的日志问题（日志里没有施工内容、也没写作业）归“未归部位”，不猜
    un = pm.get("/api/sites/detail", params={"part": part_svc.UNASSIGNED}).json()
    assert [i["text"] for i in un["items"] if i["type"] == "issue"] == ["个别工人未戴手套"]

    d = pm.get("/api/sites/detail", params={"part": "3#墩"}).json()
    types = {i["type"] for i in d["items"]}
    assert {"photo", "progress", "qc", "hazard", "issue", "note"} == types
    act = next(a for a in d["stats"]["activities"] if a["code"] == "A1010")
    assert act["qty"] == [{"qty": 36.0, "unit": "根"}] and act["log_pct"] == 30
    ph = next(i for i in d["items"] if i["type"] == "photo")["photos"][0]
    assert ph["taken_at"] == f"{D(4)}T14:20" and ph["geo"]["lat"] == 22.5
    hz = next(i for i in d["items"] if i["type"] == "hazard")
    assert hz["open"] and hz["photos"] and hz["finding_id"] == ids["hz"]
    only = pm.get("/api/sites/detail", params={"part": "3#墩", "types": "hazard,qc"}).json()
    assert {i["type"] for i in only["items"]} == {"hazard", "qc"}
    # 未确认(draft)日志的问题不进档案
    assert not any("缺料" in i.get("text", "") for i in d["items"])

    # 范围：甲只见甲，乙只见乙；乙看不到甲的照片
    t1, t2 = login("tech1"), login("tech2")
    assert [p["name"] for p in t2.get("/api/sites").json()["parts"]] == ["路基工程"]
    assert t2.get("/api/sites/detail", params={"part": "3#墩"}).status_code == 400
    assert t1.get("/api/sites/detail", params={"part": "路基工程"}).status_code == 400
    d1 = t1.get("/api/sites/detail", params={"part": "3#墩"}).json()
    assert d1["summary"] is None and d1["scope"] == "甲" and not d1["can_write"]
    assert t2.get(ph["thumb"]).status_code == 403
    assert t1.post("/api/sites/summary", json={"part": "3#墩"}).status_code == 403
    assert t1.post("/api/sites/organize").status_code == 403
    # 管理员不参与；领导只读可看
    assert login("admin").get("/api/sites").status_code == 403
    assert boss.get("/api/sites").status_code == 200 and boss.post("/api/sites/summary", json={"part": "3#墩"}).status_code == 403
    assert boss.get("/api/sites").json()["can_write"] is False


def test_summary_number_guard_and_docx(env):
    _world(env)
    pm, eng = login("pm"), login("eng")
    r = pm.post("/api/sites/summary", json={"part": "3#墩"})
    assert r.status_code == 200 and r.json()["source"] == "ai" and "3#墩" in r.json()["text"]
    d = pm.get("/api/sites/detail", params={"part": "3#墩"}).json()
    assert d["summary"]["text"] == r.json()["text"] and d["summary"]["stale"] is False
    # 数字守卫：模型自造数字 → 改用规则概况
    orig = part_svc.call_json
    part_svc.call_json = lambda kind, **kw: AIResult(data={"summary": "该部位已完成桩基 9999 根"}, run_id=1, provider="mock", model="mock", latency_ms=1)
    try:
        r2 = eng.post("/api/sites/summary", json={"part": "3#墩"}).json()
        assert r2["source"] == "rule" and "9999" not in r2["text"] and r2["guard"][0]["number"] == "9999"

        def boom(kind, **kw):
            raise AIError("不可用")
        part_svc.call_json = boom
        assert eng.post("/api/sites/summary", json={"part": "3#墩"}).json()["source"] == "rule"
    finally:
        part_svc.call_json = orig
    # 导出 docx
    rd = pm.get("/api/sites/docx", params={"part": "3#墩"})
    assert rd.status_code == 200 and rd.content[:2] == b"PK"
    from docx import Document
    doc = Document(io.BytesIO(rd.content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "部位档案：3#墩" in text and "3#墩临边防护缺失" in text and len(doc.inline_shapes) >= 1
    assert login("tech2").get("/api/sites/docx", params={"part": "3#墩"}).status_code == 400
    conn = dbmod.connect()
    actions = {r["action"] for r in conn.execute("SELECT action FROM audit").fetchall()}
    conn.close()
    assert {"site.summary", "site.export"} <= actions


def test_organize_merges_free_text_part(env):
    _world(env)
    conn = dbmod.connect()
    _mklog(conn, env, "tech2", D(2), {"works": [{"part": "路基", "process": "填筑", "desc": "填筑碾压", "activity_code": ""}], "labor": [], "machinery": [], "issues": []})
    conn.close()
    pm = login("pm")
    lst = pm.get("/api/sites").json()
    assert lst["unresolved"] == 1 and lst["unresolved_labels"] == ["路基"]
    un = pm.get("/api/sites/detail", params={"part": part_svc.UNASSIGNED}).json()
    assert any(w["desc"] == "填筑碾压" for i in un["items"] if i["type"] == "progress" for w in i["works"])
    assert pm.post("/api/sites/organize").json() == {"merged": 1, "labels": 1}
    lst = pm.get("/api/sites").json()
    assert lst["unresolved"] == 0
    un = pm.get("/api/sites/detail", params={"part": part_svc.UNASSIGNED}).json()
    assert not any(w["desc"] == "填筑碾压" for i in un["items"] if i["type"] == "progress" for w in i["works"])
    d = pm.get("/api/sites/detail", params={"part": "路基工程"}).json()
    assert any(w["desc"] == "填筑碾压" for i in d["items"] if i["type"] == "progress" for w in i["works"])


def test_claim_events_and_pack(env):
    ids = _world(env)
    pm, eng, ss = login("pm"), login("eng"), login("ss")
    evs = pm.get("/api/claims/events").json()
    cats = sorted((e["cat"], e["company"]) for e in evs["events"])
    # 个别工人未戴手套(安全/无关键词)不是索赔事件；设备故障连续两天合并为一个事件
    assert cats == [("equipment", "甲"), ("material", "甲"), ("weather", "乙")]
    eq = next(e for e in evs["events"] if e["cat"] == "equipment")
    assert eq["date_from"] == D(4) and eq["date_to"] == D(3) and eq["log_days"] == 2 and not eq["halt"]
    wt = next(e for e in evs["events"] if e["cat"] == "weather")
    mt = next(e for e in evs["events"] if e["cat"] == "material")
    assert wt["halt"] and mt["unconfirmed"]

    ev = pm.get("/api/claims/events/detail", params={"key": eq["key"]}).json()
    imp = ev["impact"]["activities"][0]
    # 基线：D6/D5/D2 三天均10根；事件日2根、4根 → 损失 0.8+0.6=1.4 天
    assert imp["code"] == "A1010" and imp["baseline"] == 10 and imp["loss_days"] == 1.4 and ev["impact"]["est_loss_days"] == 1.4
    assert [l["date"] for l in ev["logs"]] == [D(4), D(3)] and ev["logs"][0]["voice_text"].startswith("旋挖钻故障")
    assert ev["photos"][0]["taken_at"] == f"{D(4)}T14:20" and ev["photos"][0]["id"] == ids["p_log"]
    assert any(p["id"] == ids["p_note"] for p in ev["photos"])  # 事件日的随手记照片
    assert ev["labor_machinery"][0]["machinery"][0]["status"] == "故障" and ev["labor_machinery"][0]["labor"][0]["count"] == 12
    gap_items = " ".join(g["item"] for g in ev["gaps"])
    assert "设备故障报告" in gap_items and "采取的措施" in gap_items and "天气记录" not in gap_items
    # 天气事件：日志里有天气 → 不缺天气记录；没有照片以外证据 → 工期基线不足待补
    evw = pm.get("/api/claims/events/detail", params={"key": wt["key"]}).json()
    assert evw["weather"][0]["weather"] == "暴雨" and evw["impact"]["est_loss_days"] is None
    gw = " ".join(g["item"] for g in evw["gaps"])
    assert "气象" in gw and "受影响工期天数" in gw and "路基土方开挖 的工期影响天数" in gw and "天气记录" not in gw
    evm = pm.get("/api/claims/events/detail", params={"key": mt["key"]}).json()
    assert "日志确认" in " ".join(g["item"] for g in evm["gaps"]) and "事件日人机投入" in " ".join(g["item"] for g in evm["gaps"])

    # 生成证据包：AI(mock)整理，注明待核实，写audit
    r = eng.post("/api/claims/packs", json={"keys": [eq["key"], wt["key"]], "title": "9月索赔"})
    assert r.status_code == 200, r.text
    pack = r.json()
    assert pack["title"] == "9月索赔" and pack["notice"] == "AI整理，待商务/项目经理核实后使用" and pack["narrative"]["source"] == "ai"
    assert len(pack["narrative"]["events"]) == 2 and pack["events"][0]["cat"] == "equipment"
    assert pack["events"][0]["impact"]["est_loss_days"] == 1.4
    # 快照不随之后的数据变化：新增日志不改变已生成的证据包
    conn = dbmod.connect()
    _mklog(conn, env, "tech1", D(0), {"works": [], "labor": [], "machinery": [], "issues": [{"type": "设备", "severity": "低", "desc": "旋挖钻再次故障", "measure": "维修", "activity_code": "A1010"}]})
    conn.close()
    again = ss.get(f"/api/claims/packs/{pack['id']}").json()
    assert again["events"][0]["dates"] == [D(4), D(3)]
    assert len(ss.get("/api/claims/packs").json()) == 1 and "events" not in ss.get("/api/claims/packs").json()[0]

    # 导出 docx：含提示、日志原文、人机投入、照片缩略图、待补充项
    rd = pm.get(f"/api/claims/packs/{pack['id']}/docx")
    assert rd.status_code == 200 and rd.content[:2] == b"PK"
    from docx import Document
    doc = Document(io.BytesIO(rd.content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "AI整理，待商务/项目经理核实后使用" in text and "旋挖钻故障停机半天" in text
    assert "□ 设备故障报告及维修记录" in text and "钢筋工 12人" in text and "估算，待核实" in text and "暴雨" in text
    assert len(doc.inline_shapes) >= 2
    conn = dbmod.connect()
    acts = [r["action"] for r in conn.execute("SELECT action FROM audit").fetchall()]
    assert "claim.generate" in acts and "claim.export" in acts
    assert conn.execute("SELECT exported_at FROM claim_packs").fetchone()["exported_at"]
    conn.close()


def test_claim_permissions(env):
    _world(env)
    pm, boss = login("pm"), login("boss")
    key = pm.get("/api/claims/events").json()["events"][0]["key"]
    pack = pm.post("/api/claims/packs", json={"keys": [key]}).json()
    # 领导只读：能看事件与证据包，不能生成与导出
    assert boss.get("/api/claims/events").json()["can_write"] is False
    assert boss.get(f"/api/claims/packs/{pack['id']}").status_code == 200 and boss.get("/api/claims/packs").status_code == 200
    assert boss.post("/api/claims/packs", json={"keys": [key]}).status_code == 403
    assert boss.get(f"/api/claims/packs/{pack['id']}/docx").status_code == 403
    # 分包、管理员不可见
    for who in ("tech1", "tech2", "sa_a", "sa_b", "admin"):
        c = login(who)
        for path in ("/api/claims/events", "/api/claims/packs", f"/api/claims/packs/{pack['id']}", "/api/claims/events/detail?key=" + key):
            assert c.get(path).status_code == 403, (who, path)
        assert c.post("/api/claims/packs", json={"keys": [key]}).status_code == 403
    # 参数校验
    assert pm.post("/api/claims/packs", json={"keys": ["不存在"]}).status_code == 400
    assert pm.post("/api/claims/packs", json={"keys": []}).status_code == 422
    assert pm.get("/api/claims/packs/9999").status_code == 404


def test_claim_narrative_guard_and_fallback(env):
    _world(env)
    pm = login("pm")
    key = next(e["key"] for e in pm.get("/api/claims/events").json()["events"] if e["cat"] == "equipment")
    orig = claim_svc.call_json
    try:
        claim_svc.call_json = lambda kind, **kw: AIResult(data={"overview": "共损失 88 天", "events": [{"ref": "E1", "statement": "故障导致停工 77 天", "evidence": ["日志L1"], "gaps": []}]},
                                                          run_id=1, provider="mock", model="mock", latency_ms=1)
        p = pm.post("/api/claims/packs", json={"keys": [key]}).json()
        assert p["narrative"]["source"] == "rule" and "77" in p["narrative"]["note"] and "77" not in p["narrative"]["events"][0]["statement"]

        def boom(kind, **kw):
            raise AIError("超时")
        claim_svc.call_json = boom
        p = pm.post("/api/claims/packs", json={"keys": [key]}).json()
        assert p["narrative"]["source"] == "rule" and "AI 不可用" in p["narrative"]["note"] and p["notice"].startswith("AI整理")
    finally:
        claim_svc.call_json = orig

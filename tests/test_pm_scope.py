"""项目经理口径：只关注进度、质量、安全，不审核日志。简报、催办链不得把日志填报、产值带给项目经理；产值页与报表产值节只读（RJ 10-10，离线 mock 模型）。"""
from contextlib import closing

from test_facts import _add_alert, _bad_call, _gen, _setup_site

from app import db as dbmod
from app.services import brief
from app.services import nudge as nudge_svc
from test_nav_features import nav  # noqa: F401,E402  （node 跑 nav.js 的夹具）

GOOD_ITEM = {"level": "high", "area": "进度", "title": "A1020 滞后", "detail": "滞后 8 天，超出总浮时 8 天", "action": "赶工"}


def test_brief_drops_log_filling_and_output_items(env, monkeypatch):
    _setup_site(env)
    _add_alert(env["project_id"])
    payload = {"headline": "现场记录填报率不足50%，产值超计划", "items": [
        {"level": "medium", "area": "进度", "title": "现场记录填报率不足50%", "detail": "今日仅2/4技术员填报", "action": "加强填报考核"},
        {"level": "info", "area": "产值", "title": "本周产值超计划", "detail": "本周产值超计划", "action": "保持节奏"},
        {"level": "info", "area": "质量", "title": "数据审核发现重复照片", "detail": "数据审核评分偏低", "action": "核查"},
        GOOD_ITEM]}
    _bad_call(monkeypatch, "brief", "brief.daily", [payload])
    out = _gen(env)
    assert [i["title"] for i in out["items"]] == ["A1020 滞后"]  # 越界条目直接丢弃，不靠提示词自觉
    assert "填报" not in out["headline"] and "产值" not in out["headline"]  # 越界标题换成系统模板标题
    assert out["headline"].startswith("待确认进度预警")      # D-3：系统模板标题写明“待确认”
    assert all(i["area"] in brief.BRIEF_AREAS for i in out["items"])


def test_brief_facts_do_not_carry_log_filling_or_output(env):
    _setup_site(env)
    _add_alert(env["project_id"])
    with closing(dbmod.connect()) as conn:
        k = brief.kpis(conn, env["project_id"], brief.date.today())
        text = brief.facts_text(conn, env["project_id"], k)
    for word in ("产值", "填报", "现场记录", "周确认", "今日无现场记录"):
        assert word not in text
    assert "进度预警" in text and "安全隐患" in text and "质量问题" in text
    assert "产值" not in brief.SYSTEM.split("只输出JSON")[1]  # 输出格式里不再有产值类别


def test_log_nudge_chain_never_reaches_project_manager():
    users = {1: {"id": 1, "role": "tech", "company": "甲", "name": "张工"}, 2: {"id": 2, "role": "eng_sup", "company": "", "name": "郑工"},
             3: {"id": 3, "role": "manager", "company": "", "name": "项目经理"}}
    log_t = {"area": "log", "owner_id": 1}
    r, why = nudge_svc._recipients(3, log_t, users)
    assert [u["id"] for u in r] == [2]  # 日志缺报升级到工程主管为止
    r, why = nudge_svc._recipients(4, log_t, users)
    assert r == [] and "项目经理" in why
    # 隐患、质控等进度质量安全事项仍然升级到项目经理待决策
    for area in ("safety", "quality", "action"):
        r, _ = nudge_svc._recipients(4, {"area": area, "owner_id": 1}, users)
        assert [u["id"] for u in r] == [3]


# ----------------------------------------------------------------------------- 应用成效 / 周报：填报率数字不给项目经理

def _eng_sup(env) -> None:
    from app.security import hash_password
    from conftest import PW
    with closing(dbmod.connect()) as conn:
        dbmod.insert(conn, "users", {"username": "eng", "name": "郑工", "role": "eng_sup", "password_hash": hash_password(PW), "project_id": env["project_id"],
                                     "wbs_scope": "", "email": "", "active": 1, "created_at": dbmod.now_iso()})


def _submitted_report(env):
    from conftest import backdate, login, TODAY
    from test_flow import _draft_today, _import
    pm = login("pm")
    _import(pm)
    tech = login("tech1")
    log = _draft_today(tech)
    backdate(log["id"])
    tech.post(f"/api/logs/{log['id']}/submit", json={"version": log["version"], "active_seconds": 300})
    rep = pm.post("/api/reports", json={"kind": "week", "ref_date": TODAY.isoformat()})
    assert rep.status_code == 200, rep.text
    return pm, rep.json()


def test_metrics_hide_on_time_rate_from_project_manager(env):
    from conftest import login

    from app.services import metrics as metrics_svc
    _eng_sup(env)
    assert login("pm").get("/api/metrics").status_code == 403    # 应用成效整页不给项目经理
    eng_m, adm_m = (login(u).get("/api/metrics").json() for u in ("eng", "admin"))
    for m in (eng_m, adm_m):  # 工程主管、系统管理员口径不变
        assert "on_time_rate" in m["daily_log"] and "当日提交率" in m["definitions"]
    with closing(dbmod.connect()) as conn:   # 服务层的剔除仍保留作双保险
        pm_m = metrics_svc.summary(conn, env["project_id"], "manager")
    assert "on_time_rate" not in pm_m["daily_log"] and "按时" not in str(pm_m) and "应报" not in str(pm_m)
    assert {k: v for k, v in eng_m["daily_log"].items() if k != "on_time_rate"} == pm_m["daily_log"]


def test_report_hides_expected_counts_from_project_manager(env):
    from conftest import login
    _eng_sup(env)
    pm, rep = _submitted_report(env)
    hidden = ("expected", "on_time", "weeks_expected")
    assert not any(k in rep["facts"]["logs"] for k in hidden)
    assert rep["facts"]["logs"]["count"] == 1  # 依据日志份数仍在
    assert not any(k in pm.get(f"/api/reports/{rep['id']}").json()["facts"]["logs"] for k in hidden)
    edited = pm.put(f"/api/reports/{rep['id']}", json={"narrative": rep["narrative"], "status": "draft"}).json()
    assert not any(k in edited["facts"]["logs"] for k in hidden)
    eng = login("eng").get(f"/api/reports/{rep['id']}").json()
    assert all(k in eng["facts"]["logs"] for k in hidden)  # 冻结的事实没改，工程主管仍可见
    assert "应" not in report_agent_prompt(eng["facts"]).split("依据：")[1].split("\n")[0]  # 叙述起草依据里也不带应报/按时份数


def report_agent_prompt(facts: dict) -> str:
    from app.ai import report_agent
    return report_agent.build_prompt(facts)


def test_front_end_hides_when_fields_absent():
    from pathlib import Path
    js = (Path(__file__).resolve().parent.parent / "app" / "static" / "views.js").read_text(encoding="utf-8")
    assert "v-if=\"'on_time_rate' in dl\"" in js
    assert "v-if=\"f.logs.expected !== undefined\"" in js


# ----------------------------------------------------------------------------- B1 周报：项目经理只读产值节，不含自报/核定拆分

def _docx_text(content: bytes) -> str:
    import io
    from docx import Document
    d = Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    return "\n".join(parts)


def test_report_view_for_project_manager_reads_output_without_split_columns(env):
    """RJ 10-10：项目经理只读产值——报表产值节照常显示；工程量的自报/核定拆分（日志审核数据）仍不给；产值节原文不能改。"""
    from conftest import login
    _eng_sup(env)
    pm, rep = _submitted_report(env)
    for v in (rep, pm.get(f"/api/reports/{rep['id']}").json()):
        assert "output" in [s["key"] for s in v["template"]["sections"]] and "output" in v["facts"]
        assert not any(k in q for p in v["facts"]["progress"] for q in p["qty"] for k in ("self", "approved"))
        assert not any(k in p for p in v["facts"]["progress"] for k in ("log_ids_self", "log_ids_approved"))
        assert not {"self_total", "approved_total", "log_ids_self", "log_ids_approved"} & set(v["facts"]["output"])
        for name in ("progress", "output"):
            heads = v["tables"][name]["headers"] if name in v["tables"] else []
            assert not any("自报" in h or "核定" in h for h in heads)
    eng = login("eng").get(f"/api/reports/{rep['id']}").json()   # 工程主管、领导看完整版
    assert "output" in [s["key"] for s in eng["template"]["sections"]] and "output" in eng["facts"]
    assert {"其中分包自报", "其中港湾核定"} <= set(eng["tables"]["progress"]["headers"])
    # Word 导出：项目经理导出的版本含产值节、不含自报/核定；工程主管导出完整版
    pm_doc, eng_doc = (_docx_text(c.get(f"/api/reports/{rep['id']}/docx").content) for c in (pm, login("eng")))
    assert "产值完成情况" in pm_doc
    for w in ("自报", "核定"):
        assert w not in pm_doc, w
    assert "产值完成情况" in eng_doc and "其中分包自报" in eng_doc
    # 项目经理仍可审核发布；提交时改写或清空产值节一律无效，原文保留
    original = eng["narrative"]["output"]
    nar = {**pm.get(f"/api/reports/{rep['id']}").json()["narrative"], "output": "项目经理改写的产值节"}
    pub = pm.put(f"/api/reports/{rep['id']}", json={"narrative": nar, "status": "published",
                                                     "verified": [{"field": x["field"], "number": x["number"]} for x in rep["number_check"]]})
    assert pub.status_code == 200 and pub.json()["status"] == "published"
    assert login("eng").get(f"/api/reports/{rep['id']}").json()["narrative"]["output"] == original


def test_report_template_for_project_manager_has_output_section(env):
    from conftest import login
    _eng_sup(env)
    for who in ("pm", "eng"):
        assert "output" in [s["key"] for s in login(who).get("/api/reports/template/week").json()["sections"]], who


# ----------------------------------------------------------------------------- B2 产值页：项目经理只读，写接口一律拒绝

def test_output_read_only_for_project_manager(env):
    """项目经理能读产值查询接口；导入、AI 匹配、确认/否决映射、计划参数、产值分析直接调用一律 403（不只靠隐藏按钮）。"""
    from conftest import login

    from app import roles
    _eng_sup(env)
    assert not roles.can("manager", "output_ops") and roles.output_readonly("manager")
    assert roles.CAPS["output_ops"] == frozenset({"admin", "eng_sup"})
    pm = login("pm")
    for url in ("/api/output/overview", "/api/output/maps"):
        assert pm.get(url).status_code == 200, url
    ov = pm.get("/api/output/overview").json()
    assert "cumulative" in ov["report"] and len(ov["kpi10"]["items"]) == 10
    assert not any(k in w for w in ov["series"] for k in ("system_self", "system_approved", "log_ids_self", "log_ids_approved"))
    assert all("system_self" in w for w in login("eng").get("/api/output/overview").json()["series"])
    cp = pm.get("/api/cockpit").json()
    assert not any(k in w for w in cp["output_series"] for k in ("system_self", "log_ids_self")) and "week_self" not in cp["kpis"]["output"]
    for url in ("/api/output/map/suggest", "/api/output/analyze"):
        assert pm.post(url).status_code == 403, url
    assert pm.post("/api/output/maps", json={"ids": [], "status": "confirmed"}).status_code == 403
    assert pm.put("/api/output/plan", json={"contract_amount": 1, "boq_total": 1}).status_code == 403
    for kind in ("boq", "records"):
        assert pm.post(f"/api/output/{kind}", files={"file": ("x.xlsx", b"x")}).status_code == 403, kind
    eng = login("eng")
    assert eng.get("/api/output/overview").status_code == 200 and eng.get("/api/output/maps").status_code == 200


def test_output_page_write_controls_gated_by_output_ops():
    """前端：产值页的导入、匹配、确认、录入计划、分析按钮都挂在 can('output_ops') 上；项目经理没有这个能力。"""
    import re
    from pathlib import Path
    js = (Path(__file__).resolve().parent.parent / "app" / "static" / "mgr.js").read_text(encoding="utf-8")
    page = js[js.index("const OutputPage"):js.index("const AGENT_GROUPS")]
    for call in ('POST("/api/output/analyze")', "POST(`/api/output/${kind}`", 'POST("/api/output/map/suggest")', 'POST("/api/output/maps"'):
        assert call in page, call
    assert page.count("can('output_ops')") >= 3


# ----------------------------------------------------------------------------- B3 应用成效 / 智能应用

def test_metrics_api_denied_to_project_manager(env):
    from conftest import login

    from app import roles
    _eng_sup(env)
    assert not roles.can("manager", "metrics_ops")
    pm = login("pm")
    assert pm.get("/api/metrics").status_code == 403
    assert pm.post("/api/metrics/baseline", json={"kind": "report_week", "minutes": 60}).status_code == 403
    assert pm.get("/api/ai/runs").status_code == 403
    assert login("eng").get("/api/metrics").status_code == 200 and login("admin").get("/api/metrics").status_code == 200


def test_agents_page_for_project_manager_only_pm_agents(env):
    from conftest import login

    from app import roles
    _eng_sup(env)
    pm = login("pm").get("/api/agents").json()
    assert pm["agents"] and {a["key"] for a in pm["agents"]} <= roles.PM_AGENTS
    assert not any("日志" in k or "产值" in k for k in pm["penetration"])
    assert "日志" not in pm["definition"] and "产值" not in pm["definition"]
    assert "安全隐患由AI发现" in pm["penetration"]
    eng = login("eng").get("/api/agents").json()
    assert {"scribe", "dq", "output"} <= {a["key"] for a in eng["agents"]} and "日志经AI预审" in eng["penetration"]


# ----------------------------------------------------------------------------- B4 催办台：日志/周确认规则与记录不给项目经理；不出现环境变量名

def _nudge_rows(env):
    from test_roles import _team
    _team(env)
    with closing(dbmod.connect()) as conn:
        for i, (area, subject) in enumerate((("log", "张工10-05的现场日志缺报"), ("gap", "A1010 关键作业无进度上报"), ("safety", "临边防护缺失"))):
            dbmod.insert(conn, "nudge_log", {"project_id": env["project_id"], "obj_type": area, "obj_key": str(i), "stage": 3, "to_user_id": env["users"]["eng"],
                                             "status": "sent", "subject": subject, "status_text": "已超期1天", "area": area, "sent_at": dbmod.now_iso()})


def test_nudge_rules_hide_log_source_from_project_manager(env):
    from conftest import login
    _nudge_rows(env)
    pm, eng, ss = login("pm"), login("eng"), login("ss")
    # D-7：项目经理不进催办台，规则读写一律 403（不只是隐藏日志/周确认来源）
    assert pm.get("/api/nudge/rules").status_code == 403
    assert pm.put("/api/nudge/rules", json={"daily_cap": 5, "sources": {"logs": True}}).status_code == 403
    # 服务层的岗位裁剪仍在：按项目经理口径取规则时不含日志/周确认来源
    with closing(dbmod.connect()) as conn:
        scoped = nudge_svc.scope_rules(nudge_svc.load_rules(conn, env["project_id"]), "manager")
    assert "logs" not in scoped["sources"] and "logs" not in nudge_svc.scope_rules(nudge_svc.DEFAULT_RULES, "manager")["sources"]
    got = ss.get("/api/nudge/rules").json()
    assert "ZB_SCHEDULER" not in str(got)
    assert got["scheduler"] is False and "系统管理员" in got["scheduler_note"]     # 测试环境后台调度关闭：用人话提示
    assert login("admin").put("/api/nudge/rules", json={"daily_cap": 5, "sources": {"logs": False}}).json()["rules"]["sources"]["logs"] is False
    full = eng.get("/api/nudge/rules").json()["rules"]
    assert full["sources"]["logs"] is False and full["daily_cap"] == 5               # 项目经理的保存被拒，系统管理员设的值不变（规则 10-08 起归管理员）


def test_nudge_log_and_today_hide_log_items_from_project_manager(env):
    from conftest import login
    _nudge_rows(env)
    pm, eng = login("pm"), login("eng")
    # D-7：项目经理不进催办台，记录与今日动态接口 403；服务层按项目经理口径仍只剩安全事项
    assert pm.get("/api/nudge/log").status_code == 403 and pm.get("/api/nudge/today").status_code == 403
    with closing(dbmod.connect()) as conn:
        assert [r["subject"] for r in nudge_svc.recent_log(conn, env["project_id"], 60, "manager")] == ["临边防护缺失"]
        t_pm = nudge_svc.today_summary(conn, env["project_id"], role="manager")
    assert len(eng.get("/api/nudge/log").json()) == 3
    t_eng = eng.get("/api/nudge/today").json()
    assert t_pm["objects"] == 1 and [e["subject"] for e in t_pm["escalated"]] == ["临边防护缺失"]
    assert t_eng["objects"] == 3
    rep = {"sent": [{"stage": 3, "items": [{"subject": "日志缺报", "area": "log"}, {"subject": "隐患", "area": "safety"}]},
                    {"stage": 3, "items": [{"subject": "缺口", "area": "gap"}]}],
           "deferred": [{"items": [{"subject": "日志", "area": "log"}]}], "skipped": [{"subject": "日志", "area": "log"}, {"subject": "隐患2", "area": "safety"}]}
    out = nudge_svc.scope_report(rep, "manager")
    assert [[i["subject"] for i in s["items"]] for s in out["sent"]] == [["隐患"]] and out["deferred"] == []
    assert [s["subject"] for s in out["skipped"]] == ["隐患2"]
    assert nudge_svc.scope_report(rep, "eng_sup") == rep
    assert pm.post("/api/nudge/run", json={"dry": True}).status_code == 403          # D-7


def test_nudge_front_end_has_no_env_var_and_guards_log_source():
    from pathlib import Path
    js = (Path(__file__).resolve().parent.parent / "app" / "static" / "x_nudge.js").read_text(encoding="utf-8")
    assert "ZB_SCHEDULER" not in js
    assert "v-if=\"'logs' in form.sources\"" in js


def test_admin_overview_shows_scheduler_detail(env):
    from conftest import login
    sysinfo = login("admin").get("/api/admin/overview").json()["system"]
    assert sysinfo["scheduler"]["on"] is False and "ZB_SCHEDULER" in sysinfo["scheduler"]["detail"]   # 技术细节只给系统管理员


# ----------------------------------------------------------------------------- 导航：项目经理有产值页（只读）、无应用成效页

def test_nav_project_manager_has_output_but_no_metrics(nav):  # noqa: F811
    pm = nav["roles"]["manager"]
    pages = [p for x in pm["top"] + pm["bottom"] for p in x["pages"]]
    assert "agents?tab=metrics" not in pages
    assert "agents?tab=agents" in pages
    prog = next(x for x in pm["top"] if x["k"] == "prog")
    assert prog["name"] == "进度与产值" and "output" in prog["pages"]   # RJ 10-10：项目经理只读产值
    for role in ("eng_sup", "leader"):
        line = next(x for x in nav["roles"][role]["top"] if x["k"] == "prog")
        assert line["name"] == "进度与产值" and "output" in line["pages"]          # 导航编排 10-08：各岗位同组同名，产值是“进度”组里的一页
    assert "agents?tab=metrics" in [p for x in nav["roles"]["eng_sup"]["top"] for p in x["pages"]]


def test_cockpit_api_strips_log_filling_for_project_manager(env):
    """项目总览接口：项目经理拿不到日志/周确认填报数据；产值只读照常给（RJ 10-10）；工程主管照常。"""
    from conftest import login
    _eng_sup(env)
    pm = login("pm").get("/api/cockpit").json()
    assert "log_trend" not in pm and not {"logs", "weeks"} & set(pm["kpis"])
    assert {"alerts", "safety", "quality", "output"} <= set(pm["kpis"]) and "output_series" in pm
    eng = login("eng").get("/api/cockpit").json()
    assert "output_series" in eng and "output" in eng["kpis"] and "log_trend" in eng


# ----------------------------------------------------------------------------- 复核问题回归：问答/队列权限、报表其他节的产值句、催办计数

def test_ask_agents_outside_pm_scope_denied_to_project_manager(env):
    from conftest import login

    from app.services import qa
    _eng_sup(env)
    pm, eng = login("pm"), login("eng")
    for ag in ("dq", "scribe", "nope"):
        assert pm.post("/api/ask", json={"question": "本周产值多少", "agent": ag}).status_code == 403, ag
    for ag in (None, "progress", "report", "analyst"):
        assert pm.post("/api/ask", json={"question": "进度怎么样", "agent": ag}).status_code == 200, ag
    # RJ 10-10：经理主动问产值可以答（产值 Agent 或问工笔里明确问产值都带产值事实）
    for ag in ("output", None):
        r = pm.post("/api/ask", json={"question": "本周产值多少", "agent": ag})
        assert r.status_code == 200 and not r.json().get("out_of_scope"), ag
    assert qa.can_ask({"role": "manager"}, "output") and not qa.can_ask({"role": "manager"}, "dq")
    assert eng.post("/api/ask", json={"question": "本周产值多少", "agent": "output"}).status_code == 200
    assert not {"output", "dq"} & set(qa.DEFAULT_SOURCES)   # 问工笔默认来源不含产值、日志预审


def test_report_facts_for_project_manager_strip_output(env):
    from app.services import qa
    _eng_sup(env)
    _pm, rep = _submitted_report(env)
    with closing(dbmod.connect()) as conn:
        nar = dbmod.loads(conn.execute("SELECT narrative_json FROM reports WHERE id=?", (rep["id"],)).fetchone()["narrative_json"], {})
        nar["output"] = "本周产值 40.3万 USD，高于计划"
        nar["overview"] = "本周完成桩基2根。本周产值约40万；路基开挖继续。"
        conn.execute("UPDATE reports SET narrative_json=? WHERE id=?", (dbmod.dumps(nar), rep["id"]))
        pm_text = "\n".join(qa._report_facts(conn, env["project_id"], "manager", "进度怎么样"))
        pm_ask = "\n".join(qa._report_facts(conn, env["project_id"], "manager", "上周产值多少"))
        eng_text = "\n".join(qa._report_facts(conn, env["project_id"], "eng_sup"))
    assert "产值" not in pm_text and "桩基" in pm_text and "路基" in pm_text   # 没问产值时不主动带
    assert "40.3万" in pm_ask                                                 # 经理主动问产值照常引用
    assert "产值" in eng_text


def test_cockpit_kpis_strip_log_for_project_manager(env):
    """原 /api/queue 的 KPI 口径（已下线，前端不再调用）：同一口径现由 /api/cockpit 承担。"""
    from conftest import login
    _eng_sup(env)
    assert login("pm").get("/api/queue").status_code == 404
    k = login("pm").get("/api/cockpit").json()["kpis"]
    assert not {"logs", "weeks"} & set(k) and {"alerts", "safety", "quality"} <= set(k)
    assert "output" in login("eng").get("/api/cockpit").json()["kpis"]


def test_report_output_sentences_dropped_only_for_qa(env):
    """问答引用报表时，项目经理口径仍剔除产值节与其他节的产值句；报表本身对项目经理只读显示全文。"""
    from conftest import login

    from app.ai import report_agent
    from app.services import reports
    tpl = report_agent.default_template("week")
    assert "产值只写" in report_agent.system_prompt(tpl)
    nar = {"overview": "本周完成桩基2根。本周产值约40万；路基开挖继续。", "output": "产值 40万",
           "coordination": ["协调业主产值计量", "协调征地"]}
    _f, _t, out, _c = reports.pm_scope({"progress": []}, tpl, nar, [], drop_output=True)
    assert out["overview"] == "本周完成桩基2根。路基开挖继续。" and out["coordination"] == ["协调征地"] and "output" not in out
    _f, _t, out, _c = reports.pm_scope({"progress": []}, tpl, nar, [])
    assert out == nar
    _eng_sup(env)
    pm, rep = _submitted_report(env)
    with closing(dbmod.connect()) as conn:
        full = dbmod.loads(conn.execute("SELECT narrative_json FROM reports WHERE id=?", (rep["id"],)).fetchone()["narrative_json"], {})
        full["overview"] = "本周完成桩基2根。本周产值约40万；路基开挖继续。"
        conn.execute("UPDATE reports SET narrative_json=? WHERE id=?", (dbmod.dumps(full), rep["id"]))
    v = pm.get(f"/api/reports/{rep['id']}").json()
    assert "本周产值约40万" in v["narrative"]["overview"]
    assert pm.put(f"/api/reports/{rep['id']}", json={"narrative": v["narrative"], "status": "draft"}).status_code == 200
    assert "本周产值约40万" in login("eng").get(f"/api/reports/{rep['id']}").json()["narrative"]["overview"]


def test_nudge_scope_report_recounts_summary_for_project_manager():
    rep = {"dry": True, "targets": 3, "areas": {"log": 1, "gap": 1, "safety": 1},
           "reasons": {"waiting": 2, "quiet": 1}, "reasons_by_area": {"log": {"waiting": 1}, "gap": {"quiet": 1}, "safety": {"waiting": 1}},
           "sent": [], "deferred": [{"items": [{"subject": "缺口", "area": "gap"}], "why": "静默", "code": "quiet"}],
           "skipped": [{"subject": "日志", "area": "log", "code": "waiting"}, {"subject": "隐患", "area": "safety", "code": "waiting"}],
           "summary": "预览：检查到 3 件待跟进事项，将发出 0 条提醒（0 人）；暂不发送：1件在静默时段延后、2件未到升级时间。"}
    out = nudge_svc.scope_report(rep, "manager")
    assert out["targets"] == 1 and out["reasons"] == {"waiting": 1}
    assert "检查到 1 件" in out["summary"] and "静默" not in out["summary"] and "1件未到升级时间" in out["summary"]
    assert set(out["areas"]) == {"safety"} and set(out["reasons_by_area"]) == {"safety"}
    assert nudge_svc.scope_report(rep, "eng_sup") == rep


def test_nudge_run_by_project_manager_skips_log_targets(env, monkeypatch):
    from datetime import date, datetime

    from test_roles import _team
    _team(env)
    tech_id = env["users"]["tech1"]
    today = date(2026, 10, 6)

    def fake_collect(conn, pid, rules, day, users):
        return [nudge_svc._target("dailylog", f"{tech_id}:2026-10-05", "log", "张工10-05的现场日志缺报", today, today, tech_id, "home"),
                nudge_svc._target("datagap", "A1010:x", "gap", "A1010 关键作业无进度上报", today, today, tech_id, "home"),
                nudge_svc._target("finding", "99", "safety", "临边防护缺失", today, today, tech_id, "finding?id=99", 99)], []
    monkeypatch.setattr(nudge_svc, "collect_targets", fake_collect)
    now = datetime(2026, 10, 6, 10, 0)
    with closing(dbmod.connect()) as conn:
        pm = dict(conn.execute("SELECT * FROM users WHERE username='pm'").fetchone())
        rep = nudge_svc.run_project(conn, env["project_id"], now=now, actor=pm)
        areas = [r["area"] for r in conn.execute("SELECT area FROM nudge_log WHERE project_id=?", (env["project_id"],)).fetchall()]
    assert rep["targets"] == 1 and "检查到 1 件" in rep["summary"]
    assert areas == ["safety"]   # 项目经理“立即运行”不以其名义发日志缺报/数据缺口催办
    with closing(dbmod.connect()) as conn:
        eng = dict(conn.execute("SELECT * FROM users WHERE username='eng'").fetchone())
        rep = nudge_svc.run_project(conn, env["project_id"], now=now, actor=eng)
        areas = sorted(r["area"] for r in conn.execute("SELECT area FROM nudge_log WHERE project_id=?", (env["project_id"],)).fetchall())
    assert rep["targets"] == 3 and areas == ["gap", "log", "safety"]


def test_leader_sees_published_reports_only(env):
    """领导只读：周月报列表与详情只给已发布的，草稿不给（草稿还在审核修改中）。"""
    from conftest import PW, login
    from app.security import hash_password
    pm, rep = _submitted_report(env)
    with closing(dbmod.connect()) as conn:
        dbmod.insert(conn, "users", {"username": "boss", "name": "领导", "role": "leader", "password_hash": hash_password(PW), "project_id": env["project_id"],
                                     "wbs_scope": "", "email": "", "active": 1, "created_at": dbmod.now_iso()})
    boss = login("boss")
    assert rep["id"] not in [r["id"] for r in boss.get("/api/reports").json()]
    assert boss.get(f"/api/reports/{rep['id']}").status_code == 404
    assert rep["id"] in [r["id"] for r in pm.get("/api/reports").json()]           # 项目经理照常看到草稿去审核
    pm.put(f"/api/reports/{rep['id']}", json={"narrative": pm.get(f"/api/reports/{rep['id']}").json()["narrative"], "status": "published"})
    assert rep["id"] in [r["id"] for r in boss.get("/api/reports").json()]
    assert boss.get(f"/api/reports/{rep['id']}").status_code == 200


# ----------------------------------------------------------------------------- T-06 待办收敛：项目经理待办的白名单与反向白名单（AC-6.1、6.3）

PM_EXEC_SOURCES = {"candidate", "gate", "gate_more", "quick", "ticket_risk", "ticket_pending", "conceal_gap", "conceal_todo", "conceal_result", "alert",
                   "log_review", "log_low", "log_missing", "week_missing", "data_gap", "mapping"}


def test_pm_forbidden_and_allowed_sources():
    from app.services import todo as todo_svc
    assert PM_EXEC_SOURCES <= todo_svc.PM_FORBIDDEN                       # 反向白名单只增不减
    assert todo_svc.PM_ALLOWED == {"crash", "escalated", "report", "meeting", "claim_notice"}   # claim_notice：索赔通知时效剩 3 天升级（用户 10-08 定）
    assert not todo_svc.PM_ALLOWED & todo_svc.PM_FORBIDDEN
    assert "alert" not in todo_svc.PM_ALLOWED and "crash" != "alert"         # 赶工决定用新来源名，不与预警反馈重名


def test_pm_todo_guard_drops_forbidden_unknown_and_out_of_scope_text(env, monkeypatch):
    """兜底守卫：即便来源挂错，项目经理待办也只留白名单来源（含“另有 N 项”行），文字再过一遍 scope_ok。"""
    from conftest import login
    from app.services import todo as todo_svc
    _eng_sup(env)

    def fake(c):
        mk = lambda s, title: todo_svc.item(s, title, "progress", title, "原因", "状态")  # noqa: E731
        return [mk("crash", "A1020 赶工决定"), mk("crash", "填报率不足的赶工"), mk("candidate_more", "另有 3 项候选"), mk("gate", "放行缺项"),
                mk("brand_new", "未知来源"), mk("report_more", "另有 1 份报表"), mk("alert", "预警属实反馈")]
    monkeypatch.setitem(todo_svc.SOURCES, "manager", [fake])
    d = login("pm").get("/api/todo").json()
    assert sorted(i["title"] for i in d["items"]) == ["A1020 赶工决定", "另有 1 份报表"]
    # 工程主管不受这道守卫影响
    monkeypatch.setitem(todo_svc.SOURCES, "eng_sup", [fake])
    assert len(login("eng").get("/api/todo").json()["items"]) == 7


def test_pm_need_to_know_line_in_scope(env):
    from conftest import login
    from test_facts import _add_alert, _setup_site
    _setup_site(env)
    _add_alert(env["project_id"])
    d = login("pm").get("/api/todo").json()
    k = d["know"]
    assert k["n"] >= 1 and all(brief.scope_ok(p["label"]) for p in k["parts"])
    assert not any(w in p["label"] for p in k["parts"] for w in ("填报", "日志", "缺报", "产值", "记录率"))


def test_agent_cards_for_project_manager_do_not_mention_output(env):
    from conftest import login
    _eng_sup(env)
    cards = login("pm").get("/api/agents").json()
    cards = cards.get("agents", cards) if isinstance(cards, dict) else cards
    assert cards and not any("产值" in (a.get("inputs", "") + a.get("outputs", "") + a.get("duty", "")) for a in cards)


# ----------------------------------------------------------------------------- F-14：项目经理不审核日志（后端权限与数据口径）

LOG_REVIEW_KEYS = ("dq", "checks", "status", "reviewer", "reviewed_at", "review_note", "edit_seconds", "edit_ratio")  # 后两项属应用成效


def _submitted_log(env):
    from conftest import backdate, login
    from test_flow import _draft_today, _import
    _eng_sup(env)
    _import(login("admin"))
    tech = login("tech1")
    log = _draft_today(tech)
    backdate(log["id"])
    s = tech.post(f"/api/logs/{log['id']}/submit", json={"version": log["version"], "active_seconds": 300})
    assert s.status_code == 200, s.text
    return s.json()


def test_log_review_cap_excludes_project_manager():
    from app import roles as R
    assert not R.can("manager", "log_review") and R.CAPS["log_review"] == {"admin", "eng_sup"}
    assert R.can("manager", "report_publish") and R.can("manager", "report_gen")  # 周报审核、发布不受影响
    assert "log_review" not in R.caps_of("manager")


def test_project_manager_cannot_review_logs(env):
    from conftest import login
    lg = _submitted_log(env)
    pm = login("pm")
    assert pm.post(f"/api/logs/{lg['id']}/review", json={"action": "approve"}).status_code == 403
    assert pm.post(f"/api/logs/{lg['id']}/review", json={"action": "return", "note": "不符"}).status_code == 403
    assert pm.post("/api/logs/batch-approve", json={"ids": [lg["id"]]}).status_code == 403
    assert login("eng").post(f"/api/logs/{lg['id']}/review", json={"action": "approve"}).json()["status"] == "approved"


def test_log_list_denied_and_detail_stripped_for_project_manager(env):
    import json

    from conftest import login
    lg = _submitted_log(env)
    pm, eng = login("pm"), login("eng")
    assert pm.get("/api/logs").status_code == 403
    assert pm.get("/api/logs?period=week").status_code == 403
    assert eng.get("/api/logs").status_code == 200 and login("admin").get("/api/logs").status_code == 200
    d = pm.get(f"/api/logs/{lg['id']}")
    assert d.status_code == 200
    v = d.json()
    assert not set(LOG_REVIEW_KEYS) & set(v), set(LOG_REVIEW_KEYS) & set(v)
    assert v["dq_review"] is None and v["days"] is None and v["progress_rows"]["voided"] == []
    assert "数据审核" not in json.dumps(v, ensure_ascii=False) and "预审" not in json.dumps(v, ensure_ascii=False)
    assert v["content"] and v["photos"]  # 施工内容（进度、质量、安全的原始记录）照常可看
    full = eng.get(f"/api/logs/{lg['id']}").json()
    assert full["dq"] and "checks" in full and full["status"] == "submitted" and full["edit_seconds"] == 300   # 工程主管看完整版


def test_circle_pick_reason_with_data_review_hidden_from_project_manager(env):
    from conftest import login
    lg = _submitted_log(env)
    key = f"log:{lg['id']}"
    with closing(dbmod.connect()) as conn:  # 库里已有的旧理由（规则③改文案之前生成的）
        dbmod.insert(conn, "circle_picks", {"project_id": env["project_id"], "post_key": key, "reason": "周进度确认已提交，数据审核低风险（评分 97）",
                                            "source": "rule", "ai_run_id": None, "created_at": dbmod.now_iso()})
    pm, eng = login("pm"), login("eng")
    for c, seen in ((pm, False), (eng, True)):
        feed = c.get("/api/feed?limit=50").json()["posts"]
        post = next(p for p in feed if p["key"] == key)
        assert bool(post["ai_pick"]) is seen
        assert (key in [p["key"] for p in c.get("/api/feed?view=ai&limit=50").json()["posts"]]) is seen
    assert "数据审核" not in str(pm.get("/api/feed?limit=50").json())



def test_project_manager_out_of_scope_question_refused_without_ai(env, monkeypatch):
    from conftest import login

    from app.services import qa
    _eng_sup(env)

    def boom(*a, **k):
        raise AssertionError("越界问题不得调用 AI")
    monkeypatch.setattr(qa, "call_json", boom)
    pm = login("pm")
    for q in ("各分包本周日志填报率是多少，谁没有记录？", "哪些日志数据审核没通过", "周确认缺报的有谁"):
        for ag in (None, "progress"):
            r = pm.post("/api/ask", json={"question": q, "agent": ag})
            assert r.status_code == 200, r.text
            a = r.json()
            assert a["answer"].startswith("不在项目经理的查看范围") and a["refs"] == [] and a["facts_count"] == 0
            assert brief.scope_ok(a["answer"])  # 拒答本身也不复述越界词
    monkeypatch.undo()
    assert pm.post("/api/ask", json={"question": "A1010 进度怎么样"}).json()["answer"] != ""
    assert not pm.post("/api/ask", json={"question": "本月产值完成多少"}).json()["answer"].startswith("不在项目经理")   # 主动问产值可以答（RJ 10-10）
    assert brief.ask_scope_ok("本月产值完成多少") and not brief.scope_ok("本月产值完成多少")   # 简报/待办仍不推产值
    assert not login("eng").post("/api/ask", json={"question": "本周日志填报率是多少"}).json()["answer"].startswith("不在项目经理")


# ----------------------------------------------------------------------------- N-2：现场圈与问答里的整改复核只写成“AI 建议”

def test_rectify_verify_shown_as_ai_suggestion_in_feed_and_qa(env):
    from conftest import confirm_hazards, jpeg, login

    from app.services import qa
    _eng_sup(env)
    tech, pm = login("tech1"), login("pm")
    tech.post("/api/media", data={"purpose": "safety"}, files=[("files", ("hz.jpg", jpeg((200, 50, 50)), "image/jpeg"))])
    confirm_hazards()
    fid = tech.get("/api/findings?agent=safety&status=active&mine=1").json()[0]["id"]
    mine = tech.post("/api/media", data={"purpose": "rectify"}, files=[("files", ("r.jpg", jpeg((20, 210, 20)), "image/jpeg"))]).json()[0]
    assert tech.post(f"/api/findings/{fid}/rectify", json={"media_id": mine["id"]}).json()["status"] == "rectifying"
    for c in (pm, login("eng")):
        rect = next(p for p in c.get("/api/feed?limit=50").json()["posts"] if p["kind"] == "rect")
        assert rect["ai"][0]["reason"].startswith("AI 建议：已消除，待安全主管确认"), rect["ai"][0]["reason"]
    with closing(dbmod.connect()) as conn:
        facts = qa._findings_facts(conn, env["project_id"], "safety", [])
    line = next(f for f in facts if f.startswith(f"[F{fid}]"))
    assert "AI 建议：已消除，待安全主管确认" in line


# ----------------------------------------------------------------------------- F-14：项目经理日志详情页在缺字段时不报错、不显示空标签与“规则校验通过”

def test_log_editor_tolerates_pm_view_without_review_fields():
    import json as _json
    import re

    from test_frontline_f import read, run_node
    js = read("views.js")
    editor = js[js.index("const LogEditor"):js.index("const LogEditor") + 20000]
    flags = re.search(r"\n    flags\(\) \{ (return .*?;) \},", editor).group(1)
    row_flag = re.search(r"\n    rowFlag\(sec, i\) \{\n(.*?)\n    \},", editor, re.S).group(1)
    pm_log = {"id": 1, "has_ai_draft": True, "mode": "ai", "content": {}}   # 项目经理拿到的日志详情：无 checks/status/edit_*
    out = run_node(f"""
const log = {_json.dumps(pm_log)};
const flags = function () {{ {flags} }}.call({{ log }});
const rowFlag = function (sec, i) {{ {row_flag} }}.call({{ log }}, "works", 0);
console.log(JSON.stringify({{ flags, rowFlag }}));""")
    assert out == {"flags": [], "rowFlag": ""}
    tpl = editor[:editor.index("data: ()")]
    assert "v-if=\"log.checks && ((log.has_ai_draft || log.imported) || log.status !== 'draft')\"" in tpl   # 没有校验数据就不显示“规则校验通过”
    assert re.search(r"<span v-if=\"log\.status\" :class=\"\['chip', STATUS_CHIP\[log\.status\]\]\"", tpl)   # 无状态不出空标签
    assert "<span v-if=\"log.edit_ratio !== null && log.edit_ratio !== undefined\"> · 相对AI草稿修改率" in tpl
    assert "<span v-if=\"'edit_seconds' in log\"> · 编写用时" in tpl


def test_pm_overview_apis_drop_log_review_and_fill_rate(env):
    """#/eng 与分包一览的接口：项目经理拿不到日志预审队列、今日谁没记录、谁没周确认、填报数；进度、质量、产值照常返回。"""
    from conftest import login
    from test_roles import _team
    _team(env)
    eng, pm = login("eng"), login("pm")
    e_eng, e_pm = eng.get("/api/eng/overview").json(), pm.get("/api/eng/overview").json()
    assert e_eng["review"] is not None and "missing_today" in e_eng
    assert e_pm["review"] is None and e_pm["missing_today"] == [] and e_pm["week_unconfirmed"] == [] and e_pm["gaps"] == []
    assert e_pm["quality"] == e_eng["quality"] and e_pm["mapping"] == e_eng["mapping"]
    p_eng, p_pm = eng.get("/api/production/overview").json(), pm.get("/api/production/overview").json()
    assert "logs" in p_eng["kpis"] and "weeks" in p_eng["kpis"]
    assert "logs" not in p_pm["kpis"] and "weeks" not in p_pm["kpis"] and "output" in p_pm["kpis"]
    for rows in (p_pm["companies"], pm.get("/api/companies").json()):
        assert rows and all("record_today" not in c and "week_confirmed" not in c for c in rows)
        assert all("safety_open" in c for c in rows)
    assert all("record_today" in c for c in eng.get("/api/companies").json())


def test_pm_eng_page_hides_record_column():
    from pathlib import Path
    st = Path(__file__).resolve().parent.parent / "app" / "static"
    # RJ 10-10：各分包今日记录 / 本周确认表从工程主管首页移到“分包履约画像”页（x_subprofile.js 的 SubToday），只给工程主管；
    # 口径不变：接口不给 record_today（项目经理）时“今日记录”整列不出现
    sub = (st / "x_subprofile.js").read_text(encoding="utf-8")
    seg = sub[sub.index("const SubToday"):sub.index("const SubProfilePage")]
    assert '<th v-if="fill">今日记录</th>' in seg and 'c.record_today !== undefined' in seg and 'roleOf() === "eng_sup"' in seg
    assert "<sub-today></sub-today>" in sub
    assert "<th v-if=\"fill\">今日记录</th>" not in (st / "x_enghome.js").read_text(encoding="utf-8")

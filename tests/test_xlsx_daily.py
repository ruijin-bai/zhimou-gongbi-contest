"""Excel 日报导入（RJ 2026-10-09 意见 7）：按“生产日报/Daily Report”模板识别 → 预览 → 技术员逐行决定后生成当天日志草稿。

- 识别是规则解析：表头标签定位、合并单元格续行沿用上一行的内容和单位、“\\”和空白记为未填（None，不当 0）、文字数量写进 warnings；
- 只有技术员能导入；每行必须决定（计入工程量 / 只记施工内容 / 不导入），计入工程量必须选 P6 作业且有数量；
- 生成的是草稿：不提交、不改进度；并入当天已有草稿时同一作业的工程量以 Excel 为准；后台整理随手记不覆盖导入的内容；
- 已提交的那天不能导入；同一份导入只能确认一次；旧版 .xls、非 Excel 文件当场拒收。
样例工作簿是合成的（结构仿项目模板，数字为虚构），不含真实项目数据。
"""
import io
from datetime import timedelta

import openpyxl
import pytest

from app import db as dbmod
from app.services import xlsx_daily as x
from conftest import TODAY, login, sample_xer

DAY = (TODAY - timedelta(days=1)).isoformat()


def book(day=DAY, daily_bad=False) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    d = day.replace("-", ".")
    rows = [
        ["生产日报/Daily Report"],
        ["单位名称 /Section Name", "合成测试分部", None, "施工日期 /Construction date", d, "填报人 Reporter", "测试甲", "复核人 Reviewer", "测试乙"],
        ["天气情况 /Weather", "晴", None, "中方人员数量 /Number of Chinese personnel", 12, "属地员工数量 /Number of Nigerian employees", 30,
         "进场设备台（套）/Number of Equipment Onsite", "9台(挖掘机2台、装载机1台、自卸车6台)"],
        [],
        ["当日施工内容简述（形象进度）/Overview of construction content on"],
        ["标段A:"],
        ["1、3#墩桩基钻孔 2 根"],
        ["2、路基填筑 CH1+000-CH1+200 填方 300m³"],
        ["半岛段："],
        ["无"],
        ["现场进度情况/On site Progress"],
        ["序号/No", "施工区域/Construction Area", "施工内容/Construction content", "单位/unit", "设计总工程量/Total design quantity",
         "完成桩号段落/位置Complete Chainage /Location", "当日完成数量/Daily completed quantity", "累计完成数量/Accumulated completed quantity",
         "累计完成占比/Accumulated completion percentage"],
        [1, "标段A", "桩基钻孔", "根", 16, "3#墩", "约2" if daily_bad else 2, 10, 0.625],
        [2, None, "路基填筑/ subgrade", "m³", "\\", "CH1+000-CH1+200", 300, 5200, "\\"],
        [None, None, None, None, None, "CH1+300-CH1+350", None, 800, None],
        [3, "半岛段", "清表", "㎡", "\\", "", "\\", 4000, "\\"],
        ["明日施工计划/Tomorrow's Construction Plan"],
        ["1、3#墩桩基钻孔"],
        ["2、路基填筑"],
        ["现场形象进度照片On-site Progress Photograph"],
    ]
    for r in rows:
        ws.append(r)
    ws.merge_cells("B14:B15")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _q(sql, params=()):
    conn = dbmod.connect()
    try:
        return conn.execute(sql, params).fetchall()
    finally:
        conn.close()


def _p6():
    assert login("pm").post("/api/p6/import", files={"file": ("plan.xer", sample_xer(), "application/octet-stream")}).status_code == 200


def _upload(c, raw=None, name="生产日报.xlsx"):
    return c.post("/api/logs/xlsx", files={"file": (name, raw if raw is not None else book(), "application/octet-stream")})


# ============================================================================ 识别

def test_parse_template_fields_rows_and_unknowns():
    p = x.parse(book(), "生产日报.xlsx")
    assert (p["date"], p["section"], p["weather"], p["reporter"], p["reviewer"]) == (DAY, "合成测试分部", "晴", "测试甲", "测试乙")
    assert (p["staff_cn"], p["staff_local"], p["equipment_count"]) == (12, 30, 9)
    assert p["machines"] == [{"name": "挖掘机", "count": 2}, {"name": "装载机", "count": 1}, {"name": "自卸车", "count": 6}]
    assert p["overview"] == [{"group": "标段A", "text": "3#墩桩基钻孔 2 根"}, {"group": "标段A", "text": "路基填筑 CH1+000-CH1+200 填方 300m³"}]
    rows = p["rows"]
    assert [(r["area"], r["content"], r["unit"], r["daily"], r["cum"]) for r in rows] == [
        ("标段A", "桩基钻孔", "根", 2.0, 10.0),
        ("标段A", "路基填筑/ subgrade", "m³", 300.0, 5200.0),
        ("标段A", "路基填筑/ subgrade", "m³", None, 800.0),   # 合并单元格续行：内容、单位沿用上一行；当日数量空白 → 未填，不当 0
        ("半岛段", "清表", "㎡", None, 4000.0),               # “\\” → 未填
    ]
    assert rows[0]["pct_text"] == "62.5%" and rows[1]["pct_text"] == "" and rows[1]["design"] is None
    assert [t["text"] for t in p["tomorrow"]] == ["3#墩桩基钻孔", "路基填筑"]
    assert p["warnings"] == []


def test_parse_reports_text_quantity_and_missing_date():
    p = x.parse(book(daily_bad=True), "a.xlsx")
    assert p["rows"][0]["daily"] is None and any("约2" in w for w in p["warnings"])
    wb = openpyxl.load_workbook(io.BytesIO(book()))
    wb.active["E2"] = "某天"
    buf = io.BytesIO()
    wb.save(buf)
    p2 = x.parse(buf.getvalue(), "无日期.xlsx")
    assert p2["date"] is None and any("施工日期" in w for w in p2["warnings"])
    assert x.parse(buf.getvalue(), "生产日报（2026.9.17）.xlsx")["date"] == "2026-09-17"  # 表里没写，文件名里有完整日期


def test_parse_rejects_xls_and_non_excel():
    with pytest.raises(x.XlsxError, match="另存为 .xlsx"):
        x.parse(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100)
    with pytest.raises(x.XlsxError, match="只支持"):
        x.parse(b"hello")
    with pytest.raises(x.XlsxError, match="打不开"):
        x.parse(b"PK\x03\x04broken")


def test_candidates_are_suggestions_by_chainage_and_keywords():
    acts = [{"code": "C1", "name": "Subgrade Fill - CH0+950-CH1+250"}, {"code": "C2", "name": "Subgrade Fill - CH3+000-CH4+000"},
            {"code": "C3", "name": "Bored Pile Construction"}, {"code": "C4", "name": "3#墩桩基施工"}]
    road = {"area": "标段A", "content": "填砂路基/ subgrade", "location": "CH1+000-CH1+200"}
    assert [c["code"] for c in x.candidates(acts, road)][:2] == ["C1", "C2"]  # 桩号重叠的排前面
    pile = {"area": "", "content": "桩基钻孔", "location": "3#墩"}
    assert {c["code"] for c in x.candidates(acts, pile)} >= {"C3", "C4"}
    assert x.candidates(acts, {"area": "", "content": "营地板房", "location": ""}) == []


# ============================================================================ 接口：预览、确认生成草稿

def test_only_tech_can_import(env):
    assert _upload(login("pm")).status_code == 403
    r = _upload(login("tech1"), b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "old.xls")
    assert r.status_code == 400 and "另存为" in r.json()["detail"]


def test_preview_then_confirm_creates_draft_log(env):
    _p6()
    t1 = login("tech1")
    raw = book()  # openpyxl 每次保存都写当前时间，重传“同一文件”要用同一份字节
    pv = _upload(t1, raw)
    assert pv.status_code == 200, pv.text
    d = pv.json()
    assert d["parsed"]["date"] == DAY and d["day_log"] is None and d["imported_before"] is None
    assert "A1010" in [c["code"] for c in d["parsed"]["rows"][0]["candidates"]]  # 3#墩桩基施工
    assert _q("SELECT COUNT(*) n FROM logs")[0]["n"] == 0  # 预览不生成任何日志
    iid = d["import_id"]
    # 没决定完不能生成
    r = t1.post(f"/api/logs/xlsx/{iid}/confirm", json={"log_date": DAY, "decisions": [{"i": 0, "action": "skip"}]})
    assert r.status_code == 400 and "还有 3 行" in r.json()["detail"]
    # 计入工程量必须选作业、必须有数量
    dec = [{"i": 0, "action": "p6"}, {"i": 1, "action": "text"}, {"i": 2, "action": "skip"}, {"i": 3, "action": "skip"}]
    assert "请选择 P6 作业" in t1.post(f"/api/logs/xlsx/{iid}/confirm", json={"log_date": DAY, "decisions": dec}).json()["detail"]
    dec[0]["activity_code"] = "A1010"
    dec[2] = {"i": 2, "action": "p6", "activity_code": "A3020"}
    assert "没有当日完成数量" in t1.post(f"/api/logs/xlsx/{iid}/confirm", json={"log_date": DAY, "decisions": dec}).json()["detail"]
    dec[2] = {"i": 2, "action": "skip"}
    dec[0]["qty"] = 1.5  # 人改了数量，以人为准
    r = t1.post(f"/api/logs/xlsx/{iid}/confirm", json={"log_date": DAY, "decisions": dec})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["log_date"] == DAY and res["progress"] == 1
    lg = _q("SELECT * FROM logs WHERE id=?", (res["log_id"],))[0]
    assert lg["status"] == "draft" and lg["period"] == "day" and lg["weather"] == "晴" and lg["submitted_at"] is None
    c = dbmod.loads(lg["content_json"], {})
    assert c["progress"] == [{"activity_code": "A1010", "qty": 1.5, "unit": "根", "pct": None,
                              "note": "Excel：标段A 桩基钻孔（3#墩），累计 10根，累计占比 62.5%", "sources": ["X"]}]
    descs = [w["desc"] for w in c["works"]]
    assert "3#墩桩基钻孔 2 根" in descs and any("路基填筑" in s and "当日完成 300m³" in s for s in descs)
    assert not any("清表" in s for s in descs)  # 不导入的行不写
    assert c["labor"] == [{"trade": "中方人员", "count": 12, "sources": ["X"]}, {"trade": "属地员工", "count": 30, "sources": ["X"]}]
    assert [m["name"] for m in c["machinery"]] == ["挖掘机", "装载机", "自卸车"]
    assert [t["desc"] for t in c["tomorrow"]] == ["3#墩桩基钻孔", "路基填筑"]
    assert "从 Excel 日报导入" in c["notes"]
    assert not any("不是来自口述" in k["message"] for k in dbmod.loads(lg["checks_json"], []))  # Excel 数量是技术员确认过的
    assert t1.get(f"/api/logs/{res['log_id']}").json()["imported"] is True
    # 进度没有被改：日志未提交，progress_state 不计草稿
    assert _q("SELECT COUNT(*) n FROM logs WHERE status IN ('submitted','approved')")[0]["n"] == 0
    # 同一份导入只能确认一次；再传同一文件会提示导入过
    assert t1.post(f"/api/logs/xlsx/{iid}/confirm", json={"log_date": DAY, "decisions": dec}).status_code == 409
    again = _upload(t1, raw).json()
    assert again["imported_before"]["log_id"] == res["log_id"] and again["day_log"]["id"] == res["log_id"]


def test_keep_switches_and_merge_into_existing_draft(env):
    _p6()
    t1 = login("tech1")
    first = _upload(t1).json()
    dec = [{"i": 0, "action": "p6", "activity_code": "A1010"}, {"i": 1, "action": "skip"}, {"i": 2, "action": "skip"}, {"i": 3, "action": "skip"}]
    log_id = t1.post(f"/api/logs/xlsx/{first['import_id']}/confirm", json={"log_date": DAY, "decisions": dec}).json()["log_id"]
    second = _upload(t1).json()
    dec[0]["qty"] = 3
    r = t1.post(f"/api/logs/xlsx/{second['import_id']}/confirm",
                json={"log_date": DAY, "decisions": dec, "keep": {"overview": False, "tomorrow": False, "staff": False, "machinery": False}})
    assert r.json()["log_id"] == log_id
    c = dbmod.loads(_q("SELECT content_json FROM logs WHERE id=?", (log_id,))[0]["content_json"], {})
    assert [(p["activity_code"], p["qty"]) for p in c["progress"]] == [("A1010", 3.0)]  # 同一作业以后一次为准，不重复
    assert len(c["tomorrow"]) == 2  # 第二次没带明日计划，第一次的还在


def test_auto_organize_does_not_overwrite_imported_log(env):
    _p6()
    t1 = login("tech1")
    d = _upload(t1, book(day=TODAY.isoformat())).json()
    dec = [{"i": 0, "action": "p6", "activity_code": "A1010"}, {"i": 1, "action": "text"}, {"i": 2, "action": "skip"}, {"i": 3, "action": "skip"}]
    log_id = t1.post(f"/api/logs/xlsx/{d['import_id']}/confirm", json={"log_date": TODAY.isoformat(), "decisions": dec}).json()["log_id"]
    before = _q("SELECT content_json FROM logs WHERE id=?", (log_id,))[0]["content_json"]
    n = t1.post("/api/notes", json={"text": "3#墩桩基施工今天完成1根"}).json()  # 后台整理当天日志
    assert _q("SELECT content_json FROM logs WHERE id=?", (log_id,))[0]["content_json"] == before
    assert t1.delete(f"/api/notes/{n['id']}").status_code == 200  # 删随手记也不清空导入的内容
    assert _q("SELECT content_json FROM logs WHERE id=?", (log_id,))[0]["content_json"] == before


def test_cannot_import_into_submitted_day_or_future(env):
    _p6()
    t1 = login("tech1")
    d = _upload(t1).json()
    dec = [{"i": i, "action": "skip"} for i in range(4)]
    future = (TODAY + timedelta(days=1)).isoformat()
    assert "晚于今天" in t1.post(f"/api/logs/xlsx/{d['import_id']}/confirm", json={"log_date": future, "decisions": dec}).json()["detail"]
    conn = dbmod.connect()
    pid, uid = env["project_id"], env["users"]["tech1"]
    dbmod.insert(conn, "logs", {"project_id": pid, "author_id": uid, "log_date": DAY, "mode": "ai", "status": "submitted", "content_json": "{}",
                                "started_at": dbmod.now_iso(), "period": "day"})
    conn.close()
    r = t1.post(f"/api/logs/xlsx/{d['import_id']}/confirm", json={"log_date": DAY, "decisions": dec})
    assert r.status_code == 400 and "已提交" in r.json()["detail"]
    assert login("tech2").get(f"/api/logs/xlsx/{d['import_id']}").status_code == 404  # 别人的导入看不到


def test_parse_same_content_rows_inherit_missing_unit():
    """现行模板（四航局 2026-10-08）里同一施工内容分几行按桩号写，后面的行有时不填单位：沿用上一行的单位，不留空。"""
    wb = openpyxl.load_workbook(io.BytesIO(book()))
    ws = wb.active
    ws.insert_rows(17)
    for col, v in zip("ABCDEFGHI", [4, None, "清表", None, "\\", "CH2+000-CH2+050", 100, 4100, "\\"]):
        ws[f"{col}17"] = v
    buf = io.BytesIO()
    wb.save(buf)
    rows = x.parse(buf.getvalue(), "a.xlsx")["rows"]
    assert [(r["content"], r["unit"], r["daily"]) for r in rows[-2:]] == [("清表", "㎡", None), ("清表", "㎡", 100.0)]


def test_parse_rejects_other_reports_in_the_same_package():
    """周报包里同时有设备监控等其他 Excel：没有简述、进度表、明日计划的不当成日报。"""
    wb = openpyxl.Workbook()
    wb.active.append(["设备监控表 Equipment Monitoring"])
    wb.active.append(["序号", "设备名称", "台数", "状态"])
    wb.active.append([1, "挖掘机", 3, "正常"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(x.XlsxError, match="没识别出生产日报"):
        x.parse(buf.getvalue(), "Equipment_Monitoring.xlsx")

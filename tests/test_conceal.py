"""资料Agent：隐蔽工程验收自动成册。模板匹配、mock 识别 → 缺项规则、状态流转与权限、缺资料扫描、Word 成册、AI 与人工值分开。"""
import io
from datetime import timedelta
from pathlib import Path

from conftest import TODAY, jpeg, login, sample_xer
from test_roles import _add, _team

from app import db as dbmod
from app.ai import conceal_agent
from app.ai.client import AIError, AIResult
from app.services import conceal as cs

NARR = "3#墩12号桩，成孔深度28.5米，孔径1500毫米，沉渣厚度40毫米"


def _q(fn, *a, **kw):
    conn = dbmod.connect()
    try:
        return fn(conn, *a, **kw)
    finally:
        conn.close()


def _world(env):
    _team(env)
    _add(env, "boss", "leader", "领导")
    assert login("pm").post("/api/p6/import", files={"file": ("plan.xer", sample_xer(), "application/octet-stream")}).status_code == 200


def _up(c, rid, n=1, shot="", color=(100, 110, 120), taken=True, place="3#墩"):
    """水印相机上传：拍摄时间与地点随文件带来。"""
    now = f"{TODAY.isoformat()}T00:01"
    data = {"shot": shot}
    if taken:
        data["taken_at"] = [now] * n
        data["place"] = [place] * n
    r = c.post(f"/api/conceal/{rid}/photos", data=data, files=[("files", (f"p{i}.jpg", jpeg((color[0] + i * 7, color[1], color[2])), "image/jpeg")) for i in range(n)])
    assert r.status_code == 200, r.text
    return r.json()


def _new(c, tpl="pile", part="3#墩", code="A1010"):
    r = c.post("/api/conceal", json={"part": part, "template_key": tpl, "activity_code": code})
    assert r.status_code == 200, r.text
    return r.json()


def _complete_pile(c, rid):
    for k, s in enumerate(("position", "depth", "diameter", "cage", "sediment")):
        _up(c, rid, 1, s, color=(20 + k * 40, 90, 150))
    r = c.put(f"/api/conceal/{rid}/fields", json={"fields": {"pile_no": {"value": "12#"}, "depth": {"value": 28.5, "design": "≥28"}, "dia": {"value": 1500}, "sediment": {"value": 40}}})
    assert r.status_code == 200, r.text
    return r.json()


def test_templates_match_and_edit(env):
    _world(env)
    pm, eng, t1 = login("pm"), login("eng"), login("tech1")
    ts = pm.get("/api/conceal/templates").json()
    keys = {t["key"] for t in ts["templates"]}
    assert {"rebar", "formwork", "foundation_pit", "pile", "embedded", "waterproof"} <= keys and ts["can_edit"]
    assert "以设计文件、项目验收规范和监理要求为准" in ts["notice"]
    for t in ts["templates"]:
        assert t["shots"] and all(s["label"] for s in t["shots"])
        # 不编造规范条文号与允许偏差数值：模板里没有这类字段
        assert not any(k in i for i in t["items"] for k in ("tolerance", "limit", "standard", "clause"))
        assert any(i["type"] == "number" for i in t["items"]) and "kw" in t["items"][0]
    rebar = next(t for t in ts["templates"] if t["key"] == "rebar")
    assert any(s["key"] == "cover" and "垫块" in s["label"] for s in rebar["shots"]) and all(i["measure"] for i in rebar["items"] if i["type"] == "number")
    tpls = _q(cs.templates, env["project_id"])
    assert [t["key"] for t in cs.match_templates(tpls, "3#墩承台钢筋绑扎")] == ["rebar"]
    assert [t["key"] for t in cs.match_templates(tpls, "3#墩桩基施工")] == ["pile"]
    assert cs.match_templates(tpls, "路基土方开挖") == [] and cs.match_templates(tpls, "梁板预制") == []
    # 选项：按作业名与 WBS 给出部位和可选模板（技术员只看自己范围）
    opt = pm.get("/api/conceal/options").json()
    # 梁板工程：A4010“梁板预制”名称不命中模板，但放行清单按混凝土类核验钢筋/模板验收，可关联（F-11①）
    assert opt["has_plan"] and {p["name"] for p in opt["parts"]} == {"3#墩", "4#墩", "梁板工程"}
    assert [p["name"] for p in t1.get("/api/conceal/options").json()["parts"]] == ["3#墩"]
    # 改模板仅 eng_sup / manager
    body = {"name": "自定义：沉井封底", "group": "沉井", "keywords": ["沉井", "封底"], "shots": [{"key": "bottom", "label": "封底前基底", "required": True}],
            "items": [{"key": "thick", "label": "封底厚度", "type": "number", "unit": "m", "required": False, "min": 0, "max": 20}]}
    assert t1.put("/api/conceal/templates/caisson", json=body).status_code == 403
    assert login("boss").put("/api/conceal/templates/caisson", json=body).status_code == 403
    assert eng.put("/api/conceal/templates/caisson", json=body).status_code == 200
    assert eng.put("/api/conceal/templates/Bad-Key", json=body).status_code in (400, 404)
    assert eng.put("/api/conceal/templates/caisson2", json={**body, "shots": []}).status_code == 400
    assert "caisson" in {t["key"] for t in t1.get("/api/conceal/templates").json()["templates"]}
    # 内置模板改写、隐藏、恢复
    over = {**rebar, "name": "钢筋隐蔽（本项目）", "shots": rebar["shots"][:2]}
    assert pm.put("/api/conceal/templates/rebar", json=over).json()["name"] == "钢筋隐蔽（本项目）"
    assert len(next(t for t in pm.get("/api/conceal/templates").json()["templates"] if t["key"] == "rebar")["shots"]) == 2
    assert pm.delete("/api/conceal/templates/waterproof").status_code == 200
    assert "waterproof" not in {t["key"] for t in t1.get("/api/conceal/templates").json()["templates"]}
    assert pm.delete("/api/conceal/templates/rebar?restore=1").status_code == 200
    assert len(next(t for t in pm.get("/api/conceal/templates").json()["templates"] if t["key"] == "rebar")["shots"]) == 6
    assert t1.delete("/api/conceal/templates/pile").status_code == 403


def test_mock_analyze_then_missing_rules_and_human_values_separate(env):
    _world(env)
    t1 = login("tech1")
    rec = _new(t1)
    rid = rec["id"]
    assert rec["status"] == "draft" and rec["check"]["conclusion"] == "有缺项" and len(rec["check"]["missing_shots"]) == 5 and len(rec["check"]["pending_items"]) == 4
    assert rec["title"] == "3#墩桩基验收" and rec["activity_name"] == "3#墩桩基施工" and "以设计文件" in rec["notice"]
    _up(t1, rid, 1, "position", color=(30, 40, 50))
    v = _up(t1, rid, 2, "")                      # 两张未归类
    assert [p["shot"] for p in v["photos"]] == ["position", None, None]
    assert any(w["kind"] == "unassigned" for w in v["check"]["warnings"])
    # 识别前不能凭空归类：AI 只在 ai_json 里给建议
    assert t1.post(f"/api/conceal/{rid}/analyze", json={"narration": ""}).status_code == 200
    a = t1.get(f"/api/conceal/{rid}").json()
    assert a["ai"]["mock"] and "离线模拟" in a["ai"]["notice"] and a["status"] == "draft"
    sug = [p["ai"]["shot"] for p in a["photos"]]
    assert sug[0] == "position" and sug[1:] == ["depth", "diameter"] and all("离线模拟" in p["ai"]["scene"] for p in a["photos"])
    assert [p["shot"] for p in a["photos"]] == ["position", None, None]      # 人工归类未被改动
    assert {m["key"] for m in a["check"]["missing_shots"]} == {"depth", "diameter", "cage", "sediment"}
    # 口述 → 候选值；数据项仍为待补，必须人采用
    a = t1.post(f"/api/conceal/{rid}/analyze", json={"narration": NARR}).json()
    cand = {f["item"]: f for f in a["ai"]["fields"]}
    assert cand["depth"]["value"] == 28.5 and cand["dia"]["value"] == 1500 and cand["sediment"]["value"] == 40 and all(f["src"] == "口述" for f in cand.values())
    assert a["fields"] == {} and {m["key"] for m in a["check"]["pending_items"]} == {"pile_no", "depth", "dia", "sediment"}
    conn = dbmod.connect()
    row = conn.execute("SELECT * FROM conceal_records WHERE id=?", (rid,)).fetchone()
    ai_before = row["ai_json"]
    assert row["fields_json"] == "{}" and '"depth"' in ai_before and row["status"] == "draft" and row["narration"] == NARR
    # 人确认采用：写入 fields_json，ai_json 不变
    r = t1.put(f"/api/conceal/{rid}/fields", json={"fields": {"depth": {"value": 28.5, "src": "口述"}, "dia": {"value": "1500", "design": "Φ1500"}}})
    assert r.status_code == 200
    fv = r.json()["fields"]
    assert fv["depth"]["src"] == "口述" and fv["dia"]["src"] == "人工" and fv["dia"]["design"] == "Φ1500" and fv["dia"]["value"] == 1500
    row = conn.execute("SELECT * FROM conceal_records WHERE id=?", (rid,)).fetchone()
    assert row["ai_json"] == ai_before and '"dia"' in row["fields_json"]
    assert t1.put(f"/api/conceal/{rid}/fields", json={"fields": {"dia": {"value": "abc"}}}).status_code == 400
    assert t1.put(f"/api/conceal/{rid}/fields", json={"fields": {"nope": {"value": 1}}}).status_code == 400
    # 采用 AI 归类建议后才计入缺项
    r = t1.post(f"/api/conceal/{rid}/photos/accept-ai", json={}).json()
    assert r["accepted"] == 2 and [p["shot_src"] for p in r["photos"]] == ["人工", "AI建议·人确认", "AI建议·人确认"]
    assert {m["key"] for m in r["check"]["missing_shots"]} == {"cage", "sediment"}
    # 缺项快照可被其他模块直接 SQL 读取
    ms = dbmod.loads(conn.execute("SELECT missing_json FROM conceal_records WHERE id=?", (rid,)).fetchone()["missing_json"])
    assert ms["shots"] == ["cage", "sediment"] and ms["fields"] == ["pile_no", "sediment"] and ms["complete"] is False
    cols = {c["name"] for c in conn.execute("PRAGMA table_info(conceal_records)").fetchall()}
    assert {"id", "project_id", "activity_code", "part", "template_key", "title", "status", "inspection_date", "fields_json", "photos_json", "ai_json", "missing_json",
            "supervisor_name", "submitted_at", "result_note", "result_at", "created_by", "created_at"} <= cols
    # 改人工归类、取下照片
    pid = r["photos"][2]["media_id"]
    assert t1.put(f"/api/conceal/{rid}/photos/{pid}", json={"shot": "cage"}).json()["photos"][2]["shot"] == "cage"
    assert t1.put(f"/api/conceal/{rid}/photos/{pid}", json={"shot": "bogus"}).status_code == 400
    assert len(t1.delete(f"/api/conceal/{rid}/photos/{pid}").json()["photos"]) == 2
    assert conn.execute("SELECT COUNT(*) n FROM media WHERE id=?", (pid,)).fetchone()["n"] == 1   # 原件仍在
    # 没有任何照片也没口述 → 不能识别
    rid2 = _new(t1, "formwork", "3#墩", None)["id"]
    assert t1.post(f"/api/conceal/{rid2}/analyze", json={}).status_code == 400
    conn.close()


def test_time_geo_warnings_and_conclusion(env):
    _world(env)
    t1 = login("tech1")
    rid = _new(t1)["id"]
    v = _up(t1, rid, 1, "position", taken=False)               # 无拍摄时间、无定位
    kinds = {w["kind"]: w["text"] for w in v["check"]["warnings"]}
    assert "拍摄时间缺失" in kinds["no_time"] and "定位缺失" in kinds["no_geo"] and v["photos"][0]["taken_at"] is None and v["photos"][0]["geo"] is None
    v = _up(t1, rid, 1, "depth", color=(10, 200, 10))          # 有时间和地点
    kinds = {w["kind"]: w["text"] for w in v["check"]["warnings"]}
    assert kinds["no_time"].startswith("1 张") and kinds["no_geo"].startswith("1 张")
    # 有缺项绝不显示可报验
    assert v["check"]["complete"] is False and v["check"]["conclusion"] == "有缺项" and "可报验" not in v["check"]["conclusion"]
    # 数据项齐、照片齐 → 资料齐全，可报验（时间定位缺失只是提示）
    v = _complete_pile(t1, rid)
    assert v["check"]["complete"] and v["check"]["conclusion"] == "资料齐全，可报验" and v["check"]["missing_shots"] == [] and v["check"]["pending_items"] == []


def test_rules_layer_range_unit_shot_and_ai_failure(env):
    _world(env)
    t1 = login("tech1")
    rid = _new(t1)["id"]
    _up(t1, rid, 2, "")
    raw = {"photos": [{"photo": "P1", "scene": "钻孔现场", "shot": "bogus", "clear": False,
                       "readings": [{"item": "dia", "label": "探孔器", "value": 99999, "unit": "mm"}, {"item": "depth", "label": "测绳", "value": 28000, "unit": "mm"},
                                    {"item": "sediment", "label": "沉渣", "value": 35, "unit": "mm"}, {"item": None, "label": "标牌", "value": None, "text": "12#桩"}]},
                      {"photo": "P2", "scene": "钢筋笼", "shot": "cage", "clear": True, "readings": []},
                      {"photo": "P9", "scene": "不存在的照片", "shot": "cage"}],
           "fields": [{"item": "dia", "value": 99999, "unit": "mm"}, {"item": "depth", "value": 30, "unit": "m"},       # 超范围 / 口述里没有
                      {"item": "sediment", "value": 40, "unit": "cm"},                                              # 单位不一致
                      {"item": "pile_no", "value": "12#", "unit": ""}, {"item": "pile_no", "value": "99#"}, {"item": "ghost", "value": 1}]}

    def fake(kind, **kw):
        assert kind == "conceal.pack" and len(kw["images"]) == 2
        return AIResult(data=kw["validate"](raw), run_id=1, provider="jiaorong", model="vision-x", latency_ms=1)
    orig = cs.call_json
    cs.call_json = fake
    try:
        v = t1.post(f"/api/conceal/{rid}/analyze", json={"narration": "桩号12#，孔径99999，沉渣厚度40"}).json()
    finally:
        cs.call_json = orig
    p1, p2 = v["photos"][0]["ai"], v["photos"][1]["ai"]
    assert p1["shot"] is None and p1["clear"] is False and p2["shot"] == "cage"                    # 非法 shot 置空
    rd = {r["label"]: r for r in p1["readings"]}
    assert rd["探孔器"]["value"] is None and "超出常见范围" in rd["探孔器"]["doubt"]                    # 越界置空并标存疑
    assert rd["沉渣"]["value"] == 35 and rd["标牌"]["text"] == "12#桩"
    assert "单位" in rd["测绳"]["doubt"] and rd["测绳"]["value"] is None                              # 读数 28000mm 与数据项单位 m 不一致
    cand = {f["item"] + str(i): f for i, f in enumerate(v["ai"]["fields"])}
    by = {f["item"]: f for f in v["ai"]["fields"]}
    assert by["dia"]["value"] is None and "超出常见范围" in by["dia"]["doubt"]
    assert by["depth"]["value"] is None and "口述里没有" in by["depth"]["doubt"]
    assert by["sediment"]["value"] is None and "单位" in by["sediment"]["doubt"]
    pile_vals = [f for f in v["ai"]["fields"] if f["item"] == "pile_no"]
    assert pile_vals[0]["value"] == "12#" and pile_vals[1]["value"] is None and len(cand) == 5 and v["ai"]["mock"] is False and "AI 识别结果仅供参考" in v["ai"]["notice"]
    assert v["fields"] == {} and v["status"] == "draft"
    # 模型不可用：502，状态与已有结果都不变
    def boom(kind, **kw):
        raise AIError("交融不可用")
    cs.call_json = boom
    try:
        r = t1.post(f"/api/conceal/{rid}/analyze", json={})
    finally:
        cs.call_json = orig
    assert r.status_code == 502 and t1.get(f"/api/conceal/{rid}").json()["ai"]["analyzed"]
    # 直接测规则函数：口述数值必须真出现
    tpl = _q(cs.get_template, env["project_id"], "pile")
    n = conceal_agent.apply_rules(tpl, {"photos": {}, "fields": [{"item": "dia", "value": 1500.0, "unit": "毫米", "src": "口述"}]}, "孔径1500")
    assert n["fields"][0]["value"] == 1500.0
    n = conceal_agent.apply_rules(tpl, {"photos": {}, "fields": [{"item": "dia", "value": 1500.0, "unit": "mm", "src": "口述"}]}, "")
    assert n["fields"][0]["value"] is None


def test_status_flow_and_forced_ready(env):
    _world(env)
    t1, eng, pm = login("tech1"), login("eng"), login("pm")
    rid = _new(t1)["id"]
    _up(t1, rid, 1, "position")
    # 有缺项：必须写缺项说明才能强制标 ready
    assert t1.post(f"/api/conceal/{rid}/ready", json={}).status_code == 400
    assert t1.post(f"/api/conceal/{rid}/ready", json={"note": "  "}).status_code == 400
    r = t1.post(f"/api/conceal/{rid}/ready", json={"note": "钢筋笼已入孔，照片补拍中"})
    assert r.status_code == 200 and r.json()["status"] == "ready" and r.json()["forced"] and r.json()["check"]["conclusion"] == "有缺项"
    assert r.json()["ready_note"] == "钢筋笼已入孔，照片补拍中"
    # ready 后不可再改；退回补资料后可改
    assert t1.put(f"/api/conceal/{rid}/fields", json={"fields": {"dia": {"value": 1}}}).status_code == 409
    assert t1.post(f"/api/conceal/{rid}/photos/attach", json={"media_ids": [1]}).status_code == 409
    assert t1.post(f"/api/conceal/{rid}/analyze", json={}).status_code == 409
    assert t1.delete(f"/api/conceal/{rid}").status_code == 409
    assert t1.post(f"/api/conceal/{rid}/reopen").json()["status"] == "draft"
    _complete_pile(t1, rid)
    r = t1.post(f"/api/conceal/{rid}/ready", json={}).json()
    assert r["status"] == "ready" and not r["forced"] and r["check"]["conclusion"] == "资料齐全，可报验"
    # 流转次序
    assert t1.post(f"/api/conceal/{rid}/result", json={"result": "passed"}).status_code == 403          # 技术员不能录监理结果
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "passed"}).status_code == 409         # 未报验
    s = t1.post(f"/api/conceal/{rid}/submit", json={"supervisor_name": "王监理"}).json()
    assert s["status"] == "submitted" and s["submitted_at"] and s["supervisor_name"] == "王监理"
    assert t1.post(f"/api/conceal/{rid}/submit", json={}).status_code == 409
    assert t1.post(f"/api/conceal/{rid}/reopen").status_code == 409                                       # 已报验不能直接退回
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "rejected"}).status_code == 400         # 退回须写原因
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "rejected", "note": "  "}).status_code == 400
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "approved"}).status_code == 422
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "passed", "date": "2026/10/01"}).status_code == 400
    rj = eng.post(f"/api/conceal/{rid}/result", json={"result": "rejected", "note": "沉渣厚度需复测", "date": TODAY.isoformat()}).json()
    assert rj["status"] == "rejected" and rj["result_note"] == "沉渣厚度需复测" and rj["result_at"] == TODAY.isoformat()
    assert t1.post(f"/api/conceal/{rid}/reopen").json()["status"] == "draft"
    assert t1.get(f"/api/conceal/{rid}").json()["result_note"] == ""
    assert [h["action"] for h in t1.get(f"/api/conceal/{rid}").json()["history"]][-1] == "退回补资料"
    r = t1.post(f"/api/conceal/{rid}/ready", json={}).json()
    t1.post(f"/api/conceal/{rid}/submit", json={})
    ps = pm.post(f"/api/conceal/{rid}/result", json={"result": "passed", "note": "同意进入下道工序", "supervisor_name": "王监理"}).json()
    assert ps["status"] == "passed" and ps["result_at"] == TODAY.isoformat() and ps["supervisor_name"] == "王监理"
    assert t1.post(f"/api/conceal/{rid}/reopen").status_code == 409                                       # 通过后不可改
    assert t1.put(f"/api/conceal/{rid}", json={"title": "改名"}).status_code == 409
    # 列表按状态分组计数
    lst = t1.get("/api/conceal").json()
    assert lst["counts"]["passed"] == 1 and lst["records"][0]["status_name"] == "监理通过" and lst["can_write"] and not lst["can_result"]
    assert t1.get("/api/conceal", params={"status": "draft"}).json()["records"] == []
    conn = dbmod.connect()
    acts = {r["action"] for r in conn.execute("SELECT action FROM audit").fetchall()}
    conn.close()
    assert {"conceal.create", "conceal.photo", "conceal.fields", "conceal.ready", "conceal.submit", "conceal.result", "conceal.reopen"} <= acts


def test_permissions_and_scope(env):
    _world(env)
    _add(env, "tech3", "tech", "同甲技术员", "甲", "3#墩")
    t1, t2, t3, eng, pm, ss, boss, sa, adm = (login(u) for u in ("tech1", "tech2", "tech3", "eng", "pm", "ss", "boss", "sa_a", "admin"))
    rid = _new(t1)["id"]
    # 范围外：乙单位技术员、同单位但非本人的技术员都不能看/改
    for c in (t2, t3):
        assert c.get(f"/api/conceal/{rid}").status_code == 403
        assert c.put(f"/api/conceal/{rid}/fields", json={"fields": {}}).status_code == 403
        assert c.post(f"/api/conceal/{rid}/analyze", json={}).status_code == 403
        assert c.get(f"/api/conceal/{rid}/docx").status_code == 403
        assert c.get("/api/conceal").json()["records"] == []
    assert t2.post("/api/conceal", json={"part": "3#墩", "template_key": "pile", "activity_code": "A1010"}).status_code == 403   # 不在 tech2 的负责范围
    # D20：不关联作业时可填任意具体部位，只记作者（可见范围仍按作者过滤，见 test_custom_part_without_activity_d20）
    assert t2.post("/api/conceal", json={"part": "3#墩", "template_key": "pile"}).status_code == 200
    assert t2.post("/api/conceal", json={"part": "路基工程", "template_key": "foundation_pit"}).status_code == 200
    assert t1.post("/api/conceal", json={"part": "3#墩", "template_key": "nope"}).status_code == 400
    assert t1.post("/api/conceal", json={"part": "", "template_key": "pile"}).status_code == 400
    assert t1.post("/api/conceal", json={"part": "3#墩", "template_key": "pile", "activity_code": "ZZZ"}).status_code == 400
    # 分包安全员、管理员无权
    for c in (sa, adm):
        for path in ("/api/conceal", f"/api/conceal/{rid}", "/api/conceal/gaps", "/api/conceal/templates", "/api/conceal/options", "/api/conceal/docx?part=3%23墩"):
            assert c.get(path).status_code == 403, path
        assert c.post("/api/conceal", json={"part": "3#墩", "template_key": "pile"}).status_code == 403
    # 领导、安全主管只读
    for c in (boss, ss):
        assert c.get("/api/conceal").status_code == 200 and c.get(f"/api/conceal/{rid}").status_code == 200 and c.get("/api/conceal/gaps").status_code == 200
        assert c.get(f"/api/conceal/{rid}").json()["can"]["edit"] is False
        assert c.post("/api/conceal", json={"part": "3#墩", "template_key": "pile"}).status_code == 403
        assert c.put(f"/api/conceal/{rid}/fields", json={"fields": {}}).status_code == 403
        assert c.post(f"/api/conceal/{rid}/analyze", json={}).status_code == 403
        assert c.post(f"/api/conceal/{rid}/ready", json={"note": "x"}).status_code == 403
        assert c.post(f"/api/conceal/{rid}/result", json={"result": "passed"}).status_code == 403
        assert c.put("/api/conceal/templates/pile", json={"name": "x"}).status_code == 403
        assert c.delete(f"/api/conceal/{rid}").status_code == 403
    # 港湾看全项目；工程主管可代为编辑
    assert eng.get(f"/api/conceal/{rid}").json()["can"]["edit"] and len(pm.get("/api/conceal").json()["records"]) == 3
    assert eng.put(f"/api/conceal/{rid}", json={"title": "3#墩12号桩验收"}).json()["title"] == "3#墩12号桩验收"
    # 删除草稿
    assert t1.delete(f"/api/conceal/{rid}").status_code == 200 and t1.get(f"/api/conceal/{rid}").status_code == 404
    # 未登录
    from fastapi.testclient import TestClient
    from app.main import app
    anon = TestClient(app)
    assert anon.get("/api/conceal").status_code == 401


def test_gaps_scan(env):
    # 无 P6：如实返回空与说明，不造数
    _team(env)
    pm = login("pm")
    g = pm.get("/api/conceal/gaps").json()
    assert g["gaps"] == [] and g["parts"] == [] and g["has_plan"] is False and "P6" in g["note"] and g["summary"]["total"] == 0
    assert pm.get("/api/conceal/options").json()["parts"] == []
    assert pm.post("/api/p6/import", files={"file": ("plan.xer", sample_xer(), "application/octet-stream")}).status_code == 200
    g = pm.get("/api/conceal/gaps").json()
    got = {(x["activity_code"], x["template_key"]): x for x in g["gaps"]}
    # 已开始的作业：A1010(3#墩桩基) A2010(4#墩桩基)；未开始的承台钢筋 A1020 不列；里程碑/路基不匹配模板
    assert set(got) == {("A1010", "pile"), ("A2010", "pile")}
    a = got[("A1010", "pile")]
    assert a["state"] == "missing" and a["part"] == "3#墩" and "疑似" in g["note"] and "没有找到" in a["basis"] and a["p6_status"] == "进行中"
    assert g["summary"] == {"done": 0, "pending": 0, "missing": 2, "total": 2}
    assert {p["part"]: p["missing"] for p in g["parts"]} == {"3#墩": 1, "4#墩": 1} and g["groups"][0]["group"] == "桩基"
    # 技术员只看自己范围
    assert {x["activity_code"] for x in login("tech1").get("/api/conceal/gaps").json()["gaps"]} == {"A1010"}
    # 草稿 → 待补（仍列为疑似缺）；确认后 → 已齐，不再列
    t1 = login("tech1")
    rid = _new(t1)["id"]
    g = pm.get("/api/conceal/gaps").json()
    a = next(x for x in g["gaps"] if x["activity_code"] == "A1010")
    assert a["state"] == "pending" and a["draft_id"] == rid and g["summary"]["pending"] == 1
    _complete_pile(t1, rid)
    assert t1.post(f"/api/conceal/{rid}/ready", json={}).status_code == 200
    g = pm.get("/api/conceal/gaps").json()
    assert {x["activity_code"] for x in g["gaps"]} == {"A2010"} and g["summary"] == {"done": 1, "pending": 0, "missing": 1, "total": 2}
    assert _q(cs.covered, env["project_id"], activity_code="A1010") and not _q(cs.covered, env["project_id"], activity_code="A2010")
    # 同部位、同模板的记录也算覆盖（不填作业代码）
    eng = login("eng")
    r2 = eng.post("/api/conceal", json={"part": "4#墩", "template_key": "pile"}).json()["id"]
    assert {x["activity_code"] for x in pm.get("/api/conceal/gaps").json()["gaps"]} == {"A2010"}      # 草稿不算
    _complete_pile(eng, r2)
    eng.post(f"/api/conceal/{r2}/ready", json={})
    g = pm.get("/api/conceal/gaps").json()
    assert g["gaps"] == [] and g["summary"] == {"done": 2, "pending": 0, "missing": 0, "total": 2}
    # 监理退回的记录不算已齐
    t1.post(f"/api/conceal/{rid}/submit", json={})
    eng.post(f"/api/conceal/{rid}/result", json={"result": "rejected", "note": "复测"})
    g = pm.get("/api/conceal/gaps").json()
    assert [x["activity_code"] for x in g["gaps"]] == ["A1010"] and g["gaps"][0]["state"] == "pending"
    # 自定义模板关键词也参与匹配；隐藏模板后不再扫描
    assert pm.delete("/api/conceal/templates/pile").status_code == 200
    assert pm.get("/api/conceal/gaps").json()["gaps"] == []


def test_candidate_photos_and_attach(env):
    _world(env)
    t1 = login("tech1")
    rid = _new(t1)["id"]
    now, old = f"{TODAY.isoformat()}T00:01", f"{(TODAY - timedelta(days=5)).isoformat()}T08:10"

    def log_photo(taken, color):
        r = t1.post("/api/media", data={"purpose": "log", "taken_at": [taken], "place": ["3#墩"]}, files=[("files", ("a.jpg", jpeg(color), "image/jpeg"))])
        assert r.status_code == 200, r.text
        return r.json()[0]["id"]
    new_id, old_id = log_photo(now, (1, 2, 3)), log_photo(old, (9, 9, 9))
    loose = t1.post("/api/media", data={"purpose": "log", "taken_at": [now]}, files=[("files", ("b.jpg", jpeg((70, 80, 90)), "image/jpeg"))]).json()[0]["id"]
    c = t1.get("/api/conceal/candidates", params={"part": "3#墩", "date": TODAY.isoformat(), "record_id": rid}).json()
    assert [p["media_id"] for p in c["photos"]] == [new_id] and [p["media_id"] for p in c["unassigned"]] == [loose]
    assert c["photos"][0]["taken_at"] == now and not c["photos"][0]["in_record"]
    # 乙单位技术员看不到甲的候选照片
    assert login("tech2").get("/api/conceal/candidates", params={"part": "3#墩", "date": TODAY.isoformat()}).json()["photos"] == []
    v = t1.post(f"/api/conceal/{rid}/photos/attach", json={"media_ids": [new_id, old_id], "shot": "position"}).json()
    assert [p["media_id"] for p in v["photos"]] == [new_id] and v["photos"][0]["source"] == "部位已有照片" and v["photos"][0]["shot"] == "position"
    assert v["skipped"][0]["media_id"] == old_id and "前后 1 天" in v["skipped"][0]["why"]
    assert t1.post(f"/api/conceal/{rid}/photos/attach", json={"media_ids": [old_id]}).status_code == 400
    assert len(t1.post(f"/api/conceal/{rid}/photos/attach", json={"media_ids": [new_id]}).json()["photos"]) == 1       # 不重复引用
    assert t1.get("/api/conceal/candidates", params={"part": "3#墩", "date": TODAY.isoformat(), "record_id": rid}).json()["photos"][0]["in_record"]
    # 他人照片不能引用
    t2 = login("tech2")
    own2 = t2.post("/api/media", data={"purpose": "log", "taken_at": [now], "place": ["路基工程"]}, files=[("files", ("c.jpg", jpeg((5, 6, 7)), "image/jpeg")) ]).json()[0]["id"]
    assert t1.post(f"/api/conceal/{rid}/photos/attach", json={"media_ids": [own2]}).status_code == 400
    # 隐蔽验收照片不进安全识别队列
    conn = dbmod.connect()
    st = conn.execute("SELECT purpose, safety_status FROM media WHERE purpose='conceal'").fetchall()
    conn.close()
    assert st == [] or all(s["safety_status"] == "na" for s in st)


def test_docx_single_and_part(env):
    from docx import Document
    _world(env)
    t1, eng, boss = login("tech1"), login("eng"), login("boss")
    rid = _new(t1)["id"]
    _complete_pile(t1, rid)
    t1.post(f"/api/conceal/{rid}/analyze", json={"narration": NARR})
    t1.post(f"/api/conceal/{rid}/ready", json={})
    rd = t1.get(f"/api/conceal/{rid}/docx")
    assert rd.status_code == 200 and rd.content[:2] == b"PK"
    doc = Document(io.BytesIO(rd.content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "隐蔽工程验收记录" in text and "3#墩" in text and "桩基验收" in text and "沉渣厚度实测" in text and "Φ" not in text
    assert "AI 预填，经人工确认；签认以纸质/监理系统为准。" in text and "监理单位验收意见" in text and "【缺拍】" not in text and "28.5" in text
    assert len(doc.inline_shapes) >= 5 and "拍摄时间缺失" not in text
    assert doc.sections[0].footer.paragraphs[0].text.startswith("AI 预填，经人工确认")
    # 缺项如实写在册里
    rid2 = _new(eng, "rebar", "3#墩", None)["id"]
    _up(eng, rid2, 1, "overview", taken=False)
    eng.post(f"/api/conceal/{rid2}/ready", json={"note": "其余照片待补"})
    doc2 = Document(io.BytesIO(eng.get(f"/api/conceal/{rid2}/docx").content))
    text2 = "\n".join(p.text for p in doc2.paragraphs) + "\n" + "\n".join(c.text for t in doc2.tables for row in t.rows for c in row.cells)
    assert "【缺拍】" in text2 and "待补" in text2 and "其余照片待补" in text2 and "拍摄时间缺失" in text2
    # 按部位成册：默认只含已确认的记录
    rp = eng.get("/api/conceal/docx", params={"part": "3#墩"})
    assert rp.status_code == 200
    dp = Document(io.BytesIO(rp.content))
    tp = "\n".join(p.text for p in dp.paragraphs) + "\n" + "\n".join(c.text for t in dp.tables for row in t.rows for c in row.cells)
    assert "隐蔽工程验收资料册：3#墩" in tp and "共 2 份记录" in tp and "钢筋隐蔽验收" in tp and len(dp.inline_shapes) >= 5
    assert eng.get("/api/conceal/docx", params={"part": "路基工程"}).status_code == 404
    assert boss.get("/api/conceal/docx", params={"part": "3#墩"}).status_code == 200          # 领导可导出（只读）
    assert login("tech2").get("/api/conceal/docx", params={"part": "3#墩"}).status_code == 404   # 范围外看不到任何记录
    conn = dbmod.connect()
    acts = {r["action"] for r in conn.execute("SELECT action FROM audit").fetchall()}
    conn.close()
    assert {"conceal.export", "conceal.analyze"} <= acts


def test_template_snapshot_keeps_submitted_basis(env):
    _world(env)
    t1, pm = login("tech1"), login("pm")
    rid = _new(t1)["id"]
    _complete_pile(t1, rid)
    assert t1.post(f"/api/conceal/{rid}/ready", json={}).status_code == 200
    tpl = next(t for t in pm.get("/api/conceal/templates").json()["templates"] if t["key"] == "pile")
    tpl["shots"].append({"key": "extra", "label": "新增必拍项", "required": True, "hint": ""})
    assert pm.put("/api/conceal/templates/pile", json=tpl).status_code == 200
    v = t1.get(f"/api/conceal/{rid}").json()
    assert v["check"]["complete"] and len(v["template"]["shots"]) == 5          # 已确认的记录按建档/确认时的模板
    rid2 = _new(t1)["id"]
    assert any(m["key"] == "extra" for m in t1.get(f"/api/conceal/{rid2}").json()["check"]["missing_shots"])


# ----------------------------------------------------------------------------- T-14 隐蔽验收闭环（F-03 / F-04 / F-11①）

def _notes(uid, like="%"):
    conn = dbmod.connect()
    try:
        return [dict(r) for r in conn.execute("SELECT * FROM notifications WHERE user_id=? AND title LIKE ? ORDER BY id", (uid, like)).fetchall()]
    finally:
        conn.close()


def test_custom_part_without_activity_d20(env):
    """F-03 / AC-F3：不关联作业时可填任意具体部位，只记作者；关联作业时仍校验负责范围；可见范围不变。"""
    _world(env)
    t1, t2, pm = login("tech1"), login("tech2"), login("pm")
    # 走查原样：部位选“其他”，按提示填“3#墩”“T梁预制场1#台座”，都能建成
    for part in ("3#墩", "T梁预制场1#台座", "1#涵"):
        r = t2.post("/api/conceal", json={"part": part, "template_key": "rebar"})
        assert r.status_code == 200, (part, r.text)
        assert r.json()["part"] == part and r.json()["created_by"] == "李工" and r.json()["activity_code"] is None
    rid = _new(t1, "rebar", "T梁预制场1#台座", None)["id"]
    # 关联作业时仍校验负责范围，提示说清是“作业”不在范围、不关联作业可填任意部位
    r = t2.post("/api/conceal", json={"part": "3#墩", "template_key": "pile", "activity_code": "A1010"})
    assert r.status_code == 403 and "作业" in r.json()["detail"] and "不关联作业" in r.json()["detail"]
    assert t1.post("/api/conceal", json={"part": "3#墩", "template_key": "pile", "activity_code": "A1010"}).status_code == 200
    # 草稿改部位：不关联作业时同样不限
    assert t1.put(f"/api/conceal/{rid}", json={"part": "BR140 T梁"}).json()["part"] == "BR140 T梁"
    # 可见范围仍按作者：技术员互相看不到，项目部看全部
    assert t2.get(f"/api/conceal/{rid}").status_code == 403
    assert {r["created_by"] for r in t1.get("/api/conceal").json()["records"]} == {"张工"}
    assert {r["created_by"] for r in t2.get("/api/conceal").json()["records"]} == {"李工"}
    assert len(pm.get("/api/conceal").json()["records"]) == 5
    # 页面提示与规则一致：不再暗示“只能填自己范围内的部位”
    js = (Path(__file__).resolve().parent.parent / "app" / "static" / "x_conceal.js").read_text(encoding="utf-8")
    assert "不关联作业时可填写任意具体部位" in js and "关联作业须在您的负责范围内" in js


def test_submit_and_result_notify_with_link(env):
    """F-04 / AC-F4：报验 → 工程主管收到带跳转的消息（不发项目经理）；监理退回 → 技术员收到消息并写明原因；通过同样通知。"""
    _world(env)
    t1, eng, pm = login("tech1"), login("eng"), login("pm")
    uid = env["users"]
    rid = _new(t1)["id"]
    _complete_pile(t1, rid)
    assert t1.post(f"/api/conceal/{rid}/ready", json={}).status_code == 200
    before_pm = len(_notes(uid["pm"]))
    assert t1.post(f"/api/conceal/{rid}/submit", json={"supervisor_name": "王监理"}).status_code == 200
    link = {"route": "conceal", "params": {"id": str(rid)}}
    sent = eng.get("/api/notifications").json()
    n = next(x for x in sent if "隐蔽验收" in x["title"])
    assert n["link"] == link and "3#墩桩基验收" in n["title"] + n["body"] and "张工" in n["body"] and "王监理" in n["body"]
    assert "录入监理结果" in n["title"] + n["body"]
    assert len(_notes(uid["pm"])) == before_pm                       # 项目经理不收验收流程消息
    assert not _notes(uid["tech1"], "%隐蔽验收%")                    # 报验人自己不收
    # 监理退回：通知编制人，正文写明原因
    reason = "保护层厚度读数需重测"
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "rejected", "note": reason}).status_code == 200
    back = [x for x in t1.get("/api/notifications").json() if "隐蔽验收" in x["title"]]
    assert len(back) == 1 and "退回" in back[0]["title"] and reason in back[0]["body"] and back[0]["link"] == link
    assert "工程主管" in back[0]["body"] and "退回补资料" in back[0]["body"]
    assert len(_notes(uid["eng"], "%隐蔽验收%")) == 1                # 录入人自己不收
    assert len(_notes(uid["pm"])) == before_pm
    # 补正后再报验 → 工程主管再收一条；监理通过 → 技术员收到“通过”
    t1.post(f"/api/conceal/{rid}/reopen")
    t1.post(f"/api/conceal/{rid}/ready", json={})
    t1.post(f"/api/conceal/{rid}/submit", json={})
    assert len(_notes(uid["eng"], "%隐蔽验收%")) == 2
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "passed", "note": "同意进入下道工序"}).status_code == 200
    ok = [x for x in t1.get("/api/notifications").json() if "隐蔽验收" in x["title"] and "通过" in x["title"]]
    assert len(ok) == 1 and ok[0]["link"] == link and "同意进入下道工序" in ok[0]["body"]
    assert len(_notes(uid["pm"])) == before_pm
    # 失败的流转不发消息
    assert t1.post(f"/api/conceal/{rid}/submit", json={}).status_code == 409
    assert len(_notes(uid["eng"], "%隐蔽验收%")) == 2


def test_eng_self_flow_does_not_notify_self(env):
    """工程主管自己建档、报验、录入结果：不给自己发消息，也不发给项目经理。"""
    _world(env)
    eng = login("eng")
    uid = env["users"]
    rid = _new(eng, "pile", "4#墩", None)["id"]
    _complete_pile(eng, rid)
    eng.post(f"/api/conceal/{rid}/ready", json={})
    assert eng.post(f"/api/conceal/{rid}/submit", json={}).status_code == 200
    assert eng.post(f"/api/conceal/{rid}/result", json={"result": "rejected", "note": "复测"}).status_code == 200
    assert not _notes(uid["eng"], "%隐蔽验收%") and not _notes(uid["pm"], "%隐蔽验收%")


def test_options_include_unstarted_gate_activities(env):
    """F-11①：放行清单要核验隐蔽验收的作业（如承台混凝土浇筑），未开工、名称未命中模板关键词也能在新建时关联。"""
    _world(env)
    t1 = login("tech1")
    opt = t1.get("/api/conceal/options").json()
    acts = {a["code"]: a for p in opt["parts"] for a in p["activities"]}
    assert "A1030" in acts and not acts["A1030"]["started"] and set(acts["A1030"]["templates"]) == {"rebar", "formwork"}
    assert acts["A1010"]["templates"] == ["pile"]                    # 原有关键词匹配不变
    assert "A3010" not in acts                                       # 范围外仍不列
    # 关联后放行清单能核到
    rid = _new(t1, "rebar", "3#墩", "A1030")["id"]
    _up(t1, rid, 1, "overview")
    assert t1.post(f"/api/conceal/{rid}/ready", json={"note": "其余照片待补"}).status_code == 200
    t1.post(f"/api/conceal/{rid}/submit", json={})
    login("eng").post(f"/api/conceal/{rid}/result", json={"result": "passed"})
    g = login("eng").get("/api/gate/A1030").json()
    assert next(x for x in g["conditions"] if x["key"] == "rebar_conceal")["auto_state"] == "ok"

"""T-04a 回测后端（P0-2）：多期 XER、事实取最晚一期、命中率/召回率/提前量、3×3 敏感性表、期次校验、
来源标识（真实/仿真）、Word 导出、费用哨兵、权限（PM/领导 403，采用阈值只给工程主管）、不调用 AI。

回测逻辑一律用合成多期 XER；真实基线 XER 只用来验证解析与工期%换算（文件不在仓库内，缺失时跳过）。
"""
import hashlib
import io
import os
import json
import re
import sqlite3
from contextlib import closing
from datetime import date, timedelta
from pathlib import Path

import pytest

from conftest import login
from sample_xer import make_xer

from app import db as dbmod
from app.security import hash_password
from app.services import p6
from app.services import progress as prog

# ----------------------------------------------------------------------------- 合成多期 XER

S = date(2026, 1, 5)                  # 全部作业计划开始
DUR = 28                              # 原定工期 28 天 → 计划完成 2026-02-02
DD = ["2026-01-12", "2026-01-19", "2026-01-26", "2026-02-16"]

# 各类作业：S0 总浮时、各期工期%、各期 P6 预测完成推迟天数（总浮时随之减少）
#   T 判红且最终延误（TP）；F 判红但追回（FP）；N 未判红但最终延误（FN，第 3 组判红）；G 有浮时、正常（TN）
CLASSES = {
    "T": [(0, [25, 30, 40, 100], [0, 0, 6, 6])] * 3 + [(0, [25, 30, 40, 100], [0, 4, 6, 6])] * 2,
    "F": [(0, [25, 30, 75, 100], [0, 0, 0, 0])] * 5,
    "N": [(0, [25, 50, 75, 90], [0, 0, 0, 20])] * 5,
    "G": [(10, [25, 50, 75, 100], [0, 0, 0, 2])] * 5,
}


def _row(i: int, code: str, dd: str, pct, delay: int, tf: float, *, pct_type: str = "CP_Drtn", phys: str = "0",
         work: tuple | None = None, extra: dict | None = None) -> dict:
    d = date.fromisoformat(dd)
    es, ef = S, S + timedelta(days=DUR)
    done = pct is not None and pct >= 100
    started = pct is None or pct > 0
    act_end = ef + timedelta(days=delay) if done else None
    if done:
        e_s, e_f = es, act_end
    elif started:
        e_s, e_f = d, ef + timedelta(days=delay)
    else:
        e_s, e_f = es + timedelta(days=delay), ef + timedelta(days=delay)
    remain = 0 if done else DUR * (100 - (pct or 0)) / 100
    row = {
        "task_id": str(1000 + i), "proj_id": "1", "wbs_id": "W1", "clndr_id": "1",
        "phys_complete_pct": phys, "complete_pct_type": pct_type, "task_type": "TT_Task",
        "status_code": "TK_Complete" if done else ("TK_Active" if started else "TK_NotStart"),
        "task_code": code, "task_name": f"合成作业{code}",
        "total_float_hr_cnt": f"{(tf - delay) * 8:g}", "remain_drtn_hr_cnt": f"{remain * 8:g}", "target_drtn_hr_cnt": str(DUR * 8),
        "act_work_qty": "" if work is None else str(work[0]), "remain_work_qty": "" if work is None else str(work[1]),
        "act_start_date": f"{es} 08:00" if started else "", "act_end_date": f"{act_end} 17:00" if done else "",
        "early_start_date": f"{e_s} 08:00", "early_end_date": f"{e_f} 17:00",
        "target_start_date": f"{es} 08:00", "target_end_date": f"{ef} 17:00",
        "driving_path_flag": "Y" if tf - delay <= 0 else "N",
    }
    row.update(extra or {})
    return row


def _xer(dd: str, rows: list[dict], extra_tables: tuple = (), wbs_extra: dict | None = None) -> str:
    out = ["ERMHDR\t19.12\t2026-01-01\tProject\tadmin\tadmin\tdbx\tProject Management\tUSD"]

    def table(name, fields, records):
        out.append(f"%T\t{name}")
        out.append("%F\t" + "\t".join(fields))
        for r in records:
            out.append("%R\t" + "\t".join(str(r.get(f, "")) for f in fields))

    table("PROJECT", ["proj_id", "proj_short_name", "last_recalc_date"], [{"proj_id": "1", "proj_short_name": "SYN-BT", "last_recalc_date": f"{dd} 08:00"}])
    table("CALENDAR", ["clndr_id", "clndr_name", "day_hr_cnt"], [{"clndr_id": "1", "clndr_name": "7天", "day_hr_cnt": "8"}])
    wbs = [{"wbs_id": "W0", "proj_id": "1", "parent_wbs_id": "", "wbs_name": "SYN", "proj_node_flag": "Y"},
           {"wbs_id": "W1", "proj_id": "1", "parent_wbs_id": "W0", "wbs_name": "合成分部", "proj_node_flag": "N"}]
    for w in wbs:
        w.update(wbs_extra or {})
    table("PROJWBS", ["wbs_id", "proj_id", "parent_wbs_id", "wbs_name", "proj_node_flag"] + list(wbs_extra or {}), wbs)
    fields = list(dict.fromkeys(k for r in rows for k in r))
    table("TASK", fields, rows)
    table("TASKPRED", ["task_pred_id", "task_id", "pred_task_id", "pred_type", "lag_hr_cnt"], [])
    for name, fields_, records in extra_tables:
        table(name, fields_, records)
    out.append("%E")
    return "\n".join(out) + "\n"


def _period(k: int, classes: dict = CLASSES, extra_rows: list | None = None, **kw) -> str:
    rows, i = [], 0
    for prefix, specs in classes.items():
        for j, (tf, pcts, delays) in enumerate(specs, 1):
            rows.append(_row(i, f"{prefix}{j}", DD[k], pcts[k], delays[k], tf))
            i += 1
    rows += extra_rows or []
    return _xer(DD[k], rows, **kw)


def _files(texts, names=None):
    return [("files", ((names or [f"u{i}.xer" for i in range(len(texts))])[i], t.encode("utf-8"), "application/octet-stream"))
            for i, t in enumerate(texts)]


def _snaps(texts):
    from app.services import backtest as bt
    return bt.parse_files([(f"u{i}.xer", t.encode("utf-8")) for i, t in enumerate(texts)])


def _run(texts, **kw):
    from app.services import backtest as bt
    return bt.run(_snaps(texts), **kw)


def _add_user(env, username, role, project_id=None):
    conn = dbmod.connect()
    uid = dbmod.insert(conn, "users", {"username": username, "name": username, "role": role, "password_hash": hash_password("secret123"),
                                       "project_id": project_id or env["project_id"], "wbs_scope": "", "email": "", "active": 1,
                                       "created_at": dbmod.now_iso()})
    conn.close()
    return uid


# ----------------------------------------------------------------------------- 解析：完成%按类型读取（AC-2.5、G-08、D2）

GOLDEN_SAMPLE = "e42c32217b2ce69720034d58e755a9c65f39d592edbf1297938d3d804a0211ae"
GOLDEN_REAL_PHYS = "2f01a7ccc82270c77d6355b1e6d9f3c4bab8a74f01db2f32df2605e9c654eac3"
REAL_XER = Path(os.environ.get("ZB_REAL_XER", "__no_real_xer__.xer"))  # 本机真实基线 XER，不随仓库提供


def _h(x) -> str:
    return hashlib.sha256(json.dumps(x, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def test_golden_parse_without_pct_type_unchanged():
    """不带 complete_pct_type 的 XER：默认模式解析结果与改动前逐位一致；typed 模式除多一个 pct_type 外也一致。"""
    s = make_xer("2026-10-01", start="2026-09-21", progress={"M0000": 100, "A1010": 80, "A2010": 60, "A3010": 50}, shift={"A1010": 3})
    assert _h(p6.activities_from_xer(s)) == GOLDEN_SAMPLE
    assert _h(p6.activities_from_xer(s, pct_mode="phys")) == GOLDEN_SAMPLE
    meta, acts, rels = p6.activities_from_xer(s, pct_mode="typed")
    assert all(a.pop("pct_type") == "" for a in acts)
    assert _h((meta, acts, rels)) == GOLDEN_SAMPLE
    with pytest.raises(ValueError):
        p6.activities_from_xer(s, pct_mode="bogus")


def test_typed_pct_by_complete_pct_type():
    rows = [
        _row(0, "D1", "2026-01-19", 75, 0, 0, phys="5"),                                     # 工期% = (224-56)/224 = 75
        _row(1, "D2", "2026-01-19", 100, 0, 0, phys="5"),                                    # 已完工 = 100
        _row(2, "D3", "2026-01-19", 0, 0, 0, phys="5"),                                      # 未开始 = 0
        _row(3, "D4", "2026-01-19", 30, 0, 0, phys="5", extra={"remain_drtn_hr_cnt": "300"}),  # 尚需 > 原定 → 0（不出负数）
        _row(4, "U1", "2026-01-19", 40, 0, 0, pct_type="CP_Units", phys="5", work=(30, 10)),  # 数量% = 30/(30+10) = 75
        _row(5, "U2", "2026-01-19", None, 0, 0, pct_type="CP_Units", phys="5"),             # 进行中、没有数量 → 未知
        _row(6, "P1", "2026-01-19", 40, 0, 0, pct_type="CP_Phys", phys="33"),               # 实物% 照读
        _row(7, "X1", "2026-01-19", 40, 0, 0, pct_type="", phys="44"),                      # 未标注 → 实物%（与旧行为一致）
    ]
    meta, acts, _ = p6.activities_from_xer(_xer("2026-01-19", rows), pct_mode="typed")
    got = {a["code"]: (a["phys_pct"], a["pct_type"]) for a in acts}
    assert got["D1"] == (75.0, "CP_Drtn") and got["D2"] == (100.0, "CP_Drtn") and got["D3"] == (0.0, "CP_Drtn")
    assert got["D4"] == (0.0, "CP_Drtn")
    assert got["U1"] == (75.0, "CP_Units") and got["U2"] == (None, "CP_Units")
    assert got["P1"] == (33.0, "CP_Phys") and got["X1"] == (44.0, "")
    # 默认（phys）模式照旧读实物%，不加字段
    _, acts0, _ = p6.activities_from_xer(_xer("2026-01-19", rows))
    assert {a["code"]: a["phys_pct"] for a in acts0}["D1"] == 5.0 and all("pct_type" not in a for a in acts0)


def test_online_import_reads_duration_pct(env):
    """D2：线上导入路径按类型读取；P6 导入结果给出完成%类型分布。"""
    eng = _add_user(env, "eng", "eng_sup")  # noqa: F841
    rows = [_row(0, "D1", "2026-01-19", 75, 0, 0, phys="5"), _row(1, "D2", "2026-01-19", 0, 0, 0, phys="5")]
    r = login("eng").post("/api/p6/import", files={"file": ("plan.xer", _xer("2026-01-19", rows).encode(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert r.json()["pct_types"] == {"CP_Drtn": 2}
    with closing(dbmod.connect()) as conn:
        got = {a["code"]: a["phys_pct"] for a in conn.execute("SELECT code, phys_pct FROM activities").fetchall()}
    assert got == {"D1": 75.0, "D2": 0.0}


@pytest.mark.skipif(not REAL_XER.exists(), reason="真实基线 XER 不在本机")
def test_real_baseline_parse_counts():
    """真实基线（704 个作业全部未开始，完成%类型全部为工期%）：只断言计数，不回显业务内容。"""
    text = p6.decode(REAL_XER.read_bytes())
    tasks = p6.parse_xer(text)["TASK"]
    assert _h(p6.activities_from_xer(text)) == GOLDEN_REAL_PHYS  # 默认模式逐位不变
    meta, acts, rels = p6.activities_from_xer(text, pct_mode="typed")
    expected = sum(1 for t in tasks if t.get("task_type") not in ("TT_WBS", "TT_LOE") and (t.get("task_code") or "").strip())
    assert len(acts) == expected == 704
    assert meta["data_date"] == "2025-11-05"
    assert all(a["pct_type"] == "CP_Drtn" for a in acts)
    assert all(a["phys_pct"] == 0.0 for a in acts)  # 全部未开始 → 工期% 为 0
    _, acts0, _ = p6.activities_from_xer(text)
    assert [{k: v for k, v in a.items() if k != "pct_type"} for a in acts] == acts0


# ----------------------------------------------------------------------------- 回测核心（AC-2.1、2.2、2.7、2.9、2.10、2.11）

def test_multi_period_confusion_fact_from_last_period():
    texts = [_period(k) for k in range(4)]
    res = _run(list(reversed(texts)))  # 上传顺序打乱，按数据日期排序
    s = res["summary"]
    assert s["periods"] == 4 and s["data_date_from"] == DD[0] and s["data_date_to"] == DD[3] and not s["weak"]
    assert (s["samples"], s["tp"], s["fp"], s["fn"], s["tn"]) == (60, 15, 5, 10, 30)
    assert s["precision"] == {"k": 15, "n": 20, "rate": 0.75}            # 主指标：命中率 TP/(TP+FP)
    assert s["recall"] == {"k": 15, "n": 25, "rate": 0.6}
    assert s["delay_share"] == {"k": 25, "n": 60, "rate": round(25 / 60, 4)}
    assert s["accuracy_ref"] == {"k": 45, "n": 60, "rate": 0.75}         # (TP+TN)/N 只作参考
    assert s["data_gaps"] == 0 and s["fact_actual"] == 45 and s["fact_forecast"] == 15
    by = {p["judge_date"]: (p["tp"], p["fp"], p["fn"], p["tn"]) for p in s["by_period"]}
    assert by == {DD[1]: (5, 5, 5, 5), DD[2]: (5, 0, 5, 10), DD[3]: (5, 0, 0, 15)}
    # 事实取最晚一期：N1 第 3 期仍按期，最晚一期才推迟 20 天（旧算法只看下一期会判为不延误）
    n1 = next(r for r in res["rows"] if r["code"] == "N1" and r["judge_date"] == DD[1])
    assert n1["actual"] and n1["fact_kind"] == "forecast" and n1["fact_finish"] == "2026-02-22" and n1["outcome"] == "FN"
    t1 = next(r for r in res["rows"] if r["code"] == "T1" and r["judge_date"] == DD[1])
    assert t1["predicted"] and t1["fact_kind"] == "actual" and t1["outcome"] == "TP" and t1["level"] == "red"
    # 只算已完工（严格口径）与按作业去重
    st = s["strict"]
    assert (st["samples"], st["tp"], st["fp"], st["fn"], st["tn"]) == (45, 10, 5, 0, 30)
    ba = s["by_activity"]
    assert (ba["samples"], ba["tp"], ba["fp"], ba["fn"], ba["tn"]) == (20, 10, 5, 0, 5)
    assert s["tp"] + s["fp"] + s["fn"] + s["tn"] == s["samples"] == len(res["rows"])
    assert "命中率" in "".join(s["definitions"]) and any("参考" in v for v in s["definitions"].values())


def test_lead_time_vs_p6_float():
    res = _run([_period(k) for k in range(4)])
    lt = res["lead_time"]
    # 确实延误的作业：T1～T5、N1～N5。T1～T3 比 P6 总浮时转负早一期（7 天）；T4、T5 与 N 同期
    assert lt["delayed"] == 10 and lt["compared"] == 10
    assert (lt["earlier"], lt["same"], lt["later"], lt["only_ours"], lt["never"]) == (3, 7, 0, 0, 0)
    assert lt["mean_days"] == 2.1 and lt["median_days"] == 0
    item = {i["code"]: i for i in lt["items"]}
    assert item["T1"]["ours"] == DD[1] and item["T1"]["p6"] == DD[2] and item["T1"]["lead_days"] == 7


def test_lead_time_later_and_never():
    classes = {
        "L": [(0, [25, 30, 100], [2, 2, 6])],    # P6 第 1 期就转负，本规则第 2 期才判红 → 更晚 7 天
        "V": [(0, [25, 50, 75], [0, 0, 10])],    # 从未判红但最终延误 → 从未判出
        "G": [(10, [25, 50, 100], [0, 0, 0])] * 3,
    }
    texts = [_period(k, classes) for k in range(3)]
    lt = _run(texts)["lead_time"]
    item = {i["code"]: i for i in lt["items"]}
    assert item["L"+"1"]["category"] == "later" and item["L1"]["lead_days"] == -7
    assert item["V1"]["category"] == "never" and item["V1"]["ours"] is None
    assert lt["later"] == 1 and lt["never"] == 1


def test_sensitivity_grid_3x3():
    res = _run([_period(k) for k in range(4)])
    cells = res["sensitivity"]
    assert [(c["red_slip_days"], c["tolerance_days"]) for c in cells] == [(r, t) for r in (1.0, 2.0, 3.0) for t in (0.0, 3.0, 7.0)]
    cur = [c for c in cells if c["current"]]
    assert len(cur) == 1 and (cur[0]["red_slip_days"], cur[0]["tolerance_days"]) == (1.0, 0.0)
    s = res["summary"]
    assert (cur[0]["samples"], cur[0]["tp"], cur[0]["fp"]) == (s["samples"], s["tp"], s["fp"])
    assert cur[0]["precision"] == s["precision"] and cur[0]["recall"] == s["recall"]
    c17 = next(c for c in cells if c["red_slip_days"] == 1.0 and c["tolerance_days"] == 7.0)
    assert (c17["tp"], c17["fp"], c17["fn"], c17["tn"]) == (5, 15, 10, 30)   # 容差 7 天：T 推迟 6 天不再算延误
    assert all(set(c) >= {"samples", "precision", "recall"} for c in cells)


def test_two_periods_marked_weak():
    res = _run([_period(0), _period(3)])
    assert res["summary"]["weak"] is True and res["summary"]["periods"] == 2
    assert any("弱口径" in n for n in res["summary"]["notes"])


def test_gaps_and_missing_excluded_from_matrix():
    classes = {"G": [(10, [25, 50, 100], [0, 0, 0])] * 5}
    unk = [_row(90, "U1", DD[0], 25, 0, 0, pct_type="CP_Units", work=(25, 75)),
           _row(90, "U1", DD[1], None, 0, 0, pct_type="CP_Units"),                       # 第 2 期完成%未知 → 数据缺口
           _row(90, "U1", DD[2], 100, 0, 0, pct_type="CP_Units", work=(100, 0))]
    gone = [_row(91, "X1", DD[0], 25, 0, 0), _row(91, "X1", DD[1], 50, 0, 0)]          # 最晚一期没有 → 末期缺失
    texts = [_period(k, classes, extra_rows=[unk[k]] + (gone[k:k + 1] if k < 2 else [])) for k in range(3)]
    res = _run(texts)
    s = res["summary"]
    assert any(g["code"] == "U1" and g["judge_date"] == DD[1] for g in res["gap_rows"])
    assert not any(r["code"] == "U1" and r["judge_date"] == DD[1] for r in res["rows"])
    assert not any(r["code"] == "X1" for r in res["rows"])
    assert s["data_gaps"] >= 1 and s["skipped"]["missing_in_last"] >= 1
    assert s["tp"] + s["fp"] + s["fn"] + s["tn"] == s["samples"]


def test_no_alert_and_small_samples_no_percentage():
    # 全部按期：没有红色预警 → 命中率不适用、不出百分比
    classes = {"G": [(10, [25, 50, 100], [0, 0, 0])] * 3}
    s = _run([_period(k, classes) for k in range(3)])["summary"]
    assert s["precision"] == {"k": 0, "n": 0, "rate": None} and s["no_alert_samples"] is True
    assert any("无预警样本" in n for n in s["notes"])
    # 只有 2 条红色预警：分母不足 5，不给百分比
    classes = {"T": [(0, [25, 30, 40, 100], [0, 0, 6, 6])] * 1, "F": [(0, [25, 30, 75, 100], [0, 0, 0, 0])] * 1}
    s = _run([_period(k, classes) for k in range(4)])["summary"]
    assert s["precision"]["n"] < 5 and s["precision"]["rate"] is None and s["precision"]["k"] >= 1
    assert any("样本不足" in n for n in s["notes"])


def test_periods_report_and_exclusion():
    texts = [_period(k) for k in range(4)]
    rep = _run(texts)["periods_report"]
    assert [p["data_date"] for p in rep] == DD
    assert all(p["pct_types"] == {"CP_Drtn": 20} for p in rep)
    assert [p["completed"] for p in rep] == [0, 0, 0, 15]
    assert rep[0]["overlap_prev"] is None and all(p["overlap_prev"] == 1.0 for p in rep[1:])
    assert not any(p["suspect"] for p in rep) and all(p["used"] for p in rep)
    # 插入一期大范围重排（全部作业推迟 45 天）：单独标出，可以剔除
    resched = {k: [(tf, pcts[:2] + [pcts[1]] + pcts[2:], delays[:2] + [45] + delays[2:]) for tf, pcts, delays in v] for k, v in CLASSES.items()}
    rdate = "2026-01-22"
    rows = []
    for i, (code, (tf, pcts, delays)) in enumerate((f"{p}{j}", spec) for p, specs in resched.items() for j, spec in enumerate(specs, 1)):
        rows.append(_row(i, code, rdate, pcts[2], delays[2], tf))
    rtext = _xer(rdate, rows)
    res = _run(texts + [rtext])
    r = next(p for p in res["periods_report"] if p["data_date"] == rdate)
    assert r["suspect"] and any("重排" in f for f in r["flags"])
    res2 = _run(texts + [rtext], exclude=[rdate])
    r2 = {p["data_date"]: p for p in res2["periods_report"]}
    assert r2[rdate]["excluded"] and not r2[rdate]["used"] and not r2[DD[2]]["suspect"]
    assert res2["summary"]["periods"] == 4 and res2["summary"]["excluded_periods"] == [rdate]
    assert (res2["summary"]["tp"], res2["summary"]["fp"]) == (15, 5)


def test_input_validation():
    from app.services import backtest as bt
    with pytest.raises(bt.BacktestError):
        _run([_period(0)])                                      # 只有 1 期
    with pytest.raises(bt.BacktestError):
        _run([_period(0), _period(0)])                          # 数据日期重复
    with pytest.raises(bt.BacktestError):
        _run([_period(0), _period(1)], exclude=[DD[0]])         # 剔除后不足 2 期
    with pytest.raises(bt.BacktestError):
        _run([_period(0), _period(1)], exclude=["2020-01-01"])  # 剔除的期次不存在
    with pytest.raises(bt.BacktestError):
        bt.parse_files([("a.xlsx", b"x"), ("b.xer", _period(1).encode())])
    assert bt.parse_exclude("2026-01-12，2026-01-19") == ["2026-01-12", "2026-01-19"]
    assert bt.parse_exclude('["2026-01-12"]') == ["2026-01-12"] and bt.parse_exclude("") == []
    with pytest.raises(bt.BacktestError):
        bt.parse_exclude("第3期")
    assert bt.normalize_source("real") == "real" and bt.normalize_source("sim") == "test" and bt.normalize_source("test") == "test"
    with pytest.raises(bt.BacktestError):
        bt.normalize_source("")
    with pytest.raises(bt.BacktestError):
        bt.normalize_source("whatever")


def test_service_never_calls_ai():
    src = (Path(__file__).resolve().parent.parent / "app" / "services" / "backtest.py").read_text(encoding="utf-8")
    assert "from ..ai" not in src and "import ai" not in src and "call_json" not in src


# ----------------------------------------------------------------------------- 接口：保存、来源、权限、导出、采用阈值

def _team(env):
    _add_user(env, "eng", "eng_sup")
    _add_user(env, "lead", "leader")
    _add_user(env, "ss", "safety_sup")


def _post_run(client, texts, **data):
    return client.post("/api/backtest/run", files=_files(texts), data=data)


def test_api_run_saves_with_source_and_reads_back(env):
    _team(env)
    eng = login("eng")
    with closing(dbmod.connect()) as conn:
        ai0 = conn.execute("SELECT COUNT(*) n FROM ai_runs").fetchone()["n"]
    r = _post_run(eng, [_period(k) for k in range(4)], source="real", label="LPR1", tolerance_days="0")
    assert r.status_code == 200, r.text
    v = r.json()
    assert set(v) >= {"id", "summary", "periods_report", "rows", "sensitivity", "lead_time", "source", "label"}
    assert v["source"] == "real" and v["source_name"] == "真实项目数据" and v["label"] == "LPR1"
    assert v["summary"]["precision"]["rate"] == 0.75
    with closing(dbmod.connect()) as conn:
        row = conn.execute("SELECT * FROM backtests WHERE id=?", (v["id"],)).fetchone()
        assert row["project_id"] == env["project_id"] and row["source"] == "real" and row["created_by"]
        assert conn.execute("SELECT COUNT(*) n FROM ai_runs").fetchone()["n"] == ai0          # AC-2.13 不调用 AI
        assert conn.execute("SELECT COUNT(*) n FROM events WHERE kind='backtest'").fetchone()["n"] == 0
        assert conn.execute("SELECT COUNT(*) n FROM audit WHERE action='backtest.run'").fetchone()["n"] == 1
    got = eng.get(f"/api/backtest/{v['id']}").json()
    assert got["summary"] == v["summary"] and got["sensitivity"] == v["sensitivity"] and got["rows"] == v["rows"]
    assert eng.get("/api/backtest/latest?source=real").json()["id"] == v["id"]
    assert eng.get("/api/backtest/latest?source=test").json() is None
    # 来源必填；仿真/测试来源单独标注
    assert _post_run(eng, [_period(0), _period(1)]).status_code == 400
    r2 = eng.post("/api/backtest/run", files=[("files[]", (f"v{i}.xer", _period(i).encode(), "application/octet-stream")) for i in range(2)],
                  data={"source": "sim"})
    assert r2.status_code == 200, r2.text
    assert r2.json()["source"] == "test" and r2.json()["source_name"] == "仿真/测试数据" and r2.json()["summary"]["weak"]
    assert eng.get("/api/backtest/latest?source=test").json()["id"] == r2.json()["id"]
    assert eng.get("/api/backtest/latest").json()["id"] == r2.json()["id"]
    assert eng.get("/api/backtest/latest?source=real").json()["id"] == v["id"]
    # 输入错误给 400，不落库
    assert _post_run(eng, [_period(0)], source="real").status_code == 400
    assert _post_run(eng, [_period(0), _period(1)], source="real", exclude_periods="2020-01-01").status_code == 400
    bad = eng.post("/api/backtest/run", files=[("files", ("a.txt", b"x", "text/plain")), ("files", ("b.xer", _period(1).encode(), "x"))],
                   data={"source": "real"})
    assert bad.status_code == 400
    ex = _post_run(eng, [_period(k) for k in range(4)], source="real", exclude_periods=DD[2])
    assert ex.status_code == 200 and ex.json()["summary"]["excluded_periods"] == [DD[2]]


def test_api_permissions(env):
    _team(env)
    eng = login("eng")
    bid = _post_run(eng, [_period(k) for k in range(3)], source="real").json()["id"]
    for who in ("pm", "lead", "tech1", "ss"):
        c = login(who)
        assert _post_run(c, [_period(0), _period(1)], source="real").status_code == 403, who
        assert c.get("/api/backtest/latest").status_code == 403, who
        assert c.get(f"/api/backtest/{bid}").status_code == 403, who
        assert c.get(f"/api/backtest/{bid}/docx").status_code == 403, who
        assert c.post(f"/api/backtest/{bid}/adopt", json={"red_slip_days": 2, "tolerance_days": 3}).status_code == 403, who
    adm = login("admin")
    assert adm.get(f"/api/backtest/{bid}").status_code == 200
    assert adm.get(f"/api/backtest/{bid}/docx").status_code == 200
    assert _post_run(adm, [_period(0), _period(1)], source="test").status_code == 200
    # 采用阈值是业务决定：管理员不能操作
    assert adm.post(f"/api/backtest/{bid}/adopt", json={"red_slip_days": 2, "tolerance_days": 3}).status_code == 403
    # 别的项目的回测记录不可见
    conn = dbmod.connect()
    pid2 = dbmod.insert(conn, "projects", {"name": "另一项目", "code": "T2", "config_json": "{}", "created_at": dbmod.now_iso()})
    conn.close()
    _add_user(env, "eng2", "eng_sup", pid2)
    other = login("eng2")
    assert other.get(f"/api/backtest/{bid}").status_code == 404
    assert other.post(f"/api/backtest/{bid}/adopt", json={"red_slip_days": 2, "tolerance_days": 3}).status_code == 404
    assert other.get("/api/backtest/latest").json() is None


def test_adopt_threshold_writes_settings_and_audit(env):
    _team(env)
    eng = login("eng")
    bid = _post_run(eng, [_period(k) for k in range(4)], source="real").json()["id"]
    assert eng.post(f"/api/backtest/{bid}/adopt", json={"red_slip_days": 5, "tolerance_days": 0}).status_code == 400  # 不在敏感性表里
    r = eng.post(f"/api/backtest/{bid}/adopt", json={"red_slip_days": 2, "tolerance_days": 3})
    assert r.status_code == 200, r.text
    assert r.json()["adopted"]["red_slip_days"] == 2.0 and r.json()["adopted"]["tolerance_days"] == 3.0
    with closing(dbmod.connect()) as conn:
        assert prog.thresholds(conn, env["project_id"])["red_slip_days"] == 2.0
        cfg = json.loads(conn.execute("SELECT config_json FROM projects WHERE id=?", (env["project_id"],)).fetchone()["config_json"])
        assert cfg["cadence"] == "day"  # 其他设置不动
        row = conn.execute("SELECT adopted_at, adopted_by FROM backtests WHERE id=?", (bid,)).fetchone()
        assert row["adopted_at"] and row["adopted_by"]
        assert conn.execute("SELECT COUNT(*) n FROM audit WHERE action='backtest.adopt' AND target=?", (f"backtest:{bid}",)).fetchone()["n"] == 1
    assert eng.get(f"/api/backtest/{bid}").json()["adopted"]["red_slip_days"] == 2.0


def _docx_text(raw: bytes) -> tuple[str, list]:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    return text + "\n" + "\n".join("\t".join(r) for t in tables for r in t), tables


def _kn(cell: str) -> tuple[int, int]:
    m = re.search(r"(\d+)/(\d+)", cell)
    return int(m.group(1)), int(m.group(2))


def test_docx_numbers_match_api(env):
    _team(env)
    eng = login("eng")
    v = _post_run(eng, [_period(k) for k in range(4)], source="real", label="LPR1").json()
    r = eng.get(f"/api/backtest/{v['id']}/docx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    text, tables = _docx_text(r.content)
    assert "真实项目数据" in text and "LPR1" in text and "不调用 AI" in text
    main = next(t for t in tables if t[0][:2] == ["指标", "作业×期次"])
    vals = {row[0]: row[1:] for row in main[1:]}
    s = v["summary"]
    for label, key in (("样本数", "samples"), ("TP（命中）", "tp"), ("FP（误报）", "fp"), ("FN（漏报）", "fn"), ("TN", "tn")):
        assert int(vals[label][0]) == s[key], label
        assert int(vals[label][1]) == s["strict"][key] and int(vals[label][2]) == s["by_activity"][key], label
    for label, key in (("命中率（主指标）", "precision"), ("召回率", "recall"), ("实际延误占比", "delay_share")):
        assert _kn(vals[label][0]) == (s[key]["k"], s[key]["n"]), label
        assert float(re.search(r"([\d.]+)%", vals[label][0]).group(1)) == round(s[key]["rate"] * 100, 1), label
    sens = next(t for t in tables if t[0][:2] == ["红色阈值（天）", "容差（天）"])
    assert len(sens) == 1 + 9
    for row, cell in zip(sens[1:], v["sensitivity"]):
        assert float(row[0]) == cell["red_slip_days"] and float(row[1]) == cell["tolerance_days"] and int(row[2]) == cell["samples"]
        assert _kn(row[3]) == (cell["precision"]["k"], cell["precision"]["n"]) and _kn(row[4]) == (cell["recall"]["k"], cell["recall"]["n"])
    # 仿真/测试来源在 Word 里如实标注
    t = _post_run(eng, [_period(0), _period(1)], source="test").json()
    text2, _ = _docx_text(eng.get(f"/api/backtest/{t['id']}/docx").content)
    assert "仿真/测试数据" in text2 and "弱口径" in text2


def test_cost_sentinel_never_stored_or_exported(env):
    """AC-2.6：含费用哨兵值的 XER 导入并回测后，库里所有列、AI 输入、导出的 Word 都找不到哨兵值。"""
    _team(env)
    num, word = "913579.2468", "ZBCOSTSENTINEL"
    cost_fields = ["cost_per_qty", "target_cost", "act_reg_cost", "act_ot_cost", "remain_cost", "act_this_per_cost"]
    extra = (
        ("TASKRSRC", ["taskrsrc_id", "task_id", "proj_id", "rsrc_id", "rsrc_type", "target_qty", "remain_qty", "act_reg_qty"] + cost_fields,
         [{"taskrsrc_id": "1", "task_id": "1000", "proj_id": "1", "rsrc_id": "9", "rsrc_type": "RT_Mat", "target_qty": "10", "remain_qty": "5",
           "act_reg_qty": "5", **{f: num for f in cost_fields}}]),
        ("RSRC", ["rsrc_id", "rsrc_short_name", "rsrc_type", "unit_id"], [{"rsrc_id": "9", "rsrc_short_name": "MAT", "rsrc_type": "RT_Mat", "unit_id": "1"}]),
        ("RSRCRATE", ["rsrc_rate_id", "rsrc_id"] + [f"cost_per_qty{i}" for i in ("", "2", "3", "4", "5")],
         [{"rsrc_rate_id": "1", "rsrc_id": "9", **{f"cost_per_qty{i}": num for i in ("", "2", "3", "4", "5")}}]),
        ("PROJCOST", ["proj_cost_id", "proj_id", "task_id", "cost_name", "target_cost", "act_cost"],
         [{"proj_cost_id": "1", "proj_id": "1", "task_id": "1000", "cost_name": word, "target_cost": num, "act_cost": num}]),
    )
    wbs_extra = {"orig_cost": num, "indep_remain_total_cost": word}

    def period(k):
        rows, i = [], 0
        for prefix, specs in CLASSES.items():
            for j, (tf, pcts, delays) in enumerate(specs, 1):
                rows.append(_row(i, f"{prefix}{j}", DD[k], pcts[k], delays[k], tf, extra={"target_cost": num, "cost_note": word}))
                i += 1
        return _xer(DD[k], rows, extra_tables=extra, wbs_extra=wbs_extra)

    texts = [period(k) for k in range(4)]
    eng = login("eng")
    assert eng.post("/api/p6/import", files={"file": ("plan.xer", texts[2].encode(), "application/octet-stream")}).status_code == 200
    r = _post_run(eng, texts, source="real", label="SENT")
    assert r.status_code == 200, r.text
    blobs = [r.text, eng.get("/api/backtest/latest").text, _docx_text(eng.get(f"/api/backtest/{r.json()['id']}/docx").content)[0]]
    for b in blobs:
        assert num not in b and "913579" not in b and word not in b
    with closing(sqlite3.connect(str(dbmod.settings.db_path))) as raw:
        tables = [t for (t,) in raw.execute("SELECT name FROM sqlite_master WHERE type='table'")]
        assert "ai_runs" in tables and "backtests" in tables
        hits = []
        for t in tables:
            cols = [c[1] for c in raw.execute(f"PRAGMA table_info({t})")]
            for row in raw.execute(f"SELECT * FROM {t}"):
                for c, val in zip(cols, row):
                    sv = val.decode("utf-8", "ignore") if isinstance(val, bytes) else str(val)
                    if "913579" in sv or word in sv:
                        hits.append(f"{t}.{c}")
        assert hits == []


def test_metrics_shows_only_latest_real_backtest(env):
    """AC-2.12、AC-2.14（metrics 部分）：应用成效只引用最新一条“真实”回测，仿真/测试的不计入。"""
    _team(env)
    eng, adm = login("eng"), login("admin")
    _post_run(eng, [_period(k) for k in range(3)], source="test")
    assert adm.get("/api/metrics").json()["backtest"] is None
    real = _post_run(eng, [_period(k) for k in range(4)], source="real", label="LPR1").json()
    _post_run(eng, [_period(k) for k in range(3)], source="test")
    m = adm.get("/api/metrics").json()["backtest"]
    assert m["id"] == real["id"] and m["source"] == "real" and m["label"] == "LPR1"
    assert m["samples"] == real["summary"]["samples"] and m["precision"] == real["summary"]["precision"]["rate"]
    assert m["periods"] == 4 and m["data_date_from"] == DD[0] and m["data_date_to"] == DD[3]
    assert m["source_note"] == "真实数据，不在仿真时间线内"


def test_old_endpoint_removed_and_run_does_not_write_events(env):
    """旧接口 POST /api/p6/backtest 已下线（前端不再调用）；原 AC-2.14 的口径由 /api/backtest/run 承担：
    结果写 backtests、必须标来源；不再写 events；项目经理、领导 403；库里残留的旧 events(kind='backtest') 不进入应用成效。"""
    _team(env)
    files = _files([_period(k) for k in range(4)])
    assert login("eng").post("/api/p6/backtest", files=files, data={"source": "test"}).status_code in (404, 405)
    for who in ("pm", "lead"):
        assert login(who).post("/api/backtest/run", files=files, data={"source": "test"}).status_code == 403, who
    conn = dbmod.connect()
    dbmod.event(conn, "backtest", project_id=env["project_id"], data={"samples": 99, "accuracy": 0.9})  # 旧版遗留记录
    conn.close()
    eng = login("eng")
    assert eng.post("/api/backtest/run", files=files).status_code == 400  # 不再默认记为测试来源，必须选来源
    r = eng.post("/api/backtest/run", files=files, data={"source": "test"})
    assert r.status_code == 200, r.text
    v = r.json()
    s = v["summary"]
    assert v["source"] == "test" and s["samples"] == 60 and s["precision"]["rate"] == 0.75 and not s["weak"]
    with closing(dbmod.connect()) as c:
        row = c.execute("SELECT source FROM backtests WHERE id=?", (v["id"],)).fetchone()
        assert row["source"] == "test"
        assert c.execute("SELECT COUNT(*) n FROM events WHERE kind='backtest'").fetchone()["n"] == 1  # 只有上面那条遗留
    assert login("admin").get("/api/metrics").json()["backtest"] is None
    r2 = eng.post("/api/backtest/run", files=files, data={"source": "real"})
    assert r2.status_code == 200 and r2.json()["source"] == "real"
    m = login("admin").get("/api/metrics").json()["backtest"]
    assert m["id"] == r2.json()["id"] and m["link"] == f"#/backtest?id={m['id']}"
    assert login("admin").post("/api/backtest/run", files=[files[0]], data={"source": "test"}).status_code == 400


def test_adoption_definition_includes_onspot(env):
    """T-05 遗留：采纳口径写明含当场纠正（D6）；应用成效页的口径说明直接渲染接口里的 definitions。"""
    d = login("admin").get("/api/metrics").json()["ai_adoption"]["definitions"]["采纳率/驳回率"]
    assert "含当场纠正（D6）" in d
    mgr = (Path(__file__).resolve().parent.parent / "app" / "static" / "mgr.js").read_text(encoding="utf-8")
    assert 'v-for="(v, k) in a.definitions"' in mgr and "<adoption-table></adoption-table>" in mgr


# ----------------------------------------------------------------------------- 命令行（离线，不写库）

def test_cli_backtest_offline(env, capsys):
    import tempfile

    from app.__main__ import main
    folder = Path(tempfile.mkdtemp(prefix="zb_bt_cli_"))
    paths = []
    for k in range(4):
        p = folder / f"u{k}.xer"
        p.write_text(_period(k), encoding="utf-8")
        paths.append(str(p))
    assert main(["backtest", *paths]) == 0
    out = json.loads(capsys.readouterr().out)
    assert out["summary"]["samples"] == 60 and out["summary"]["precision"]["rate"] == 0.75 and "rows" not in out
    assert len(out["sensitivity"]) == 9 and len(out["periods_report"]) == 4
    assert main(["backtest", *paths, "--exclude", DD[2], "--rows"]) == 0
    out2 = json.loads(capsys.readouterr().out)
    assert out2["summary"]["excluded_periods"] == [DD[2]] and out2["rows"]
    assert main(["backtest", paths[0]]) == 2
    with closing(dbmod.connect()) as conn:
        assert conn.execute("SELECT COUNT(*) n FROM backtests").fetchone()["n"] == 0
